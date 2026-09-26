"""The PRICE box's bullets and indents: Kyle's REBID layout by default, the ribbon on every price
line, and the editor and the customer's document agreeing line for line.

Hanz, 2026-09-26: "fix the indents and the bullets now". Before it, on staging: "I couldnt add an
indendt and bullet" and "also the bullet point and the indent are not working in the pricing box".
Two things were true at once: the ribbon let go of every price line the page composes (it was a
channel the ribbon could not reach), and the render stripped every PRICE-list bullet whatever
anybody set (`_flatten_price_bullets`, PR #132, the 2026-07-16 "no bullets in the pricing" rule).

Hanz reversed that rule on 2026-09-25, with it in front of him: the price box reads like Kyle's
hand-made "Nickell RC Sustainment REBID" proposal. A red square on every money line (the base, its
tax rows and Total, each option and its own rows, the manual and combo lines), the hollow "o" on an
option's sub-lines, nothing on a heading or a blank line — and what the estimator sets with the
ribbon on any line beats that default and prints.

WHAT THIS FILE PROVES, BY RUNNING BOTH HALVES:

  (1) THE RULE, in both languages (price_rules.py / price-lines-core.js), over one matrix.
  (2) THE SCREEN AND THE PAPER AGREE, line for line: js/price-bullets-harness.js builds the editor's
      PRICE box through the page's own code over the real template's blocks, presses the ribbon's
      buttons, and hands back what it shows and the payload it would send; the payload goes
      through the real renderer and the .docx's PRICE box is read back paragraph by paragraph.
      Every Direct template (epoxy, polish, combo, budget), two GC files and Gyp.
  (3) THE PAPER ITSELF — the four ways the first REBID attempt would have gone wrong (the
      format-verdict's REBID hazards), each as a docx-level check.
"""
import copy
import io
import itertools
import json
import pathlib
import shutil
import subprocess

import docx
import pytest
from docx.oxml.ns import qn
from starlette.requests import Request

import main
import price_rules
import proposal_writer as pw

FRONTEND = pathlib.Path(__file__).resolve().parents[2] / "frontend"
HARNESS = pathlib.Path(__file__).resolve().parent / "js" / "price-bullets-harness.js"
CORE = FRONTEND / "js" / "price-lines-core.js"

needs_node = pytest.mark.skipif(shutil.which("node") is None, reason="node is not installed")

W = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"
MC = "{http://schemas.openxmlformats.org/markup-compatibility/2006}"


# ── (1) ONE RULE, TWO LANGUAGES ──────────────────────────────────────────────────────────────────
KEYS = ["base", "sales_tax", "total", "option:Alt1", "option:Alt1:total", "manual:0",
        "combo:epoxy.flooring", "heading_base", "heading_options", "alt_name", "alt_total"]
POSITIONS = [None, "before", "after"]
TEXTS = ["$22,600 – Resinous flooring", "", "   "]
OVERRIDES = [None, {}, {"bullet": False}, {"bullet": False, "indent": 576}, {"bullet": True},
             {"bullet": True, "level": 1}, {"level": 1}, {"bullet": True, "level": 0},
             {"bullet": False, "level": 1}, {"indent": 99999}, {"indent": -5},
             {"bullet": "yes", "level": 3, "indent": "12"}, {"bullet": True, "level": True}, "garbage"]


def _node(cases):
    p = subprocess.run(["node", "-e", """
      const P = require(process.argv[1]);
      const cases = JSON.parse(require("fs").readFileSync(0, "utf8"));
      console.log(JSON.stringify(cases.map(c => {
        if (c.kind === "clean") return P.cleanLineProps(c.o);
        if (c.kind === "default") return P.lineDefault(c.key, c.pos);
        if (c.kind === "intent") return P.lineIntent(c.key, c.pos, c.o);
        if (c.kind === "resolve") return P.resolveLineProps(c.key, c.pos, c.text, c.o);
        if (c.kind === "step") return P.paraStep(c.it, c.action);
        if (c.kind === "store") return P.overrideFor(c.key, c.pos, c.it);
        if (c.kind === "consts") return { left: P.LEVEL_LEFT, hang: P.LEVEL_HANG, step: P.INDENT_STEP,
                                          max: P.INDENT_MAX, heads: P.HEADING_KEYS };
      })));
    """, str(CORE)], input=json.dumps(cases), capture_output=True, text=True, encoding="utf-8")
    assert p.returncode == 0, p.stderr
    return json.loads(p.stdout)


@needs_node
def test_the_line_rule_is_the_same_rule_in_both_languages():
    cases = []
    for o in OVERRIDES:
        cases.append({"kind": "clean", "o": o})
    for k, pos in itertools.product(KEYS, POSITIONS):
        cases.append({"kind": "default", "key": k, "pos": pos})
        for o in OVERRIDES:
            cases.append({"kind": "intent", "key": k, "pos": pos, "o": o})
            for t in TEXTS:
                cases.append({"kind": "resolve", "key": k, "pos": pos, "text": t, "o": o})
    js = _node(cases)
    for c, j in zip(cases, js):
        if c["kind"] == "clean":
            py = price_rules.clean_line_props(c["o"])
        elif c["kind"] == "default":
            py = price_rules.line_default(c["key"], c["pos"])
        elif c["kind"] == "intent":
            py = price_rules.line_intent(c["key"], c["pos"], c["o"])
        else:
            py = price_rules.resolve_line_props(c["key"], c["pos"], c["text"], c["o"])
        assert j == py, c
    consts = _node([{"kind": "consts"}])[0]
    consts["heads"] = sorted(consts["heads"])
    assert consts == {"left": list(price_rules.LEVEL_LEFT_TW), "hang": list(price_rules.LEVEL_HANG_TW),
                      "step": price_rules.LINE_INDENT_STEP_TW, "max": price_rules.LINE_INDENT_MAX_TW,
                      "heads": sorted(price_rules.HEADING_LINE_KEYS)}


def test_the_rebid_default_is_kyles_layout():
    """Money lines a square, sub-lines the "o", headings nothing — and a blank line never a bullet,
    whatever anybody set (a lone square on an empty line is what Kyle's test click would print)."""
    assert price_rules.resolve_line_props("base", None, "$1 – x") == {"bullet": True, "level": 0}
    assert price_rules.resolve_line_props("option:A:total", None, "$1 – Total") == {"bullet": True, "level": 0}
    assert price_rules.resolve_line_props("option:A", "after", "Notes: x") == {"bullet": True, "level": 1}
    assert price_rules.resolve_line_props("heading_options", None, "Options:") == {"bullet": False, "indent": 0}
    assert price_rules.resolve_line_props("heading_options", "before", "typed on the gap") == {
        "bullet": False, "indent": 0}
    assert price_rules.resolve_line_props("alt_name", None, "ALTERNATE") == {"bullet": False, "indent": 0}
    for o in (None, {"bullet": True}, {"bullet": True, "level": 1}):
        assert price_rules.resolve_line_props("base", "after", "", o)["bullet"] is False
        assert price_rules.resolve_line_props("base", "after", "  ", o)["bullet"] is False
    # Off sticks, and the words stay where the bullet had them.
    assert price_rules.resolve_line_props("base", None, "$1 – x", {"bullet": False}) == {
        "bullet": False, "indent": 288}
    assert price_rules.resolve_line_props("option:A", "after", "x", {"bullet": False}) == {
        "bullet": False, "indent": 1440}


@needs_node
def test_every_ribbon_press_stores_the_smallest_override_that_says_it():
    """paraStep is one press; overrideFor is what the draft keeps. For every line kind, every
    starting state and every press, storing the result and reading it back through the rule gives
    exactly the stepped state — and a press that lands back on the default stores nothing."""
    keys = ["base", "option:Alt1", "heading_options"]
    starts = [None, {"bullet": False}, {"bullet": True, "level": 1}, {"bullet": False, "indent": 864},
              {"bullet": False, "indent": 0}, {"bullet": False, "indent": 2880}]
    cases, meta = [], []
    for k, pos, o in itertools.product(keys, POSITIONS, starts):
        it = price_rules.line_intent(k, pos, o)
        for action in ("bullet", "indent", "outdent"):
            cases.append({"kind": "step", "it": it, "action": action})
            meta.append((k, pos, o, it, action))
    steps = _node(cases)
    stores = _node([{"kind": "store", "key": k, "pos": pos, "it": s}
                    for (k, pos, _o, _it, _a), s in zip(meta, steps) if s is not None])
    si = iter(stores)
    for (k, pos, o, it, action), s in zip(meta, steps):
        if action == "bullet":
            assert s is not None and s["bullet"] is (not it["bullet"]), (k, pos, o)
            if not it["bullet"]:
                assert s["indent"] == price_rules.LEVEL_LEFT_TW[s["level"]]
            else:
                assert s["indent"] == it["indent"], "switching a bullet off moved the words"
        if action == "indent" and it["bullet"]:
            assert (s is None) == (it["level"] == 1) and (s is None or s["level"] == 1), (k, pos, o)
        if action == "outdent" and it["bullet"]:
            assert (s is None) == (it["level"] == 0) and (s is None or s["level"] == 0), (k, pos, o)
        if not it["bullet"] and action in ("indent", "outdent"):
            want = it["indent"] + (288 if action == "indent" else -288)
            if 0 <= want <= 2880:
                assert s is not None and s["indent"] == want and s["bullet"] is False
            else:
                assert s is None
        if s is None:
            continue
        stored = next(si)
        back = price_rules.line_intent(k, pos, stored)
        assert (back["bullet"], back["indent"]) == (s["bullet"], s["indent"]), (k, pos, o, action, stored)
        if s["bullet"]:
            assert back["level"] == s["level"]
        if (s["bullet"], s["indent"]) == (price_rules.line_intent(k, pos, None)["bullet"],
                                          price_rules.line_intent(k, pos, None)["indent"]) and (
                not s["bullet"] or s["level"] == price_rules.line_intent(k, pos, None)["level"]):
            assert stored is None, ("a press back to the default stored an override", k, pos, stored)


def test_every_template_defines_the_price_list_levels_the_rule_assumes():
    """The rule places a bulleted line by its list level alone: square at 0 / text at 288 twips on
    level 0, "o" at 1080 / text at 1440 on level 1. True of numId 3 in every shipped template."""
    for rel in sorted(set(pw.TEMPLATE_PICKER.values())):
        d = docx.Document(str(pw.TEMPLATES_ROOT / rel))
        lv = pw._numbering_levels(d)
        for level in (0, 1):
            got = lv.get(("3", str(level))) or {}
            assert got.get("fmt") == "bullet", (rel, level, got)
            ind = got.get("ind") or {}
            assert int(ind.get("left")) == price_rules.LEVEL_LEFT_TW[level], (rel, level, ind)
            assert int(ind.get("hanging")) == price_rules.LEVEL_HANG_TW[level], (rel, level, ind)
        assert lv[("3", "1")]["text"] == "o", rel


# ── (2) THE SCREEN AND THE PAPER, LINE FOR LINE ──────────────────────────────────────────────────
def _template_json(work_type, audience):
    req = Request({"type": "http", "method": "GET", "path": "/api/proposal-template", "headers": [],
                   "query_string": b""})
    return json.loads(main.api_proposal_template(req, work_type=work_type, audience=audience).body)


def _price_box(tj):
    """The PRICE text box's blocks, in the editor's order — the box holding the base line (or, on
    the budget sheet, its per-sf lines)."""
    idx = next(b["txbx"] for b in tj["blocks"]
               if b.get("txbx") is not None and ("base_bid_formatted" in (b["text"] or "")
                                                 or "/ sf" in (b["text"] or "")))
    return [b for b in tj["blocks"] if b.get("txbx") == idx]


def _mark_sz(p_elem):
    """The paragraph mark's size in half-points — what an EMPTY line is as tall as."""
    ppr = p_elem.find(W + "pPr")
    sz = ppr.find(W + "rPr/" + W + "sz") if ppr is not None else None
    return int(sz.get(W + "val")) if sz is not None else None


# Kyle's files carry 1pt spacer paragraphs between some rows (sz 2); on Gyp one sits INSIDE the
# {{#has_options}} region, where the editor draws the region's own lines and not its spacer. They
# are not lines anybody reads or types on, so both sides leave them out of the comparison.
HAIRLINE_SZ = 8


def _style_left(d, sid):
    for _ in range(10):
        if not sid:
            return None
        s = next((x for x in d.styles.element.findall(W + "style") if x.get(W + "styleId") == sid), None)
        if s is None:
            return None
        ind = s.find(W + "pPr/" + W + "ind")
        if ind is not None:
            for a in ("start", "left"):
                if ind.get(W + a) is not None:
                    return int(float(ind.get(W + a)))
        b = s.find(W + "basedOn")
        sid = b.get(W + "val") if b is not None else None
    return None


def _doc_lines(blob):
    """The rendered PRICE box, paragraph by paragraph, as Word lays it out: the bullet (a numbering
    level that prints one), its level, where the text starts (the paragraph's own indent, else its
    list level's, else its STYLE's — Word walks the style, so this does too), and how far the first
    line starts from there."""
    d = docx.Document(io.BytesIO(blob))
    levels = pw._numbering_levels(d)
    box = None
    for tx in d.element.body.iter(W + "txbxContent"):
        if any(True for _ in tx.iterancestors(MC + "Fallback")):
            continue
        txt = [pw._own_text(p).strip() for p in tx.findall(W + "p")]
        if any(t.startswith(("Base Bid", "Floor Budget")) or (t.startswith("$") and "Option 1" in t)
               for t in txt):
            box = tx
            break
    assert box is not None, "no PRICE box in the rendered document"
    out = []
    for p in box.findall(W + "p"):
        text = pw._own_text(p).rstrip()
        blank = not text.strip()
        ref = pw._para_num_ref(p)
        lv = levels.get(ref) if ref else None
        bullet = bool(lv and lv.get("fmt") == "bullet")
        ppr = p.find(W + "pPr")
        ind = ppr.find(W + "ind") if ppr is not None else None
        own = {a: int(float(ind.get(W + a))) for a in ("left", "start", "hanging", "firstLine")
               if ind is not None and ind.get(W + a) is not None}
        lvl_ind = {k: int(v) for k, v in ((lv or {}).get("ind") or {}).items()} if bullet else {}
        if "start" in own or "left" in own:
            left = own.get("start", own.get("left"))
        elif "left" in lvl_ind or "start" in lvl_ind:
            left = lvl_ind.get("start", lvl_ind.get("left"))
        else:
            ps = ppr.find(W + "pStyle") if ppr is not None else None
            left = (_style_left(d, ps.get(W + "val")) if ps is not None else None) or 0
        first = 0 if bullet else own.get("firstLine", 0) - own.get("hanging", 0)
        out.append({"text": text, "blank": blank, "bullet": bullet and not blank,
                    "level": int(ref[1]) if bullet and not blank else None,
                    "indent_tw": left, "first_tw": first, "numbered": ref is not None,
                    "sz": _mark_sz(p), "p": p})
    return out


def _render(body):
    req = Request({"type": "http", "method": "POST", "path": "/t", "headers": [], "query_string": b""})
    return main._render_documents(body, req, want_estimate=False)["docx"]["content"]


def _state(work_type, audience, **over):
    """Kyle's REBID figures: a $24,700 base ($22,600 + $350 material sales tax + $1,750 remodel),
    one priced option ($85,500, taxable, no remodel) with three lines typed under it — two sub-lines
    and a blank line between them, the way his REBID lists Notes / Schedule under Alternate #1."""
    st = {
        "work_type": work_type, "audience": audience, "project_name": "Nickell RC", "job_name": "Nickell RC",
        "proposal_lump_sum": 24700, "proposal_sales_tax": 350, "proposal_remodel_tax": 1750,
        "proposal_taxable": True, "proposal_remodel_on": True, "tax_layout": "BROKEN_OUT",
        "priced_tabs": [], "cell_values": {}, "sheet_area": {"epoxy_sf": 885, "polish_sf": 885, "cove_lf": 100},
        "rooms": [
            {"id": "Epoxy", "name": "Epoxy", "is_base": True,
             "bid": {"total": 24700, "sales_tax": 350, "remodel": 1750}},
            {"id": "Alt1", "name": "Alt 1", "is_base": False, "show": True, "price_mode": "total",
             "system_desc": "Alternate #1: Add Restrooms", "option_desc": "Alternate #1: Add Restrooms",
             "base_total": 24700,
             "bid": {"total": 85500, "sales_tax": 500, "remodel": 0, "taxable": True, "remodel_on": False}},
        ],
        "price_overrides": {"after": {"option:Alt1": ["Notes: Areas per Schedule Note 1", "",
                                                      "Schedule: same 2 mobs/phases"]}},
    }
    st.update(over)
    return st


def _typed(pos, idx):
    return {"key": "option:Alt1", "kind": "extra", "pos": pos, "idx": idx}


# Per template: the draft it starts from and the presses made on it. Every template gets a bullet
# switched OFF, a line moved to the "o" (indent), and a line brought back (outdent); every file
# with typed lines gets an "o" sub-line and a blank typed line, and every file with an Options
# heading its gap.
CASES = {
    ("epoxy", "Direct"): {
        "state": {"price_overrides": {
            "after": {"option:Alt1": ["Notes: Areas per Schedule Note 1", "", "Schedule: same 2 mobs/phases"]},
            # A line typed ON the gap above "Options:" — a heading's line, so no bullet.
            "before": {"heading_options": ["Pricing valid 30 days"]}, "options_gap": 1}},
        # Through the RIBBON here — the caret lands on the line, the page's focusin aims the
        # ribbon, the button's click handler runs (test_the_ribbon_acts_on_the_price_line_it_is_on).
        "actions": [
            {"ribbon": {"key": "base"}, "click": "bullet"},                # off: sticks
            {"ribbon": {"key": "sales_tax"}, "click": "indent"},           # square -> "o"
            {"on": _typed("after", 2), "press": "outdent"},                # "o" -> square
            {"ribbon": _typed("after", 0), "click": "bullet"},             # "o" off, words stay
            {"ribbon": _typed("after", 0), "click": "indent"},             # ...and in one more step
            {"ribbon": {"key": "remodel"}, "click": "indent"},             # moved...
            {"ribbon": {"key": "remodel"}, "click": "reset"},              # ...and Reset: the default
            {"ribbon": {"key": "heading_options"}, "click": "bullet"},     # a heading gets one
            {"enter_after": {"key": "total"}, "type": "Includes 1 gallon patch"},   # Enter + type
        ],
    },
    ("polish", "Direct"): {
        "actions": [
            {"on": {"starts": "$22,600"}, "press": "indent"},              # the TEMPLATE base line
            {"on": {"key": "total"}, "press": "bullet"},
            {"on": _typed("after", 2), "press": "outdent"},
            {"enter_after": {"key": "option:Alt1:total"}, "type": "Exclusions: wall coating"},
        ],
    },
    ("combo", "Direct"): {
        "state": {"rooms": [], "priced_tabs": [
            {"id": "Epoxy", "role": "epoxy", "kind": "base", "total": 24700, "sales_tax": 350,
             "remodel": 1750, "taxable": True, "remodel_on": True},
            {"id": "Polish", "role": "polish", "kind": "base", "total": 5000, "sales_tax": 50,
             "remodel": 20, "taxable": True, "remodel_on": True}],
            "price_overrides": {"after": {"combo:epoxy.flooring": ["Includes moisture mitigation", ""]}}},
        "actions": [
            {"on": {"key": "combo:polish.total"}, "press": "bullet"},
            {"on": {"key": "combo:epoxy.sales_tax"}, "press": "indent"},
            {"on": {"key": "combo:epoxy.flooring", "kind": "extra", "pos": "after", "idx": 0}, "press": "outdent"},
            {"enter_after": {"key": "combo:polish.flooring"}, "type": "Grind and seal, 2 passes"},
        ],
    },
    ("budget", "Direct"): {
        "state": {"rooms": [], "price_overrides": {}},
        "actions": [
            {"on": {"starts": "$1.50"}, "press": "bullet"},
            {"on": {"starts": "$2.25"}, "press": "indent"},
            {"on": {"starts": "$5.75"}, "press": "indent"},
            {"on": {"starts": "$5.75"}, "press": "outdent"},
            {"on": {"starts": "Add On"}, "press": "bullet"},               # a heading joins the list
        ],
    },
    ("epoxy", "GC"): {
        "actions": [
            # The FIRST row of the box to the "o" — then a heading's bullet switched on. The
            # document puts a new bullet on its siblings' level, so without the level riding the
            # patch the heading would print an "o" the editor never showed.
            {"on": {"starts": "$22,600"}, "press": "indent"},
            {"on": {"starts": "Options & Unit Prices"}, "press": "bullet"},
            {"on": {"starts": "$350"}, "press": "bullet"},
            {"on": {"starts": "If a different"}, "press": "bullet"},
            {"on": {"starts": "If a different"}, "press": "indent"},
        ],
    },
    ("polish", "GC"): {
        "actions": [
            # Kyle's own "o" row, authored at left 1170 / hanging 180: its bullet off must not
            # leave the first line hanging left of the rest.
            {"on": {"starts": "Note: Additional"}, "press": "bullet"},
            {"on": {"starts": "$1,100"}, "press": "indent"},
            {"on": {"starts": "($x)"}, "press": "indent"},
            {"on": {"starts": "($x)"}, "press": "outdent"},
        ],
    },
    ("gyp", "Direct"): {
        "actions": [
            {"on": {"starts": "$350"}, "press": "indent"},
            {"on": {"key": "option:Alt1"}, "press": "bullet"},
            {"on": _typed("after", 2), "press": "outdent"},
        ],
    },
}


def _case(work_type, audience, spec=None, name=None):
    spec = CASES[(work_type, audience)] if spec is None else spec
    tj = _template_json(work_type, audience)
    st = _state(work_type, audience)
    for k, v in (spec.get("state") or {}).items():
        st[k] = copy.deepcopy(v)
    return {"name": name or f"{work_type}/{audience}", "work_type": work_type, "audience": audience,
            "blocks": _price_box(tj), "options_heading_ids": tj.get("options_heading_ids") or [],
            "state": st, "actions": spec.get("actions") or []}


def _run_cases(cases):
    """The editor (harness) and the document (renderer) for each case, in order."""
    p = subprocess.run(["node", str(HARNESS), str(FRONTEND)], input=json.dumps(cases),
                       capture_output=True, text=True, encoding="utf-8")
    assert p.returncode == 0, "the harness itself failed:\n" + p.stderr
    out = []
    for c, r in zip(cases, json.loads(p.stdout)):
        pl = r["payload"]
        body = {"work_type": c["work_type"], "audience": c["audience"], "values": pl["values"],
                "price_overrides": pl["price_overrides"],
                "paragraph_overrides": pl["paragraph_overrides"],
                "combo_options": pl["combo_options"], "rooms": c["state"]["rooms"],
                "alternate_computed_bid": c["state"].get("alternate_computed_bid"),
                "alternate_label": (c["state"].get("alternate") or {}).get("label", "")}
        blob = _render(body)
        out.append({"case": c, "editor": r, "doc": _doc_lines(blob), "blob": blob, "body": body})
    return out


_RUNS: dict = {}


@pytest.fixture(scope="module")
def runs():
    """Every case: the editor (harness) and the document (renderer), once."""
    if shutil.which("node") is None:
        pytest.skip("node is not installed")
    if _RUNS:
        return _RUNS
    for run in _run_cases([_case(wt, aud) for wt, aud in CASES]):
        _RUNS[(run["case"]["work_type"], run["case"]["audience"])] = run
    return _RUNS


def _editor_sz(case, line):
    """A template block's paragraph-mark size, off the pristine template (the editor draws a blank
    block at the size of its mark); None for the lines the page composes (9pt price rows)."""
    if line.get("id") is None:
        return None
    d = docx.Document(str(pw.pick_template(case["work_type"], case["audience"])))
    for i, _k, p, _b, _t, _x in pw.iter_editable_blocks(d):
        if i == line["id"]:
            return _mark_sz(p)
    return None


def _cmp(line):
    """One line, as compared: the words, blank or not, bullet, level, where the words start."""
    if line["blank"]:
        return ("", "blank")
    return (line["text"].rstrip(), line["bullet"], line["level"], line["indent_tw"],
            0 if line["bullet"] else line["first_tw"])


# The one WORDING difference the price box still has, older than this fix and about words, not
# bullets, so it is named here rather than hidden: the editor's Options heading reads "Options:" on
# every file, and the Polish Direct file's own heading reads "Options " (the heading is painted
# from the draft once he types one; see fix 5-6). Its bullet and indent are still compared.
_KNOWN_WORDING = {("polish", "Direct"): {"Options:": "Options"}}


def _shown(run):
    case = run["case"]
    hair = {}
    lines = []
    for ln in run["editor"]["lines"]:
        if ln["blank"] and ln.get("id") is not None:
            sz = hair.setdefault(ln["id"], _editor_sz(case, ln))
            if sz is not None and sz < HAIRLINE_SZ:
                continue
        c = _cmp(ln)
        fix = _KNOWN_WORDING.get((case["work_type"], case["audience"]), {})
        if not ln["blank"] and c[0] in fix:
            c = (fix[c[0]],) + c[1:]
        lines.append(c)
    return lines


def _printed(run):
    return [_cmp(ln) for ln in run["doc"]
            if not (ln["blank"] and ln["sz"] is not None and ln["sz"] < HAIRLINE_SZ)]


@needs_node
@pytest.mark.parametrize("work_type,audience", list(CASES))
def test_the_editor_and_the_document_agree_line_for_line(runs, work_type, audience):
    """THE PROPERTY: what the estimator sees in the price box is what the customer reads — the
    words, the bullet, which bullet, where the words start, and the blank lines between them."""
    run = runs[(work_type, audience)]
    assert all(p["done"] for p in run["editor"]["pressed"] if "done" in p), run["editor"]["pressed"]
    shown, printed = _shown(run), _printed(run)
    assert not [ln for ln in run["editor"]["lines"] if ln["unstyled"]], (
        "a shown line carries no geometry of its own, so the stylesheet decides where it sits")
    assert shown == printed, "\n".join(
        f"{'  ' if a == b else '!!'} screen {a!r}\n{'  ' if a == b else '!!'} paper  {b!r}"
        for a, b in itertools.zip_longest(shown, printed))


@needs_node
@pytest.mark.parametrize("work_type,audience", list(CASES))
def test_no_blank_line_carries_a_bullet(runs, work_type, audience):
    """Word prints a lone square on an EMPTY list paragraph; the editor draws none. Not on the
    template's own spacers, not on a blank line he typed, not on the Options gap."""
    run = runs[(work_type, audience)]
    stray = [i for i, ln in enumerate(run["doc"]) if ln["blank"] and ln["numbered"]]
    assert not stray, f"blank paragraph(s) {stray} still carry numbering"
    assert not [ln for ln in run["editor"]["lines"] if ln["blank"] and ln["bullet"]]


@needs_node
@pytest.mark.parametrize("work_type,audience", list(CASES))
def test_what_the_ribbon_set_survives_a_reload(runs, work_type, audience):
    """A new page, drawn from nothing but the saved draft, shows the same box — and composes the
    same payload, so a reload cannot lose it on the next generate either."""
    r = runs[(work_type, audience)]["editor"]
    assert r["reload"] == r["lines"]
    assert r["reloadPayload"]["paragraph_overrides"] == r["payload"]["paragraph_overrides"]
    for k in ("line_props", "before_props", "after_props", "before", "after"):
        assert (r["reloadPayload"]["price_overrides"].get(k) or {}) == (r["payload"]["price_overrides"].get(k) or {}), k


@needs_node
def test_the_ribbon_acts_on_the_price_line_it_is_on(runs):
    """"I couldnt add an indendt and bullet": the ribbon used to let go of a price line, so every
    button did nothing there. Now the caret landing on one aims the ribbon at it; Bullet / Indent /
    Outdent / Reset are live and press it, and Bold / Italic / Underline / the size are switched
    OFF, because the line's channel stores words, not runs — a bold would show and print plain."""
    r = runs[("epoxy", "Direct")]["editor"]
    assert len(r["ribbon"]) == 7
    for press in r["ribbon"]:
        aimed = press["aimed"]
        for k in ("bold", "italic", "underline", "size"):
            assert aimed[k]["disabled"] is True, (press["on"], k)
        assert aimed["reset"]["disabled"] is False
        for k in ("bullet", "outdent", "indent"):
            assert aimed[k]["visibility"] == "", (press["on"], k)
    by = [(p["on"].get("key"), p["on"].get("idx"), p["click"]) for p in r["ribbon"]]
    base_off, tax_in = r["ribbon"][0], r["ribbon"][1]
    assert by[0] == ("base", None, "bullet")
    # The bullet reads PRESSED on a money line before the press, and not after it.
    assert base_off["aimed"]["bullet"]["pressed"] == "true" and base_off["after"]["bullet"]["pressed"] == "false"
    assert json.loads(base_off["pl"]) == {"bullet": False, "indent": 288}
    # On the square, Outdent has nowhere to go and Indent moves it to the "o"; on the "o", the reverse.
    assert tax_in["aimed"]["outdent"]["disabled"] is True and tax_in["aimed"]["indent"]["disabled"] is False
    assert tax_in["after"]["outdent"]["disabled"] is False and tax_in["after"]["indent"]["disabled"] is True
    assert json.loads(tax_in["pl"]) == {"bullet": True, "level": 1}
    # Reset puts the line back on the default and stores nothing for it.
    reset = r["ribbon"][5]
    assert (reset["on"]["key"], reset["click"]) == ("remodel", "reset")
    assert reset["before"] is not None and reset["pl"] is None
    lp = r["payload"]["price_overrides"]["line_props"]
    assert "remodel" not in lp
    assert lp == {"base": {"bullet": False, "indent": 288}, "sales_tax": {"bullet": True, "level": 1},
                  "heading_options": {"bullet": True, "level": 0}}
    assert r["payload"]["price_overrides"]["after_props"] == {
        "option:Alt1": [{"bullet": False, "indent": 1728}, None, {"bullet": True, "level": 0}]}


def _sq(run):
    """The printed box in brief: (first words, "sq" | "o" | "-", where the words start)."""
    return [(ln["text"][:14], ("o" if ln["level"] == 1 else "sq") if ln["bullet"] else "-", ln["indent_tw"])
            for ln in run["doc"] if not ln["blank"]]


# What the customer's document prints, written out by hand from Kyle's REBID layout and the presses
# in CASES. The comparison above proves the two halves agree; this proves they agree on the RIGHT
# answer — two halves that regressed to flush together would pass the comparison alone.
EXPECTED = {
    ("epoxy", "Direct"): [
        ("Base Bid", "-", 0),
        ("$22,600 – Epox", "-", 288),          # bullet off; the words stay where it had them
        ("$350 – Materia", "o", 1440),         # indent: the "o"
        ("$1,750 – Remod", "sq", 288),
        ("$24,700 – Tota", "sq", 288),
        ("Includes 1 gal", "o", 1440),         # Enter + typed under the Total: a sub-line
        ("Pricing valid ", "-", 0),            # typed on the gap: a heading's line
        ("Options:", "sq", 288),               # a heading given a bullet
        ("$85,000 – Alte", "sq", 288),
        ("Notes: Areas p", "-", 1728),         # "o" off (words stay at 1440) then indent a step
        ("Schedule: same", "sq", 288),         # outdent: "o" back to the square
        ("$500 – Materia", "sq", 288),
        ("$85,500 – Tota", "sq", 288),
    ],
    ("polish", "Direct"): [
        ("Base Bid", "-", 0),
        ("$22,600 – Poli", "o", 1440),
        ("$350 – Materia", "sq", 288),
        ("$1,750 – Remod", "sq", 288),
        ("$24,700 – Tota", "-", 288),
        ("Options", "-", 0),
        ("$85,000 – Alte", "sq", 288),
        ("Notes: Areas p", "o", 1440),
        ("Schedule: same", "sq", 288),
        ("$500 – Materia", "sq", 288),
        ("$85,500 – Tota", "sq", 288),
        ("Exclusions: wa", "o", 1440),
    ],
    ("combo", "Direct"): [
        ("$22,600 – Opti", "sq", 288),
        ("Includes moist", "sq", 288),
        ("$350 – Materia", "o", 1440),
        ("$1,750 – Remod", "sq", 288),
        ("$24,700 – Tota", "sq", 288),
        ("$4,930 – Optio", "sq", 288),
        ("Grind and seal", "o", 1440),
        ("$50 – Material", "sq", 288),
        ("$20 – Remodel ", "sq", 288),
        ("$5,000 – Total", "-", 288),
    ],
    ("budget", "Direct"): [
        ("Floor Budget P", "-", 0),
        ("$1.15 / sf – C", "sq", 288),
        ("$1.50 / sf – G", "-", 288),
        ("$2.25 / sf – G", "o", 1440),
        ("$5.75 / sf – I", "sq", 288),
        ("$6.75 / sf – U", "sq", 288),
        ("Add On’s", "sq", 288),
        ("$3.50 / lf – C", "sq", 288),
        ("$6.50 / lf – E", "sq", 288),
    ],
    ("epoxy", "GC"): [
        ("Base Bid", "-", 0),
        ("$22,600 – Resi", "o", 1440),
        ("$350 – Materia", "-", 288),
        ("$1,750 – Remod", "sq", 288),
        ("$24,700 – Tota", "sq", 288),
        ("Options & Unit", "sq", 288),         # level 0, though the box's first row is on the "o"
        ("If a different", "-", 576),
        ("$x – Add for", "sq", 288),
        ("$4,200 – Add f", "sq", 288),
        ("$x – Add for T", "sq", 288),         # the WORK-list rows Kyle put in this box
        ("Repairs and ma", "-", 288),
        ("$3.00 $3.25 $2", "sq", 288),
        ("$400 /day or $", "sq", 288),
    ],
    ("polish", "GC"): [
        ("Base Bid", "-", 0),
        ("$22,600 – Poli", "sq", 288),
        ("$350 – Materia", "sq", 288),
        ("$1,750 – Remod", "sq", 288),
        ("$24,700 – Tota", "sq", 288),
        ("Options & Unit", "-", 0),
        ("$x – Add for", "sq", 288),
        ("$x – Add for u", "sq", 288),
        ("($x) – Deduct ", "sq", 288),
        ("Note: Addition", "-", 1170),
        ("$1,100 – Add f", "o", 1440),
        ("$x – Add for T", "sq", 288),
        ("Repairs and ma", "-", 0),
        ("$400 /day or $", "sq", 288),
    ],
    ("gyp", "Direct"): [
        ("Base Bid", "-", 0),
        ("$22,600 – Gyps", "sq", 288),
        ("$350 – Materia", "o", 1440),
        ("$1,750 – Kansa", "sq", 288),
        ("$24,700 – Tota", "sq", 288),
        ("1 Mobilization", "sq", 288),
        ("Options:", "-", 0),
        ("$85,000 – Alte", "-", 288),
        ("Notes: Areas p", "o", 1440),
        ("Schedule: same", "sq", 288),
        ("$500 – Materia", "sq", 288),
        ("$85,500 – Tota", "sq", 288),
    ],
}


@needs_node
@pytest.mark.parametrize("work_type,audience", list(CASES))
def test_the_price_box_prints_kyles_rebid_layout_and_every_press(runs, work_type, audience):
    got = _sq(runs[(work_type, audience)])
    assert got == EXPECTED[(work_type, audience)], "\n".join(map(repr, got))


# ── (3) THE PAPER: the REBID hazards the flatten was hiding ──────────────────────────────────────
def _vals(**over):
    v = {"job_name": "Fmt QA", "project_name": "Fmt QA", "city_state": "Olathe, KS",
         "bid_date_formatted": "7/15/26", "system_name": "MACRO", "texture": "OP", "epoxy_sf": "12,000",
         "cove_lf": "250", "disposal": "d", "schedule_notes": "~5d", "scope_notes": "scope",
         "total_formatted": "$63,801", "state_name": "Kansas", "base_bid_formatted": "$58,523",
         "material_tax_formatted": "$2,639", "tax_amount_formatted": "$2,639", "estimator_name": "Kyle",
         "site_visit_phrase": "per site visit on 7/15", "base_tax_phrase": "", "exclusions": "std",
         "total_label": "$63,801 – Total", "mobilizations_line": "1 Mobilization to Site.",
         "work_description": "per plans"}
    v.update(over)
    return v


def _price_list_paras(blob):
    d = docx.Document(io.BytesIO(blob))
    return d, [p for p in d.element.body.iter(qn("w:p"))
               if (pw._para_num_ref(p) or ("",))[0] == "3"
               and not any(True for _ in p.iterancestors(MC + "Fallback"))]


@pytest.mark.parametrize("work_type,audience", sorted(set(pw.TEMPLATE_PICKER)))
def test_hazard_1_no_price_row_is_tucked_into_the_margin(work_type, audience):
    """Kyle's Direct files put every PRICE row on the list with `w:ind left=0` — the trick that hid
    the square in the margin. Kept, the square prints OUTSIDE the text column (at 0 - 288 twips).
    On every template, every bulleted PRICE row is placed by its level, or states a real indent."""
    blob = pw.fill_proposal(work_type=work_type, audience=audience, values=_vals(),
                            price_lines=[{"amount_formatted": "$2,500", "label": "Add VE"}],
                            has_options=True)
    d, rows = _price_list_paras(blob)
    assert rows, "no PRICE-list row printed"
    levels = pw._numbering_levels(d)
    for p in rows:
        ind = p.find(qn("w:pPr")).find(qn("w:ind"))
        ref = pw._para_num_ref(p)
        hang = int((levels[ref].get("ind") or {}).get("hanging") or 0)
        left = None
        if ind is not None:
            for a in ("w:start", "w:left"):
                if ind.get(qn(a)) is not None:
                    left = int(ind.get(qn(a)))
                    break
            if ind.get(qn("w:hanging")) is not None:
                hang = int(ind.get(qn("w:hanging")))
        if left is None:
            left = int((levels[ref].get("ind") or {}).get("left"))
        assert left - hang >= 0, (work_type, audience, pw._own_text(p)[:40], left, hang)
    # ...and the pricing now carries its bullets at all (what #132 took out).
    money = [p for p in rows if pw._own_text(p).strip().startswith(("$", "($"))]
    assert money, (work_type, audience)


def test_hazard_2_kyles_saved_click_on_an_empty_line_prints_no_square():
    """The beta's saved override on the epoxy file — paragraph 163, a blank line under the options,
    `bullet: true` and no words. With the flatten gone it would print a lone red square on the
    customer's document; the editor draws none (`.tw-block.tw-empty.tw-li::before`)."""
    blob = pw.fill_proposal(work_type="epoxy", audience="Direct", values=_vals(),
                            paragraph_overrides=[{"id": 163, "text": "", "para": {"bullet": True}}])
    d, rows = _price_list_paras(blob)
    assert not [p for p in rows if not pw._own_text(p).strip()], "an empty PRICE-list paragraph kept its bullet"


def _gc_polish_row(d, starts):
    for i, _k, p, _b, t, _x in pw.iter_editable_blocks(d):
        if (t or "").startswith(starts):
            return i
    raise AssertionError(starts)


def _printed_row(blob, starts):
    d = docx.Document(io.BytesIO(blob))
    for p in d.element.body.iter(qn("w:p")):
        if any(True for _ in p.iterancestors(MC + "Fallback")):
            continue
        if pw._own_text(p).startswith(starts):
            ppr = p.find(qn("w:pPr"))
            ind = ppr.find(qn("w:ind"))
            return (pw._para_num_ref(p),
                    {k.split("}")[1]: v for k, v in (ind.attrib.items() if ind is not None else [])})
    raise AssertionError(starts)


def test_hazard_3_the_o_is_a_list_level_the_contract_carries_backend_first():
    """`level` is part of the paragraph-props contract, sanitised on the server: the "o" is the
    PRICE list's level 1, which {bullet, indent} alone could never ask for (two indent presses give
    a square at 288 with the words at 576). And a props entry saved BEFORE `level` existed means what
    it always meant: its row stays on the level the template gave it."""
    assert pw.sanitize_para_props({"bullet": True, "indent": 1440, "level": 1}) == {
        "bullet": True, "indent": 1440, "level": 1}
    for bad in (True, -1, 9, "1", 1.0):
        assert "level" not in pw.sanitize_para_props({"level": bad}), bad
    assert main._sanitize_paragraph_overrides(
        [{"id": 128, "para": {"bullet": True, "indent": 1440, "level": 1}}])[0]["para"] == {
            "bullet": True, "indent": 1440, "level": 1}

    tpl = docx.Document(str(pw.pick_template("polish", "GC")))
    base_id = _gc_polish_row(tpl, "{{base_bid_formatted}}")
    note_id = _gc_polish_row(tpl, "Note: Additional")
    # The new shape: the base row to the "o", Kyle's "o" row to the square.
    blob = pw.fill_proposal(work_type="polish", audience="GC", values=_vals(), paragraph_overrides=[
        {"id": base_id, "para": {"bullet": True, "indent": 1440, "level": 1}},
        {"id": note_id, "para": {"bullet": True, "indent": 288, "level": 0}}])
    assert _printed_row(blob, "$58,523") == (("3", "1"), {}), "the base row is not on the \"o\" level"
    assert _printed_row(blob, "Note: Additional") == (("3", "0"), {}), "Kyle's \"o\" row did not come back"
    # The OLD shape, as a draft saved last week holds it: {bullet, indent}, no level.
    old = pw.fill_proposal(work_type="polish", audience="GC", values=_vals(), paragraph_overrides=[
        {"id": base_id, "para": {"bullet": True, "indent": 576}},
        {"id": note_id, "para": {"bullet": True, "indent": 1170}}])
    ref, ind = _printed_row(old, "$58,523")
    assert ref == ("3", "0") and ind.get("left") == "576", (ref, ind)
    ref, ind = _printed_row(old, "Note: Additional")
    assert ref == ("3", "1") and ind == {"left": "1170", "hanging": "180"}, (ref, ind)


def test_hazard_3_a_bullet_off_kyles_o_row_does_not_hang_its_first_line():
    """Kyle authored the GC Polish "Note:" row at left 1170 / hanging 180. Switching its bullet off
    kept the hanging, so the first line printed 180 twips LEFT of where the editor draws the words."""
    tpl = docx.Document(str(pw.pick_template("polish", "GC")))
    note_id = _gc_polish_row(tpl, "Note: Additional")
    blob = pw.fill_proposal(work_type="polish", audience="GC", values=_vals(), paragraph_overrides=[
        {"id": note_id, "para": {"bullet": False, "indent": 1170}}])
    ref, ind = _printed_row(blob, "Note: Additional")
    assert ref is None and ind.get("left") == "1170" and "hanging" not in ind, (ref, ind)


def test_hazard_4_blank_price_lines_are_the_price_rows_own_paragraphs():
    """A blank line typed next to a price line and the Options gap's blank lines print in the price
    box's own paragraph and run properties (9pt Zetta Serif Book #404040) — never a bare <w:p/>,
    which prints at the document's 12pt and is taller than the editor's line."""
    body = {"work_type": "epoxy", "audience": "Direct", "values": _vals(tax_layout="BROKEN_OUT",
                                                                        price_taxable=True),
            "price_overrides": {"after": {"total": ["", "typed"]}, "options_gap": 2},
            "price_lines": [{"label": "Add VE", "amount": 2500}]}
    blob = _render(body)
    lines = _doc_lines(blob)
    at = next(i for i, ln in enumerate(lines) if ln["text"].endswith("– Total"))
    blanks = [lines[at + 1]] + [ln for ln in lines[at + 3:at + 5]]
    assert [ln["text"] for ln in lines[at + 1:at + 6]] == ["", "typed", "", "", "Options:"], (
        [ln["text"] for ln in lines[at:at + 7]])
    row_rpr = lines[at]["p"].find(W + "pPr/" + W + "rPr")
    for ln in blanks:
        assert ln["blank"] and not ln["numbered"]
        rpr = ln["p"].find(W + "pPr/" + W + "rPr")
        assert rpr is not None, "a bare blank paragraph: it prints at the document default"
        assert rpr.find(W + "sz").get(W + "val") == "18"
        for tag in ("rFonts", "color"):
            want = row_rpr.find(W + tag)
            got = rpr.find(W + tag)
            assert want is not None and got is not None and dict(got.attrib) == dict(want.attrib), tag


def test_the_alternate_system_block_prints_a_heading_and_three_money_lines():
    """The ALTERNATE SYSTEM block, on the Direct files: its name is a heading on the PRICE list in
    Kyle's file (so it would print a square) and prints none; its three rows are money lines; and a
    line_props override on one of them prints. (The screen-against-paper comparison of the block,
    words included, is section (4): the editor now draws the flooring row's tax wording the way
    each file prints it.)"""
    for wt in ("epoxy", "polish", "combo"):
        body = {"work_type": wt, "audience": "Direct", "values": _vals(),
                "alternate_computed_bid": {"alternate_full_bid": {"total_base_bid": 30000, "remodel_tax": 1000},
                                           "alternate": {"label": "MACRO Flake"}},
                "alternate_label": "MACRO Flake",
                "price_overrides": {"line_props": {"alt_total": {"bullet": False},
                                                   "alt_remodel": {"bullet": True, "level": 1}}}}
        lines = _doc_lines(_render(body))
        at = next(i for i, ln in enumerate(lines) if ln["text"].startswith("ALTERNATE SYSTEM"))
        got = [(ln["text"][:10], ln["bullet"], ln["level"], ln["indent_tw"]) for ln in lines[at:at + 4]]
        assert got == [("ALTERNATE ", False, None, 0), ("$29,000 – ", True, 0, 288),
                       ("$1,000 – R", True, 1, 1440), ("$30,000 – ", False, None, 288)], (wt, got)


def test_the_payload_keeps_line_props_and_drops_garbage():
    """`_sanitize_price_overrides` keeps the three new buckets — cleaned by the same rule the editor
    uses — and a draft saved before them sanitises to empty buckets (the default layout)."""
    got = main._sanitize_price_overrides({
        "line_props": {"base": {"bullet": False, "indent": 576}, "total": {"bullet": "x"},
                       "option:A": {"level": 1, "junk": 1}},
        "before_props": {"heading_options": [None, {"bullet": True}], "x": "nope"},
        "after_props": {"option:A": [{"bullet": False}, {}, {"level": 7}]}})
    assert got["line_props"] == {"base": {"bullet": False, "indent": 576}, "option:A": {"level": 1}}
    assert got["before_props"] == {"heading_options": [None, {"bullet": True}]}
    assert got["after_props"] == {"option:A": [{"bullet": False}, None, None]}
    old = main._sanitize_price_overrides({"lines": {"base": "x"}})
    assert old["line_props"] == {} and old["before_props"] == {} and old["after_props"] == {}


# ── (4) THE REVIEW OF FIX 7: seven findings on the first cut, each run through both halves ──────
# The same comparison as (2) — the editor's box through the page's own code, the payload through the
# real renderer — over the cases the first cut's matrix never reached: the ALTERNATE SYSTEM block,
# a template price row taken below its hanging, the keys (Backspace, Tab), and the caret moving
# without a focus event.
_ALT = {"alternate_computed_bid": {"alternate_full_bid": {"total_base_bid": 30000, "remodel_tax": 1000},
                                   "alternate": {"label": "MACRO Flake"}},
        "alternate": {"label": "MACRO Flake"}}
_COMBO_TABS = CASES[("combo", "Direct")]["state"]["priced_tabs"]


def _alt_state(**over):
    st = copy.deepcopy(_ALT)
    st.update(over)
    return st


def _block_id(work_type, audience, starts):
    for b in _price_box(_template_json(work_type, audience)):
        if (b.get("text") or "").startswith(starts):
            return b["id"]
    raise AssertionError((work_type, audience, starts))


REVIEW = {
    # FINDING 1. A keystroke anywhere in the box — here on the Base Bid heading, its words unchanged —
    # and a ribbon press on the alternate's flooring line: the sweep stored that line with its tax
    # wording turned into a marker that resolved to nothing, and the PDF printed "…as described
    # above" while the screen still showed "(material sales tax INCLUDED)".
    "alt-keystroke/epoxy": ("epoxy", "Direct", {"state": _alt_state(), "actions": [
        {"input_on": {"key": "heading_base"}},
        {"ribbon": {"key": "alt_flooring"}, "click": "indent"}]}),
    "alt-keystroke/epoxy-one-line": ("epoxy", "Direct", {"state": _alt_state(tax_layout="ONE_LINE"),
                                                         "actions": [{"input_on": {"key": "base"}}]}),
    # Polish and Combo print the BASE's wording there ({{base_tax_phrase}}), not Epoxy's literal: the
    # editor showed "(material sales tax INCLUDED)" on both whatever the base said.
    "alt-keystroke/polish-one-line": ("polish", "Direct", {"state": _alt_state(tax_layout="ONE_LINE"),
                                                           "actions": [{"input_on": {"key": "option:Alt1"}}]}),
    "alt-keystroke/polish-broken": ("polish", "Direct", {"state": _alt_state(),
                                                         "actions": [{"input_on": {"key": "total"}}]}),
    "alt-keystroke/combo": ("combo", "Direct", {"state": _alt_state(rooms=[], priced_tabs=_COMBO_TABS),
                                                "actions": [{"input_on": {"key": "combo:epoxy.flooring"}}]}),
    # A draft the bug already reached: the stored empty marker now resolves to the wording again.
    "alt-healed/epoxy": ("epoxy", "Direct", {"state": _alt_state(price_overrides={
        "lines2": {"alt_flooring": "$29,000 – Flooring as described above ⟦tax⟧"}})}),
    # FINDING 4. Lines typed around the ALTERNATE rows were drawn (with an "o") and never printed.
    "alt-typed/epoxy": ("epoxy", "Direct", {"state": _alt_state(price_overrides={
        "after": {"alt_name": ["Upgrade for the showroom"]}}), "actions": [
        {"enter_after": {"key": "alt_total"}, "type": "Includes 2 coats of urethane"},
        {"enter_after": {"key": "alt_flooring"}, "type": "Moisture test by others"},
        # ...and an empty one opened under the remodel row goes again with Backspace, as on any
        # other price line (the alternate rows were left out of the join, so it stayed as a blank).
        {"enter_after": {"key": "alt_remodel"}, "type": ""},
        {"key": "Backspace", "on": {"key": "alt_remodel", "kind": "extra", "pos": "after", "idx": 0}}]}),
    "alt-typed/polish": ("polish", "Direct", {"state": _alt_state(), "actions": [
        {"enter_after": {"key": "alt_total"}, "type": "Includes 2 coats of urethane"}]}),
    # FINDINGS 2 AND 5. A template price row with its bullet off and then outdented (or Backspace
    # twice at its start) was drawn 0.2in in over a money line the PDF printed flush.
    "below-hanging/gc": ("epoxy", "GC", {"actions": [
        {"on": {"starts": "$22,600"}, "press": "bullet"}, {"on": {"starts": "$22,600"}, "press": "outdent"},
        {"on": {"starts": "$350"}, "press": "bullet"}, {"on": {"starts": "$350"}, "press": "outdent"},
        # An unbulleted row Kyle authored with a first-line indent, moved and moved back: its state
        # is the template's again, so it keeps the template's own geometry, as the PDF does.
        {"on": {"starts": "Repairs and ma"}, "press": "indent"},
        {"on": {"starts": "Repairs and ma"}, "press": "outdent"}]}),
    "below-hanging/gyp": ("gyp", "Direct", {"actions": [
        {"on": {"starts": "$350"}, "press": "bullet"}, {"on": {"starts": "$350"}, "press": "outdent"}]}),
    "below-hanging/polish-backspace": ("polish", "Direct", {"actions": [
        {"key": "Backspace", "on": {"starts": "$22,600"}}, {"key": "Backspace", "on": {"starts": "$22,600"}}]}),
    "below-hanging/budget-saved": ("budget", "Direct", {"state": {"rooms": [], "price_overrides": {}}}),
    # FINDING 3. The caret moved onto a blank line of the Options gap inside a box that already had
    # focus: the ribbon stayed aimed at the Total, and Bullet took the Total's square off.
    "gap-aim/epoxy": ("epoxy", "Direct", {"actions": [
        {"select": {"key": "total"}}, {"select": {"gap": 0}}, {"press_ribbon": "bullet"},
        {"select": {"key": "total"}}]}),
    # FINDING 6. Backspace and Tab on the lines the page composes.
    "keys/epoxy": ("epoxy", "Direct", {"actions": [
        {"key": "Backspace", "on": {"key": "total"}},            # the square off, the words stay
        {"key": "Backspace", "on": {"key": "total"}},            # then the indent, to the margin
        {"key": "Backspace", "on": {"key": "total"}},            # then refused: a row is never merged
        {"key": "Backspace", "on": _typed("after", 0)},          # the "o" off, NOT glued to the line above
        {"key": "Tab", "on": {"key": "base"}},                   # a square to the "o"
        {"key": "Tab", "on": {"key": "base"}},                   # nowhere deeper: handed back
        {"key": "Tab", "shift": True, "on": {"key": "sales_tax"}},   # already the square: handed back
        {"key": "Tab", "shift": True, "on": _typed("after", 2)},     # the "o" back to the square
        {"key": "Backspace", "on": _typed("after", 1)},          # a blank typed line: it goes
    ]}),
}
# The budget row's saved {bullet: false, indent: 0}, restored through setParaState as a reload does
# — what the old editor's Bullet press stored on a row whose indent it read as the tucked 0.
REVIEW["below-hanging/budget-saved"][2]["saved_paragraph_overrides"] = [
    {"id": _block_id("budget", "Direct", "$1.50"), "para": {"bullet": False, "indent": 0}}]

_REVIEW_RUNS: dict = {}


@pytest.fixture(scope="module")
def review_runs():
    if shutil.which("node") is None:
        pytest.skip("node is not installed")
    if not _REVIEW_RUNS:
        cases = []
        for name, (wt, aud, spec) in REVIEW.items():
            c = _case(wt, aud, spec, name)
            c["saved_paragraph_overrides"] = spec.get("saved_paragraph_overrides") or []
            cases.append(c)
        for run in _run_cases(cases):
            _REVIEW_RUNS[run["case"]["name"]] = run
    return _REVIEW_RUNS


@needs_node
@pytest.mark.parametrize("name", list(REVIEW))
def test_review_cases_the_editor_and_the_document_agree_line_for_line(review_runs, name):
    """Every review case, the ALTERNATE SYSTEM block included: the words, the bullet, which one,
    where the words start, and the blank lines between — the screen and the paper."""
    run = review_runs[name]
    shown, printed = _shown(run), _printed(run)
    assert not [ln for ln in run["editor"]["lines"] if ln["unstyled"]]
    assert shown == printed, name + "\n" + "\n".join(
        f"{'  ' if a == b else '!!'} screen {a!r}\n{'  ' if a == b else '!!'} paper  {b!r}"
        for a, b in itertools.zip_longest(shown, printed))


def _line(run, starts, side="doc"):
    lines = run["doc"] if side == "doc" else run["editor"]["lines"]
    hits = [ln for ln in lines if ln["text"].startswith(starts)]
    assert len(hits) == 1, (starts, [ln["text"] for ln in lines])
    return hits[0]


@needs_node
def test_a_keystroke_or_a_press_in_the_box_keeps_the_alternate_flooring_wording(review_runs):
    """Finding 1: the wording is the line's PHRASE now, on both halves, so the sweep's marker goes
    back to it — nothing is stored for a line nobody re-worded, and the PDF keeps the wording."""
    for name, want in (("alt-keystroke/epoxy", price_rules.PHRASE_MATERIAL),
                       ("alt-keystroke/epoxy-one-line", price_rules.PHRASE_MATERIAL),
                       ("alt-keystroke/polish-one-line", price_rules.PHRASE_BOTH),
                       ("alt-keystroke/polish-broken", ""),
                       ("alt-keystroke/combo", ""),
                       ("alt-healed/epoxy", price_rules.PHRASE_MATERIAL)):
        run = review_runs[name]
        pov = run["editor"]["payload"]["price_overrides"]
        if not name.startswith("alt-healed"):
            assert "alt_flooring" not in (pov.get("lines2") or {}), (name, pov.get("lines2"))
        paper = _line(run, "$29,000 – Flooring")["text"]
        screen = _line(run, "$29,000 – Flooring", "editor")["text"]
        assert paper == screen == ("$29,000 – Flooring as described above" + (" " + want if want else "")), (
            name, screen, paper)
    # ...and the ribbon press on it stored its bullet, not its words.
    lp = review_runs["alt-keystroke/epoxy"]["editor"]["payload"]["price_overrides"]["line_props"]
    assert lp == {"alt_flooring": {"bullet": True, "level": 1}}


@needs_node
def test_lines_typed_around_the_alternate_rows_print(review_runs):
    """Finding 4: a line typed under "$30,000 – Total" (Enter, then words), under the flooring row
    and under the alternate's name, each prints where the editor draws it, with its bullet."""
    run = review_runs["alt-typed/epoxy"]
    pov = run["editor"]["payload"]["price_overrides"]
    assert pov["after"] == {"alt_name": ["Upgrade for the showroom"],
                            "alt_total": ["Includes 2 coats of urethane"],
                            "alt_flooring": ["Moisture test by others"]}
    doc = [ln for ln in run["doc"] if not ln["blank"]]
    at = next(i for i, ln in enumerate(doc) if ln["text"].startswith("ALTERNATE SYSTEM"))
    got = [(ln["text"][:14], ln["bullet"], ln["level"], ln["indent_tw"]) for ln in doc[at:]]
    assert got == [("ALTERNATE SYST", False, None, 0),
                   ("Upgrade for th", False, None, 0),        # under a heading: a heading's line
                   ("$29,000 – Floo", True, 0, 288),
                   ("Moisture test ", True, 1, 1440),         # under a money line: its "o"
                   ("$1,000 – Remod", True, 0, 288),
                   ("$30,000 – Tota", True, 0, 288),
                   ("Includes 2 coa", True, 1, 1440)], got
    polish = [ln["text"] for ln in review_runs["alt-typed/polish"]["doc"] if ln["text"].strip()]
    assert polish[-2:] == ["$30,000 – Total", "Includes 2 coats of urethane"], polish[-4:]


def test_a_line_typed_under_an_alternate_row_goes_with_the_block_when_there_is_none():
    """No alternate, no ALTERNATE block — and no stray line: the typed lines are cloned inside
    {{#alternate}}, which the writer removes whole."""
    body = {"work_type": "epoxy", "audience": "Direct",
            "values": _vals(tax_layout="BROKEN_OUT", price_taxable=True),
            "price_overrides": {"after": {"alt_total": ["Orphan typed line"], "alt_name": ["Orphan head"]},
                                "before": {"alt_flooring": ["Orphan above"]}}}
    lines = [ln["text"] for ln in _doc_lines(_render(body))]
    # The base's own Total prints (broken out): a row these lines could wrongly attach to is there.
    assert any(t.endswith("– Total") for t in lines), lines
    text = "\n".join(lines)
    assert "Orphan" not in text and "ALTERNATE SYSTEM" not in text


@needs_node
def test_a_template_price_row_below_its_hanging_is_drawn_where_it_prints(review_runs):
    """Findings 2 and 5: {bullet: false, indent: 0} — Bullet off then Outdent, Backspace twice, or
    a draft saved by the old editor — prints the words at the margin, and the editor now draws them
    there instead of one hanging (0.2in) further in."""
    for name, starts, pid_expected in (("below-hanging/gc", "$22,600", True), ("below-hanging/gc", "$350", True),
                                       ("below-hanging/gyp", "$350", True),
                                       ("below-hanging/polish-backspace", "$22,600", True),
                                       ("below-hanging/budget-saved", "$1.50", True)):
        run = review_runs[name]
        paper = _line(run, starts)
        screen = _line(run, starts, "editor")
        assert (paper["bullet"], paper["indent_tw"]) == (False, 0), (name, starts, paper["indent_tw"])
        assert (screen["bullet"], screen["indent_tw"], screen["first_tw"]) == (False, 0, 0), (name, starts, screen)
    po = review_runs["below-hanging/polish-backspace"]["editor"]["payload"]["paragraph_overrides"]
    assert [o["para"] for o in po] == [{"bullet": False, "indent": 0}]
    # A row put BACK on its template state keeps the template's own geometry (Reset draws as before).
    gyp = review_runs["below-hanging/gyp"]
    assert _line(gyp, "$22,600", "editor")["indent_tw"] == _line(gyp, "$22,600")["indent_tw"] == 288


@needs_node
def test_the_ribbon_lets_go_when_the_caret_moves_onto_a_blank_gap_line(review_runs):
    """Finding 3: arrowing from the Total onto the blank line above "Options:" fires no focusin; the
    selectionchange listener now idles the ribbon, so Bullet cannot reach the line the caret left."""
    run = review_runs["gap-aim/epoxy"]
    r = run["editor"]["ribbon"]
    assert r[0]["aimedAt"]["key"] == "total" and r[0]["bar"]["bullet"]["disabled"] is False
    assert r[1]["aimedAt"] is None, r[1]["aimedAt"]
    assert r[1]["bar"]["bullet"]["disabled"] is True and r[1]["bar"]["indent"]["disabled"] is True
    assert r[2]["disabled"] is True and r[2]["aimedAt"] is None
    assert r[3]["aimedAt"]["key"] == "total"            # back on a line: aimed again
    assert not (run["editor"]["payload"]["price_overrides"].get("line_props") or {})
    total = _line(run, "$24,700 – Total")
    assert (total["bullet"], total["level"], total["indent_tw"]) == (True, 0, 288)


@needs_node
def test_backspace_and_tab_act_on_the_lines_the_page_composes(review_runs):
    """Finding 6, Hanz 2026-08-25: "When I back space, it doesnt remove the bullet point." The
    template rows follow Word; the page-built lines now do too — Backspace at the start takes the
    bullet, then the indent, then (a row) refuses; a typed "o" line loses its "o" instead of being
    glued onto the money line above; Tab and Shift+Tab move a line between the square and the "o",
    and a Tab that cannot move it is handed back to the browser."""
    run = review_runs["keys/epoxy"]
    k = [p for p in run["editor"]["pressed"] if "key" in p]
    j = lambda p: json.loads(p["pl"]) if p["pl"] else None   # noqa: E731
    assert [(p["prevented"], j(p)) for p in k[:3]] == [
        (True, {"bullet": False, "indent": 288}), (True, {"bullet": False, "indent": 0}),
        (True, {"bullet": False, "indent": 0})]
    assert all(p["text"] == "$24,700 – Total" for p in k[:3])
    # The ribbon follows the key: aimed at the line, its Bullet no longer pressed.
    assert k[0]["aimedAt"]["key"] == "total" and k[0]["bar"]["bullet"]["pressed"] == "false"
    assert k[4]["aimedAt"]["key"] == "base" and k[4]["bar"]["indent"]["disabled"] is True
    assert k[4]["bar"]["outdent"]["disabled"] is False
    assert (k[3]["prevented"], j(k[3]), k[3]["text"]) == (
        True, {"bullet": False, "indent": 1440}, "Notes: Areas per Schedule Note 1")
    assert [(p["prevented"], j(p)) for p in k[4:8]] == [
        (True, {"bullet": True, "level": 1}), (False, {"bullet": True, "level": 1}),
        (False, None), (True, {"bullet": True, "level": 0})]
    assert k[8]["prevented"] is True and k[8]["attached"] is False     # the blank typed line went
    pov = run["editor"]["payload"]["price_overrides"]
    assert "option:Alt1" not in (pov.get("lines2") or {}), "a typed line was glued into the option's line"
    assert pov["after"]["option:Alt1"] == ["Notes: Areas per Schedule Note 1", "Schedule: same 2 mobs/phases"]
    assert pov["line_props"] == {"total": {"bullet": False, "indent": 0}, "base": {"bullet": True, "level": 1}}
    assert pov["after_props"] == {"option:Alt1": [{"bullet": False, "indent": 1440}, {"bullet": True, "level": 0}]}
    assert (_line(run, "$24,700 – Total")["bullet"], _line(run, "$24,700 – Total")["indent_tw"]) == (False, 0)
    assert _line(run, "$22,600 – Epoxy")["level"] == 1


@needs_node
def test_the_alternate_flooring_phrase_is_the_same_rule_in_both_languages():
    rows = ["{{alternate.lump_sum_formatted}} – Flooring as described above (material sales tax INCLUDED)",
            "{{alternate.lump_sum_formatted}} – Flooring as described above {{base_tax_phrase}}",
            "{{alternate.lump_sum_formatted}} – Flooring as described above {{ base_tax_phrase }}",
            "{{alternate.lump_sum_formatted}} – Flooring (Remodel Tax AND material sales tax INCLUDED)",
            "{{alternate.lump_sum_formatted}} – Flooring (tax exempt)",
            "{{alternate.lump_sum_formatted}} – Flooring", "", None]
    bases = ["", price_rules.PHRASE_BOTH, price_rules.PHRASE_NONE, None]
    cases = [{"row": r, "base": b} for r in rows for b in bases]
    p = subprocess.run(["node", "-e", """
      const P = require(process.argv[1]);
      const cases = JSON.parse(require("fs").readFileSync(0, "utf8"));
      console.log(JSON.stringify(cases.map(c => P.altFlooringPhrase(c.row, c.base))));
    """, str(CORE)], input=json.dumps(cases), capture_output=True, text=True, encoding="utf-8")
    assert p.returncode == 0, p.stderr
    js = json.loads(p.stdout)
    for c, j in zip(cases, js):
        assert j == price_rules.alt_flooring_phrase(c["row"], c["base"]), c
    assert price_rules.alt_flooring_phrase(rows[0], price_rules.PHRASE_NONE) == price_rules.PHRASE_MATERIAL
    assert price_rules.alt_flooring_phrase(rows[1], price_rules.PHRASE_BOTH) == price_rules.PHRASE_BOTH
    assert price_rules.alt_flooring_phrase(rows[3], "") == price_rules.PHRASE_BOTH
    for wt in ("epoxy", "polish", "combo"):
        row = pw.template_alt_flooring_row(wt, "Direct")
        assert "alternate.lump_sum_formatted" in row, wt
