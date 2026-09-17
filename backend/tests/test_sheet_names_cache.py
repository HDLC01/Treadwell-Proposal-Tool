"""`/api/sheets` stops re-reading the workbook, and still notices when it changes.

WHY THIS ENDPOINT AND NOT A HEAVIER ONE. It returns 236 bytes. It is also the FIRST thing
estimate-review fetches and the one every other fetch waits behind: the tab bar cannot render
without the tab names, and no grid is requested until it does. Measured 2026-09-17 on the dev
box, five runs through the real ASGI stack: 193-229 ms to produce those 236 bytes, because
`list_sheet_names` opened the whole .xlsx every single call. `read_only=True` does not save
it — openpyxl still parses workbook.xml and the shared-string table before it can name a
sheet.

WHAT THE CACHE MUST NOT COST. Kyle swaps the estimate template by dropping a new file in, and
that has always taken effect without restarting the server — every other cache in
estimate_writer is keyed by mtime for exactly that reason. So the second half of this module
is the invalidation, executed against a real file whose mtime moves.

The staleness test runs on a COPY of the template in tmp_path. frontend/ and the committed
template are read by other modules in this suite at the same moment, and proving invalidation
needs a file we are allowed to rewrite mid-test.
"""
import shutil

import pytest
from fastapi.testclient import TestClient

import estimate_writer
import main

client = TestClient(main.app)


@pytest.fixture
def count_opens(monkeypatch):
    """Count workbook opens by wrapping the name `list_sheet_names` actually calls."""
    calls = []
    real = estimate_writer.load_workbook

    def counting(*a, **k):
        calls.append(a[0] if a else k.get("filename"))
        return real(*a, **k)

    monkeypatch.setattr(estimate_writer, "load_workbook", counting)
    return calls


def test_the_tab_list_is_read_off_disk_once_however_often_it_is_asked_for(count_opens):
    """THE ONE THAT MATTERS. Not "it is fast" — a timing assertion on a shared dev box is a
    flake — but the thing that made it slow: one workbook open per request."""
    estimate_writer.list_sheet_names()          # may or may not be a hit already; either way
    count_opens.clear()                         # only the calls AFTER this point are counted
    for _ in range(5):
        estimate_writer.list_sheet_names()
    assert count_opens == [], (
        "list_sheet_names opened the workbook %d more times after it was already cached"
        % len(count_opens))


def test_the_endpoint_still_names_every_tab_in_the_workbook():
    """The cache is worthless if it caches the wrong thing. Sixteen tabs, and the three the
    warm loop prioritises have to be among them or _WARM_FIRST is naming tabs that do not
    exist."""
    names = client.get("/api/sheets").json()["sheets"]
    assert len(names) >= 16, "only %d tabs: %r" % (len(names), names)
    for wanted in main._WARM_FIRST:
        assert wanted in names, "_WARM_FIRST names %r, which is not a tab in the template" % wanted


def test_two_requests_return_the_same_tab_list():
    """Belt and braces on the copy below: a cache that handed back a mutated list would show up
    as the tab bar changing between page loads."""
    assert client.get("/api/sheets").json() == client.get("/api/sheets").json()


def test_a_caller_cannot_reorder_the_tab_bar_for_the_whole_process():
    """The cache holds a list, and a list handed out by reference is a list somebody can sort.
    Today's only caller drops it straight into a JSON response, so this guards a future one —
    and it is cheap enough to be worth not having to think about again."""
    first = estimate_writer.list_sheet_names()
    original = list(first)
    first.reverse()
    first.append("a tab that does not exist")
    assert estimate_writer.list_sheet_names() == original, (
        "mutating the returned list changed what every later caller sees")


# ── invalidation ─────────────────────────────────────────────────────────────
def test_a_swapped_in_template_is_noticed_without_a_restart(tmp_path):
    """Kyle drops a new estimate sheet in and it takes effect. Every cache in estimate_writer
    is mtime-keyed for this, and a cache that is not is a cache that serves last month's tabs
    until somebody redeploys — with nothing on screen to say so."""
    book = tmp_path / "sheet.xlsx"
    shutil.copy(estimate_writer.TEMPLATE_PATH, book)

    before = estimate_writer.list_sheet_names(path=book)
    assert len(before) >= 16

    from openpyxl import load_workbook as _open
    wb = _open(book)
    wb.create_sheet("A Tab Kyle Added")
    wb.save(book)

    after = estimate_writer.list_sheet_names(path=book)
    assert "A Tab Kyle Added" in after, (
        "the tab list is stale after the file changed — %r" % (after,))


def test_the_invalidation_test_above_could_have_failed(tmp_path):
    """The counterexample. `test_a_swapped_in_template_is_noticed` proves nothing if this
    function never caches at all, so: leave the file alone and the SAME call must not re-read
    it. One test says the cache invalidates, this one says there is a cache to invalidate."""
    book = tmp_path / "sheet.xlsx"
    shutil.copy(estimate_writer.TEMPLATE_PATH, book)
    estimate_writer.list_sheet_names(path=book)

    real = estimate_writer.load_workbook
    opened = []

    def counting(*a, **k):
        opened.append(1)
        return real(*a, **k)

    estimate_writer.load_workbook = counting
    try:
        estimate_writer.list_sheet_names(path=book)
    finally:
        estimate_writer.load_workbook = real
    assert opened == [], "this path re-reads on every call, so the test above is vacuous"


def test_two_workbooks_do_not_share_one_entry(tmp_path):
    """The cache is keyed by PATH as well as mtime, like _WB_CACHE and _TEMPLATE_BYTES above
    it, and for the same reason those are: this module serves the Project Info Sheet's workbook
    too. A path-less key hands one workbook's tab list to the other's request, which reads as
    correct data right up until somebody notices the tabs are from the wrong file."""
    other = tmp_path / "other.xlsx"
    shutil.copy(estimate_writer.TEMPLATE_PATH, other)
    from openpyxl import load_workbook as _open
    wb = _open(other)
    wb.create_sheet("Only In The Other Book")
    wb.save(other)

    assert "Only In The Other Book" in estimate_writer.list_sheet_names(path=other)
    assert "Only In The Other Book" not in estimate_writer.list_sheet_names(), (
        "the real template came back holding a tab that only exists in a different file")
