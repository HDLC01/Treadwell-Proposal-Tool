"""An edited price line keeps the estimator's WORDS; its amount and tax wording follow the estimate.

Hanz, 2026-09-25, on the "Hanz Fix" test project on staging. He typed "THis is a test send to Hanz"
under the base line and "Test again 123" under an option; both lines got a warning mark and FROZE
the line's dollar amount, because the typed text had been stored as part of the price line itself.
The downloaded PDF printed them as line breaks inside that one paragraph. Across real revisions the
same mechanism left options quoting last week's figures (Omakase $46,945 against $51,916) and made
Kyle retype every amount on Carson Ross rev 2 by hand.

Now:
  * an edited line stores his words with today's amount and tax wording as MARKERS wherever they
    were still there verbatim, and both the editor and the document put today's values back;
  * a line with a DIFFERENT dollar figure keeps it, is marked in the editor, and Send asks
    ("This line says $X but the estimate says $Y — send anyway?") — Hanz: warn, then let him send;
  * lines typed above and below a price line are their OWN lines: stored beside it, printed as
    their own paragraphs in the price row's own formatting, never freezing it;
  * one keystroke in the box no longer freezes the hidden "$0 – Total" rows nobody touched, and no
    longer deletes a saved Options heading;
  * every review highlight is gone from the customer document.

EXECUTED: the editor half runs the page's own functions and key handlers (js/price-lines-harness.js),
the document half runs the real renderer and reads the .docx back.
"""
import io
import json
import pathlib
import shutil
import subprocess

import docx
import pytest
from docx.oxml.ns import qn
from starlette.requests import Request

import main
import proposal_writer as pw

FRONTEND = pathlib.Path(__file__).resolve().parents[2] / "frontend"
HARNESS = pathlib.Path(__file__).resolve().parent / "js" / "price-lines-harness.js"
needs_node = pytest.mark.skipif(shutil.which("node") is None, reason="node is not installed")
_MC = "{http://schemas.openxmlformats.org/markup-compatibility/2006}"
AMT, TAX = "⟦amount⟧", "⟦tax⟧"


@pytest.fixture(scope="module")
def ran():
    if shutil.which("node") is None:
        pytest.skip("node is not installed")
    p = subprocess.run(["node", str(HARNESS), str(FRONTEND)], capture_output=True, text=True,
                       encoding="utf-8", timeout=120)
    assert p.returncode == 0, "the harness itself failed:\n" + p.stderr
    return json.loads(p.stdout.strip().splitlines()[-1])


# ── the editor ──────────────────────────────────────────────────────────────────────────────────
@needs_node
def test_hanzs_saved_project_is_migrated_into_lines_of_their_own(ran):
    """His staging draft, exactly as saved: the typed text inside the two price lines' frozen
    overrides. Drawn once, the price lines are the computed ones again (no mark, no frozen amount)
    and what he typed is its own lines, blank ones included, in the order he typed them."""
    h = ran["hanzFix"]
    pov = h["pov"]
    assert pov["lines"] == {} and pov["lines2"] == {}, pov
    assert pov["after"] == {"base": ["", "THis is a test send to Hanz"],
                            "option:Copy1": ["Test again 123"]}, pov
    assert pov["before"] == {"option:Copy1": ["", ""]}, pov
    base = [l for l in h["lines"] if l["key"] == "base"]
    assert [l["kind"] for l in base] == ["line", "extra", "extra"]
    assert base[0]["text"] == "$7,447 – Epoxy flooring as described above (material sales tax INCLUDED)"
    assert "tw-overridden" not in base[0]["cls"] and not base[0]["title"], base[0]
    assert h["saves"] >= 1, "the migration was not saved"
    # And it follows the estimate now: both amounts move, the typed lines stay.
    moved = {(l["key"], l["kind"]): l["text"] for l in ran["hanzFixRepriced"] if l["kind"] == "line"}
    assert moved[("base", "line")].startswith("$8,000 – ")
    assert moved[("option:Copy1", "line")].startswith("$8,100 – ")
    assert [l["text"] for l in ran["hanzFixRepriced"] if l["kind"] == "extra"] == [
        "", "THis is a test send to Hanz", "", "", "Test again 123"]


@needs_node
def test_an_edit_that_keeps_the_amount_follows_the_estimate(ran):
    k = ran["keptAmount"]
    assert k["stored"] == f"{AMT} – Epoxy flooring in the warehouse only as described above {TAX}"
    assert k["repriced"] == ("$9,100 – Epoxy flooring in the warehouse only as described above "
                             "(material sales tax INCLUDED)")
    # His words are marked as his (the amber wash), but no warning: the money is the estimate's.
    assert "tw-po-live" in k["cls"] and "tw-money-off" not in k["cls"]
    # Broken out: the bracket goes, the pre-tax figure comes, the words stay.
    assert k["broken"] == "$9,004 – Epoxy flooring in the warehouse only as described above"


@needs_node
def test_a_hand_typed_figure_is_kept_marked_and_asked_about_at_send(ran):
    h = ran["handTyped"]
    assert h["stored"].startswith("$9,999 – ") and AMT not in h["stored"]
    assert h["shown"].startswith("$9,999 – ")
    assert "tw-money-off" in h["cls"]
    assert "does not follow the estimate" in h["title"]
    assert h["warnings"] == [{"key": "base", "says": "$9,999", "estimate": "$9,100"}]
    assert h["ask"] == "This line says $9,999 but the estimate says $9,100 — send anyway?"


@needs_node
def test_a_keystroke_in_the_box_freezes_no_row_nobody_touched(ran):
    """money-verdict M3: in one line the tax rows and the Total are hidden and still read "$0 – …";
    the box-wide sweep used to store all three as edits on any keystroke."""
    assert ran["hiddenRows"]["stored"] == ["base"], ran["hiddenRows"]
    assert ran["hiddenRows"]["legacy"] == []


@needs_node
def test_a_saved_options_heading_is_shown_and_survives_a_keystroke(ran):
    """David Dyer rev 3 lost Kyle's typed heading to one keystroke elsewhere in the box."""
    assert ran["heading"] == {"shown": "Options & Alternates:", "kept": "Options & Alternates:"}


@needs_node
def test_enter_makes_a_line_of_its_own_and_backspace_takes_an_empty_one_away(ran):
    e = ran["enter"]
    assert e["prevented"] and e["addedKind"] == "extra" and e["caretInNew"]
    assert e["after"] == {"base": [""]} and e["baseStored"] is None
    assert e["typed"] == {"base": ["THis is a test send to Hanz"]} and e["baseStillLive"]
    assert e["twoLines"] == {"base": ["THis is a test send to Hanz", ""]}
    assert e["backspace"] == {"prevented": True, "after": {"base": ["THis is a test send to Hanz"]}}
    assert [(l["kind"], l["text"]) for l in e["repainted"]] == [
        ("line", "$7,447 – Epoxy flooring as described above (material sales tax INCLUDED)"),
        ("extra", "THis is a test send to Hanz")]
    assert e["above"] == {"option:Copy1": [""]}


@needs_node
def test_options_are_itemised_off_their_own_tab_and_an_edited_one_keeps_following_it(ran):
    assert [(l["key"], l["text"]) for l in ran["brokenOptions"]] == [
        ("option:Copy1", "$7,597 – Treadwell 3/16\" Urethne Cement With Shop Floor and Armor Top as described above"),
        ("option:Copy1:sales_tax", "$99 – Material Sales Tax"),
        ("option:Copy1:total", "$7,696 – Total"),
    ]
    ed = ran["editedOption"]
    assert ed["stored"].startswith(AMT) and ed["stored"].endswith(TAX)
    assert ed["lines"][0]["text"].startswith("$8,597 – ") and ", shop floor" in ed["lines"][0]["text"]
    assert ed["lines"][2]["text"] == "$8,696 – Total"


@needs_node
def test_a_break_inside_a_price_line_becomes_a_line_of_its_own_and_is_kept(ran):
    """Enter at a caret is handled (splitPriceLine), but a break can still land inside a price line
    another way -- Enter over a selection that spans lines puts one there. The text after it is a
    line of its own under the price line, which stays the computed one, and the same keystroke's
    re-read of the typed lines does not wipe it. Mutation: write the spill straight into `after`
    (the re-read then clears it)."""
    b = ran["breakInside"]
    assert b["after"] == {"base": ["Typed under it"]}, b
    assert b["base"] is None, b
    assert b["repainted"] == [
        ["line", "$7,447 – Epoxy flooring as described above (material sales tax INCLUDED)"],
        ["extra", "Typed under it"]], b
    assert b["next"] == {"base": ["Typed under it!"]}, b


# The shapes real drafts hold today (prod, read 2026-09-26): each saved line with the parts the
# page computes for it now. Product wording only -- no project, customer or contact detail.
_REAL_SHAPES = [
    # A base line whose figure the preview froze with cents, the tax wording taken off by hand.
    ("cents", "$1,870.00 – Epoxy flooring as described above",
     {"amount": "$1,870", "phrase": "(material sales tax INCLUDED)", "slot": True, "candidates": ["$1,870"]},
     {"main": f"{AMT} – Epoxy flooring as described above", "before": [], "after": [], "drop": False}),
    # Kyle's own wording on an exempt job: kept, and not a tax phrase this tool printed.
    ("exempt words", "$20,901 – Flake Broadcast Epoxy flooring as described above (material sales tax EXCLUDED)",
     {"amount": "$20,901", "phrase": "(tax exempt)", "slot": True},
     {"main": f"{AMT} – Flake Broadcast Epoxy flooring as described above (material sales tax EXCLUDED)",
      "before": [], "after": [], "drop": False}),
    ("blank lines round an option",
     "\n$23,115 – Treadwell 3/16\" Urethne Cement Hybrid System (material sales tax EXCLUDED) — Includes 6\" Cove Base\n\n",
     {"amount": "$23,115", "phrase": "(tax exempt)", "slot": True},
     {"main": f"{AMT} – Treadwell 3/16\" Urethne Cement Hybrid System (material sales tax EXCLUDED) — Includes 6\" Cove Base",
      "before": [""], "after": ["", ""], "drop": False}),
    # Broken out: the base line froze the TOTAL the day it was typed under one line.
    ("total frozen on a broken-out base", "$21,260 – Epoxy flooring as described above",
     {"amount": "$19,360", "phrase": "", "slot": True, "candidates": ["$21,260"]},
     {"main": f"{AMT} – Epoxy flooring as described above {TAX}", "before": [], "after": [], "drop": False}),
    # A figure that is neither today's nor the total: his, kept -- and a line typed under it.
    ("a figure of its own", "\n\n$31,054 – Sealed Concrete Option Alternate\nTest Test option: $5000\n\n",
     {"amount": "$31,348", "phrase": "", "slot": True, "candidates": ["$34,470"]},
     {"main": f"$31,054 – Sealed Concrete Option Alternate {TAX}", "before": ["", ""],
      "after": ["Test Test option: $5000", "", ""], "drop": False}),
    ("Add in front of a total", "Add $8,865 – Sealed Concrete in painted floor surface areas per drawing (tax exempt)",
     {"amount": "$8,865", "phrase": "(tax exempt)", "slot": True},
     {"main": f"Add {AMT} – Sealed Concrete in painted floor surface areas per drawing {TAX}",
      "before": [], "after": [], "drop": False}),
    ("an add/deduct line", "Add $3,189 – Upgrade to decorative quartz broadcast in lieu of flake floor system",
     {"amount": "Add $3,189", "phrase": "", "slot": False},
     {"main": f"{AMT} – Upgrade to decorative quartz broadcast in lieu of flake floor system",
      "before": [], "after": [], "drop": False}),
    ("Kyle's alignment spaces",
     "  $24,086 – Polish in Weekend Phases (3 max) as described above (material sales tax INCLUDED) ",
     {"amount": "$24,086", "phrase": "(material sales tax INCLUDED)", "slot": True},
     {"main": f"  {AMT} – Polish in Weekend Phases (3 max) as described above {TAX} ",
      "before": [], "after": [], "drop": False}),
    ("the sweep's phantom Total", "$0 – Total", {"amount": "$6,307", "zeroIsPhantom": True},
     {"main": None, "before": [], "after": [], "drop": True}),
    ("a real $0 row", "$0 – Remodel Tax", {"amount": "$0", "zeroIsPhantom": True},
     {"main": f"{AMT} – Remodel Tax", "before": [], "after": [], "drop": False}),
    # A tax wording on a line with no tax wording of its own is his (review, 2026-09-26): the old
    # page never printed one there, so migrating it into the marker deleted his words for good.
    ("a manual line with a wording of his", "$1,200 – Add for moisture mitigation (material sales tax INCLUDED)",
     {"amount": "$1,200", "phrase": "", "slot": False},
     {"main": f"{AMT} – Add for moisture mitigation (material sales tax INCLUDED)",
      "before": [], "after": [], "drop": False}),
    ("an add line with a wording of his", "Add $3,189 – Upgrade to decorative quartz broadcast (tax exempt)",
     {"amount": "Add $3,189", "phrase": "", "slot": False},
     {"main": f"{AMT} – Upgrade to decorative quartz broadcast (tax exempt)",
      "before": [], "after": [], "drop": False}),
    # His price first, today's figure quoted after it: the price stays his (review, 2026-09-26).
    ("his price quoting today's", "$5,800 – Epoxy flooring as described above (discounted from $6,307)",
     {"amount": "$6,307", "phrase": "(material sales tax INCLUDED)", "slot": True},
     {"main": "$5,800 – Epoxy flooring as described above (discounted from $6,307)",
      "before": [], "after": [], "drop": False}),
    # A combo "Option N" line the old page drew as Total less remodel under "Included".
    ("combo line frozen at total less remodel",
     "$10,000 – Option 1: Epoxy flooring in the kitchen as described above (material sales tax INCLUDED)",
     {"amount": "$10,500", "phrase": "(Remodel Tax AND material sales tax INCLUDED)", "slot": True,
      "candidates": ["$10,500", "$10,000", "$9,900"]},
     {"main": f"{AMT} – Option 1: Epoxy flooring in the kitchen as described above {TAX}",
      "before": [], "after": [], "drop": False}),
]


@needs_node
def test_the_lines_real_drafts_hold_migrate_without_losing_a_word():
    """TWPrice.migrateLine, run over the shapes Kyle's saved drafts actually hold: today's figure
    (in either money style) becomes the marker, a tax wording this tool printed becomes the other,
    his own words and figures stay exactly as typed, the blank lines and typed lines round a line
    become lines of their own, and only the box sweep's "$0 – Total" phantom is dropped. Mutation:
    no same-amount match (the cents-styled "$1,870.00" stays frozen and is flagged as his)."""
    core = FRONTEND / "js" / "price-lines-core.js"
    script = ("const P = require(process.argv[1]); let s = ''; process.stdin.on('data', d => s += d);"
              "process.stdin.on('end', () => { const cases = JSON.parse(s);"
              "console.log(JSON.stringify(cases.map(c => P.migrateLine(c[0], c[1])))); });")
    p = subprocess.run(["node", "-e", script, str(core)],
                       input=json.dumps([[t, parts] for _n, t, parts, _w in _REAL_SHAPES]),
                       capture_output=True, text=True, encoding="utf-8", timeout=60)
    assert p.returncode == 0, p.stderr
    got = json.loads(p.stdout)
    for (name, _t, _parts, want), g in zip(_REAL_SHAPES, got):
        assert g == want, (name, g)


# ── the document ────────────────────────────────────────────────────────────────────────────────
def _render(payload):
    req = Request({"type": "http", "method": "POST", "path": "/t", "headers": [], "query_string": b""})
    return main._render_documents(payload, req, want_estimate=False)["docx"]["content"]


def _paras(blob):
    d = docx.Document(io.BytesIO(blob))
    return d, [p for p in d.element.xpath("//w:p")
               if not any(True for _ in p.iterancestors(f"{_MC}Fallback"))
               and p.find(".//" + qn("w:txbxContent")) is None]


def _text(p):
    return pw._own_text(p)


def _payload(**over):
    values = {
        "project_name": "Hanz Fix", "job_name": "Hanz Fix", "city_state": "Iloilo City, KS",
        "texture": "Smooth", "system_name": "Treadwell 3/16\" Urethne Cement", "scope_notes": "s",
        "schedule_notes": "s", "exclusions": "e", "estimator_name": "Hanz", "bid_date_formatted": "9/25/26",
        "total_formatted": "$7,447", "material_tax_formatted": "$96", "tax_amount_formatted": "$0",
        "base_bid_formatted": "$7,447", "base_tax_phrase": "", "tax_layout": "ONE_LINE",
        "price_taxable": True, "price_remodel_on": False, "epoxy_sf": "99", "sqft": "99",
    }
    values.update(over.pop("values", {}))
    p = {"work_type": "epoxy", "audience": "Direct", "values": values, "remodel": [],
         "rooms": [
             {"id": "Epoxy", "name": "Epoxy", "is_base": True, "bid": {"total": 7447, "sales_tax": 96, "remodel": 0}},
             {"id": "Copy1", "name": "Epoxy copy", "is_base": False, "show": True, "price_mode": "total",
              "option_desc": "Treadwell 3/16\" Urethane Cement", "base_total": 7447,
              "bid": {"total": 7696, "sales_tax": 99, "remodel": 0, "taxable": True, "remodel_on": False}}]}
    p.update(over)
    return p


def test_typed_lines_print_as_their_own_paragraphs_in_the_price_rows_formatting():
    """What Hanz typed, as its own paragraphs right below the line it was typed under — cloned from
    that row, so 9pt Zetta Serif Book #404040 like the rows around it (a bare run would print at the
    document's 12pt default) — and the price line above stays the computed, live one."""
    blob = _render(_payload(price_overrides={"after": {
        "base": ["", "THis is a test send to Hanz"], "option:Copy1": ["Test again 123"]}}))
    _d, paras = _paras(blob)
    texts = [_text(p) for p in paras]
    i = texts.index("$7,447 – Epoxy flooring as described above (material sales tax INCLUDED)")
    assert texts[i + 1] == "" and texts[i + 2] == "THis is a test send to Hanz", texts[i:i + 4]
    j = texts.index("$7,696 – Treadwell 3/16\" Urethane Cement as described above (material sales tax INCLUDED)")
    assert texts[j + 1] == "Test again 123", texts[j:j + 3]
    for p in (paras[i + 2], paras[j + 1]):
        runs = [r for r in p.findall(qn("w:r")) if "".join(t.text or "" for t in r.iter(qn("w:t")))]
        assert runs, _text(p)
        rpr = runs[0].find(qn("w:rPr"))
        assert rpr is not None, f"{_text(p)!r} printed in a bare run"
        assert rpr.find(qn("w:rFonts")).get(qn("w:ascii")) == "Zetta Serif Book"
        assert rpr.find(qn("w:sz")).get(qn("w:val")) == "18"
        assert rpr.find(qn("w:color")).get(qn("w:val")) == "404040"
        assert rpr.find(qn("w:highlight")) is None and rpr.find(qn("w:u")) is None
    # One paragraph each: no line breaks inside the price lines any more.
    assert not any(list(p.iter(qn("w:br"))) for p in (paras[i], paras[j]))


def test_an_edited_line_prints_todays_amount_and_wording_and_a_typed_figure_as_typed():
    """Markers resolve to TODAY's figures in the document too. Broken out here: the base line has
    no bracket, the Total is the re-priced bid, and the option he re-worded quotes its new price."""
    blob = _render(_payload(
        values={"total_formatted": "$9,100", "tax_layout": "BROKEN_OUT", "material_tax_formatted": "$96"},
        rooms=[{"id": "Epoxy", "name": "Epoxy", "is_base": True, "bid": {"total": 9100, "sales_tax": 96, "remodel": 0}},
               {"id": "Copy1", "name": "Epoxy copy", "is_base": False, "show": True, "price_mode": "total",
                "option_desc": "Treadwell 3/16\" Urethane Cement", "base_total": 9100,
                "bid": {"total": 8696, "sales_tax": 99, "remodel": 0, "taxable": True, "remodel_on": False}}],
        price_overrides={"lines2": {
            "base": f"{AMT} – Epoxy flooring in the warehouse only as described above {TAX}",
            "option:Copy1": f"{AMT} – Treadwell 3/16\" Urethane Cement, shop floor {TAX}",
            "option:Copy1:total": "$8,000 – Total, as agreed",
        }}))
    _d, paras = _paras(blob)
    texts = [_text(p) for p in paras]
    assert "$9,004 – Epoxy flooring in the warehouse only as described above" in texts, texts
    assert "$8,597 – Treadwell 3/16\" Urethane Cement, shop floor" in texts
    assert "$8,000 – Total, as agreed" in texts          # his own figure, kept
    assert "$9,100 – Total" in texts


def test_a_line_saved_before_the_live_shape_still_prints_as_it_was_saved():
    """A payload frozen before this change: its whole-line override prints verbatim, as it always
    did — including the typed breaks inside it."""
    blob = _render(_payload(price_overrides={"lines": {
        "base": "$7,447 – Epoxy flooring as described above (material sales tax INCLUDED)\n\nold note"}}))
    _d, paras = _paras(blob)
    hit = [p for p in paras if _text(p).startswith("$7,447 – Epoxy flooring as described above")]
    assert hit and "old note" in _text(hit[0]) and list(hit[0].iter(qn("w:br")))


def test_a_legacy_exempt_payload_reads_its_wording_off_the_sheet():
    """"Exempt" used to be a label the estimator picked; an exempt job is now one whose sheet says
    No. A payload saved with the old EXEMPT on a job that carried sales tax prints what the sheet
    says; on a real exempt job (no sales tax) it still prints "(tax exempt)"."""
    taxed = _render(_payload(values={"tax_layout": None, "tax_inclusion": "EXEMPT",
                                     "price_taxable": None, "price_remodel_on": None}))
    exempt = _render(_payload(values={"tax_layout": None, "tax_inclusion": "EXEMPT", "price_taxable": None,
                                      "price_remodel_on": None, "material_tax_formatted": "$0"},
                              rooms=[]))
    t1 = [_text(p) for p in _paras(taxed)[1]]
    t2 = [_text(p) for p in _paras(exempt)[1]]
    assert "$7,447 – Epoxy flooring as described above (material sales tax INCLUDED)" in t1
    assert "$7,447 – Epoxy flooring as described above (tax exempt)" in t2


@pytest.mark.parametrize("work_type,audience", [("epoxy", "Direct"), ("polish", "Direct"), ("combo", "Direct"),
                                                ("epoxy", "GC"), ("polish", "GC"), ("sealer", "GC"),
                                                ("gyp", "Direct")])
def test_no_review_highlight_reaches_the_customer_on_any_template(work_type, audience):
    """Hanz: "Strip them all." The Direct files' yellow Remodel Tax row and the GC files' yellow /
    green / cyan "-or-" markers are Kyle's review marks, not customer content."""
    d = docx.Document(str(pw.pick_template(work_type, audience)))
    authored = len(list(d.element.body.iter(qn("w:highlight"))))
    blob = _render({"work_type": work_type, "audience": audience,
                    "values": _payload()["values"] | {"tax_layout": "BROKEN_OUT", "price_remodel_on": True,
                                                      "tax_amount_formatted": "$120"},
                    "remodel": [{"amount_formatted": "$120"}]})
    out = docx.Document(io.BytesIO(blob))
    assert len(list(out.element.body.iter(qn("w:highlight")))) == 0, (
        f"{work_type}/{audience}: {authored} authored highlights, some reached the document")


# ── the review of 2026-09-26: his words, his price, a combo line, and screen == document ─────────
_OPT_DESC = "Treadwell 3/16\" Urethne Cement With Shop Floor and Armor Top"


def _screen_payload(**over):
    """The harness's own draft as a payload: the same rooms, the same manual line."""
    return _payload(
        rooms=[{"id": "Epoxy", "name": "Epoxy", "is_base": True, "bid": {"total": 7447, "sales_tax": 96, "remodel": 0}},
               {"id": "Copy1", "name": "Epoxy copy", "is_base": False, "show": True, "price_mode": "total",
                "system_desc": _OPT_DESC, "option_desc": _OPT_DESC, "base_total": 7447,
                "bid": {"total": 7696, "sales_tax": 99, "remodel": 0, "taxable": True, "remodel_on": False}}],
        price_lines=[{"label": "Add for moisture mitigation", "amount": 1200}], **over)


def _price_box(blob, first, last):
    """The PRICE box as printed: its non-empty paragraphs from `first` through `last`."""
    t = [x for x in (_text(p) for p in _paras(blob)[1]) if x.strip()]
    i = t.index(first)
    j = t.index(last, i)
    return t[i:j + 1]


@needs_node
def test_a_tax_wording_of_his_own_is_kept_on_screen(ran):
    """Only TODAY's wording on a line with a place for tax wording becomes the live marker. A known
    wording on a line with none of its own (a manual line, the Total) used to become the marker and
    show as nothing; "(tax exempt)" typed over a taxable job's wording used to resolve back to the
    computed wording, the override deleted as equal to it. Mutations: every known wording a marker
    on every line again (the pre-review captureLine); the marker added behind his own wording."""
    o = ran["ownPhrase"]
    assert o["stored"] == {
        "manual:0": f"{AMT} – Add for moisture mitigation (material sales tax INCLUDED)",
        "base": f"{AMT} – Epoxy flooring as described above (tax exempt)"}, o["stored"]
    assert o["manual"] == "$1,200 – Add for moisture mitigation (material sales tax INCLUDED)"
    assert o["base"] == "$7,447 – Epoxy flooring as described above (tax exempt)"
    assert o["brokenStored"]["total"] == f"{AMT} – Total (material sales tax INCLUDED)"
    # A wording of his in the base line's empty (Broken out) slot: no marker added behind it, so
    # a switch back to one line does not print a second wording after his.
    assert o["brokenStored"]["base"] == f"{AMT} – Epoxy flooring as described above (material sales tax EXCLUDED)"
    assert o["backToOneLine"] == "$7,447 – Epoxy flooring as described above (material sales tax EXCLUDED)"
    # Saved before the live shape: migrated with his wording in it, not deleted from the draft.
    assert o["migrated"] == {"lines2": {"manual:0": f"{AMT} – Add for moisture mitigation (tax exempt)"},
                             "lines": {}, "shown": "$1,200 – Add for moisture mitigation (tax exempt)"}


@needs_node
@pytest.mark.parametrize("layout", ["ONE_LINE", "BROKEN_OUT"])
def test_a_tax_wording_of_his_own_prints_as_the_screen_shows_it(ran, layout):
    """Screen == document: what the editor stored for those lines, through the real renderer, is
    the PRICE box the editor showed, line for line."""
    o = ran["ownPhrase"]
    stored, screen = ((o["stored"], o["oneLine"]) if layout == "ONE_LINE"
                      else (o["brokenStored"], o["broken"]))
    blob = _render(_screen_payload(values={"tax_layout": layout},
                                   price_overrides={"lines2": stored}))
    want = [l["text"] for l in screen if l["text"].strip()]
    assert _price_box(blob, "Base Bid", want[-1]) == want


@needs_node
def test_his_price_that_quotes_todays_figure_is_marked_and_asked_about(ran):
    """"$5,800 – … (discounted from $7,447)": the price is the line's first figure. Today's figure
    further along is his quote, and used to become the live amount -- no mark, no Send prompt, the
    $5,800 printed as if it followed the estimate while the quote moved with every re-price.
    Mutations: the amount marker anywhere in the line (amountIndex for priceIndex); moneyOff back
    to "any marker means it follows"."""
    q = ran["quoted"]
    assert q["stored"] == "$5,800 – Epoxy flooring as described above (discounted from $7,447)"
    assert "tw-money-off" in q["cls"] and "does not follow the estimate" in q["title"]
    assert q["warnings"] == [{"key": "base", "says": "$5,800", "estimate": "$7,447"}]
    assert q["ask"] == "This line says $5,800 but the estimate says $7,447 — send anyway?"
    assert q["repriced"] == "$5,800 – Epoxy flooring as described above (discounted from $7,447)"
    # A line the previous build stored with the quote as the marker is still his price.
    assert "tw-money-off" in q["previousBuild"]["cls"]
    assert q["previousBuild"]["warnings"] == [{"key": "base", "says": "$5,800", "estimate": "$7,447"}]


@needs_node
def test_a_combo_line_reworded_under_included_migrates_live_and_adds_up(ran):
    """Under "Included" the old page drew a combo Option line as Total less remodel ($10,000) with
    the Remodel Tax and Total rows under it. Migrated with that figure frozen, one line printed
    "$10,000 … (Remodel Tax AND material sales tax INCLUDED)" on a $10,500 bid, no rows under it.
    Now the old figure is recognised as the estimate's and follows it; his words stay. Mutations:
    no `candidates` on the combo line; not passed on by the painter; not passed on by the payload."""
    c = ran["comboLegacy"]
    assert c["lines2"] == {"combo:epoxy.flooring":
                           f"{AMT} – Option 1: Epoxy flooring in the kitchen as described above {TAX}"}
    line = ("$10,500 – Option 1: Epoxy flooring in the kitchen as described above "
            "(Remodel Tax AND material sales tax INCLUDED)")
    assert [x[1] for x in c["lines"]] == [line]
    assert "tw-money-off" not in c["lines"][0][2] and c["warnings"] == []
    # Built into the payload before the page ever drew it: the same migration, the same line.
    assert c["payloadFirst"] == [{"label": line, "amount_formatted": ""}]
    assert c["broken"] == [
        ["combo:epoxy.flooring", "$9,900 – Option 1: Epoxy flooring in the kitchen as described above"],
        ["combo:epoxy.sales_tax", "$100 – Material Sales Tax"],
        ["combo:epoxy.remodel", "$500 – Remodel Tax"],
        ["combo:epoxy.total", "$10,500 – Total"]]
    # The document: the same lines, and Broken out they add up.
    for layout, payload_lines, screen in (("ONE_LINE", c["payload"], [line]),
                                          ("BROKEN_OUT", c["brokenPayload"], [x[1] for x in c["broken"]])):
        blob = _render({"work_type": "combo", "audience": "Direct", "combo_options": payload_lines,
                        "remodel": [{"amount_formatted": "$500"}],
                        "values": _payload()["values"] | {
                            "tax_layout": layout, "price_remodel_on": True, "total_formatted": "$10,500",
                            "material_tax_formatted": "$100", "tax_amount_formatted": "$500"}})
        assert _price_box(blob, screen[0], screen[-1]) == screen, layout
    rows = [int(x[1].split(" – ")[0].replace("$", "").replace(",", "")) for x in c["broken"]]
    assert rows[0] + rows[1] + rows[2] == rows[3]
