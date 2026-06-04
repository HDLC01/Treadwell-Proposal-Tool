"""Combo proposal fill + non-ASCII file download.

Guards two bugs found while generating a combo proposal:
  1. /api/file 500'd when the project name had a non-ASCII char (em-dash) because
     the Content-Disposition header is latin-1 only.
  2. The combo template's WORK + PRICE blocks weren't tokenized, so combos came
     out with dummy "$xx,xxx" / "Square Feet" placeholders.
"""
import io

from fastapi.testclient import TestClient
from docx import Document

import main
import proposal_writer

client = TestClient(main.app)

MC = "{http://schemas.openxmlformats.org/markup-compatibility/2006}"


def _rendered_text(docx_bytes: bytes) -> str:
    """Text of paragraphs Word actually renders (the mc:Choice/drawing copy),
    excluding the legacy mc:Fallback duplicates."""
    doc = Document(io.BytesIO(docx_bytes))
    out = []
    for p in doc.element.xpath("//w:p"):
        if any(True for _ in p.iterancestors(f"{MC}Fallback")):
            continue
        out.append("".join(t.text or "" for t in p.xpath(".//w:t")))
    return "\n".join(out)


# ── /api/file non-ASCII filename ───────────────────────────────────────
def test_file_download_handles_non_ascii_filename():
    token = main._cache_file(b"hello", "Combo Demo — Olathe é.xlsx",
                             "application/octet-stream")
    r = client.get(f"/api/file/{token}")
    assert r.status_code == 200
    assert r.content == b"hello"
    cd = r.headers["content-disposition"]
    assert "filename*=UTF-8''" in cd          # RFC 5987 form for modern browsers
    assert cd.split(";")[0] == "attachment"
    # the plain filename= fallback must be pure ASCII (no raw em-dash)
    assert "—" not in cd.split("filename*=")[0]


def test_file_download_404_for_unknown_token():
    assert client.get("/api/file/does-not-exist").status_code == 404


# ── combo proposal fills (combined single lump) ────────────────────────
COMBO_VALUES = {
    "job_name": "Combo QA", "city_state": "Olathe, KS", "bid_date_formatted": "6/15/26",
    "system_name": "Treadwell MACRO Flake Single Broadcast", "texture": "Orange Peel",
    "epoxy_sf": "12,000", "cove_lf": "250", "polish_sf": "8,000",
    "disposal": "a dumpster", "schedule_notes": "~5 days",
    "lump_sum_formatted": "$61,162.00", "tax_amount_formatted": "$2,639.00",
    "state_name": "Kansas", "total_formatted": "$63,801.00",
}


def test_combo_proposal_fills_price_and_areas():
    data = proposal_writer.fill_proposal(work_type="combo", audience="Direct", values=COMBO_VALUES)
    text = _rendered_text(data)
    for needle in ["$61,162.00", "$2,639.00", "$63,801.00", "Kansas Remodel Tax",
                   "12,000 Square Feet", "8,000 Square Feet",
                   "Epoxy & Polished Concrete flooring", "6/15/26"]:
        assert needle in text, f"combo proposal missing {needle!r}"


def test_combo_proposal_has_no_unfilled_placeholders():
    data = proposal_writer.fill_proposal(work_type="combo", audience="Direct", values=COMBO_VALUES)
    text = _rendered_text(data)
    import re
    assert not re.search(r"\{\{[^}]+\}\}", text), "unsubstituted {{token}} left in combo output"
    for dummy in ("$xx,xxx", "Square Feet, Lineal Feet", "xx/xx/26"):
        assert dummy not in text, f"combo dummy placeholder {dummy!r} not replaced"
