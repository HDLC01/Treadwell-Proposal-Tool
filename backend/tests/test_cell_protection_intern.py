"""_apply_cell_protection writes the protection style index, not the object.

Setting `cell.protection = Protection(locked=False)` runs openpyxl's
StyleDescriptor, which calls `wb._protections.add(value)` for EVERY cell — a
hash + compare of a Serialisable just to hand back an index it already knew.
On Kyle's template that is 84,670 cells across the 11 protected sheets, and it
was ~470ms of the ~500ms the whole protection pass cost, on every /api/generate,
every portal PDF, every revision download and every To-Dropbox re-file.

The pass now interns the two Protection records once per sheet and writes the
index straight into each cell's StyleArray. That is the same write the
descriptor performs, so the risk is not "is it faster" — it is "did the
document change". These tests answer both:

  * equivalence — replay the ORIGINAL descriptor loop over an identical
    workbook and demand every resulting StyleArray match, cell for cell;
  * the perf contract itself — count the interning calls, so a well-meaning
    revert to `cell.protection = unlocked` goes red instead of quietly
    costing half a second a generate again.

Both guard a single function, so they share one expensive template parse.
"""
import pytest
from openpyxl.styles import Protection
from openpyxl.utils.indexed_list import IndexedList

import estimate_writer as ew


# The protection pass only touches sheets that carry locks or ship protected.
# Everything here is measured against that same set.
def _layouts(wb):
    return ew._resolve_ws_layouts(wb, None, {}, {})


def _reference_apply(ws_layouts):
    """The pre-optimisation loop, verbatim, so equivalence is checked against
    the real openpyxl descriptor rather than against a paraphrase of it."""
    unlocked = Protection(locked=False)
    locked = Protection(locked=True)
    wb = ws_layouts[0][0].parent
    for style in wb._named_styles:
        if style.name == "Normal":
            style.protection = Protection(locked=False)
    for ws, addrs in ws_layouts:
        if not addrs:
            ws.protection.sheet = False
            ws.protection.disable()
            continue
        for cell in list(ws._cells.values()):
            cell.protection = unlocked
        for addr in addrs:
            ws[addr].protection = locked
        ws.protection.sheet = True
        ws.protection.formatCells = False
        ws.protection.formatColumns = False
        ws.protection.formatRows = False
        ws.protection.enable()


class _CountingProtections(IndexedList):
    """An IndexedList that remembers how often it was asked to intern."""

    def add(self, value):
        self.adds += 1
        return super().add(value)


@pytest.fixture(scope="module")
def reference():
    """A workbook protected the OLD way — the thing the new code must match."""
    wb = ew._fresh_template()
    _reference_apply(_layouts(wb))
    return wb


@pytest.fixture(scope="module")
def optimised():
    """The same workbook protected by the shipping code, plus the intern count."""
    wb = ew._fresh_template()
    counting = _CountingProtections(wb._protections)
    counting.adds = 0
    wb._protections = counting
    ew._apply_cell_protection(_layouts(wb))
    return wb


def test_every_cell_gets_the_same_style_array_as_the_descriptor(reference, optimised):
    """Cell for cell, the direct index write must land what openpyxl's
    descriptor landed — not just the locked flag, the WHOLE StyleArray, so a
    stray write to the font or number-format slot can't slip through."""
    compared = 0
    for name in reference.sheetnames:
        a, b = reference[name], optimised[name]
        assert set(a._cells) == set(b._cells), f"{name}: cell set differs"
        for key, cell_a in a._cells.items():
            cell_b = b._cells[key]
            sa = None if cell_a._style is None else list(cell_a._style)
            sb = None if cell_b._style is None else list(cell_b._style)
            assert sa == sb, f"{name}!{cell_a.coordinate}: style {sa} vs {sb}"
            compared += 1
    # A comparison that walked an empty workbook would pass too. Kyle's template
    # carries ~171k cells in total; anything near zero means the fixture broke,
    # not that the code is right.
    assert compared > 50_000, f"only {compared} cells compared — fixture is hollow"


def test_protection_settings_and_the_normal_style_match(reference, optimised):
    """The flags around the per-cell write — which sheets end up protected, what
    Excel still allows on them, and the Normal-style unlock that keeps blank
    cells typeable — have to survive the change as well."""
    for name in reference.sheetnames:
        a, b = reference[name].protection, optimised[name].protection
        for attr in ("sheet", "formatCells", "formatColumns", "formatRows",
                     "insertRows", "deleteRows", "password"):
            assert getattr(a, attr, None) == getattr(b, attr, None), \
                f"{name}.protection.{attr}"
    for wb in (reference, optimised):
        normal = next(s for s in wb._named_styles if s.name == "Normal")
        assert normal.protection.locked is False


def test_protection_is_interned_per_sheet_not_per_cell(optimised):
    """The perf contract. The budget is two interns per protected sheet (one
    for locked, one for unlocked) plus one for the Normal-style rebind, which
    re-interns the whole style when its protection is reassigned. The
    descriptor version spent one per CELL, which is what made this ~470ms."""
    layouts = _layouts(optimised)
    touched = sum(len(ws._cells) for ws, addrs in layouts if addrs)
    # Without a workload this assertion proves nothing — make the gap explicit.
    assert touched > 50_000, f"only {touched} cells touched — fixture is hollow"
    budget = 2 * len(layouts) + 1
    assert optimised._protections.adds <= budget, (
        f"interned {optimised._protections.adds} times (budget {budget}) for "
        f"{len(layouts)} sheets and {touched} cells — the per-cell lookup is back"
    )


def test_interning_added_no_new_protection_records(optimised):
    """Kyle's template already carries both Protection records, so hoisting the
    intern cannot shift any index: `add` finds them and appends nothing. This is
    why the hoist is safe to do up front instead of in call order — if a future
    template ever arrived without one of them, this goes red and the ordering
    argument has to be re-made."""
    prots = [(p.locked, p.hidden) for p in optimised._protections]
    assert prots == [(True, False), (False, False)], prots
