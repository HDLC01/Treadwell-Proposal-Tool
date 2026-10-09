"""Phase 7b of the v2 estimating program: the condition cells respect a split draft, a condition with several
cells is read from the first one that holds an answer, and every literal flag cell on the template is either
written or listed with the phase that will write it.

WHAT REVIEW OF PHASE 7 FOUND. js/work-types.js made Taxable four cells, which fixed a tax-exempt Leveling or
Gypsum option charging 9.475% on a v2 bid. It also made the v2 save write all four on EVERY draft, including a
draft the estimate screen has split per sheet (`tax_flags_per_sheet`), where each sheet keeps its own Taxable and
Remodel answer and the live intake refuses to restate them. A v2 test copy of such a project dropped the mark
as well, so it arrived looking unsplit. And the read-back looked at the first cell of a condition only: "Reno" in
Polish!B10 with Epoxy!B10 blank was read as nothing and saved back as "New".

WHAT THIS FILE PINS, all of it EXECUTED:
  1. the rule (js/work-types.js writeCellsFor), for every condition and job type, split and not;
  2. the v2 model's writes and read-back through it, including the first cell that holds an answer;
  3. the test copy (js/polish-sandbox.js buildCopy) of a split project;
  4. THE JOURNEY, end to end: a split project is copied, opened in v2 and saved, through the modules, through
     the REAL Takeoff step and through the REAL v2 intake. The option sheets' answers come out as they went in;
  5. the live intake, which asks the same table the same question (and still drops an address that is no cell);
  6. a scan of every priced sheet of Kyle's template for its own literal flag cells (Local, Taxable, Prevailing
     wage, Remodel tax, New or Reno): each is written by the table or on KNOWN_GAPS, with the phase that will
     write it and the reason, and a new unlisted one fails;
  7. the architecture document, which said a save "loses or changes none".

RED WITHOUT THE CHANGE. Every check is a plain function of what a harness answered, and BREAKS at the bottom
breaks one line of one file in a scratch copy, runs the harness on it, and requires the named check to fail. A
check no break turns red proves nothing, so the last tests require every check to be named there. The journey
harnesses use only calls that exist in the tree before this phase as well, so they can be pointed at that tree.

Run under node; a missing node FAILS under CI (tests/_node.py).
"""
from __future__ import annotations

import json
import pathlib
import re
import shutil
import warnings

import pytest

from _golden_support import FIXTURES, FRONTEND, break_source, copy_unmodified
from _node import last_json_line, require_node, run_node

import test_intake_conditions as live_intake_tests

TESTS = pathlib.Path(__file__).resolve().parent
ROOT = TESTS.parents[1]
JS = TESTS / "js"
MODULE_HARNESS = JS / "condition-cells-harness.js"
TAKEOFF_HARNESS = JS / "condition-cells-takeoff-harness.js"
INTAKE_HARNESS = JS / "condition-cells-intake-harness.js"
BETA_HARNESS = JS / "beta-routing-harness.js"
TEMPLATE = ROOT / "backend" / "templates" / "estimate_sheet_5.7.xlsx"
DOC = ROOT / "docs" / "v2-architecture.md"

# What each page harness reads out of a frontend directory (they do not fall back to the real one), so a
# scratch copy that breaks ONE file can be completed with the rest.
TAKEOFF_FILES = ["polish-estimate.html", "js/polish-estimate.js", "js/bid-model.js", "js/library-core.js",
                 "js/work-types.js", "js/polish-sandbox.js", "js/library-picker.js"]
INTAKE_FILES = ["polish-intake.html", "js/polish-intake.js", "js/bid-model.js", "js/work-types.js",
                "js/polish-sandbox.js"]
BETA_FILES = ["js/index.js", "index.html", "shared.js", "js/projects.js", "js/county-picker.js",
              "js/address-lookup.js", "js/work-types.js", "js/intake-scope.js"]

GYP = 'Gyp (USG 1-8")'
JOBS = ["epoxy", "polish", "combo", "gyp"]
# The four literals a job's Taxable answer fans out to while the draft is not split.
FAN_OUT = ["Epoxy!B6", "Leveling!B6", GYP + "!B8", "Gyp (FR)!B8"]
# Where a SPLIT draft keeps the job's two answers, per job type: the base sheets' own cells.
SPLIT_TAXABLE = {"epoxy": ["Epoxy!B6"], "polish": ["Polish!B6"], "combo": ["Epoxy!B6", "Polish!B6"],
                 "gyp": [GYP + "!B8"]}
SPLIT_REMODEL = {"epoxy": ["Epoxy!D6"], "polish": ["Polish!D6"], "combo": ["Epoxy!D6", "Polish!D6"],
                 "gyp": [GYP + "!D8"]}
# What the harnesses' split project holds in every tax cell BUT the base's own (Polish!B6 and Polish!D6): the
# Leveling and Gypsum (FR) sheets are tax-exempt, Epoxy is taxable. A save that honours the split leaves each
# of these exactly as it found it.
OPTIONS = {"Epoxy!B6": "Yes", "Epoxy!D6": "No", "Leveling!B6": "No", "Gyp (FR)!B8": "No",
           GYP + "!B8": "Yes", GYP + "!D8": "Yes"}


def _run(harness, frontend=FRONTEND):
    require_node()
    proc = run_node(harness, frontend)
    assert proc.returncode == 0, (
        "the harness itself failed. Read this before assuming a product bug:\n" + proc.stderr)
    return last_json_line(proc.stdout)


@pytest.fixture(scope="module")
def ran():
    return _run(MODULE_HARNESS)


@pytest.fixture(scope="module")
def takeoff():
    return _run(TAKEOFF_HARNESS)


@pytest.fixture(scope="module")
def intake():
    return _run(INTAKE_HARNESS)


@pytest.fixture(scope="module")
def live():
    return _run(BETA_HARNESS)["conditions"]


@pytest.fixture(scope="module")
def book():
    return _template()


def _template():
    import openpyxl
    with warnings.catch_warnings():
        warnings.simplefilter("ignore")
        return openpyxl.load_workbook(TEMPLATE)


def _options_as_they_were(cells):
    """The tax cells a split draft's options hold, read off a save: every one is what the source had."""
    got = {k: cells[k] for k in OPTIONS}
    assert got == OPTIONS, "a save changed an option sheet's own tax answer: %r" % (got,)


# ══ 1. the rule ══════════════════════════════════════════════════════════════════════════════════════
def test_the_table_marks_exactly_the_two_tax_conditions_as_per_sheet(ran):
    """Taxable and Remodel tax are the two answers a sheet can own. They carry the base sheet's own cell for
    every tab a job can be priced on, and for no option-only tab (Seal and Leveling are never a base)."""
    table = ran["rule"]["table"]
    marked = {c["key"]: c["perSheet"] for c in table if c["perSheet"]}
    assert set(marked) == {"taxable", "remodel_tax"}
    for key, per in marked.items():
        assert set(per) == {"epoxy", "polish", "gyp"}, key
    assert ran["rule"]["isPerSheet"] == {c["key"]: {"ok": c["key"] in marked} for c in table}
    assert marked["taxable"] == {"epoxy": "Epoxy!B6", "polish": "Polish!B6", "gyp": GYP + "!B8"}
    assert marked["remodel_tax"] == {"epoxy": "Epoxy!D6", "polish": "Polish!D6", "gyp": GYP + "!D8"}


def test_each_per_sheet_cell_is_that_sheets_flag_cell_in_the_backend_and_in_the_template(ran, book):
    """An independent source on each side: backend/estimate_writer.py's table of every flag-block sheet and its
    (Taxable?, Remodel Tax?) cells, and the labels in Kyle's own template."""
    import estimate_writer as ew
    base = {"epoxy": "Epoxy", "polish": "Polish", "gyp": ew.GYP_SHEET}
    marked = {c["key"]: c["perSheet"] for c in ran["rule"]["table"] if c["perSheet"]}
    found = _flag_cells(book, ran["sheets"])
    for tab, sheet in base.items():
        taxable, remodel = ew.FLAG_BLOCK_CELLS[sheet]
        assert marked["taxable"][tab] == "%s!%s" % (sheet, taxable), tab
        assert marked["remodel_tax"][tab] == "%s!%s" % (sheet, remodel), tab
        assert found[(sheet, "taxable")][0] == taxable and found[(sheet, "remodel_tax")][0] == remodel, sheet


def test_an_unsplit_draft_writes_every_cell_of_the_row_and_a_split_one_the_base_sheets_own(ran):
    """The rule, for every condition and job type. Not split, or a condition no sheet owns: the condition's
    `cells`, all of them. Split, and Taxable or Remodel tax: the base sheets' own cell, one per tab of the job
    type (a combo job is priced on Epoxy and Polish). The expectations are written out here, not read from the
    table."""
    table = {c["key"]: c for c in ran["rule"]["table"]}
    assert len(ran["rule"]["grid"]) == len(table)
    for key, row in ran["rule"]["grid"].items():
        for job in JOBS:
            answers = row[job]
            assert answers["unsplit"] == {"ok": table[key]["cells"]}, (key, job)
            expected = {"taxable": SPLIT_TAXABLE, "remodel_tax": SPLIT_REMODEL}.get(key, {}).get(job, table[key]["cells"])
            assert answers["split"] == {"ok": expected}, (key, job)


def test_a_job_type_is_priced_on_the_first_sheet_of_each_of_its_tabs(ran):
    assert ran["rule"]["baseSheets"] == {
        "epoxy": {"ok": ["Epoxy"]}, "polish": {"ok": ["Polish"]}, "combo": {"ok": ["Epoxy", "Polish"]},
        "gyp": {"ok": [GYP]}}


def test_a_caller_that_knows_the_base_better_replaces_the_tables_cells_and_only_real_addresses_survive(ran):
    """The live intake hands in the estimate screen's snapshot of each tab's own cells, which follows a copied
    tab or a moved row where the table's template addresses cannot. On a split draft that list IS the answer.
    Only a real "Sheet!A1" survives (an address becomes a key in cell_values and can come out of a draft), an
    empty list means none is known, and it is ignored on a draft that is not split and for a condition no sheet
    owns."""
    own = ran["rule"]["own"]
    assert own["real"] == {"ok": ["Polish!B6", "Epoxy!B6"]}, "the list did not replace the table's cell, in order"
    assert own["filtered"] == {"ok": ["Polish!B6", "Gyp (FR)!B8"]}, "an address that is no cell got through"
    assert own["empty"] == {"ok": []}
    assert own["ignoredWhenUnsplit"] == {"ok": FAN_OUT}
    assert own["ignoredForAConditionNoSheetOwns"] == {"ok": ["Epoxy!B4", "Polish!B4"]}
    assert own["notAList"] == {"ok": ["Polish!B6"]}, "something that is not a list was read as one"


def test_the_rule_refuses_what_it_was_not_asked_properly(ran):
    """Job type and split are REQUIRED, and a tab is not a job type: a default is how a gypsum job once read
    the Polish defaults, and a guess at `split` is how a split draft gets restated."""
    r = ran["rule"]["refusals"]
    for name, answer in r.items():
        assert "threw" in answer, "%s answered %r" % (name, answer)
        assert answer["threw"].startswith("work-types.js: "), (name, answer)
    for name in ("noSplit", "splitOne", "splitYes", "splitNull"):
        assert "not true or false" in r[name]["threw"], name
    assert "a tab is not a job type" in r["aTab"]["threw"] and "a tab is not a job type" in r["baseSheetsATab"]["threw"]
    for name in ("noJob", "aLookalike", "baseSheetsNothing"):
        assert "not a job type" in r[name]["threw"], name
    for name in ("noCondition", "notACondition", "isPerSheetNotACondition"):
        assert "not a job condition" in r[name]["threw"], name


def test_the_draft_is_split_when_it_carries_the_estimate_screens_mark_and_for_no_other_reason(ran):
    got = [a["ok"] for a in ran["rule"]["isSplit"]]
    # undefined, null, {}, mark true, mark false, mark 1, mark "", "abc", 5, [], mark undefined
    assert got == [False, False, False, True, False, True, False, False, False, False, False]


def test_a_caller_that_edits_what_it_was_handed_edits_nothing_the_next_caller_reads(ran):
    assert ran["rule"]["own"]["isACopy"] == {"ok": {
        "unsplit": FAN_OUT, "split": ["Polish!B6"], "baseSheets": ["Epoxy", "Polish"]}}


# ══ 2. the v2 model, on cells a split draft holds ════════════════════════════════════════════════════
def test_a_save_on_a_draft_that_is_not_split_writes_the_four_taxable_cells_as_before(ran):
    """Phase 7's behaviour, unchanged: every literal follows the one answer, and the base sheets' own cells are
    not written at all."""
    w = ran["model"]["writes"]
    assert {k: w["unsplitOn"][k] for k in FAN_OUT} == dict.fromkeys(FAN_OUT, "Yes")
    assert {k: w["unsplitOff"][k] for k in FAN_OUT} == dict.fromkeys(FAN_OUT, "No")
    assert w["unsplitOn"]["Epoxy!D6"] == "Yes" and w["unsplitOff"]["Epoxy!D6"] == "No"
    for state in ("unsplitOn", "unsplitOff"):
        assert [w[state][k] for k in ("Polish!B6", "Polish!D6", GYP + "!D8")] == [None, None, None], state


def test_a_save_on_a_split_draft_writes_the_base_sheets_own_cells_and_no_other_tax_cell(ran):
    """THE DEFECT. Leveling!B6 and Gyp (FR)!B8 say No, Epoxy is taxable, and the base (Polish) is what the job
    answers for. Every other sheet's cell comes out as it went in, in both directions."""
    w = ran["model"]["writes"]
    for state, base in (("splitOn", ("Yes", "Yes")), ("splitOff", ("No", "No"))):
        _options_as_they_were(w[state])
        assert (w[state]["Polish!B6"], w[state]["Polish!D6"]) == base, state
    # a split draft that holds no cell yet: the base's two are written, and nothing else tax is invented
    assert w["splitEmpty"] == {"Epoxy!B6": None, "Leveling!B6": None, GYP + "!B8": None, "Gyp (FR)!B8": None,
                               "Polish!B6": "No", "Epoxy!D6": None, "Polish!D6": "Yes", GYP + "!D8": None}
    # every other condition writes the same cells whether or not the draft is split
    assert w["otherCells"]["same"] is True
    assert w["otherCells"]["keys"] == ["Epoxy!B10", "Epoxy!B4", "Epoxy!D5", "Polish!B10", "Polish!B4",
                                       "Polish!E25", "Polish!E29", "Polish!F29"]


def test_a_save_is_idempotent_and_leaves_its_inputs_alone(ran):
    assert ran["model"]["writes"]["twice"] is True
    assert ran["model"]["writes"]["inputsUntouched"] is True


def test_the_read_back_on_a_split_draft_is_the_base_sheets_own_cell(ran):
    """The cell wins over the model, and on a split draft the cell is the BASE sheet's own: Polish's answer,
    not Epoxy's (which on a polish job is not the base) and not an option's. Each case is posed so reading the
    wrong cell gives the other answer."""
    r = ran["model"]["reads"]
    assert r["split"] == {"taxable": True, "remodel_tax": False}, "the base's own answers were not read"
    assert r["splitBaseSaysNo"] == {"taxable": False, "remodel_tax": False}
    assert r["splitEpoxyOnly"] == {"taxable": True, "remodel_tax": False}, "Epoxy's cell was read on a split draft"
    assert r["splitOptionsOnly"] == {"taxable": True}, "an option sheet's cell was read as the job's answer"
    assert r["unsplit"] == {"taxable": True, "remodel_tax": False} == r["unsplitExplicit"], (
        "a draft that is not split reads Epoxy's cell, as it always did")


def test_a_condition_with_several_cells_is_read_from_the_first_one_that_holds_an_answer(ran):
    r = ran["model"]["reads"]
    assert r["localSecond"] == {"local": False}, "Polish!B4 alone was read as nothing"
    assert r["localFirst"] == {"local": True}, "the first cell no longer wins when it holds an answer"
    assert r["taxableSecond"] == {"taxable": False}


def test_renovation_survives_a_save_whichever_of_its_two_cells_holds_it(ran):
    """The model has no renovation answer, so a save carries the cells' own. A "Reno" in Polish!B10 with
    Epoxy!B10 blank used to be written back as "New" in both (the first cell was blank, so the default
    applied): a blank B10 takes the Reno branch of Polish!C17 and triples the patch material rate, so a silent
    flip the other way is an underbid with nothing on screen to show it."""
    r = ran["model"]["reno"]
    both = lambda a, b: {"Epoxy!B10": a, "Polish!B10": b}  # noqa: E731
    assert r["polishOnly"] == both("Reno", "Reno")
    assert r["epoxyOnly"] == both("Reno", "Reno")
    assert r["bothNew"] == both("New", "New")
    assert r["firstWins"] == both("New", "New"), "the first cell that holds an answer still wins"
    assert r["firstIsSpaces"] == both("Reno", "Reno"), "whitespace is no answer for Renovation"
    assert r["secondIsSpelledOddly"] == both("Reno", "Reno")
    assert r["blank"] == both("New", "New"), "a blank is New, never blank"
    assert r["survivesTwoSaves"] == both("Reno", "Reno")


# ══ 3. the test copy ═════════════════════════════════════════════════════════════════════════════════
def test_a_copy_of_a_split_project_arrives_split_with_its_tax_answers(ran):
    """buildCopy keeps `tax_flags_per_sheet` (it is no price: it says how the copied tax cells are read) and
    the base sheets' own cells beside the live intake's, and still leaves the sheet's working and everything
    derived behind. A source that is not split gives a copy that is not marked."""
    s = ran["sandbox"]
    assert s["marker"] is True and s["copyableKeysHasMarker"] is True
    assert s["unsplitHasNoMarker"] is True, "a copy of a draft that is not split was marked split"
    assert s["taxCells"] == {"Epoxy!B6": "Yes", "Leveling!B6": "No", GYP + "!B8": "Yes", "Gyp (FR)!B8": "No",
                             "Polish!B6": "Yes", "Epoxy!D6": "No", "Polish!D6": "No", GYP + "!D8": "Yes"}
    assert s["sheetCellDropped"] is True, "a cell that is the sheet's own working was copied"
    assert s["derivedLeftBehind"] == [], s["derivedLeftBehind"]
    assert s["sourceUntouched"] is True
    assert ran["rule"]["copyable"]["ok"] == s["copyableCells"], "the copy keeps something other than the table's cells"
    assert {"Polish!B6", "Polish!D6", GYP + "!D8"} <= set(s["copyableCells"])
    assert len(s["copyableCells"]) == len(set(s["copyableCells"]))


# ══ 4. THE JOURNEY: copy a split project, open it in v2, save, look at the option sheets ═════════════
def _base(cells, taxable, remodel):
    assert (cells["Polish!B6"], cells["Polish!D6"]) == (taxable, remodel), (cells["Polish!B6"], cells["Polish!D6"])


def test_a_split_project_copied_and_saved_in_v2_keeps_its_option_sheets_tax_answers(ran):
    """Through the REAL buildCopy and a v2 save built the way the pages build it. The task's own case: Leveling
    and Gypsum (FR) say No, Epoxy says Yes, the draft is split. Saved twice, as the page does on every edit.
    The same answers price the same on a split draft as on one that is not."""
    j = ran["journey"]
    for name, taxable, remodel in (("split", True, False), ("baseExempt", False, False), ("flipped", False, True)):
        r = j[name]
        assert r["copyKeys"].count("tax_flags_per_sheet") == 1, (name, "the copy lost the mark")
        assert (r["taxable"], r["remodel"]) == (taxable, remodel), name
        _options_as_they_were(r["cells"])
        _options_as_they_were(r["cellsAgain"])
        _base(r["cells"], "Yes" if taxable else "No", "Yes" if remodel else "No")
        assert r["cellsAgain"] == r["cells"], name
    assert j["split"]["total"] == j["unsplit"]["total"], "the draft being split moved a price"
    assert j["baseExempt"]["total"] != j["split"]["total"], "the exempt answer moved no price: the case proves nothing"
    # a draft that is not split keeps the fan-out: every literal follows the one answer
    assert "tax_flags_per_sheet" not in j["unsplit"]["copyKeys"]
    assert {k: j["unsplit"]["cells"][k] for k in FAN_OUT} == dict.fromkeys(FAN_OUT, "Yes")


def test_a_v2_copy_of_a_split_project_that_is_not_polish_opens_at_the_source_bases_own_tax_answer(ran):
    """v2 prices a POLISH bid, so on a split draft it reads and writes Polish!B6 and Polish!D6. A project whose
    job is priced on other sheets had its real answer changed on the grid after the split stamped Polish's, so
    the copy has to start with that answer in the Polish cells, or the beta price opens taxed (Epoxy exempt,
    Polish stale Yes) and no longer matches the live bid it was copied to compare against. Saved twice."""
    j = ran["journey"]
    # epoxy job: Epoxy!B6 No and Epoxy!D6 Yes are the base's answer; Polish!B6 was stamped Yes and Polish!D6 No
    for name in ("epoxyBased", "epoxyBasedNoPolish"):
        r = j[name]
        assert (r["taxable"], r["remodel"]) == (False, True), name
        for cells in (r["cells"], r["cellsAgain"]):
            assert (cells["Polish!B6"], cells["Polish!D6"]) == ("No", "Yes"), (name, cells)
            # the source's own cells and every option's are exactly as the source had them
            assert (cells["Epoxy!B6"], cells["Epoxy!D6"]) == ("No", "Yes"), name
            assert (cells["Leveling!B6"], cells["Gyp (FR)!B8"], cells[GYP + "!B8"], cells[GYP + "!D8"]) == \
                ("No", "No", "Yes", "Yes"), name
        assert r["cellsAgain"] == r["cells"], name
    # gyp job: the gyp base's own B8 and D8
    g = j["gypBased"]
    assert (g["taxable"], g["remodel"]) == (False, True)
    assert (g["cells"]["Polish!B6"], g["cells"]["Polish!D6"]) == ("No", "Yes")
    assert (g["cells"][GYP + "!B8"], g["cells"][GYP + "!D8"]) == ("No", "Yes"), "the gyp base's own cells moved"
    assert g["cellsAgain"] == g["cells"]
    # a combo job has a Polish base of its own, and its answer is the one v2 reads
    c = j["comboBased"]
    assert (c["taxable"], c["remodel"]) == (True, False)
    assert (c["cells"]["Polish!B6"], c["cells"]["Polish!D6"]) == ("Yes", "No")
    # a job the vocabulary does not know is read as the polish job v2 prices, and a draft that is not split
    # keeps the fan-out untouched
    u = j["unknownJob"]
    assert (u["cells"]["Polish!B6"], u["cells"]["Polish!D6"]) == ("Yes", "No")
    e = j["epoxyUnsplit"]
    assert {k: e["cells"][k] for k in FAN_OUT} == dict.fromkeys(FAN_OUT, "No")


def _page_journey(r, name, base_taxable, base_remodel):
    """What a page's save did to a copy of the split project: the options untouched, the base sheet's own two
    cells holding the answers the page had (the same on the second save), and nothing written on opening."""
    _options_as_they_were(r["cells"])
    _options_as_they_were(r["cellsAgain"])
    _base(r["cells"], base_taxable, base_remodel)
    assert r["cellsAgain"] == r["cells"], name
    assert r["savesOnOpen"] == 0, "opening a copy wrote to the draft"
    assert r["marker"] is True


def test_the_takeoff_step_saves_a_split_copy_without_touching_its_options(takeoff):
    """The REAL Takeoff step, on a copy the REAL sandbox made. `taxableRead` is what the page took off the
    draft when it opened: the base sheet's own answer, not Epoxy's Yes."""
    t = takeoff
    _page_journey(t["split"], "split", "Yes", "No")
    assert (t["split"]["taxableRead"], t["split"]["remodelRead"]) == (True, False)
    # Polish is tax-exempt in the draft while Epoxy says Yes: the page must read THAT, and a save keeps it
    _page_journey(t["baseExempt"], "baseExempt", "No", "No")
    assert (t["baseExempt"]["taxableRead"], t["baseExempt"]["remodelRead"]) == (False, False), (
        "the page read Epoxy's answer on a split draft")
    # flipped in v2: the base's answers move, the options' do not
    _page_journey(t["flipped"], "flipped", "No", "Yes")
    # a draft that is not split: the fan-out, every literal follows the one answer
    u = t["unsplit"]
    assert u["marker"] is None and u["savesOnOpen"] == 0
    assert {k: u["cells"][k] for k in FAN_OUT} == dict.fromkeys(FAN_OUT, "No")
    assert u["taxableRead"] is True, "an unsplit draft is read from Epoxy's cell"
    # Renovation in the second cell only: it survives opening and two saves
    assert t["reno"]["reno"] == t["reno"]["renoAgain"] == {"Epoxy!B10": "Reno", "Polish!B10": "Reno"}
    _options_as_they_were(t["reno"]["cells"])
    # a copy that already holds a v2 estimate: the model says taxable and no remodel, the base sheet's own cells
    # say the opposite, Epoxy's say the opposite of those. The cell wins and it is Polish's.
    s = t["savedV2"]
    _page_journey(s, "savedV2", "No", "Yes")
    assert (s["taxableRead"], s["remodelRead"]) == (False, True), (
        "the page read Epoxy's answers, or the model's, over the base sheet's own cells")


def test_the_v2_intake_shows_and_saves_a_split_copy_without_touching_its_options(intake):
    """The REAL v2 intake, on a copy the REAL sandbox made, saved because the estimator touched an UNRELATED
    switch. `shown` is what its switches say on arrival: the base sheet's own answers."""
    i = intake
    assert i["split"]["shown"] == {"prevailing_wage": False, "taxable": True, "remodel_tax": False, "bond": False}
    _page_journey(i["split"], "split", "Yes", "No")
    # the base sheet says tax-exempt and remodel: the switches must say so, and a save puts them back
    assert i["baseOwn"]["shown"] == {"prevailing_wage": False, "taxable": False, "remodel_tax": True, "bond": False}, (
        "the switches showed Epoxy's answers on a split draft")
    _page_journey(i["baseOwn"], "baseOwn", "No", "Yes")
    assert (i["baseOwn"]["taxable"], i["baseOwn"]["remodel"]) == (False, True)
    # the estimator flips Taxable here: the base's answer moves, the options' do not
    assert i["flipped"]["shown"]["taxable"] is True and i["flipped"]["taxable"] is False
    _page_journey(i["flipped"], "flipped", "No", "No")
    # a draft that is not split: the fan-out
    u = i["unsplit"]
    assert u["marker"] is None and u["savesOnOpen"] == 0
    assert {k: u["cells"][k] for k in FAN_OUT} == dict.fromkeys(FAN_OUT, "No")
    # Renovation in the second cell only
    assert i["reno"]["reno"] == i["reno"]["renoAgain"] == {"Epoxy!B10": "Reno", "Polish!B10": "Reno"}
    _options_as_they_were(i["reno"]["cells"])
    # a copy that already holds a v2 estimate (nothing is seeded onto it): the model says taxable and no
    # remodel, the base sheet's own cells say the opposite, Epoxy's say the opposite of those. The cell wins.
    s = i["savedV2"]
    assert s["shown"] == {"prevailing_wage": False, "taxable": False, "remodel_tax": True, "bond": False}, (
        "the switches showed Epoxy's answers, or the model's, over the base sheet's own cells")
    _page_journey(s, "savedV2", "No", "Yes")


# ══ 5. the live intake asks the same table the same question ═════════════════════════════════════════
def test_the_live_intake_still_writes_a_split_drafts_base_cell_and_nothing_else(live):
    """test_intake_conditions.py's two split tests, which pin the live intake's behaviour, run on this tree
    through the shared rule: the same switches, the same cells, the same work-type changes."""
    live_intake_tests.test_on_a_split_draft_the_two_tax_switches_are_the_bases_own_cells(live)
    live_intake_tests.test_a_work_type_change_on_a_split_draft_rereads_the_new_bases_own_cells(live)


def test_the_live_intake_writes_both_of_a_combo_jobs_base_sheets_and_no_other_tax_cell(live):
    """A combo job is priced on Epoxy and Polish, which `baseSheets` of the table says. Polish starts on Yes
    here (the older combo case starts it on No, the answer it ends on, so it cannot see a lost second sheet);
    flipping Taxable has to write No into BOTH halves, and leave Leveling and the gyp base as they were."""
    assert live["comboBothHalves"]["cells"] == {
        "Epoxy!B6": "No", "Polish!B6": "No", "Epoxy!D6": "No", "Polish!D6": "No", "Leveling!B6": "Yes",
        GYP + "!B8": "No"}


def test_the_live_intake_drops_an_address_that_is_not_a_cell(live):
    """priced_tabs[].flag_cells comes out of a draft and becomes a KEY in cell_values, so only a real
    "Sheet!A1" may be written. The check moved from the page into the shared rule; every address the
    snapshot offers here is wrong, so flipping the two switches writes no tax cell and invents no key."""
    h = live["hostileFlagCells"]
    assert h["invented"] == [], h["invented"]
    assert h["taxCells"] == {"Epoxy!B6": "Yes", "Polish!B6": "No", "Polish!D6": "No"}


# ══ 6. every literal flag cell on the template is written or listed ══════════════════════════════════
# The labels of the flag block, in column A (the value is beside it, in B) and column C (the value is in D).
# Normalised: whitespace collapsed and lower case. Hard Bid? is scanned too, and is NEVER written (Hanz,
# 2026-10-03): it is accounted for below by its own rule, not by KNOWN_GAPS.
LABELS = {"local?": "local", "hard bid?": "hard_bid", "taxable?": "taxable", "prevailingwage?": "prevailing_wage",
          "remodel tax?": "remodel_tax", "new / reno:": "reno"}
KINDS = sorted(LABELS.values())


def _normal(text):
    return re.sub(r"\s+", " ", str(text)).strip().lower()


def _flag_cells(book, sheets):
    """{(sheet, kind): (cell, value)}: the cell beside each flag label in the project block of each priced
    sheet, the sheets being every sheet the vocabulary files under a tab (the harness reports them). Every
    (sheet, kind) must be found exactly once, so a template that moves a row says so loudly."""
    found, seen = {}, {}
    for sheet in sheets:
        ws = book[sheet]
        for row in range(1, 15):
            for col, target in (("A", "B"), ("C", "D")):
                label = ws["%s%d" % (col, row)].value
                kind = LABELS.get(_normal(label)) if isinstance(label, str) else None
                if kind is None:
                    continue
                seen[(sheet, kind)] = seen.get((sheet, kind), 0) + 1
                found[(sheet, kind)] = ("%s%d" % (target, row), ws["%s%d" % (target, row)].value)
    problems = [k for k, n in seen.items() if n != 1]
    missing = [(s, k) for s in sheets for k in KINDS if (s, k) not in found]
    assert not problems and not missing, (
        "the flag block moved in Kyle's template: rows seen twice %r, rows not found %r. Re-derive this scan, "
        "do not loosen it." % (problems, missing))
    return found


def _is_formula(value):
    return isinstance(value, str) and value.startswith("=")


def _own_literals(book, sheets):
    """Every flag cell that is the sheet's OWN input rather than a mirror of another sheet's: a literal, or a
    blank input (New / Reno: is blank on the sheets that read it), as "Sheet!A1". Hard Bid? is left out."""
    return {"%s!%s" % (sheet, cell)
            for (sheet, kind), (cell, value) in _flag_cells(book, sheets).items()
            if kind != "hard_bid" and not _is_formula(value)}


def _hard_bid_cells(book, sheets):
    return {"%s!%s" % (sheet, cell)
            for (sheet, kind), (cell, _v) in _flag_cells(book, sheets).items() if kind == "hard_bid"}


def _readers(book, sheet, cell):
    """The formulas ON THE SAME SHEET that read `cell` (another sheet's mirror of it is the template's own
    wiring and is not counted)."""
    col, row = re.match(r"([A-Z]+)(\d+)$", cell).groups()
    pattern = re.compile(r"(?<![A-Za-z0-9_!])\$?%s\$?%s(?![0-9])" % (re.escape(col), re.escape(row)))
    return sorted(c.coordinate for r in book[sheet].iter_rows() for c in r
                  if isinstance(c.value, str) and c.value.startswith("=") and c.coordinate != cell
                  and pattern.search(c.value))


NOT_WRITTEN_LIVE = " The live intake does not write it either."
# cell -> (the phase that will write it, whether a formula on its sheet reads it today, why it is not written)
KNOWN_GAPS = {
    "Seal!B4": ("Phase 15", True,
                "Local? is Seal's own literal ('Seal (+Jnts)'!B4 mirrors it) and Seal's travel line reads it. "
                "No Seal option is priced in v2 yet." + NOT_WRITTEN_LIVE),
    "Seal!B10": ("Phase 15", True,
                 "New / Reno: is a blank input on Seal ('Seal (+Jnts)'!B10 mirrors it) and Seal's patch material "
                 "rate reads it. No Seal option is priced in v2 yet." + NOT_WRITTEN_LIVE),
    "Leveling!B4": ("Phase 16", True,
                    "Local? is Leveling's own literal and its travel line reads it. No Leveling option is priced "
                    "in v2 yet." + NOT_WRITTEN_LIVE),
    "Leveling!B10": ("Phase 16", False,
                     "New / Reno: is a blank input on Leveling that no formula on the sheet reads today. No "
                     "Leveling option is priced in v2 yet." + NOT_WRITTEN_LIVE),
    GYP + "!B5": ("Phase 14", True,
                  "Local? is the gyp base's own literal (three of the other four gyp sheets mirror it) and the "
                  "sheet's travel and markup lines read it. Gyp is not priced in v2 yet." + NOT_WRITTEN_LIVE),
    "Gyp (FR)!B5": ("Phase 14", True,
                    "Local? is Gyp (FR)'s own literal, not a mirror of the gyp base's, and its travel and markup "
                    "lines read it. Gyp is not priced in v2 yet." + NOT_WRITTEN_LIVE),
    GYP + "!B12": ("Phase 14", False,
                   "New / Reno: is the gyp base's own 'New' and no formula on the sheet reads it today. Gyp is "
                   "not priced in v2 yet." + NOT_WRITTEN_LIVE),
    "Gyp (FR)!B12": ("Phase 14", False,
                     "New / Reno: is Gyp (FR)'s own 'New' and no formula on the sheet reads it today. Gyp is not "
                     "priced in v2 yet." + NOT_WRITTEN_LIVE),
}
PHASE_OF_SHEET = {"Seal": "Phase 15", "Leveling": "Phase 16", "Gyp": "Phase 14"}


def _written(ran):
    """Every cell the table writes in any state of a draft: each condition's `cells` and `perSheet`."""
    out = set()
    for c in ran["rule"]["table"]:
        out.update(c["cells"])
        out.update((c["perSheet"] or {}).values())
    return out


def _unlisted(book, sheets, written, gaps):
    return sorted(_own_literals(book, sheets) - written - set(gaps))


def _stale(book, sheets, written, gaps):
    """Gaps that are no longer gaps: written now, or not a literal flag cell of the template any more."""
    return sorted((set(gaps) & written) | (set(gaps) - _own_literals(book, sheets)))


def test_every_literal_flag_cell_of_every_priced_sheet_is_written_by_the_table_or_a_known_gap(ran, book):
    """THE COMPLETENESS SCAN. Each priced sheet's own Local, Taxable, Prevailing wage, Remodel tax and New or
    Reno input is either written by the work-types table or on KNOWN_GAPS, with the phase that will write it.
    A new literal flag cell that is on neither fails here, by name."""
    unlisted = _unlisted(book, ran["sheets"], _written(ran), KNOWN_GAPS)
    assert unlisted == [], (
        "literal flag cell(s) of the template that nothing writes and nobody has listed: %r. Write them from the "
        "table, or add them to KNOWN_GAPS with the phase that will and why." % unlisted)


def test_a_known_gap_is_a_real_literal_flag_cell_nothing_writes_not_even_the_live_intake(ran, book):
    """The list cannot go stale: every entry is still a literal flag cell of the template, the table still does
    not write it, and neither does the live intake (the page's own literal, lifted and evaluated). The phase is
    the one that owns the sheet, and whether a formula reads the cell is checked against the workbook, so the
    reason is not a guess that Kyle's next edit can make wrong."""
    assert _stale(book, ran["sheets"], _written(ran), KNOWN_GAPS) == [], (
        "a known gap is not a gap any more: take it off the list")
    live_cells = set(ran["live"]["cells"])
    for cell, (phase, read, reason) in KNOWN_GAPS.items():
        sheet = cell.rsplit("!", 1)[0]
        assert phase == PHASE_OF_SHEET[sheet.split(" ")[0].split("(")[0].strip()], (cell, phase)
        assert cell not in live_cells, "%s is written by the live intake, so it is not a gap" % cell
        assert bool(_readers(book, sheet, cell.rsplit("!", 1)[1])) is read, (
            "%s: a formula on its sheet reads it is %s now, and the reason on KNOWN_GAPS says %s" % (
                cell, not read, read))
        assert "live intake does not write it either" in reason and phase in ("Phase 14", "Phase 15", "Phase 16")


def test_the_scan_finds_all_six_flag_rows_on_all_eleven_priced_sheets(ran, book):
    import estimate_writer as ew
    assert sorted(ran["sheets"]) == sorted(ew.FLAG_BLOCK_CELLS), (
        "the vocabulary and the backend disagree about which sheets carry a flag block")
    found = _flag_cells(book, ran["sheets"])
    assert len(found) == 66 and {s for s, _k in found} == {
        "Epoxy", "Epoxy blank", "Polish", "Seal", "Seal (+Jnts)", "Leveling", GYP, "Gyp (USG N12ULTRA)",
        'Gyp (USG N25 1-4")', "Gyp (GWorx SC190)", "Gyp (FR)"}
    # the cells the table already writes are among the literals it found: it is reading the right rows
    assert found[("Seal", "local")] == ("B4", "Yes") and found[("Seal", "reno")] == ("B10", None)
    assert found[("Seal (+Jnts)", "local")][1] == "=Seal!B4", "a mirror is not an input"
    assert found[(GYP, "taxable")] == ("B8", "Yes") and found[(GYP, "remodel_tax")] == ("D8", "=Epoxy!D6")


def test_the_table_writes_no_hard_bid_cell_and_every_one_of_them_is_accounted_for(ran, book):
    """Hanz, 2026-10-03: nothing writes a Hard Bid? cell, on any sheet. They are literals in the template, so
    they are accounted for here by that one rule, and not by KNOWN_GAPS."""
    hard = _hard_bid_cells(book, ran["sheets"])
    assert len(hard) == 11
    assert not hard & _written(ran), sorted(hard & _written(ran))
    import estimate_writer as ew
    assert {"%s!%s" % (s, a) for s, a in ew.HARD_BID_FLAG_CELLS.items()} <= hard


def test_the_scan_can_fail_a_new_unlisted_literal_and_a_gap_that_has_closed(ran):
    """The scan is only worth having if it goes red. A mirror that becomes a literal ('Seal (+Jnts)'!B4 typed
    over) is an unlisted literal flag cell. A gap the table now writes (here, Seal!B4 added to what it writes)
    is a stale entry. The real template is clean, and each of these is not."""
    written, sheets = _written(ran), ran["sheets"]
    assert _unlisted(_template(), sheets, written, KNOWN_GAPS) == []
    typed_over = _template()
    typed_over["Seal (+Jnts)"]["B4"] = "Yes"
    assert _unlisted(typed_over, sheets, written, KNOWN_GAPS) == ["Seal (+Jnts)!B4"]
    new_row = _template()
    new_row["Leveling"]["D6"] = "No"          # Remodel Tax? on Leveling was a mirror of Epoxy's; now its own
    assert _unlisted(new_row, sheets, written, KNOWN_GAPS) == ["Leveling!D6"]
    assert _stale(_template(), sheets, written | {"Seal!B4"}, KNOWN_GAPS) == ["Seal!B4"]
    assert _stale(_template(), sheets, written, {**KNOWN_GAPS, "Epoxy blank!B4": ("Phase 99", False, "x")}) == [
        "Epoxy blank!B4"]
    # and a template that moves a flag row is refused rather than scanned wrong
    moved = _template()
    moved["Polish"]["A4"] = "Nearby?"
    with pytest.raises(AssertionError, match="the flag block moved"):
        _flag_cells(moved, sheets)


# ══ 7. the architecture document ═════════════════════════════════════════════════════════════════════
def _flat(text):
    return re.sub(r"\s+", " ", text)


def test_the_architecture_doc_no_longer_says_a_save_loses_or_changes_none():
    """It said: "A save of an existing v2 estimate gains exactly those five cells and loses or changes none".
    That was wrong for a split draft and for Renovation in the second cell. The sentence is gone as a claim,
    and the document says what a save does now, names the rule and the test, and names the gaps' phases."""
    doc = _flat(DOC.read_text(encoding="utf-8"))
    assert ('A save of an existing v2 estimate gains exactly those five cells and loses or changes none, and a '
            '"Reno" already in B10 survives.') not in doc
    for must in ("Phase 7b", "`writeCellsFor(condition, jobType, split, own)`", "tax_flags_per_sheet",
                 "the first cell that holds an answer", "test_v2_condition_cells.py", "Phase 14", "Phase 15",
                 "Phase 16", "a draft that is not split gains the five cells above and changes no other",
                 "On a split draft it writes the base sheet's own Taxable and Remodel cell"):
        assert must in doc, must


def test_the_vectors_the_doc_names_are_in_the_golden():
    golden = json.loads((FIXTURES / "polish_chain_golden.json").read_text(encoding="utf-8"))
    ids = {v["id"] for v in golden["vectors"]}
    assert "cond/fromCells/polishB4IsReadWhenEpoxyB4IsBlank" in ids
    assert "cond/fromCells/polishB4IsNotRead" not in ids
    added = [i for i in ids if i.startswith(("cond/cells/split/", "cond/cells/reno/", "cond/fromCells/split/"))
             or i == "cond/fromCells/epoxyB4WinsOverPolishB4"
             or re.match(r"cond/fromCells/taxable/(?!\d)", i)]
    # one renamed (and moved) vector, nineteen added, none removed: 2,228 - 1 + 1 + 19 = 2,247
    assert len(added) == 19 and len(ids) == 2247, "the doc says nineteen vectors were added to 2,228"


# ══ 8. red without the change ════════════════════════════════════════════════════════════════════════
MODEL = "js/bid-model.js"
TYPES = "js/work-types.js"
SANDBOX = "js/polish-sandbox.js"
ESTIMATE = "js/polish-estimate.js"
INTAKE = "js/polish-intake.js"

# Each entry breaks ONE line of ONE file in a scratch copy: (the harness to run, the file, the line, what it
# becomes, files that must come along, the check that must go red on what the broken tree answers).
RUNNERS = {"module": (MODULE_HARNESS, None), "takeoff": (TAKEOFF_HARNESS, TAKEOFF_FILES),
           "intake": (INTAKE_HARNESS, INTAKE_FILES), "beta": (BETA_HARNESS, BETA_FILES)}

BASE_CELL = '      perSheet: { epoxy: "Epoxy!B6", polish: "Polish!B6", gyp: \'Gyp (USG 1-8")!B8\' },'
SPLIT_GATE = "    if (!split || !c.perSheet) return c.cells.slice();"
OWN_FILTER = ("      return own.filter(function (cell) { return typeof cell === \"string\" && "
              "CELL_ADDRESS.test(cell); });")
BASE_SHEETS = '    return jobRecord(jobTypeKey, "baseSheets").tabs.map(function (t) { return tabByKey.get(t).sheets[0]; });'
COPY_PER_SHEET = "      if (c.perSheet) Object.keys(c.perSheet).forEach(function (tabKey) { add(c.perSheet[tabKey]); });"
WRITE_LINE = "      var to = types.writeCellsFor(key, MODEL_JOB, !!split);"
READ_LINE = "      var cell = firstFilled(types.writeCellsFor(key, MODEL_JOB, !!split), cv, unanswered);"
PATCH_LINE = "      cell_values: conditionCellWrites(model.conditions, cells, ctx.library, types.isSplit(state)),"
RENO_LINE = "      var first = firstFilled(carried.cells, out, isBlank);"
MARK_LINE = '    "polish_estimate", "cell_values", "tax_flags_per_sheet",'
COPY_BASE_LINE = "    if (WT.isSplit(srcData) && WT.isJobType(srcData.work_type)) {"
ESTIMATE_ADOPT = "    M.conditions = B.conditionsFromCells(M.conditions, state.cell_values, T.isSplit(state));"
ESTIMATE_SEEDED = "        B.seedConditionDefaults(M.conditions, condRows), state.cell_values, T.isSplit(state));"
INTAKE_ADOPT = "    M.conditions = B.conditionsFromCells(M.conditions, state.cell_values, T.isSplit(state));"
INTAKE_SAVE = "    return B.conditionCellWrites(M.conditions, draft.cell_values, undefined, T.isSplit(draft));"
INTAKE_SEEDED = "        state.cell_values, T.isSplit(state));"


def _both_live_split_checks(live):
    live_intake_tests.test_on_a_split_draft_the_two_tax_switches_are_the_bases_own_cells(live)
    live_intake_tests.test_a_work_type_change_on_a_split_draft_rereads_the_new_bases_own_cells(live)


# name -> (runner, file, anchor line, what it becomes, also, the check)
BREAKS = {
    "the table forgets the polish base's own Taxable cell": (
        "module", TYPES, BASE_CELL,
        '      perSheet: { epoxy: "Epoxy!B6", gyp: \'Gyp (USG 1-8")!B8\' },', [MODEL],
        test_an_unsplit_draft_writes_every_cell_of_the_row_and_a_split_one_the_base_sheets_own),
    "a draft that is not split is treated as split": (
        "module", TYPES, SPLIT_GATE, "    if (!c.perSheet) return c.cells.slice();", [MODEL],
        test_an_unsplit_draft_writes_every_cell_of_the_row_and_a_split_one_the_base_sheets_own),
    "a split draft is treated as unsplit": (
        "module", TYPES, SPLIT_GATE, "    return c.cells.slice();", [MODEL],
        test_an_unsplit_draft_writes_every_cell_of_the_row_and_a_split_one_the_base_sheets_own),
    "an address that is no cell is written": (
        "module", TYPES, OWN_FILTER, "      return own.slice();", [MODEL],
        test_a_caller_that_knows_the_base_better_replaces_the_tables_cells_and_only_real_addresses_survive),
    "the live intake writes an address that is no cell": (
        "beta", TYPES, OWN_FILTER, "      return own.slice();", [],
        test_the_live_intake_drops_an_address_that_is_not_a_cell),
    "a combo job loses its second base sheet": (
        "module", TYPES, BASE_SHEETS,
        '    return jobRecord(jobTypeKey, "baseSheets").tabs.slice(0, 1).map(function (t) { '
        'return tabByKey.get(t).sheets[0]; });', [MODEL],
        test_a_job_type_is_priced_on_the_first_sheet_of_each_of_its_tabs),
    "the live intake loses its combo's second base sheet": (
        "beta", TYPES, BASE_SHEETS,
        '    return jobRecord(jobTypeKey, "baseSheets").tabs.slice(0, 1).map(function (t) { '
        'return tabByKey.get(t).sheets[0]; });', [],
        test_the_live_intake_writes_both_of_a_combo_jobs_base_sheets_and_no_other_tax_cell),
    "the live intake prices every job on Epoxy's sheet": (
        "beta", TYPES, BASE_SHEETS,
        '    return jobRecord(jobTypeKey, "baseSheets").tabs.map(function (t) { return "Epoxy"; });', [],
        _both_live_split_checks),
    "the copy forgets the base sheets' own tax cells": (
        "module", TYPES, COPY_PER_SHEET, "      void 0;", [MODEL],
        test_a_copy_of_a_split_project_arrives_split_with_its_tax_answers),
    "the model's save ignores whether the draft is split": (
        "module", MODEL, WRITE_LINE, "      var to = types.writeCellsFor(key, MODEL_JOB, false);", [],
        test_a_save_on_a_split_draft_writes_the_base_sheets_own_cells_and_no_other_tax_cell),
    "the model's read-back ignores whether the draft is split": (
        "module", MODEL, READ_LINE,
        "      var cell = firstFilled(types.writeCellsFor(key, MODEL_JOB, false), cv, unanswered);", [],
        test_the_read_back_on_a_split_draft_is_the_base_sheets_own_cell),
    "the one save both pages share forgets the draft's mark": (
        "module", MODEL, PATCH_LINE,
        "      cell_values: conditionCellWrites(model.conditions, cells, ctx.library, false),", [],
        test_a_split_project_copied_and_saved_in_v2_keeps_its_option_sheets_tax_answers),
    "the model reads only the first renovation cell": (
        "module", MODEL, RENO_LINE, "      var first = firstFilled(carried.cells.slice(0, 1), out, isBlank);", [],
        test_renovation_survives_a_save_whichever_of_its_two_cells_holds_it),
    "the copy drops the mark": (
        "module", SANDBOX, MARK_LINE, '    "polish_estimate", "cell_values",', [MODEL],
        test_a_split_project_copied_and_saved_in_v2_keeps_its_option_sheets_tax_answers),
    "the copy forgets the source base's own tax answer": (
        "module", SANDBOX, COPY_BASE_LINE, "    if (false && WT.isSplit(srcData) && WT.isJobType(srcData.work_type)) {", [MODEL],
        test_a_v2_copy_of_a_split_project_that_is_not_polish_opens_at_the_source_bases_own_tax_answer),
    "the Takeoff step reads the draft as unsplit": (
        "takeoff", ESTIMATE, ESTIMATE_ADOPT,
        "    M.conditions = B.conditionsFromCells(M.conditions, state.cell_values);", [],
        test_the_takeoff_step_saves_a_split_copy_without_touching_its_options),
    "the Takeoff step's seeded read ignores the mark": (
        "takeoff", ESTIMATE, ESTIMATE_SEEDED,
        "        B.seedConditionDefaults(M.conditions, condRows), state.cell_values);", [],
        test_the_takeoff_step_saves_a_split_copy_without_touching_its_options),
    "the v2 intake reads the draft as unsplit": (
        "intake", INTAKE, INTAKE_ADOPT,
        "    M.conditions = B.conditionsFromCells(M.conditions, state.cell_values);", [],
        test_the_v2_intake_shows_and_saves_a_split_copy_without_touching_its_options),
    "the v2 intake saves as if the draft were unsplit": (
        "intake", INTAKE, INTAKE_SAVE,
        "    return B.conditionCellWrites(M.conditions, draft.cell_values);", [],
        test_the_v2_intake_shows_and_saves_a_split_copy_without_touching_its_options),
    "the v2 intake's seeded read ignores the mark": (
        "intake", INTAKE, INTAKE_SEEDED, "        state.cell_values);", [],
        test_the_v2_intake_shows_and_saves_a_split_copy_without_touching_its_options),
}


def _on_broken_tree(tmp_path, runner, rel, old, new, also):
    harness, files = RUNNERS[runner]
    frontend = break_source(tmp_path, rel, old, new, also=also)
    if files:
        copy_unmodified(frontend, files)
    if runner == "intake":
        # the intake harness parses the county table out of backend/reference_tax.py, beside the frontend
        (tmp_path / "backend").mkdir(exist_ok=True)
        shutil.copyfile(ROOT / "backend" / "reference_tax.py", tmp_path / "backend" / "reference_tax.py")
    answer = _run(harness, frontend)
    return answer["conditions"] if runner == "beta" else answer


@pytest.mark.parametrize("name", sorted(BREAKS))
def test_the_checks_go_red_when_the_code_changes(tmp_path, name):
    """One line of one file is broken in a scratch copy, its harness runs on it, and the named check has to
    fail on what it answers. The harness runs OUTSIDE the try, so a break that merely crashes it is an error and
    not a pass, and break_source refuses an anchor that is not there exactly once."""
    require_node()
    runner, rel, old, new, also, check = BREAKS[name]
    broken = _on_broken_tree(tmp_path, runner, rel, old, new, also)
    with pytest.raises(AssertionError):
        check(broken)


def test_the_unbroken_answers_pass_every_check_the_table_names(ran, takeoff, intake, live):
    """The other half: the same checks, called the same way, are green on the real code."""
    real = {"module": ran, "takeoff": takeoff, "intake": intake, "beta": live}
    for runner, rel, old, new, also, check in BREAKS.values():
        check(real[runner])


def test_every_check_a_break_names_is_a_real_test_and_each_part_of_the_fix_has_a_break():
    named = {check.__name__ for *_, check in BREAKS.values()}
    for must in (test_an_unsplit_draft_writes_every_cell_of_the_row_and_a_split_one_the_base_sheets_own,
                 test_a_save_on_a_split_draft_writes_the_base_sheets_own_cells_and_no_other_tax_cell,
                 test_the_read_back_on_a_split_draft_is_the_base_sheets_own_cell,
                 test_renovation_survives_a_save_whichever_of_its_two_cells_holds_it,
                 test_a_copy_of_a_split_project_arrives_split_with_its_tax_answers,
                 test_a_split_project_copied_and_saved_in_v2_keeps_its_option_sheets_tax_answers,
                 test_a_v2_copy_of_a_split_project_that_is_not_polish_opens_at_the_source_bases_own_tax_answer,
                 test_the_takeoff_step_saves_a_split_copy_without_touching_its_options,
                 test_the_v2_intake_shows_and_saves_a_split_copy_without_touching_its_options,
                 test_the_live_intake_drops_an_address_that_is_not_a_cell):
        assert must.__name__ in named, "no break turns %s red" % must.__name__
