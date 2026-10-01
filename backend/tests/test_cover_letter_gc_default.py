"""The cover letter's GC default, with BOTH scripts the Proposal step loads.

Hanz, 2026-10-02: "general contractor projects should also have cover letter default on which
means that the toggle button for the cover letter is always on for all GC projects." The first
build passed every test and was dead on the page: coverletter-editor.js loads after
proposal-review.js, read the absent flag as false and WROTE it on load, so every new GC project
saved "off" before anybody touched the box (staging walk, 2026-10-02). These run both scripts
together (cover-letter-default-harness.js), as the browser does.

Also here: the letterhead pinned to the page in the letter (LibreOffice drew it 121.5pt to the
right), the editor's signature contact line, and the Dropbox listings waiting for sign-in.
"""
import json
import pathlib
import re
import shutil
import subprocess
import zipfile

import pytest

BACKEND = pathlib.Path(__file__).resolve().parents[1]
FRONTEND = BACKEND.parent / "frontend"
HARNESS = pathlib.Path(__file__).resolve().parent / "js" / "cover-letter-default-harness.js"
needs_node = pytest.mark.skipif(shutil.which("node") is None, reason="node is not installed")


@pytest.fixture(scope="module")
def ran():
    proc = subprocess.run(["node", str(HARNESS), str(FRONTEND)], capture_output=True, text=True,
                          encoding="utf-8", timeout=60)
    assert proc.returncode == 0, "the harness itself failed:\n" + proc.stderr
    return json.loads(proc.stdout)


def _letter_writes(writes):
    return [w for w in writes if "cover_letter_enabled" in w]


@needs_node
def test_a_fresh_gc_project_opens_with_the_letter_on_and_saves_nothing(ran):
    """Mutation: make coverletter-editor.js's init call setEnabled (the writing half) again."""
    for key in ("gcFresh", "gcPadded"):
        r = ran[key]
        assert r["threw"] is None, r["threw"]
        assert r["onLoad"]["checked"] is True, (key, r["onLoad"])
        assert r["onLoad"]["surfaceHidden"] is False, (key, r["onLoad"])
        assert _letter_writes(r["onLoad"]["writes"]) == [], (
            "loading the page wrote the cover-letter flag: %r" % r["onLoad"]["writes"])


@needs_node
def test_the_estimators_choice_still_wins_both_ways(ran):
    assert ran["gcUntickedBefore"]["onLoad"]["checked"] is False
    assert ran["gcTickedBefore"]["onLoad"]["checked"] is True
    assert ran["directFresh"]["onLoad"]["checked"] is False
    r = ran["gcFreshThenUntick"]
    assert _letter_writes(r["after"]) == [{"cover_letter_enabled": False}], r["after"]


@needs_node
def test_without_the_resolver_the_saved_flag_is_used_as_before(ran):
    r = ran["gcNoResolver"]
    assert r["threw"] is None, r["threw"]
    assert r["onLoad"]["checked"] is False
    assert _letter_writes(r["onLoad"]["writes"]) == []


# ── the letterhead, pinned to the page ───────────────────────────────────────────────────────
@pytest.mark.parametrize("audience", ["GC", "Direct"])
def test_the_letters_full_page_art_lands_at_the_page_edge(audience):
    """Kyle's letters anchor the full-page letterhead -0.69" from the COLUMN of a continuous
    section whose margin is 0.69"; LibreOffice measured from the first section's 2.375" margin and
    drew it 121.5pt to the right. fill_cover_letter now writes the same position relative to the
    page: 990 twips x 635 - 628650 EMU = 0.

    Mutation: return 0 from pin_column_anchors_to_page without rewriting anything."""
    import cover_letter_writer as cl
    data = cl.fill_cover_letter(work_type="epoxy", audience=audience,
                                values={"job_name": "J", "project_name": "J",
                                        "estimator_name": "E", "estimator_email": "e@x.com"})
    xml = zipfile.ZipFile(__import__("io").BytesIO(data)).read("word/document.xml").decode("utf8")
    big = [a for a in re.findall(r"<wp:anchor\b.*?</wp:anchor>", xml, re.S)
           if re.search(r'<wp:extent cx="7772\d{3}"', a)]
    assert big, "no full-page letterhead image in the letter"
    for a in big:
        pos = re.search(r'<wp:positionH relativeFrom="([^"]*)">\s*<wp:posOffset>(-?\d+)', a)
        assert pos and pos.group(1) == "page" and int(pos.group(2)) == 0, pos and pos.groups()


def test_pinning_skips_multi_column_sections_and_keeps_vertical_positions():
    import io
    import docx
    import cover_letter_writer as cl
    from docx.oxml import parse_xml
    d = docx.Document()
    W = 'xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main"'
    WP = 'xmlns:wp="http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing"'
    anchor = ('<w:r %s %s><w:drawing><wp:anchor behindDoc="1">'
              '<wp:positionH relativeFrom="column"><wp:posOffset>-100</wp:posOffset></wp:positionH>'
              '<wp:positionV relativeFrom="paragraph"><wp:posOffset>-200</wp:posOffset></wp:positionV>'
              '</wp:anchor></w:drawing></w:r>' % (W, WP))
    p = d.add_paragraph()
    p._p.append(parse_xml(anchor))
    sect = d.element.body.find('{http://schemas.openxmlformats.org/wordprocessingml/2006/main}sectPr')
    mar = sect.find('{http://schemas.openxmlformats.org/wordprocessingml/2006/main}pgMar')
    left = int(mar.get('{http://schemas.openxmlformats.org/wordprocessingml/2006/main}left'))
    assert cl.pin_column_anchors_to_page(d) == 1
    xml = d.element.body.xml
    assert re.search(r'<wp:positionH relativeFrom="page">\s*<wp:posOffset>%d</wp:posOffset>'
                     % (left * 635 - 100), xml), xml
    assert re.search(r'<wp:positionV relativeFrom="paragraph">\s*<wp:posOffset>-200</wp:posOffset>', xml), xml
    # A two-column section is left alone: "column" is not the margin there.
    d2 = docx.Document()
    p2 = d2.add_paragraph()
    p2._p.append(parse_xml(anchor))
    s2 = d2.element.body.find('{http://schemas.openxmlformats.org/wordprocessingml/2006/main}sectPr')
    cols = s2.find('{http://schemas.openxmlformats.org/wordprocessingml/2006/main}cols')
    cols.set('{http://schemas.openxmlformats.org/wordprocessingml/2006/main}num', "2")
    assert cl.pin_column_anchors_to_page(d2) == 0


# ── the editor's signature contact line ──────────────────────────────────────────────────────
@needs_node
def test_the_editors_letter_shows_the_contact_line_the_pdf_prints():
    """The editor showed a raw {{estimator_contact_line}} (staging walk, 2026-10-02): only the
    server built it. clTokens now fills it with the server's rule.

    Mutation: drop the backfill in clTokens."""
    src = (FRONTEND / "js" / "coverletter-editor.js").read_text(encoding="utf-8")
    def lift(name):
        m = re.search(r"\n  function " + name + r"\(\) \{", src)
        assert m, name
        depth, i = 1, m.end()
        while depth:
            depth += {"{": 1, "}": -1}.get(src[i], 0)
            i += 1
        return src[m.start():i]
    g = re.search(r"\n  const g = .*?;\n", src).group(0)
    script = (
        "const window = { computeTokenValues: (m) => ({ estimator_email: m.estimator_email }) };"
        "const document = { getElementById: () => null };"
        "let STATE = {};"
        + g + lift("clState").replace("(window.TW && TW.getState())", "STATE") + lift("clTokens") +
        "const out = {};"
        "STATE = { estimator_email: 'hanz@wetreadwell.com' }; out.withEmail = clTokens().estimator_contact_line;"
        "STATE = {}; out.without = clTokens().estimator_contact_line;"
        "console.log(JSON.stringify(out));")
    proc = subprocess.run(["node", "-e", script], capture_output=True, text=True, encoding="utf-8",
                          timeout=60)
    assert proc.returncode == 0, proc.stderr
    got = json.loads(proc.stdout.strip().splitlines()[-1])
    assert got == {"withEmail": "hanz@wetreadwell.com | wetreadwell.com", "without": "wetreadwell.com"}


# ── the Dropbox listings wait for sign-in ────────────────────────────────────────────────────
def test_both_dropbox_listings_wait_for_sign_in():
    """A fresh Done load fetched /api/dropbox/folders and /project-folders before the bearer token
    existed and took 401s (staging walk, 2026-10-02) -- the #124 race. Each fetch must be preceded,
    in the same function, by an await of TWAuth.ready."""
    src = (FRONTEND / "js" / "dropbox.js").read_text(encoding="utf-8")
    for url in ('"/api/dropbox/folders"', '"/api/dropbox/project-folders"'):
        i = src.index(url)
        before = src[max(0, i - 600):i]
        assert "await TWAuth.ready" in before, url
