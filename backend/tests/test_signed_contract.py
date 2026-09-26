"""The signed contract: the proposal, signed on its own acceptance block.

WHAT THIS FEATURE IS FOR. A customer accepts a proposal in the portal by typing
their name and ticking a consent box. Under ESIGN (15 U.S.C. § 7001) and Kansas
UETA (K.S.A. 16-1601 et seq.) that is a signature; what makes it defensible is
the audit trail — who, when, from where, under which wording, and against WHICH
EXACT DOCUMENT. All of that is still recorded and still validated here; what
changed on 2026-09-19 is where it is PRINTED.

THE CERTIFICATE PAGE IS GONE. Hanz read one and said three things: "we dont need
another page for the signatory", "I think the signature goes here?" — pointing at
the ACCEPTANCE / SIGNATURE / DATE / PRINTED NAME / TOTAL row Kyle already prints
on the proposal form — and "this is too much information... Just the basic
information is what we need". So `/api/admin/signed-contract` now writes four
values onto that row and returns a document with the SAME NUMBER OF PAGES it was
given. No IP address, no browser string, no consent paragraph, no digests appear
on the paper. `certificate_writer` still validates every one of them on the way
in, and the portal still stores the lot on its approval row.

THE ASSERTION THAT MATTERS MOST is still byte-identity of the proposal's own
pages. The portal stores the SHA-256 of the PDF the customer opened. If this
route ever re-rendered or re-flowed those pages, the document we keep would not
hash to the value recorded beside it, and it would be evidence AGAINST us. Every
page but the signed one comes through object-for-object, and the signed page
keeps its own content stream object untouched with the overlay appended after it
— `test_every_other_page_is_byte_identical` and
`test_the_signed_page_keeps_its_own_content_stream` are the two to break first
when this handler changes.

THE PAGE IS NOT THE LAST PAGE. A Direct Epoxy proposal is four pages: the FORM is
page 1 and pages 2-4 are Terms & Conditions on blank letterhead. Signing "the
last page" would put the customer's name on blank paper at the back of the Terms
and a test that only looked at `pages[-1]` would stay green through it. Every
assertion below therefore names the page by INDEX, and
`test_the_last_page_is_not_the_acceptance_page` states the trap outright.

THE FIXTURE CARRIES KYLE'S REAL ARTWORK. The form and the blank letterhead are
read straight out of the shipped templates, so "this page is the acceptance page
and that one is not" is decided in the tests by exactly the bytes that decide it
in production. The page geometry is the real one too — 612 x 791 at y=1.75 for
the form page, y=0.5 for the letterhead pages, measured off a genuine render —
so the coordinate assertions are about where the ink actually lands.

NO LIBREOFFICE ANYWHERE. Signing is pure pypdf now; nothing in this file needs a
renderer. The `certificate_writer` fill tests that remain cover the module that
still shapes and validates the payload. Its docx/PDF render half is no longer
reached by any route — kept, not deleted, pending Hanz's call on removing it.
"""
import hashlib
import io
import json
import os
import re
import zipfile

import pypdf
import pytest
from fastapi.testclient import TestClient
from pypdf.generic import ContentStream

import acceptance_signature as asig
import certificate_writer as cw
import drafts
import main
import proposal_writer

client = TestClient(main.app)
PATH = "/api/admin/signed-contract"
TOKEN = "svc-test"
TEMPLATE = proposal_writer.TEMPLATES_ROOT / cw.TEMPLATE_NAME

PROPOSAL_SHA = "3f2a1b" + "c" * 58
REVISION_SHA = "ab" * 32

CERT = {
    "project_name": "Olathe Distribution Center — Phase 2",
    "proposal_id": "b7f1c2de-4a90-4f21-9f3a-2c1d8e5a6b70",
    "revision_no": 3,
    "signer_name": "Dana Whitfield",
    "signer_title": "Director of Facilities",
    "signer_email": "dwhitfield@olathedc.example.com",
    "signed_at_utc": "2026-09-18T15:42:07Z",
    "signed_at_central": "September 18, 2026 at 10:42 AM CDT",
    "ip_address": "198.51.100.42",
    "user_agent": ("Mozilla/5.0 (Macintosh; Intel Mac OS X 10_15_7) AppleWebKit/537.36 "
                   "(KHTML, like Gecko) Chrome/141.0.0.0 Safari/537.36"),
    "options_summary": "Base bid accepted: epoxy, 24,800 SF. Alternate declined.",
    "total": 187450.75,
    "deposit_amount": 46862.69,
    "consent_text": ("By typing my name and checking this box I agree that this is my legal "
                     "electronic signature and that it binds me to this proposal and its "
                     "Terms and Conditions."),
    "consent_version": "tw-consent-2026-09-01",
    "proposal_pdf_sha256": PROPOSAL_SHA,
    "revision_sha256": REVISION_SHA,
}


# ── Kyle's real artwork, out of the shipped templates ────────────────────────
def _artwork() -> tuple:
    """(form, letterhead) raster bytes, identified rather than assumed.

    The letterhead is the raster the COVER LETTER template also carries — they
    are byte-identical, which is the fact that makes "the acceptance page is the
    page whose artwork appears nowhere else" true even when a cover letter is
    prepended. Asserting it here means a template re-export that broke the
    invariant would fail this file rather than mis-sign a contract.
    """
    def media(rel):
        with zipfile.ZipFile(proposal_writer.TEMPLATES_ROOT / rel) as z:
            return {n: z.read(n) for n in z.namelist() if n.startswith("word/media/")}

    letter = set(media("CoverLetter/Direct/Epoxy.docx").values())
    proposal = media(proposal_writer.TEMPLATE_PICKER[("epoxy", "Direct")])
    shared = [b for b in proposal.values() if b in letter]
    unique = [b for b in proposal.values() if b not in letter]
    assert len(shared) == 1 and len(unique) == 1, (
        "the Direct Epoxy template should carry exactly two rasters — the form, which is "
        "its own, and the blank letterhead, which the cover letter shares. Found %d shared "
        "and %d unique." % (len(shared), len(unique)))
    return unique[0], shared[0]


FORM_ART, LETTERHEAD_ART = _artwork()

# The real placements, measured off a genuine 4-page Direct Epoxy render: the
# artwork is 612 x 791 on a 612 x 792 page and sits a point differently on the
# form page than on the Terms pages. Both are used, because a handler that
# hard-coded one offset would be wrong on the other.
FORM_CM = (612.0, 0.0, 0.0, 791.0, 0.0, 1.75)
LETTERHEAD_CM = (612.0, 0.0, 0.0, 791.0, 0.1, 0.5)


def _esc(s: str) -> str:
    return s.replace("\\", r"\\").replace("(", r"\(").replace(")", r"\)")


def make_pdf(pages) -> bytes:
    """A genuine multi-page PDF. `pages` is [(lines, artwork_bytes_or_None, cm), ...].

    Hand-built rather than `PdfWriter.add_blank_page()`: a blank page has NO
    content stream, so a byte-identity test over blank pages would compare
    nothing against nothing and pass whatever the handler did to them.

    The image streams carry the real PNG bytes under an /Image XObject with the
    artwork's real 1275x1649 dimensions. They are not decodable as PDF image
    samples and nothing here renders them — every consumer under test reads them
    as a fingerprint and a size, which is exactly what production reads them as
    when it decides which page holds the acceptance block.
    """
    objs = {1: None, 2: None}

    def add(body: bytes) -> int:
        num = max(objs) + 1
        objs[num] = body
        return num

    font_id = add(b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>")
    art_ids: dict = {}
    page_ids = []
    for lines, art, cm in pages:
        body = ""
        res_xobject = ""
        if art is not None:
            key = hashlib.sha256(art).hexdigest()
            if key not in art_ids:
                head = ("<< /Type /XObject /Subtype /Image /Width %d /Height %d /ColorSpace "
                        "/DeviceGray /BitsPerComponent 8 /Length %d >>\nstream\n"
                        % (asig.ARTWORK_PX[0], asig.ARTWORK_PX[1], len(art)))
                art_ids[key] = add(head.encode("latin-1") + art + b"\nendstream")
            res_xobject = "/XObject << /Im0 %d 0 R >> " % art_ids[key]
            body += "q\n%s cm\n/Im0 Do\nQ\n" % " ".join(("%g" % v) for v in cm)
        body += "BT /F1 9 Tf 36 750 Td 12 TL\n"
        body += "".join("(%s) Tj T*\n" % _esc(line) for line in lines)
        body += "ET"
        stream = body.encode("latin-1", "replace")
        content_id = add(b"<< /Length %d >>\nstream\n" % len(stream) + stream + b"\nendstream")
        page_ids.append(add(
            ("<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] /Resources << /Font "
             "<< /F1 %d 0 R >> %s>> /Contents %d 0 R >>"
             % (font_id, res_xobject, content_id)).encode("latin-1")))

    objs[1] = b"<< /Type /Catalog /Pages 2 0 R >>"
    objs[2] = ("<< /Type /Pages /Kids [%s] /Count %d >>"
               % (" ".join("%d 0 R" % p for p in page_ids), len(page_ids))).encode()

    out = bytearray(b"%PDF-1.4\n")
    offsets = {}
    for num in sorted(objs):
        offsets[num] = len(out)
        out += b"%d 0 obj\n" % num + objs[num] + b"\nendobj\n"
    xref, top = len(out), max(objs) + 1
    out += b"xref\n0 %d\n0000000000 65535 f \n" % top
    for num in range(1, top):
        out += b"%010d 00000 n \n" % offsets.get(num, 0)
    out += b"trailer\n<< /Size %d /Root 1 0 R >>\nstartxref\n%d\n%%%%EOF\n" % (top, xref)
    return bytes(out)


TERMS = [
    (["TERMS AND CONDITIONS", "1. Agreement."], LETTERHEAD_ART, LETTERHEAD_CM),
    (["10. Mutual Waiver of Consequential Damages."], LETTERHEAD_ART, LETTERHEAD_CM),
    (["27. Governing law."], LETTERHEAD_ART, LETTERHEAD_CM),
]
FORM_PAGE = (["TREADWELL PROPOSAL", "Olathe Distribution Center - Phase 2",
              "Scope: 24,800 SF epoxy"], FORM_ART, FORM_CM)
COVER_PAGE = (["Dear Dana,", "Thank you for the opportunity."], LETTERHEAD_ART, LETTERHEAD_CM)

# A Direct Epoxy proposal: the form is page 1, then three pages of Terms.
PROPOSAL_PDF = make_pdf([FORM_PAGE] + TERMS)
ACCEPTANCE_INDEX = 0
# The same proposal with the cover letter prepended into its own bytes. The form
# — and the acceptance block with it — is page 2.
PROPOSAL_PDF_WITH_LETTER = make_pdf([COVER_PAGE, FORM_PAGE] + TERMS)
ACCEPTANCE_INDEX_WITH_LETTER = 1


def _draft_row(work_type="epoxy", audience="Direct", cover_letter=False):
    return {"id": CERT["proposal_id"],
            "data": {"proposal_payload": {
                "work_type": work_type, "audience": audience,
                "cover_letter_enabled": cover_letter,
                "values": {"job_name": "Olathe Distribution Center — Phase 2"}}}}


def _docx_lines(docx_bytes: bytes) -> list:
    """Every paragraph of a .docx, in order — what a reader would see."""
    from docx import Document
    d = Document(io.BytesIO(docx_bytes))
    return [p.text for p in d.paragraphs if p.text.strip()]


@pytest.fixture
def wired(monkeypatch):
    """A token, and a draft that says "Direct epoxy, no cover letter".

    The draft is what tells the handler WHICH PAGE carries the acceptance block:
    the same payload `/api/admin/proposal-pdf` renders from, read at the same
    revision. Stubbed because this suite is hermetic — it never reaches Supabase.
    """
    monkeypatch.setenv("SERVICE_TOKEN", TOKEN)
    monkeypatch.setattr(drafts, "get_revision", lambda pid, rev: _draft_row())
    monkeypatch.setattr(drafts, "load_draft", lambda pid: _draft_row())


def _post(pdf_bytes=PROPOSAL_PDF, cert=None, token=TOKEN, raw_cert=None, omit_pdf=False):
    files = {} if omit_pdf else {"proposal_pdf": ("proposal.pdf", pdf_bytes, "application/pdf")}
    data = {}
    if raw_cert is not None:
        data["certificate"] = raw_cert
    elif cert is not False:
        data["certificate"] = json.dumps(CERT if cert is None else cert)
    headers = {} if token is None else {"X-Service-Token": token}
    return client.post(PATH, files=files or None, data=data or None, headers=headers)


def _pages(pdf_bytes):
    return list(pypdf.PdfReader(io.BytesIO(pdf_bytes)).pages)


def _text(page) -> str:
    return page.extract_text() or ""


# ── the certificate payload: still required, still validated, no longer printed ─
# These cover `certificate_writer`, which shapes and validates the field set the
# portal sends. The endpoint still calls `field_problems` and `certificate_fields`;
# the module's docx render half is no longer reached by any route.
def test_the_template_ships():
    assert TEMPLATE.is_file(), (
        "Signature_Certificate.docx must ship in backend/templates/ — regenerate it with "
        "backend/prepare_signature_certificate_template.py and commit it.")


def test_the_template_offers_a_slot_for_every_field():
    """A field with nowhere to print is evidence silently lost."""
    from docx import Document
    text = "\n".join(p.text for p in Document(str(TEMPLATE)).paragraphs)
    for name in cw.FIELDS:
        assert "{{" + name + "}}" in text, name + " has no placeholder on the certificate"


def test_the_fill_leaves_no_placeholder_behind():
    """`_ensure_value_aliases` exists because a customer-facing document must never
    show a literal {{token}}. A certificate is the last place that can be allowed."""
    lines = _docx_lines(cw.build_certificate_docx(CERT, TEMPLATE))
    leftover = [t for line in lines for t in re.findall(r"\{\{[a-z0-9_]+\}\}", line)]
    assert leftover == [], leftover


def test_every_value_reaches_the_page():
    text = "\n".join(_docx_lines(cw.build_certificate_docx(CERT, TEMPLATE)))
    for value in cw.certificate_fields(CERT).values():
        assert value in text, value


def test_the_label_formatting_survives_the_fill():
    """The value is written into the run that owns the token, not over the whole
    paragraph — so the bold label beside it keeps its formatting. Collapsing the
    paragraph into one run (the shortcut `invoice_writer` can take, because its
    placeholders ARE whole paragraphs) would lose it silently."""
    from docx import Document
    d = Document(io.BytesIO(cw.build_certificate_docx(CERT, TEMPLATE)))
    row = [p for p in d.paragraphs if p.text.startswith("Name\t")][0]
    assert row.runs[0].text == "Name" and row.runs[0].bold is True
    assert "Dana Whitfield" in row.text


def test_an_empty_title_prints_a_phrase_not_a_blank():
    """A blank beside a label reads as "the fill failed", which is the one thing
    this page must never be ambiguous about."""
    f = cw.certificate_fields({**CERT, "signer_title": ""})
    assert f["signer_title"] == cw.NO_TITLE
    assert cw.certificate_fields({**CERT, "deposit_amount": None})["deposit_amount"] == cw.NO_DEPOSIT


def test_money_and_revision_are_formatted_for_a_reader():
    f = cw.certificate_fields(CERT)
    assert f["total"] == "$187,450.75" and f["deposit_amount"] == "$46,862.69"
    assert f["revision_no"] == "3"


def test_the_central_timestamp_is_taken_verbatim():
    """No timezone maths here, ever. The portal knows when the signature was taken;
    this box's clock runs ~13 hours off Central."""
    f = cw.certificate_fields(CERT)
    assert f["signed_at_central"] == CERT["signed_at_central"]
    assert f["signed_at_utc"] == CERT["signed_at_utc"]


# ── what the writer refuses ──────────────────────────────────────────────────
def test_a_missing_required_field_is_named():
    for name in ("signer_email", "consent_text", "proposal_pdf_sha256"):
        short = {k: v for k, v in CERT.items() if k != name}
        assert any(name in p for p in cw.field_problems(short)), name


def test_a_blank_required_field_counts_as_missing():
    assert any("signer_name" in p for p in cw.field_problems({**CERT, "signer_name": "   "}))


def test_the_optional_fields_really_are_optional():
    short = {k: v for k, v in CERT.items() if k not in ("signer_title", "deposit_amount")}
    assert cw.field_problems(short) == []


def test_a_mis_shaped_digest_is_refused():
    """Still validated even though it is no longer printed: the digest is what
    pins WHICH document was signed, and the portal stores it against the
    approval row. A mis-shaped one is a broken record either way."""
    for bad in ("3f2a1b", "not-a-hash", "z" * 64):
        probs = cw.field_problems({**CERT, "proposal_pdf_sha256": bad})
        assert any("proposal_pdf_sha256" in p for p in probs), bad


def test_a_non_numeric_total_is_refused_by_name():
    assert any("total" in p for p in cw.field_problems({**CERT, "total": "about twenty grand"}))


def test_the_writer_shouts_if_the_template_loses_a_row(tmp_path):
    """A template edited to drop a row would otherwise print a certificate that
    quietly omits a fact — the module raises instead."""
    from docx import Document
    d = Document(str(TEMPLATE))
    for p in d.paragraphs:
        if "{{signer_email}}" in p.text:
            for t in p._p.iter("{http://schemas.openxmlformats.org/wordprocessingml/2006/main}t"):
                t.text = (t.text or "").replace("{{signer_email}}", "")
    hacked = tmp_path / "hacked.docx"
    d.save(str(hacked))
    with pytest.raises(RuntimeError) as exc:
        cw.build_certificate_docx(CERT, hacked)
    assert "signer_email" in str(exc.value)


@pytest.mark.skipif(not (os.environ.get("TW_RENDER_TESTS") or __import__("shutil").which("soffice")),
                    reason="needs LibreOffice; CI and the dev box render nothing")
def test_the_real_render_is_one_page():
    """The certificate PAGE is no longer produced by any route, but the renderer
    that built it is still on disk, so its page budget is still checked wherever
    soffice exists. Delete this with `build_certificate_pdf` if the certificate
    half of the module goes."""
    pdf = cw.build_certificate_pdf(CERT, TEMPLATE)
    assert len(pypdf.PdfReader(io.BytesIO(pdf)).pages) == 1


# ── the geometry, on its own ─────────────────────────────────────────────────
def test_the_artwork_pixel_to_point_map_inverts_the_image_axis():
    """A PDF image's FIRST row of samples lies along the TOP of the unit square,
    so the pixel row has to be flipped before the CTM is applied. Get that wrong
    and the signature lands the same distance BELOW the rule as it should have
    been above it — which still renders, still extracts, and is still wrong.

    The expected values are worked out from the placement independently of the
    module: y = 1.75 + 791 * (1 - 1431/1649), x = 0 + 612 * (470/1275).
    """
    x, y = asig.artwork_point(FORM_CM, 470, asig.RULE_Y_SIGNATURE)
    assert round(x, 2) == round(612.0 * 470 / 1275, 2) == 225.60
    assert round(y, 2) == round(1.75 + 791.0 * (1 - 1431 / 1649), 2) == 106.32

    # The top-left corner of the artwork maps to the TOP of the placement, not
    # the bottom. This is the assertion the flip exists for.
    assert asig.artwork_point(FORM_CM, 0, 0)[1] > asig.artwork_point(FORM_CM, 0, 1649)[1]


def test_a_value_too_wide_for_its_rule_is_shrunk_before_it_is_drawn():
    wide = "Wolfeschlegelsteinhausenbergerdorff-Montgomery-Fitzwilliam"
    small = asig.fit_size(wide, 198.0, 14.0, italic=True)
    assert small is not None and small < 14.0
    assert asig.text_width(wide, small, italic=True) <= 198.0


def test_a_value_that_will_not_fit_at_a_legible_size_is_refused_not_smudged():
    """The alternative is a contract nobody can read, or a name printed over the
    DATE label. Both are worse than a named refusal."""
    assert asig.fit_size("x" * 4000, 198.0, 14.0, italic=False) is None
    with pytest.raises(asig.AcceptanceError) as exc:
        asig.build_overlay(FORM_CM, {"signature": "x" * 4000}, "R", "I")
    assert "signature" in str(exc.value)


def test_the_date_on_the_rule_is_the_portal_stamp_without_the_time():
    """The DATE rule is ~108pt wide and the portal's stamp is half as wide again.
    The date is what a contract's DATE line asks for; the hour is still on the
    portal's approval row, which is where a dispute would read it from. NO
    timezone maths — the string is sliced, never recomputed."""
    assert asig.acceptance_date("September 18, 2026 at 10:42 AM CDT") == "September 18, 2026"
    # An unfamiliar format prints whole rather than printing nothing.
    assert asig.acceptance_date("2026-09-18") == "2026-09-18"
    assert asig.acceptance_date(None) == ""


def test_only_the_direct_form_carries_an_acceptance_block():
    """Verified by artwork hash across the shipped templates: Direct Epoxy /
    Polish / Combo / Budget all paint the same form, which has the ACCEPTANCE
    row; the three GC templates share another that has none, and Gyp has its own,
    also without."""
    for wt in ("epoxy", "polish", "combo", "budget"):
        assert asig.template_has_acceptance_block(wt, "Direct") is True, wt
    for wt in ("epoxy", "polish", "combo", "sealer"):
        assert asig.template_has_acceptance_block(wt, "GC") is False, wt
    assert asig.template_has_acceptance_block("gyp", None) is False
    assert asig.template_has_acceptance_block("gyp", "Direct") is False


def test_an_unmapped_work_type_raises_rather_than_inheriting_the_epoxy_fallback():
    """`proposal_writer.pick_template` falls back to (epoxy, Direct) so a proposal
    always renders. Inheriting that here would answer "yes, it has a block" for a
    combination nobody has ever seen and then draw a signature onto whatever page
    came back."""
    with pytest.raises(asig.AcceptanceError) as exc:
        asig.template_has_acceptance_block("terrazzo", "Direct")
    assert "terrazzo" in str(exc.value)


# ── finding the page ─────────────────────────────────────────────────────────
def test_the_last_page_is_not_the_acceptance_page():
    """THE BUG THIS MODULE EXISTS TO PREVENT. A Direct Epoxy proposal is four
    pages and the last three are Terms & Conditions on blank letterhead. Signing
    `pages[-1]` would put the customer's name on the back of the Terms, on blank
    paper. The check refuses it because that page's artwork is shared with its
    neighbours."""
    reader = pypdf.PdfReader(io.BytesIO(PROPOSAL_PDF))
    last = len(reader.pages) - 1
    assert last != ACCEPTANCE_INDEX
    with pytest.raises(asig.AcceptanceError) as exc:
        asig.verify_acceptance_page(reader, last)
    assert "same artwork" in str(exc.value)

    # And the right page passes, so the test above is not green for the trivial
    # reason that nothing ever passes.
    asig.verify_acceptance_page(reader, ACCEPTANCE_INDEX)


def test_a_cover_letter_moves_the_acceptance_page_to_index_one():
    reader = pypdf.PdfReader(io.BytesIO(PROPOSAL_PDF_WITH_LETTER))
    asig.verify_acceptance_page(reader, ACCEPTANCE_INDEX_WITH_LETTER)
    with pytest.raises(asig.AcceptanceError):
        asig.verify_acceptance_page(reader, 0)


def test_a_page_with_no_full_page_artwork_is_refused():
    plain = make_pdf([(["no artwork here"], None, None)])
    with pytest.raises(asig.AcceptanceError) as exc:
        asig.verify_acceptance_page(pypdf.PdfReader(io.BytesIO(plain)), 0)
    assert "does not carry the full-page Treadwell form" in str(exc.value)


def test_an_index_past_the_end_is_refused_by_name():
    with pytest.raises(asig.AcceptanceError) as exc:
        asig.verify_acceptance_page(pypdf.PdfReader(io.BytesIO(PROPOSAL_PDF)), 9)
    assert "9" in str(exc.value) or "10" in str(exc.value)


def test_artwork_at_another_resolution_is_refused():
    """Every number in RULES is a pixel of the 1275x1649 form. Re-export it at 300
    DPI and they all point somewhere else — a signature half an inch below the
    line is not a typo, it is a contract that looks unsigned."""
    reader = pypdf.PdfReader(io.BytesIO(PROPOSAL_PDF))
    page = reader.pages[ACCEPTANCE_INDEX]
    image = page["/Resources"]["/XObject"]["/Im0"].get_object()
    image[pypdf.generic.NameObject("/Height")] = pypdf.generic.NumberObject(3298)
    with pytest.raises(asig.AcceptanceError) as exc:
        asig.verify_acceptance_page(reader, ACCEPTANCE_INDEX)
    assert "3298" in str(exc.value)


# ── the endpoint: auth and the request ───────────────────────────────────────
def test_no_token_is_401(wired):
    assert _post(token=None).status_code == 401


def test_a_wrong_token_is_401(wired):
    r = _post(token="not-the-token")
    assert r.status_code == 401 and r.json()["detail"] == "unauthorized"


def test_an_unset_server_token_cannot_be_matched_by_an_empty_header(monkeypatch):
    """An unconfigured SERVICE_TOKEN must lock the door, not open it to everybody
    who sends no header."""
    monkeypatch.delenv("SERVICE_TOKEN", raising=False)
    assert _post(token="").status_code == 401


def test_malformed_certificate_json_is_400(wired):
    r = _post(raw_cert="{not json at all")
    assert r.status_code == 400
    body = r.json()
    assert body["ok"] is False and "certificate" in body["error"]


def test_a_missing_field_is_named_in_the_400(wired):
    for name in ("signer_email", "signed_at_central"):
        short = {k: v for k, v in CERT.items() if k != name}
        r = _post(cert=short)
        assert r.status_code == 400, name
        assert r.json()["ok"] is False
        assert name in r.json()["error"], (name, r.json()["error"])


def test_a_missing_proposal_pdf_is_400(wired):
    r = _post(omit_pdf=True)
    assert r.status_code == 400 and "proposal_pdf" in r.json()["error"]


def test_an_unreadable_proposal_pdf_is_400_not_502(wired):
    """The caller's fault and our fault get paged very differently."""
    r = _post(pdf_bytes=b"this is not a pdf")
    assert r.status_code == 400 and "proposal_pdf" in r.json()["error"]


def test_the_route_is_public_to_the_auth_gate_and_gated_by_the_token():
    """It is server-to-server: the portal has no Google session. The SERVICE_TOKEN
    check inside the handler is the whole gate, so the path has to be in the
    public set AND the handler has to check."""
    assert PATH in main._AUTH_PUBLIC_PATHS
    assert main._auth_is_public(PATH, "POST") is True


# ── the endpoint: the document that comes back ───────────────────────────────
def test_a_signed_contract_comes_back_as_a_pdf(wired):
    r = _post()
    assert r.status_code == 200
    assert r.headers["content-type"] == "application/pdf"
    assert r.content.startswith(b"%PDF")


def test_the_output_has_exactly_the_pages_the_input_had(wired):
    """No appended page. Hanz: "we dont need another page for the signatory"."""
    assert len(_pages(_post().content)) == len(_pages(PROPOSAL_PDF)) == 4


def test_the_four_values_land_on_the_acceptance_page_by_index(wired):
    """BY INDEX, never "the last page" — see this file's header."""
    pages = _pages(_post().content)
    signed = _text(pages[ACCEPTANCE_INDEX])
    assert "Dana Whitfield" in signed
    assert "September 18, 2026" in signed
    assert "$187,450.75" in signed
    # The printed name is a second, separate line of ink, not the same run read
    # twice: the block asks for both a signature and a printed name.
    assert signed.count("Dana Whitfield") == 2

    for i, page in enumerate(pages):
        if i == ACCEPTANCE_INDEX:
            continue
        other = _text(page)
        assert "Dana Whitfield" not in other, i
        assert "$187,450.75" not in other, i


def test_a_cover_letter_proposal_signs_page_two(wired, monkeypatch):
    """The letter is prepended into the proposal's own bytes, so the form — and
    the acceptance block with it — moves to index 1. A handler that assumed index
    0 would sign the letterhead of the cover letter."""
    monkeypatch.setattr(drafts, "get_revision",
                        lambda pid, rev: _draft_row(cover_letter=True))
    pages = _pages(_post(pdf_bytes=PROPOSAL_PDF_WITH_LETTER).content)
    assert len(pages) == 5
    assert "Dana Whitfield" in _text(pages[ACCEPTANCE_INDEX_WITH_LETTER])
    assert "Dana Whitfield" not in _text(pages[0])


def test_nothing_but_the_four_values_reaches_the_paper(wired):
    """Hanz: "this is too much information... Just the basic information is what
    we need". The IP address, the browser string, the consent paragraph and both
    digests are still validated, still stored on the portal's approval row, and
    printed nowhere.

    Checked against the raw PDF bytes as well as the extracted text: a value drawn
    in a font pypdf cannot decode would vanish from `extract_text` and still be on
    the page."""
    out = _post().content
    text = "\n".join(_text(p) for p in _pages(out))
    for secret in (CERT["ip_address"], CERT["user_agent"], CERT["consent_text"],
                   CERT["consent_version"], PROPOSAL_SHA, REVISION_SHA,
                   CERT["signer_email"], CERT["options_summary"]):
        assert secret not in text, secret
        assert secret.encode("latin-1", "replace") not in out, secret
    assert b"15 U.S.C" not in out and b"7001" not in out


def _page_fingerprint(page):
    """What must not change about a page we are only passing through.

    The DECODED content stream alone is not enough, and that is not a guess: the
    first version of this test compared only `get_data()` and stayed GREEN when the
    handler was mutated to call `compress_content_streams()` on every page. A
    re-encode leaves the drawing operators identical while changing the bytes we
    store — so the stream's FILTER and the page geometry are pinned beside it.
    """
    contents = page["/Contents"].get_object()
    return (contents.get("/Filter"), contents.get_data(),
            str(page.get("/MediaBox")), str(page.get("/Rotate")))


def test_every_other_page_is_byte_identical(wired):
    """THE guarantee, kept from the certificate-page version of this route. The
    portal stores the SHA-256 of the PDF the customer opened; if these pages were
    re-rendered, re-encoded or re-flowed, the document we keep would not hash to
    the value recorded beside it, and it would be evidence against us instead of
    for us."""
    before = _pages(PROPOSAL_PDF)
    after = _pages(_post().content)
    assert len(after) == len(before)
    checked = 0
    for i in range(len(before)):
        if i == ACCEPTANCE_INDEX:
            continue
        assert _page_fingerprint(after[i]) == _page_fingerprint(before[i]), (
            "page %d was rewritten" % i)
        checked += 1
    assert checked == 3, "this test compared %d pages, not the three it means to" % checked


def test_the_signed_page_keeps_its_own_content_stream(wired):
    """The signed page differs from the input ONLY by the overlay.

    Its original content stream object is not re-encoded, not decompressed and
    not rewritten: `q` and `Q` go in as their own entries around it and the
    overlay follows, so the proposal's own drawing bytes come back out exactly as
    they went in."""
    was = _pages(PROPOSAL_PDF)[ACCEPTANCE_INDEX]["/Contents"].get_object()
    parts = _pages(_post().content)[ACCEPTANCE_INDEX]["/Contents"]
    assert isinstance(parts, pypdf.generic.ArrayObject) and len(parts) == 4
    kept = parts[1].get_object()
    assert kept.get_data() == was.get_data()
    assert kept.get("/Filter") == was.get("/Filter")
    assert parts[0].get_object().get_data().strip() == b"q"
    assert parts[2].get_object().get_data().strip() == b"Q"
    assert b"Tj" in parts[3].get_object().get_data()


def _overlay_ops(page) -> list:
    """(font, size, x, y, text) for every string the overlay draws."""
    stream = ContentStream(page["/Contents"][-1].get_object(), page.pdf)
    out, font, size, pos = [], None, None, None
    for operands, op in stream.operations:
        if op == b"Tf":
            font, size = str(operands[0]), float(operands[1])
        elif op == b"Tm":
            pos = (float(operands[4]), float(operands[5]))
        elif op == b"Tj":
            out.append((font, size, pos[0], pos[1], operands[0]))
    return out


def test_the_ink_lands_on_kyles_rules_at_the_measured_coordinates(wired):
    """The coordinate assertion, executed rather than asserted about source text.

    Expected values are derived here from the placement and the rule's own pixel
    row, independently of the module: the signature baseline sits LIFT above the
    rule at 1.75 + 791 * (1 - 1431/1649), and its ink starts INSET in from the
    rule's left end at 612 * 470/1275.
    """
    ops = _overlay_ops(_pages(_post().content)[ACCEPTANCE_INDEX])
    drawn = {str(text): (font, size, x, y) for font, size, x, y, text in ops}
    assert set(drawn) == {"Dana Whitfield", "September 18, 2026", "$187,450.75"}

    sig = [o for o in ops if str(o[4]) == "Dana Whitfield"]
    assert len(sig) == 2, "the signature and the printed name are two separate marks"
    sig_line = max(sig, key=lambda o: o[3])            # the upper of the two rules
    printed_line = min(sig, key=lambda o: o[3])

    assert round(sig_line[2], 2) == round(612.0 * 470 / 1275 + asig.INSET_PT, 2)
    assert round(sig_line[3], 2) == round(
        1.75 + 791.0 * (1 - asig.RULE_Y_SIGNATURE / 1649) + asig.LIFT_PT, 2)
    assert round(printed_line[3], 2) == round(
        1.75 + 791.0 * (1 - asig.RULE_Y_PRINTED / 1649) + asig.LIFT_PT, 2)

    # The signature is the only italic mark; the printed name below it is ordinary
    # type. Same string, two faces, which is what the block asks for.
    assert sig_line[0] != printed_line[0]
    assert sig_line[1] > printed_line[1]

    date = [o for o in ops if str(o[4]) == "September 18, 2026"][0]
    total = [o for o in ops if str(o[4]) == "$187,450.75"][0]
    assert round(date[2], 2) == round(612.0 * 966 / 1275 + asig.INSET_PT, 2)
    assert round(date[3], 2) == round(sig_line[3], 2)      # same row as the signature
    assert round(total[2], 2) == round(612.0 * 976 / 1275 + asig.INSET_PT, 2)
    assert round(total[3], 2) == round(printed_line[3], 2)  # same row as the printed name


def test_the_signature_is_an_italic_serif_that_needs_no_font_file(wired):
    """There is no script font in the container — the Dockerfile installs Carlito and
    Liberation, the host mounts Zetta Serif, nothing else — so naming a cursive face would
    silently substitute and render one way in CI and another in production. The
    PDF base-14 faces need no font file on the box at all, so there is one code
    path, not two."""
    page = _pages(_post().content)[ACCEPTANCE_INDEX]
    fonts = {str(k): v.get_object() for k, v in page["/Resources"]["/Font"].items()}
    ours = {name: f for name, f in fonts.items() if str(f.get("/BaseFont")).startswith("/Times")}
    assert {str(f["/BaseFont"]) for f in ours.values()} == {"/Times-Italic", "/Times-Roman"}
    for f in ours.values():
        assert str(f["/Encoding"]) == "/WinAnsiEncoding"
        assert "/FontFile" not in f and "/FontFile2" not in f and "/FontFile3" not in f
    # The page's own Helvetica is still there — the overlay added resources, it
    # did not replace them.
    assert any(str(f.get("/BaseFont")) == "/Helvetica" for f in fonts.values())


# ── the endpoint: what it refuses ────────────────────────────────────────────
def test_a_gc_proposal_is_refused_and_no_pdf_comes_back(wired, monkeypatch):
    """Kyle's GC form has no acceptance row. Until he adds one those proposals
    approve exactly as they always did and no signed contract is produced — the
    same treatment Budget Pricing already gets for having no Terms and
    Conditions. The portal decides this before it calls; this is the guard that
    stops a stale portal getting a signature drawn on blank letterhead."""
    monkeypatch.setattr(drafts, "get_revision", lambda pid, rev: _draft_row(audience="GC"))
    r = _post()
    assert r.status_code == 400
    assert not r.content.startswith(b"%PDF")
    assert r.json()["ok"] is False
    assert "no signature line" in r.json()["error"]


def test_a_gyp_proposal_is_refused_and_no_pdf_comes_back(wired, monkeypatch):
    monkeypatch.setattr(drafts, "get_revision",
                        lambda pid, rev: _draft_row(work_type="gyp", audience="Direct"))
    r = _post()
    assert r.status_code == 400 and not r.content.startswith(b"%PDF")
    assert "no signature line" in r.json()["error"]


def test_a_draft_we_do_not_have_is_a_400_naming_it(wired, monkeypatch):
    monkeypatch.setattr(drafts, "get_revision", lambda pid, rev: None)
    r = _post()
    assert r.status_code == 400
    assert "no proposal on file" in r.json()["error"]


def test_a_database_failure_is_a_502_with_a_sentence_not_a_guess(wired, monkeypatch):
    """A draft we cannot READ is not a draft that says "page 0". Refusing is the
    only safe answer — the alternative is a signature on a guessed page."""
    def boom(pid, rev):
        raise RuntimeError("connection reset by peer")
    monkeypatch.setattr(drafts, "get_revision", boom)
    r = _post()
    assert r.status_code == 502
    assert "could not look up" in r.json()["error"]
    assert "connection reset" not in r.json()["error"], (
        "the library's own message must not reach the caller")


def test_a_draft_that_never_generated_is_a_502_with_a_sentence(wired, monkeypatch):
    monkeypatch.setattr(drafts, "get_revision", lambda pid, rev: {"id": "x", "data": {}})
    r = _post()
    assert r.status_code == 502 and "no generated document" in r.json()["error"]


def test_a_proposal_whose_pages_are_not_what_we_expected_is_refused(wired):
    """The draft said "no cover letter", so the handler looked at page 1 — and
    found the cover letter's letterhead there instead of the form. It refuses
    rather than signing the wrong page, and the refusal says which page."""
    r = _post(pdf_bytes=PROPOSAL_PDF_WITH_LETTER)
    assert r.status_code == 502
    assert not r.content.startswith(b"%PDF")
    assert "page 1" in r.json()["error"]


def test_a_signature_that_will_not_fit_refuses_rather_than_printing_a_smudge(wired):
    r = _post(cert={**CERT, "signer_name": "Bartholomew " * 120})
    assert r.status_code == 502
    assert "acceptance block" in r.json()["error"]


# ── a caller-controlled value cannot forge a log line or leak an exception ────
# CodeQL flagged this route: `cert` is arbitrary JSON off an authenticated but
# still external POST, and three log calls used to print `cert["proposal_id"]`
# straight into a %s slot -- an embedded CRLF would let anyone holding
# SERVICE_TOKEN write a fabricated log line under a real one. Two other spots
# echoed str(exc) straight into the 400 response body, which is an information
# exposure risk for a library exception whose message is not meant for a caller.
def test_log_safe_strips_control_characters_and_caps_length():
    """The unit itself: repr() is the actual mechanism (CodeQL recognizes it as a
    sanitizer for py/log-injection; a hand-rolled character replace was tried
    first and kept getting flagged), so what lands in a log line is the escaped
    two-character sequence, never the raw byte -- and a value past the cap is
    truncated rather than flooding a log line."""
    forged = "real-id\r\n2026-01-01 ERROR fake admin login succeeded"
    safe = main._log_safe(forged)
    assert "\r" not in safe and "\n" not in safe
    # The escape sequences themselves show up as visible text -- that is the
    # point, not an accident: a log reader sees exactly what was sent.
    assert "\\r\\n" in safe
    # The forged line's own text survives (it is not secret), just unable to
    # start a new log record -- \r and \n are the only bytes that matter here.
    assert "fake admin login succeeded" in safe

    assert main._log_safe(None) == "''"
    long_value = "x" * 500
    capped = main._log_safe(long_value, limit=50)
    # 50 chars + the ellipsis mark, then repr()'s own quoting around all of it.
    assert capped.startswith("'") and capped.endswith("…'")
    assert len(capped) == 50 + len("…") + 2  # + the two quote characters repr() adds


def test_a_forged_proposal_id_cannot_inject_a_second_log_line(wired, caplog):
    """DRIVEN THROUGH THE REAL ENDPOINT, not the helper in isolation -- a sanitizer
    that exists but is never called on the actual request path protects nothing.

    Uploads a cover-letter proposal against a draft that says there is none, so
    the signing step refuses and this actually reaches the "refused to sign
    proposal %s" site -- one of the places that logs the caller's proposal_id."""
    caplog.set_level("WARNING")
    forged_id = "legit-1\r\nCRITICAL admin_override=true"
    resp = _post(pdf_bytes=PROPOSAL_PDF_WITH_LETTER,
                 cert={**CERT, "proposal_id": forged_id})
    assert resp.status_code == 502
    assert any("refused to sign" in r.getMessage() for r in caplog.records), (
        "this test did not actually reach the log line it means to check")
    for record in caplog.records:
        assert "\r" not in record.getMessage()
        assert "\n" not in record.getMessage()


def test_a_bad_certificate_json_does_not_echo_the_parser_exception(wired):
    """400, but the body says what is wrong in one plain sentence -- not whatever
    text json.JSONDecodeError happened to compose (position, line/column, quoted
    fragments of the bad input)."""
    resp = _post(raw_cert="{not json")
    assert resp.status_code == 400
    body = resp.json()
    assert body["error"] == "certificate is not valid JSON"
    assert "line" not in body["error"] and "column" not in body["error"]


def test_an_unreadable_proposal_pdf_does_not_echo_the_reader_exception(wired):
    """Same guarantee for the other exception-message site: pypdf's own error text
    (which can quote raw bytes) never reaches the response body."""
    resp = _post(pdf_bytes=b"not a pdf at all")
    assert resp.status_code == 400
    body = resp.json()
    assert body["error"] == "proposal_pdf is not a readable PDF"


# ── against a real production render ─────────────────────────────────────────
SAMPLE = r"C:/Users/Admin/Documents/Treadwell Main/sample-epoxy-proposal.pdf"


@pytest.mark.skipif(not os.path.isfile(SAMPLE),
                    reason="the real Direct Epoxy render is customer material and is not "
                           "committed; it lives on the dev box only")
def test_against_the_real_direct_epoxy_render():
    """The one thing a fixture cannot prove: that a genuine LibreOffice render of
    Kyle's template — JPEG artwork, its own placements, four pages — is read the
    same way. The form is page 1 and the Terms are pages 2-4, so a handler that
    reached for the last page would be caught here too."""
    raw = open(SAMPLE, "rb").read()
    reader = pypdf.PdfReader(io.BytesIO(raw))
    assert len(reader.pages) == 4
    asig.verify_acceptance_page(reader, 0)
    for wrong in (1, 2, 3):
        with pytest.raises(asig.AcceptanceError):
            asig.verify_acceptance_page(reader, wrong)

    signed = asig.sign_acceptance_page(raw, 0, signer_name="Dana Whitfield",
                                       date_text="September 18, 2026",
                                       total_text="$187,450.75")
    out = pypdf.PdfReader(io.BytesIO(signed))
    assert len(out.pages) == 4
    assert "Dana Whitfield" in (out.pages[0].extract_text() or "")
    for i in (1, 2, 3):
        assert "Dana Whitfield" not in (out.pages[i].extract_text() or "")
