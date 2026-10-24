"""ONE WORKBOOK FOR A PROJECT OF SEVERAL SCOPES (build spec T2, first slice: no Dropbox).

RJ's picture: a project split into scopes (Breakroom polish, Lab 1 epoxy, Lab 2 epoxy) is one
workbook with one tab per scope. estimate_writer.fill_group_estimate builds it in one fill_estimate
pass; GET /api/scope-group/{gid}/workbook hands it to staff as a one-hour /api/file link.

What these pin, each against the BUILT .xlsx (openpyxl reads it back), never the source:
  * one tab per scope, named for it, first, in the group's order; a later scope on the same sheet is
    a Copy<N> of it, and a scope's tab comes from its model SECTION's `sheet`, not base_tab_id;
  * cross-sheet formulas follow the renamed tabs (`=Epoxy!B1` becomes `='Lab 1'!B1`);
  * every tab shows ITS OWN scope's answers: tab 2 and tab 3 never show tab 1's;
  * the rate locks are kept on every scope tab, copies included;
  * 12 scopes build, 13 are refused; colliding names are told apart;
  * a scope that cannot be put on a tab of its own is refused by name, never guessed at;
  * a second Gyp scope is its own tab, its square feet its own, with the note;
  * every tab, scope tab or template tab, carries the "Job info and conditions only" note;
  * the route reads only LIVE members and answers in sentences.
"""
import copy
import io
import warnings

import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook

import drafts
import estimate_writer as ew
import main

client = TestClient(main.app)
GID = "grp_workbook0001"


def _load(content: bytes):
    with warnings.catch_warnings():
        warnings.simplefilter("ignore")
        return load_workbook(io.BytesIO(content))


_TEMPLATE = []


def _template():
    """The pristine template, read once (the tests only read it)."""
    if not _TEMPLATE:
        with warnings.catch_warnings():
            warnings.simplefilter("ignore")
            _TEMPLATE.append(ew._fresh_template())
    return _TEMPLATE[0]


def _section(layout, sheet=None, kind="base"):
    s = {"layout": layout, "kind": kind, "profile": "%s@1" % layout}
    if sheet:
        s["sheet"] = sheet
    return s


def epoxy(label, base="Epoxy", **data):
    """An Epoxy scope. `base` "Copy1" is a split scope whose base section is a Copy id."""
    sec = _section("epoxy", "Epoxy" if base != "Epoxy" else None)
    d = {"project_name": "Acme HQ", "work_type": "epoxy", "base_tab_id": base,
         "polish_estimate": {"version": 2, "kind": "off", "tabs": {base: sec}}}
    d.update(data)
    return {"id": "id-" + label, "label": label, "data": d}


def polish(label, **data):
    d = {"project_name": "Acme HQ", "work_type": "polish", "base_tab_id": "Polish",
         "polish_estimate": {"version": 2}}
    d.update(data)
    return {"id": "id-" + label, "label": label, "data": d}


def gyp(label, sheet=ew.GYP_SHEET, **data):
    d = {"project_name": "Acme HQ", "work_type": "gyp", "base_tab_id": sheet,
         "polish_estimate": {"version": 2, "kind": "off", "tabs": {sheet: _section("gyp")}}}
    d.update(data)
    return {"id": "id-" + label, "label": label, "data": d}


# Tab 1 answers everything; tabs 2 and 4 (Lab 2 a split scope on Copy1, Lab 3 a plain Epoxy one)
# answer less, so anything they INHERITED from tab 1 would show.
MEMBERS = [
    epoxy("Lab 1", drawings_dated="2026-01-05", system_1_sf=5000,
          cell_values={"Epoxy!B4": "No", "Epoxy!B10": "Reno", "Polish!B4": "No"}),
    epoxy("Lab 2", base="Copy1", system_1_sf=1200, cell_values={"Epoxy!B4": "Yes"}),
    polish("Breakroom", cell_values={"Polish!B4": "Yes"}),
    epoxy("Lab 3", system_1_sf=800),
]


@pytest.fixture(scope="module")
def built():
    return _load(ew.fill_group_estimate(copy.deepcopy(MEMBERS)))


# ── the tabs ─────────────────────────────────────────────────────────────────────────────────
def test_one_tab_per_scope_named_for_it_first_in_the_groups_order(built):
    assert built.sheetnames[:4] == ["Lab 1", "Lab 2", "Breakroom", "Lab 3"]
    # The template's other tabs stay behind them (their formulas are read by the scope tabs).
    assert len(built.sheetnames) == len(_template().sheetnames) + 2
    assert "Seal" in built.sheetnames and "Epoxy" not in built.sheetnames


def test_a_later_scope_on_the_same_sheet_is_a_copy_and_the_sheet_comes_from_the_section():
    plan = ew.plan_group_workbook(copy.deepcopy(MEMBERS))
    assert [(t["tab"], t["sheet"]) for t in plan["tabs"]] == [
        ("Epoxy", "Epoxy"), ("Copy1", "Epoxy"), ("Polish", "Polish"), ("Copy2", "Epoxy")]
    assert plan["tab_copies"] == [{"id": "Copy1", "source": "Epoxy"},
                                  {"id": "Copy2", "source": "Epoxy"}]
    # Lab 2's base_tab_id is "Copy1", a section id: its sheet is the section's own `sheet`.
    assert ew.group_member_sheet(MEMBERS[1]["data"]) == ("Copy1", "Epoxy")
    assert ew.group_member_sheet(gyp("FR", sheet="Gyp (FR)")["data"]) == ("Gyp (FR)", "Gyp (FR)")


def test_mirror_formulas_follow_the_renamed_tabs(built):
    # A tab no scope is on (Seal) reads the project name off Epoxy, which is now "Lab 1"; the Polish
    # tab (Breakroom) still reads its contact block off it. (Its job-info block is its OWN scope's
    # literals now, finding 23: see test_a_polish_scopes_job_info_is_its_own.)
    assert built["Seal"]["B1"].value == "='Lab 1'!B1"
    assert built["Breakroom"]["G2"].value == "='Lab 1'!G2"
    for name in ("Breakroom", "Lab 2", "Seal"):
        stale = [c.coordinate for c in built[name]._cells.values()
                 if isinstance(c.value, str) and c.value.startswith("=") and "Epoxy!" in c.value
                 and "Epoxy blank" not in c.value]
        assert stale == [], (name, stale[:5])


def test_every_tab_shows_its_own_scopes_answers_not_tab_ones(built):
    tpl = _template()["Epoxy"]
    one, two, three = built["Lab 1"], built["Lab 2"], built["Lab 3"]
    assert (one["E20"].value, two["E20"].value, three["E20"].value) == (5000, 1200, 800)
    assert (one["B4"].value, two["B4"].value) == ("No", "Yes")
    assert three["B4"].value == tpl["B4"].value          # Lab 3 never answered: the template's
    # Only tab 1 answered these; tabs 2 and 3 must not show tab 1's answer.
    assert one["B9"].value == "2026-01-05" and one["B10"].value == "Reno"
    for tab in (two, three):
        assert tab["B9"].value == tpl["B9"].value and tab["B10"].value == tpl["B10"].value
    assert two["B1"].value == "Acme HQ" and three["B1"].value == "Acme HQ"
    # Each scope's own cells land on its own tab: Lab 1's Polish!B4 is not the Breakroom's answer.
    assert built["Breakroom"]["B4"].value == "Yes"


# ── a Polish (or Seal) scope's job info is its own (finding 23) ───────────────────────────────
JOB = ("B1", "B2", "B3", "C3", "B4", "D5", "B6", "D6", "B7", "B8", "B9", "B10")


def seal(label, **data):
    d = {"project_name": "Acme HQ", "work_type": "seal", "base_tab_id": "Seal",
         "polish_estimate": {"version": 2, "kind": "off", "tabs": {"Seal": _section("seal")}}}
    d.update(data)
    return {"id": "id-" + label, "label": label, "data": d}


BREAKROOM = polish("Breakroom", bid_date="2026-10-01", architect="Arch B", drawings_dated="BR-dwg",
                   cell_values={"Epoxy!D5": "No"})
LAB_1 = epoxy("Lab 1", bid_date="2026-10-02", architect="Arch A", drawings_dated="Lab1-dwg",
              address="1 Main St", cell_values={"Epoxy!D5": "Yes"})


def _no_mirror_into_another_tab(ws):
    return [(a, ws[a].value) for a in JOB
            if isinstance(ws[a].value, str) and ws[a].value.startswith("=") and "!" in ws[a].value]


@pytest.mark.parametrize("order", ["polish first", "polish second"])
def test_a_polish_scopes_job_info_is_its_own(order):
    """The review's own build: Breakroom (polish, bid 2026-10-01) beside Lab 1 (epoxy, 2026-10-02,
    drawings 'Lab1-dwg') and Lab 2 (a split Epoxy scope on Copy1). Breakroom's tab used to read
    B2 = ='Lab 1'!B2, and its own date appeared nowhere."""
    lab2 = epoxy("Lab 2", base="Copy1", bid_date="2026-10-03")
    members = ([BREAKROOM, LAB_1, lab2] if order == "polish first" else [LAB_1, BREAKROOM, lab2])
    wb = _load(ew.fill_group_estimate(copy.deepcopy(members)))
    br, one = wb["Breakroom"], wb["Lab 1"]
    assert (br["B2"].value, br["B8"].value, br["B9"].value) == ("2026-10-01", "Arch B", "BR-dwg")
    assert (one["B2"].value, one["B8"].value, one["B9"].value) == ("2026-10-02", "Arch A", "Lab1-dwg")
    assert br["D5"].value == "No" and one["D5"].value == "Yes"
    assert br["B1"].value == "Acme HQ"
    # Lab 1's address is Lab 1's: Breakroom never answered one, so its cell is empty, not Lab 1's.
    assert one["B3"].value == "1 Main St" and br["B3"].value is None
    assert _no_mirror_into_another_tab(br) == [], "Breakroom still reads another scope's answers"
    assert wb["Lab 2"]["B2"].value == "2026-10-03"


def test_a_polish_tab_keeps_the_scopes_own_typed_flag_over_the_job_answer():
    """The scope's own typed cell on its own sheet wins, as it always has (FLAG_BLOCK_CELLS)."""
    br = polish("Breakroom", taxable="Yes", cell_values={"Polish!B6": "No", "Epoxy!D6": "Yes"})
    wb = _load(ew.fill_group_estimate([copy.deepcopy(LAB_1), br]))
    assert wb["Breakroom"]["B6"].value == "No" and wb["Breakroom"]["D6"].value == "Yes"
    assert wb["Lab 1"]["D6"].value == _template()["Epoxy"]["D6"].value


def test_a_seal_scopes_job_info_is_its_own_too():
    """Seal mirrors Epoxy (and its B6 mirrors Polish, which mirrors Epoxy): followed to the end."""
    sl = seal("Hallway", bid_date="2026-10-05", architect="Arch S", taxable="No")
    wb = _load(ew.fill_group_estimate([copy.deepcopy(LAB_1), copy.deepcopy(BREAKROOM), sl]))
    hw = wb["Hallway"]
    assert (hw["B2"].value, hw["B8"].value, hw["B6"].value) == ("2026-10-05", "Arch S", "No")
    assert hw["B9"].value is None and hw["D5"].value == _template()["Epoxy"]["D5"].value
    assert _no_mirror_into_another_tab(hw) == []


def test_the_plan_names_the_cells_it_empties():
    plan = ew.plan_group_workbook([copy.deepcopy(LAB_1), copy.deepcopy(BREAKROOM)])
    assert plan["cell_values"]["Polish!B2"] == "2026-10-01"
    assert "B3" in plan["blank_cells"]["Polish"] and "Epoxy" not in plan["blank_cells"]


def test_the_rate_locks_are_kept_on_every_scope_tab(built):
    tpl = _template()["Epoxy"]
    for name in ("Lab 1", "Lab 2", "Lab 3", "Breakroom"):
        assert built[name].protection.sheet is True, name
    for name in ("Lab 1", "Lab 2", "Lab 3"):
        ws = built[name]
        assert all(ws[a].protection.locked for a in ew.LOCK_MAP["Epoxy"]), name
        assert ws["E20"].protection.locked == tpl["E20"].protection.locked is False


def _note_cells(ws):
    return [c.coordinate for row in ws.iter_rows(min_row=1, max_row=1) for c in row
            if c.value == ew.GROUP_WORKBOOK_NOTE]


def test_every_tab_carries_the_note_scope_tabs_and_template_tabs_alike(built):
    # Spec T2: "Every tab gets a visible note cell". The template tabs behind the scope tabs carry
    # no pricing either, and the one no scope is on shows the first scope's answers.
    for name in ("Lab 1", "Lab 2", "Breakroom", "Lab 3", "Seal", "Takeoff", "Gyp (FR)"):
        ws = built[name]
        assert ws["D1"].value == ew.GROUP_WORKBOOK_NOTE, name
        assert "D1:F1" in {str(r) for r in ws.merged_cells.ranges}, name
    assert [ws.title for ws in built.worksheets if _note_cells(ws) != []] == built.sheetnames
    assert all(len(_note_cells(ws)) == 1 for ws in built.worksheets)


def test_a_tab_whose_d1_is_merged_gets_the_note_beside_it_not_over_it(built):
    # Specs+Dwgs+Addn merges A1:D1 ("Drawings"): the note goes in the next free span, E1:G1, and
    # the template's own heading and merge stay as they were.
    ws = built["Specs+Dwgs+Addn"]
    assert ws["A1"].value == _template()["Specs+Dwgs+Addn"]["A1"].value == "Drawings"
    assert _note_cells(ws) == ["E1"]
    merged = {str(r) for r in ws.merged_cells.ranges}
    assert {"A1:D1", "E1:G1"} <= merged


def test_the_tab_no_scope_is_on_carries_the_note():
    # A Polish-only project: the Epoxy tab is filled with the first scope's answers, and says so.
    wb = _load(ew.fill_group_estimate([polish("Breakroom"), polish("Lobby")]))
    assert wb["Epoxy"]["D1"].value == ew.GROUP_WORKBOOK_NOTE


def test_a_sheet_no_scope_is_on_gets_the_first_scopes_answers_and_only_its():
    # A Polish-only project: the Epoxy tab is no scope's, but the Polish tabs read it (B1, the
    # project's name, and the job's flags), so it carries the FIRST scope's answers.
    wb = _load(ew.fill_group_estimate([
        polish("Breakroom", cell_values={"Epoxy!D5": "Yes"}),
        polish("Lobby", cell_values={"Epoxy!D6": "Yes", "Polish!B4": "No"}),
    ]))
    tpl = _template()["Epoxy"]
    assert wb.sheetnames[:2] == ["Breakroom", "Lobby"]
    assert wb["Epoxy"]["B1"].value == "Acme HQ" and wb["Epoxy"]["D5"].value == "Yes"
    assert wb["Epoxy"]["D6"].value == tpl["D6"].value      # Lobby is not the first scope
    assert wb["Lobby"]["B4"].value == "No" and wb["Lobby"]["G2"].value == "=Epoxy!G2"
    # ...while each Polish tab's own job info is its own (finding 23): Lobby's remodel answer is not
    # the Epoxy tab's, which is Breakroom's.
    assert wb["Lobby"]["D6"].value == "Yes" and wb["Breakroom"]["D6"].value == tpl["D6"].value
    assert wb["Breakroom"]["D5"].value == "Yes" and wb["Lobby"]["D5"].value == tpl["D5"].value


# ── the ceiling and the names ────────────────────────────────────────────────────────────────
def test_twelve_scopes_build_and_thirteen_are_refused():
    twelve = [epoxy("Room %d" % i) for i in range(1, 13)]
    wb = _load(ew.fill_group_estimate(twelve))
    assert wb.sheetnames[:12] == ["Room %d" % i for i in range(1, 13)]
    assert all(wb["Room %d" % i]["D1"].value == ew.GROUP_WORKBOOK_NOTE for i in range(1, 13))
    with pytest.raises(ew.GroupWorkbookError) as exc:
        ew.plan_group_workbook(twelve + [epoxy("Room 13")])
    assert exc.value.code == "too_many_scopes"


def test_colliding_scope_names_are_told_apart():
    long_a = "North wing second floor breakroom A"
    long_b = "North wing second floor breakroom B"      # the same first 31 characters
    wb = _load(ew.fill_group_estimate([epoxy(long_a), epoxy(long_b), polish("Takeoff")]))
    cut = long_a[:31].strip()
    assert wb.sheetnames[:3] == [cut, cut[:29] + " 2", "Takeoff 2"]
    # ...and the formulas still name the tab that is there.
    assert wb["Takeoff 2"]["G2"].value == "='%s'!G2" % cut


# ── refusals ─────────────────────────────────────────────────────────────────────────────────
@pytest.mark.parametrize("key,value,words", [
    ("extras", [{"label": "Super Stick", "qty": 1, "unit_price": 50}], "extra material lines"),
    ("rooms", [{"name": "Room A"}], "rooms"),
    ("tab_copies", [{"id": "Copy1", "source": "Epoxy"}], "copied sheets"),
    ("tab_structs", [{"sheet": "Epoxy", "kind": "insert_row", "at": 5, "count": 1}],
     "inserted or deleted rows"),
    ("lock_overrides", {"Epoxy": {"unlock": ["B73"]}}, "locked or unlocked cells"),
    ("alternate_computed_bid", {"alternate_full_bid": {"total_base_bid": 31000}},
     "an alternate system"),
])
def test_a_scope_that_cannot_be_put_on_a_tab_of_its_own_is_refused_by_name(key, value, words):
    members = copy.deepcopy(MEMBERS)
    members[1]["data"][key] = value
    with pytest.raises(ew.GroupWorkbookError) as exc:
        ew.fill_group_estimate(members)
    assert exc.value.code == "member_not_supported"
    assert "“Lab 2”" in exc.value.message and words in exc.value.message


@pytest.mark.parametrize("mutate,code", [
    (lambda d: d.pop("polish_estimate"), "member_not_v2"),
    (lambda d: d.pop("work_type"), "member_not_estimated"),
    # A Combo priced as one bid: two bases, base_tab_id None. The tool does not guess which tab.
    (lambda d: d.update(base_tab_id=None, polish_estimate={
        "version": 2, "tabs": {"Epoxy": _section("epoxy")}}), "member_no_base"),
])
def test_a_scope_whose_tab_cannot_be_read_is_refused(mutate, code):
    members = copy.deepcopy(MEMBERS)
    mutate(members[2]["data"])
    with pytest.raises(ew.GroupWorkbookError) as exc:
        ew.plan_group_workbook(members)
    assert exc.value.code == code and "“Breakroom”" in exc.value.message


def test_an_unrefused_member_with_empty_lists_is_built():
    members = copy.deepcopy(MEMBERS)
    members[1]["data"].update(extras=[], rooms=[], tab_structs=[], lock_overrides={},
                              alternate_computed_bid={"alternate_full_bid": {}})
    assert [t["label"] for t in ew.plan_group_workbook(members)["tabs"]] == \
        ["Lab 1", "Lab 2", "Breakroom", "Lab 3"]


# ── Gyp ──────────────────────────────────────────────────────────────────────────────────────
def test_a_second_gyp_scope_is_its_own_tab_with_its_own_square_feet_and_the_note():
    wb = _load(ew.fill_group_estimate([
        gyp("Level 1", gyp_soft_sf=4000, gyp_hard_sf=100),
        gyp("Level 2", gyp_soft_sf=2500, project_name="Acme HQ"),
    ]))
    one, two = wb["Level 1"], wb["Level 2"]
    tpl = _template()
    assert wb.sheetnames[:2] == ["Level 1", "Level 2"]
    assert (one["G9"].value, two["G9"].value) == (4000, 2500)
    # Only Level 1 answered the hard square feet: Level 2 shows the template's, not Level 1's 100.
    assert one["I9"].value == 100 and two["I9"].value == tpl[ew.GYP_SHEET]["I9"].value != 100
    assert two["B2"].value == "Acme HQ"
    for ws in (one, two):
        assert ws["D1"].value == ew.GROUP_WORKBOOK_NOTE
        assert ws.protection.sheet is True
        assert all(ws[a].protection.locked for a in ew.LOCK_MAP["Gyp"])
    # The other four gyp sheets are no scope's: Level 1's square feet are not copied onto them,
    # as fill_estimate does for a one-gyp bid.
    for name in ew.GYP_SHEETS[1:]:
        assert wb[name]["G9"].value == tpl[name]["G9"].value, name


# ── the route ────────────────────────────────────────────────────────────────────────────────
def _row(member, created, gid=GID, **extra):
    data = copy.deepcopy(member["data"])
    data["scope_group"] = {"id": gid, "label": member["label"]}
    data.update(extra)
    return {"id": member["id"], "data": data, "owner_email": "kyle@wetreadwell.com",
            "created_at": created, "updated_at": created, "deleted_at": None}


def _store(fake_supabase, monkeypatch, rows):
    store = {"drafts": rows, "events": []}
    monkeypatch.setattr(drafts, "get_client", lambda: fake_supabase(store))
    return store


def test_the_route_builds_the_live_members_workbook(fake_supabase, monkeypatch):
    rows = [_row(MEMBERS[0], "2026-10-01"), _row(MEMBERS[2], "2026-10-02"),
            _row(epoxy("Old lab"), "2026-10-03", archived=True),
            _row(epoxy("Trashed"), "2026-10-04"),
            _row(epoxy("Elsewhere"), "2026-10-05", gid="grp_other00001")]
    rows[3]["deleted_at"] = "2026-10-06"
    _store(fake_supabase, monkeypatch, rows)
    r = client.get("/api/scope-group/%s/workbook" % GID)
    assert r.status_code == 200, r.text
    j = r.json()
    assert [t["label"] for t in j["tabs"]] == ["Lab 1", "Breakroom"]
    assert j["filename"] == "Acme_HQ_scopes_estimate.xlsx"
    token = j["xlsx_download_url"].rsplit("/", 1)[-1]
    entry = main._FILE_CACHE[token]
    assert entry["filename"] == j["filename"]
    assert _load(entry["content"]).sheetnames[:2] == ["Lab 1", "Breakroom"]


def test_the_route_never_puts_a_test_scope_beside_real_ones(fake_supabase, monkeypatch):
    _store(fake_supabase, monkeypatch, [
        _row(MEMBERS[0], "2026-10-01"), _row(MEMBERS[2], "2026-10-02", is_test=True)])
    monkeypatch.setattr(ew, "fill_group_estimate", lambda members: b"xlsx")
    j = client.get("/api/scope-group/%s/workbook" % GID).json()
    assert [t["label"] for t in j["tabs"]] == ["Lab 1"]


def test_the_route_refuses_in_sentences(fake_supabase, monkeypatch):
    bad = copy.deepcopy(MEMBERS[1])
    bad["data"]["extras"] = [{"label": "x"}]
    _store(fake_supabase, monkeypatch, [_row(MEMBERS[0], "2026-10-01"), _row(bad, "2026-10-02")])
    r = client.get("/api/scope-group/%s/workbook" % GID)
    assert r.status_code == 409
    assert r.json()["code"] == "member_not_supported" and "“Lab 2”" in r.json()["error"]
    assert client.get("/api/scope-group/grp_nobody00001/workbook").status_code == 404
    r = client.get("/api/scope-group/bad%20id!/workbook")
    assert r.status_code == 400 and r.json()["code"] == "invalid_scope_group"
