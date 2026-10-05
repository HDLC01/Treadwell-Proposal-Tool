"""The Terms & Conditions pages print ONE letterhead each wherever the renderer breaks them, and a
WORK line that would print nothing but its label is left out.

THE AUDIT (2026-10-02, a Direct Epoxy proposal on staging, the 4-page LibreOffice PDF): clause 9
broke mid-sentence ("...a reasonable opportunity" / "to" / "inspect the alleged..."), clause 18
printed in three pieces round a one-inch hole, the last page's letterhead sat at y=42..833pt with
no footer and no red bar, the Terms printed in Liberation Sans, and page 1's WORK box printed a
bare "Texture:" line and a bare "Notes:" bullet.

THE CAUSE. Each Terms page's full-page letterhead PNG is anchored to an empty paragraph Kyle put at
the top of each page in Word, and he split clauses by hand where Word broke the page, so that each
host landed on a page top. Only a renderer that breaks every line where Word does keeps that, and
LibreOffice does not -- with Caladea (Cambria's metric twin) installed either, measured in the
production image. So `proposal_writer._rebuild_terms_pages` makes the Terms a section of their
own whose HEADER carries the letterhead, page-relative and behind the text, takes the per-page art
and Kyle's padding out, and rejoins the hand-split clauses (see its note).

WHAT THIS FILE HOLDS, on the real fill path and the real templates:
  * the structure: two sections; page 1's setup, art and footer untouched; the Terms section's
    header draws Kyle's own Terms art and nothing else; its footer is empty; no full-page art is
    left anchored to a paragraph; no clause is left split; no word is lost or added;
  * the package is one Word can open: unique drawing ids across every part, every relationship
    resolves, the section and paragraph properties in schema order, no working mark leaks;
  * ONE PLAN: `render_adjustments` -- what /api/proposal-template hands the editor -- applied to
    the template's own block texts reproduces exactly the Terms paragraphs the document prints;
  * edits: an override still lands on its clause, and a line the estimator typed between the two
    halves of a split keeps both halves apart and itself in place;
  * the WORK lines: a blank Texture / Notes line goes, a filled one prints, Kyle's own words after
    a label always print, a line that heads printing sub-items stays.

The real render (one letterhead per page at y=0..792, the red bar on every page, no holes, the
Terms in Caladea) cannot run in CI: backend/ops/terms_render_proof.py does it in the production
image and prints the table.
"""
import io
import pathlib
import re
import zipfile

import docx
import pytest
from docx.oxml.ns import qn
from fastapi.testclient import TestClient
from lxml import etree

import cover_letter_writer as clw
import docx_merge
import main
import proposal_writer as pw

client = TestClient(main.app)

ROOT = pathlib.Path(__file__).resolve().parents[2]

W = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"
WP = "{http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing}"
A = "{http://schemas.openxmlformats.org/drawingml/2006/main}"
R = "{http://schemas.openxmlformats.org/officeDocument/2006/relationships}"
PR = "{http://schemas.openxmlformats.org/package/2006/relationships}"
MC = "{http://schemas.openxmlformats.org/markup-compatibility/2006}"

# Every template with a Terms section (Budget has none), and its numbered-clause count.
TERMS = [("epoxy", "Direct", 27), ("polish", "Direct", 27), ("combo", "Direct", 27),
         ("polish", "GC", 27), ("epoxy", "GC", 27), ("sealer", "GC", 27), ("gyp", None, 28)]
IDS = ["%s-%s" % (wt, aud) for wt, aud, _n in TERMS]

VALUES = {
    "job_name": "Release Check Warehouse", "project_name": "Release Check Warehouse",
    "city_state": "Olathe, KS", "bid_date_formatted": "9/26/26",
    "site_visit_phrase": "per site visit on 9/26/26", "estimator_name": "Kyle Loseke",
    "system_name": "Treadwell MACRO Flake Single Broadcast", "epoxy_system_name": "Epoxy System",
    "texture": "", "work_notes": "",
    "epoxy_sf": "1,000", "polish_sf": "1,000", "sqft": "1,000", "cove_lf": "50",
    "area_description": "~1,000 sf of epoxy flooring",
    "scope_notes": "Prep and coat.", "schedule_notes": "One week.", "exclusions": "Moving furniture.",
    "base_bid_formatted": "$5,569", "material_tax_formatted": "$11", "total_formatted": "$5,580",
    "lump_sum_formatted": "$5,580", "tax_amount_formatted": "$0", "state_name": "Kansas",
    "base_tax_phrase": "(material sales tax INCLUDED)",
}
SYSTEM = {"prefix": "System:", "name": "Epoxy System", "texture": "", "sqft": "1,000",
          "lf_clause": " and 50 LF of 6\" epoxy cove base"}

# Kyle's hand splits, by the words on either side: the four the audit saw on every template.
SPLITS = (("reasonable opportunity to", "inspect the alleged"),
          ("and filed with", "AAA. The Party filing"),
          ("before the date when", "the filing of a lawsuit"),
          ("in any proceeding between", "the Parties arising out of"))

_SENTENCE_END = re.compile(pw._TERMS_SENTENCE_END)


def _fill(work_type, audience, values=None, **kw):
    v = dict(VALUES)
    v.update(values or {})
    if work_type == "epoxy" and audience == "Direct" and "systems" not in kw:
        kw["systems"] = [dict(SYSTEM, texture=v["texture"])]
    return pw.fill_proposal(work_type=work_type, audience=audience, values=v, **kw)


def _template(work_type, audience):
    return docx.Document(str(pw.pick_template(work_type, audience)))


def _tops(d):
    return [c for c in d.element.body if c.tag == W + "p"]


def _text(p):
    return pw._own_text(p)


def _heading_index(tops):
    return next(i for i, p in enumerate(tops) if _text(p).strip().upper() == "TERMS AND CONDITIONS")


def _break_index(tops):
    """The paragraph page 1's section ends on, in a FILLED document."""
    got = [i for i, p in enumerate(tops) if p.find(W + "pPr/" + W + "sectPr") is not None]
    assert len(got) == 1, got
    return got[0]


def _numbered(p):
    return p.find(W + "pPr/" + W + "numPr") is not None


def _full_page_anchors(root, page=(612 * 12700, 792 * 12700)):
    out = []
    for a in root.iter(WP + "anchor"):
        if a.find(".//" + A + "blip") is None:
            continue
        ext = a.find(WP + "extent")
        if int(ext.get("cx")) >= 0.9 * page[0] and int(ext.get("cy")) >= 0.9 * page[1]:
            out.append(a)
    return out


def _parts(raw):
    with zipfile.ZipFile(io.BytesIO(raw)) as z:
        return {n: z.read(n) for n in z.namelist()}


def _rels_of(parts, part):
    """{rId: (type, target)} of one part, from its .rels."""
    d, b = part.rsplit("/", 1)
    raw = parts.get(d + "/_rels/" + b + ".rels")
    if raw is None:
        return {}
    root = etree.fromstring(raw)
    return {r.get("Id"): (r.get("Type"), r.get("Target"), r.get("TargetMode"))
            for r in root.iter(PR + "Relationship")}


def _media_bytes(parts, part, rid):
    _typ, target, _mode = _rels_of(parts, part)[rid]
    base = part.rsplit("/", 1)[0]
    import posixpath
    return parts[posixpath.normpath(posixpath.join(base, target))]


def _terms_art_bytes(work_type, audience):
    """The bytes of the PNG Kyle anchors on his own Terms pages, read off the template."""
    path = pw.pick_template(work_type, audience)
    d = docx.Document(str(path))
    tops = _tops(d)
    h = _heading_index(tops)
    host = next(p for p in reversed(tops[max(0, h - 3):h + 1]) if _full_page_anchors(p))
    rid = _full_page_anchors(host)[0].find(".//" + A + "blip").get(R + "embed")
    return _media_bytes(_parts(path.read_bytes()), "word/document.xml", rid)


def _words(texts):
    return " ".join(" ".join(texts).split())


# ══ the structure ═══════════════════════════════════════════════════════════════════════════════
@pytest.mark.parametrize("work_type,audience,_n", TERMS, ids=IDS)
def test_the_terms_are_a_section_whose_header_draws_the_letterhead(work_type, audience, _n):
    """Two sections. The second is the Terms: its header draws ONE picture -- Kyle's own Terms art,
    page-relative at (0, 0), behind the text -- its footer draws nothing, and its top margin
    clears the logo. AND NO FULL-PAGE ART IS LEFT ANCHORED TO A PARAGRAPH after the Terms start:
    the old render left three (four on Gyp), and that is what moved with the text."""
    raw = _fill(work_type, audience)
    d = docx.Document(io.BytesIO(raw))
    assert len(d.sections) == 2
    terms = d.sections[1]
    assert not terms.header.is_linked_to_previous and not terms.footer.is_linked_to_previous
    assert terms.top_margin.pt >= pw._TERMS_TOP_MARGIN_TW / 20.0
    hdr = terms.header._element
    anchors = list(hdr.iter(WP + "anchor"))
    assert len(anchors) == 1 and len(list(hdr.iter(W + "drawing"))) == 1
    a = anchors[0]
    assert a.get("behindDoc") == "1"
    assert a.find(WP + "wrapNone") is not None
    for tag in ("positionH", "positionV"):
        pos = a.find(WP + tag)
        assert pos.get("relativeFrom") == "page", tag
        assert pos.find(WP + "posOffset").text == "0", tag
    assert not "".join(t.text or "" for t in hdr.iter(W + "t")).strip()
    rid = a.find(".//" + A + "blip").get(R + "embed")
    parts = _parts(raw)
    hdr_part = next(n for n in parts if n.startswith("word/header") and n.endswith(".xml")
                    and b"<wp:anchor" in parts[n])
    assert _media_bytes(parts, hdr_part, rid) == _terms_art_bytes(work_type, audience)
    ftr = terms.footer._element
    assert not list(ftr.iter(W + "drawing")) and not list(ftr.iter(W + "pict"))
    assert not "".join(t.text or "" for t in ftr.iter(W + "t")).strip()

    tops = _tops(d)
    after = tops[_break_index(tops) + 1:]
    assert _text(after[0]).strip() and _heading_index(after) <= 1, "the Terms do not start the section"
    left = [x for p in after for x in _full_page_anchors(p)]
    assert not left, "full-page art is still anchored to a Terms paragraph"


@pytest.mark.parametrize("work_type,audience,_n", TERMS, ids=IDS)
def test_page_one_keeps_its_own_setup_art_and_footer(work_type, audience, _n):
    """The section page 1 ends on is a COPY of the template's own section properties (page size,
    margins, columns, grid; Gyp's real footer too), and page 1's full-page art is Kyle's anchor,
    unchanged."""
    tpl = _template(work_type, audience)
    want = tpl.element.body.find(W + "sectPr")
    raw = _fill(work_type, audience)
    d = docx.Document(io.BytesIO(raw))
    tops = _tops(d)
    got = tops[_break_index(tops)].find(W + "pPr/" + W + "sectPr")
    for tag in ("pgSz", "pgMar", "cols", "docGrid"):
        assert dict(got.find(W + tag).attrib) == dict(want.find(W + tag).attrib), tag
    assert got.find(W + "type") is None
    # Page 1 keeps its letterhead where Kyle put it.
    first_art = _full_page_anchors(_tops(tpl)[0])
    assert len(first_art) == 1
    mine = _full_page_anchors(tops[0])
    assert len(mine) == 1
    assert etree.tostring(mine[0], method="c14n") == etree.tostring(first_art[0], method="c14n")
    # Gyp's page-1 footer is Kyle's (the EMF letterhead), still on page 1 and only there.
    kyle_ftr = want.findall(W + "footerReference")
    if kyle_ftr:
        rid = got.find(W + "footerReference").get(R + "id")
        target = _rels_of(_parts(raw), "word/document.xml")[rid][1]
        assert b"<wp:anchor" in _parts(raw)["word/" + target]


@pytest.mark.parametrize("work_type,audience,n_clauses", TERMS, ids=IDS)
def test_no_clause_is_left_split_and_no_word_is_lost(work_type, audience, n_clauses):
    """Every hand split is rejoined, clause numbers untouched, and the Terms read word for word
    what Kyle's template reads -- the joins move words, they never add or drop one."""
    tpl_tops = _tops(_template(work_type, audience))
    h = _heading_index(tpl_tops)
    want_words = _words(_text(p) for p in tpl_tops[h:])

    d = docx.Document(io.BytesIO(_fill(work_type, audience)))
    tops = _tops(d)
    terms = tops[_break_index(tops) + 1:]
    assert _words(_text(p) for p in terms) == want_words
    assert sum(1 for p in terms if _numbered(p)) == n_clauses

    first = next(i for i, p in enumerate(terms) if _numbered(p))
    printing = [p for p in terms[first:] if _text(p).strip()]
    for a, b in zip(printing, printing[1:]):
        assert _SENTENCE_END.search(_text(a).strip()) or _numbered(b), (
            "still split: %r / %r" % (_text(a)[-40:], _text(b)[:40]))
    joined = "\n".join(_text(p) for p in terms)
    for left, right in SPLITS:
        assert re.search(re.escape(left) + r" +" + re.escape(right), joined), (left, right)
    # Nothing Kyle wrote as its own paragraph was folded into another: clause 9's disclaimer.
    assert any(_text(p).startswith("TO THE EXTENT PERMITTED BY APPLICABLE LAW") for p in terms)


@pytest.mark.parametrize("work_type,audience,_n", TERMS, ids=IDS)
def test_no_hole_is_left_in_the_terms(work_type, audience, _n):
    """Kyle's padding round each page top is gone; his own spacing under the heading stays."""
    d = docx.Document(io.BytesIO(_fill(work_type, audience)))
    tops = _tops(d)
    terms = tops[_break_index(tops) + 1:]
    run = longest = 0
    for p in terms:
        run = run + 1 if not _text(p).strip() else 0
        longest = max(longest, run)
    assert longest <= 3, "a run of %d empty lines is left in the Terms" % longest
    assert _text(terms[-1]).strip(), "the Terms end on an empty line"


# ══ the package ═════════════════════════════════════════════════════════════════════════════════
_SECTPR_ORDER = ["headerReference", "footerReference", "footnotePr", "endnotePr", "type", "pgSz",
                 "pgMar", "paperSrc", "pgBorders", "lnNumType", "pgNumType", "cols", "formProt",
                 "vAlign", "noEndnote", "titlePg", "textDirection", "bidi", "rtlGutter", "docGrid",
                 "printerSettings", "sectPrChange"]
_PPR_ORDER = ["pStyle", "keepNext", "keepLines", "pageBreakBefore", "framePr", "widowControl",
              "numPr", "suppressLineNumbers", "pBdr", "shd", "tabs", "suppressAutoHyphens",
              "kinsoku", "wordWrap", "overflowPunct", "topLinePunct", "autoSpaceDE", "autoSpaceDN",
              "bidi", "adjustRightInd", "snapToGrid", "spacing", "ind", "contextualSpacing",
              "mirrorIndents", "suppressOverlap", "jc", "textDirection", "textAlignment",
              "textboxTightWrap", "outlineLvl", "divId", "cnfStyle", "rPr", "sectPr", "pPrChange"]


def _in_order(el, order):
    seen = [etree.QName(c).localname for c in el if isinstance(c.tag, str)]
    rank = [order.index(n) for n in seen if n in order]
    if seen and seen[0] in ("headerReference", "footerReference"):
        rank = [0 if r == 1 else r for r in rank]          # the two reference kinds interleave
    return rank == sorted(rank) and all(n in order for n in seen), seen


@pytest.mark.parametrize("work_type,audience,_n", TERMS, ids=IDS)
def test_the_package_is_one_word_can_open(work_type, audience, _n):
    """What can be checked without Word: every XML part parses; drawing ids are unique across
    EVERY part (a duplicate is what makes Word offer a repair); every relationship id a part uses
    resolves in that part's .rels; every header and footer is declared in [Content_Types].xml;
    both sections' properties and the paragraph that carries the break are in schema order; and
    none of the fill's working marks reaches the file."""
    raw = _fill(work_type, audience)
    parts = _parts(raw)
    docpr = []
    for name, data in parts.items():
        if not name.endswith(".xml") and not name.endswith(".rels"):
            continue
        root = etree.fromstring(data)
        if not name.startswith("word/") or name.endswith(".rels"):
            continue
        docpr += [e.get("id") for e in root.iter(WP + "docPr")]
        rels = _rels_of(parts, name)
        for el in root.iter():
            for attr in (R + "embed", R + "id", R + "link"):
                v = el.get(attr)
                if v is not None:
                    assert v in rels, "%s points at %s, which its .rels does not declare" % (name, v)
        for attr in ("twBareLine", "twOptionsHeading", "twTypedLine", "twLineProps"):
            assert attr.encode() not in data, (name, attr)
    assert len(docpr) == len(set(docpr)), "duplicate wp:docPr id: %s" % sorted(docpr)
    # The header's letterhead has an id of its own, past every id Kyle's file uses -- not the id of
    # the anchor it was copied from, which would collide if that anchor ever stayed in the body.
    hdr_ids = [e.get("id") for n, data in parts.items() if re.fullmatch(r"word/header\d+\.xml", n)
               for e in etree.fromstring(data).iter(WP + "docPr")]
    tpl_ids = [int(e.get("id")) for e in etree.fromstring(_parts(
        pw.pick_template(work_type, audience).read_bytes())["word/document.xml"]).iter(WP + "docPr")]
    assert len(hdr_ids) == 1 and int(hdr_ids[0]) > max(tpl_ids), (hdr_ids, max(tpl_ids))
    ct = parts["[Content_Types].xml"].decode("utf-8")
    for name in parts:
        if re.fullmatch(r"word/(header|footer)\d+\.xml", name):
            assert 'PartName="/%s"' % name in ct, name

    d = docx.Document(io.BytesIO(raw))
    tops = _tops(d)
    brk = tops[_break_index(tops)]
    for sect in (brk.find(W + "pPr/" + W + "sectPr"), d.element.body.find(W + "sectPr")):
        ok, seen = _in_order(sect, _SECTPR_ORDER)
        assert ok, seen
    for p in (brk, tops[_break_index(tops) + 1]):
        ok, seen = _in_order(p.find(W + "pPr"), _PPR_ORDER)
        assert ok, seen
    # And python-docx, which is what every later reader of these bytes uses, reads both sections.
    assert [s.start_type for s in d.sections] == [s.start_type for s in d.sections[:1]] * 2


# ══ ONE plan: what the editor is told is what the render does ═══════════════════════════════════
def _apply_plan(work_type, audience, overrides=None):
    """`render_adjustments` applied to the template's own block texts, the way the editor applies
    it: drop the blank, then join each pair across what is gone. Returns the Terms paragraphs."""
    d = _template(work_type, audience)
    plan = pw.render_adjustments(d)["terms"]
    blocks = {i: (t, p) for i, _k, p, _b, t, x in pw.iter_editable_blocks(d) if x is None}
    texts = {i: t for i, (t, _p) in blocks.items()}
    texts.update(overrides or {})
    numbered = {i for i, (_t, p) in blocks.items() if _numbered(p)}
    seq = [i for i in sorted(blocks) if i >= plan["first_id"]]
    drop = set(plan["drop_if_blank_ids"])
    seq = [i for i in seq if not (i in drop and not texts[i].strip())]
    end = re.compile(plan["sentence_end"])
    into_of = {}
    for a, b in plan["join_pairs"]:
        into = into_of.get(a, a)
        if into not in seq or b not in seq or seq.index(b) != seq.index(into) + 1:
            continue
        ta, tb = texts[into], texts[b]
        if not ta.strip() or not tb.strip() or end.search(ta.strip()) or b in numbered:
            continue
        sep = "" if (ta[-1:].isspace() or tb[:1].isspace()) else " "
        texts[into] = ta + sep + tb
        seq.remove(b)
        into_of[b] = into
    return [texts[i] for i in seq]


@pytest.mark.parametrize("work_type,audience,_n", TERMS, ids=IDS)
def test_the_plan_the_editor_gets_reproduces_the_printed_terms(work_type, audience, _n):
    d = docx.Document(io.BytesIO(_fill(work_type, audience)))
    tops = _tops(d)
    printed = [_text(p) for p in tops[_break_index(tops) + 1:]]
    assert printed == _apply_plan(work_type, audience)


@pytest.mark.parametrize("work_type,audience,_n", TERMS, ids=IDS)
def test_the_plan_names_the_template_ids_and_kyles_art(work_type, audience, _n):
    d = _template(work_type, audience)
    adj = pw.render_adjustments(d)
    assert adj["version"] == pw.RENDER_ADJUSTMENTS_VERSION
    t = adj["terms"]
    tops = _tops(d)
    h = _heading_index(tops)
    # Top-level body paragraphs come first in the walk, so a Terms block id IS its body index.
    assert t["first_id"] in (h - 2, h - 1, h)
    assert t["section_break_after_id"] == t["first_id"] - 1
    assert all(_full_page_anchors(tops[i]) for i in t["art_host_ids"])
    assert len(t["art_host_ids"]) == (4 if work_type == "gyp" else 3)
    assert all(not _text(tops[i]).strip() or i in t["art_host_ids"] for i in t["drop_if_blank_ids"])
    assert len(t["join_pairs"]) == 4
    for (a, b), (left, right) in zip(t["join_pairs"], SPLITS):
        assert _text(tops[a]).rstrip().endswith(left.split()[-1]), (a, left)
    assert t["art"]["name"] and t["art"]["w_pt"] == 612.0 and abs(t["art"]["h_pt"] - 791.0) < 1
    assert t["page"]["margin"]["top"] == pw._TERMS_TOP_MARGIN_TW / 20.0
    re.compile(t["sentence_end"])


@pytest.mark.parametrize("work_type,audience", sorted(pw.TEMPLATE_PICKER, key=lambda k: (k[0], str(k[1]))),
                         ids=lambda v: str(v))
def test_the_editor_is_served_the_writers_plan(work_type, audience):
    r = client.get("/api/proposal-template", params={"work_type": work_type,
                                                     "audience": audience or "Direct"})
    assert r.status_code == 200, r.text
    assert r.json()["render_adjustments"] == pw.template_render_adjustments(work_type, audience or "Direct")


def test_budget_has_no_terms_and_is_left_alone():
    adj = pw.template_render_adjustments("budget", "Direct")
    assert adj["terms"] is None and adj["lines"] == []
    d = docx.Document(io.BytesIO(pw.fill_proposal(work_type="budget", audience="Direct",
                                                  values=dict(VALUES))))
    assert len(d.sections) == 1


# ══ edits ══════════════════════════════════════════════════════════════════════════════════════
def _ids(work_type="epoxy", audience="Direct"):
    d = _template(work_type, audience)
    return d, pw.render_adjustments(d)["terms"]


def test_an_edit_still_lands_on_its_clause_and_on_a_continuation():
    d, plan = _ids()
    clause10 = next(i for i, _k, p, _b, t, x in pw.iter_editable_blocks(d)
                    if x is None and t.startswith("Mutual Waiver"))
    cont = plan["join_pairs"][0][1]                     # "inspect the alleged improper..."
    raw = _fill("epoxy", "Direct", paragraph_overrides=[
        {"id": clause10, "text": "Mutual Waiver of Consequential Damages.  Reworded by Kyle."},
        {"id": cont, "text": "inspect the work first."}])
    tops = _tops(docx.Document(io.BytesIO(raw)))
    terms = [_text(p) for p in tops[_break_index(tops) + 1:]]
    assert "Mutual Waiver of Consequential Damages.  Reworded by Kyle." in terms
    assert any(t.endswith("reasonable opportunity to inspect the work first.") for t in terms)


def test_a_line_typed_between_two_halves_stays_and_keeps_them_apart():
    """The estimator types into one of Kyle's padding lines inside clause 18's split. It prints,
    in place; the two halves round it stay two paragraphs; every other split is still joined. And
    the plan the editor gets says the same."""
    d, plan = _ids()
    a, b = plan["join_pairs"][2]                       # "...before the date when" / "the filing..."
    typed = next(i for i in plan["drop_if_blank_ids"] if a < i < b)
    raw = _fill("epoxy", "Direct", paragraph_overrides=[{"id": typed, "text": "Typed by Kyle."}])
    tops = _tops(docx.Document(io.BytesIO(raw)))
    terms = [_text(p) for p in tops[_break_index(tops) + 1:]]
    k = terms.index("Typed by Kyle.")
    assert terms[k - 1].endswith("before the date when ") and terms[k + 1].startswith("the filing")
    assert any(t.endswith("reasonable opportunity to inspect the alleged improper or defective work "
                          "prior to any removal or alteration by Customer. Instead of repair or "
                          "replacement, Treadwell may, in its sole discretion, fully and completely "
                          "satisfy its obligations under this warranty by refunding to Customer the "
                          "Contract Price paid by Customer.") for t in terms)
    assert terms == _apply_plan("epoxy", "Direct", {typed: "Typed by Kyle."})


def test_a_clause_ended_by_hand_is_not_joined():
    """Give the split's first half a full stop and it is a sentence: the two print apart."""
    d, plan = _ids()
    a, _b = plan["join_pairs"][0]
    head = next(t for i, _k, _p, _bl, t, x in pw.iter_editable_blocks(d) if i == a)
    raw = _fill("epoxy", "Direct", paragraph_overrides=[{"id": a, "text": head + " it."}])
    tops = _tops(docx.Document(io.BytesIO(raw)))
    terms = [_text(p) for p in tops[_break_index(tops) + 1:]]
    k = next(i for i, t in enumerate(terms) if t.endswith("opportunity to it."))
    assert terms[k + 1].startswith("inspect the alleged")


# ══ the cover letter in front of a two-section proposal ═════════════════════════════════════════
def test_a_merged_cover_letter_cannot_take_the_letterheads_drawing_id():
    """The Terms letterhead is now a drawing in a HEADER part, numbered after the body's last id.
    `docx_merge` shifts the letter's ids past every part of the host, not only its body: with a
    letter whose drawing id would land exactly on the header's, the ids stay unique."""
    proposal = _fill("epoxy", "Direct")
    letter = clw.fill_cover_letter(work_type="epoxy", audience="Direct", values=dict(VALUES))
    parts = _parts(letter)
    xml = re.sub(rb'(<wp:docPr [^>]*\bid=")\d+(")', rb"\g<1>0\g<2>", parts["word/document.xml"])
    assert xml != parts["word/document.xml"]
    out = io.BytesIO()
    with zipfile.ZipFile(io.BytesIO(letter)) as zin, zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as zo:
        for item in zin.infolist():
            zo.writestr(item, xml if item.filename == "word/document.xml" else zin.read(item.filename))
    merged = docx_merge.prepend_cover_letter(proposal, out.getvalue())
    ids = []
    for name, data in _parts(merged).items():
        if name.startswith("word/") and name.endswith(".xml"):
            ids += [e.get("id") for e in etree.fromstring(data).iter(WP + "docPr")]
    assert len(ids) == len(set(ids)), sorted(ids, key=int)
    n_letter = len(docx.Document(io.BytesIO(letter)).sections)
    assert len(docx.Document(io.BytesIO(merged)).sections) == n_letter + 2


# ══ the fonts ══════════════════════════════════════════════════════════════════════════════════
def test_the_terms_ask_for_cambria_and_the_image_installs_its_metric_twin():
    """Every Terms run names Cambria, and the REGARDS name inherits the theme's minor font, which is
    Cambria too. The container has no Cambria; LibreOffice maps it to Caladea when Caladea is
    installed (fontconfig's metric aliases), and to Liberation Sans when it is not -- which is what
    the audit's PDF printed. So the Dockerfile's LibreOffice layer must keep fonts-crosextra-caladea."""
    d = _template("epoxy", "Direct")
    tops = _tops(d)
    fonts = {rf.get(W + "ascii") for p in tops[_heading_index(tops):] for rf in p.iter(W + "rFonts")}
    assert fonts - {None} <= {"Cambria", "MS Mincho"} and "Cambria" in fonts, fonts
    with zipfile.ZipFile(str(pw.pick_template("epoxy", "Direct"))) as z:
        theme = z.read("word/theme/theme1.xml").decode("utf-8")
    assert re.search(r'<a:minorFont>\s*<a:latin typeface="Cambria"', theme)
    code = re.sub(r"\\\n", " ", "\n".join(
        ln for ln in (ROOT / "Dockerfile").read_text(encoding="utf-8").splitlines()
        if not ln.lstrip().startswith("#")))
    layer = next(ln for ln in code.splitlines() if "libreoffice-writer" in ln)
    for pkg in ("libreoffice-writer", "fonts-crosextra-carlito", "fonts-crosextra-caladea",
                "fonts-liberation"):
        assert re.search(r"(?<![\w-])" + re.escape(pkg) + r"(?![\w-])", layer), pkg


# ══ the WORK lines that would print only their label ═══════════════════════════════════════════
LABEL_TEMPLATES = [("epoxy", "Direct", ("Texture:", "Notes:")),
                   ("polish", "Direct", ("Notes:",)),
                   ("combo", "Direct", ("Texture:", "Notes:"))]


def _box_lines(raw, choice_only=False):
    """Every text-box paragraph's text, BOTH copies (mc:Choice and the VML mc:Fallback) -- or the
    mc:Choice copy alone, which is the only one an editor override reaches (the walk skips the
    Fallback twin, as it always has)."""
    d = docx.Document(io.BytesIO(raw))
    return [pw._own_text(p) for t in d.element.body.iter(W + "txbxContent") for p in t.iter(W + "p")
            if not (choice_only and pw._is_fallback_paragraph(p))]


def _bare(lines, label):
    return [t for t in lines if re.fullmatch(r"\s*" + re.escape(label) + r"[\s.]*", t)]


@pytest.mark.parametrize("work_type,audience,labels", LABEL_TEMPLATES, ids=lambda v: str(v))
def test_a_blank_texture_or_note_leaves_its_line_out(work_type, audience, labels):
    lines = _box_lines(_fill(work_type, audience))
    for label in labels:
        assert not _bare(lines, label), label
        assert not any(t.strip().startswith(label) for t in lines), label
    # The lines round them are still there.
    assert any(t.strip().startswith("Exclusions:") for t in lines)


@pytest.mark.parametrize("work_type,audience,labels", LABEL_TEMPLATES, ids=lambda v: str(v))
def test_a_filled_texture_or_note_prints(work_type, audience, labels):
    lines = _box_lines(_fill(work_type, audience, {"texture": "Orange Peel",
                                                   "work_notes": "Owner moves the racking"}))
    if "Texture:" in labels:
        assert sum(1 for t in lines if "Texture:" in t and "Orange Peel" in t) == 2   # both copies
    assert sum(1 for t in lines if t.strip().startswith("Notes:") and "Owner moves the racking" in t) == 2


def test_each_priced_system_loses_its_own_bare_texture_row():
    systems = [dict(SYSTEM, prefix="Option 1:", name="MACRO"), dict(SYSTEM, prefix="Option 2:", name="Quartz")]
    lines = _box_lines(_fill("epoxy", "Direct", systems=systems))
    assert not any("Texture:" in t for t in lines)
    assert sum(1 for t in lines if t.strip().startswith("Area:")) == 4                    # 2 x 2 copies
    systems[1]["texture"] = "Smooth"
    lines = _box_lines(_fill("epoxy", "Direct", systems=systems))
    assert [t.split("Texture:")[1].strip() for t in lines if "Texture:" in t] == ["Smooth", "Smooth"]


def test_a_rewritten_texture_row_prints_and_a_relabelled_bare_one_does_not():
    kept = _box_lines(_fill("epoxy", "Direct", systems=[dict(SYSTEM, texture_line="Texture:  Owner's pick")]))
    assert sum(1 for t in kept if t.strip() == "Texture:  Owner's pick") == 2
    gone = _box_lines(_fill("epoxy", "Direct", systems=[dict(SYSTEM, texture_label="Finish:")]))
    assert not any(t.strip().startswith(("Finish:", "Texture:")) for t in gone)


def test_a_note_typed_over_the_line_prints_even_with_the_field_blank():
    d = _template("epoxy", "Direct")
    nid = next(i for i, _k, _p, _b, t, _x in pw.iter_editable_blocks(d) if t.startswith("Notes:"))
    lines = _box_lines(_fill("epoxy", "Direct", paragraph_overrides=[
        {"id": nid, "text": "Notes:  Bring the long ladder"}]), choice_only=True)
    assert sum(1 for t in lines if t.strip() == "Notes:  Bring the long ladder") == 1


def test_gyps_notes_line_stays_while_it_heads_kyles_note_and_goes_when_that_is_emptied():
    lines = _box_lines(_fill("gyp", None))
    assert len(_bare(lines, "Notes:")) == 2                       # it heads "Floor Leveling is NOT..."
    # A heading of Kyle's that ends in a colon is not one of these lines: only the template's own
    # "Label: {{token}}" rows are candidates.
    assert len(_bare(lines, "System & Scope:")) == 2
    d = _template("gyp", None)
    sub = next(i for i, _k, _p, _b, t, _x in pw.iter_editable_blocks(d)
               if t.startswith("Floor Leveling is NOT included"))
    lines = _box_lines(_fill("gyp", None, paragraph_overrides=[{"id": sub, "text": ""}]),
                       choice_only=True)
    assert not _bare(lines, "Notes:")
    lines = _box_lines(_fill("gyp", None, {"work_notes": "Lift on site"}))
    assert sum(1 for t in lines if t.strip() == "Notes: Lift on site") == 2


@pytest.mark.parametrize("work_type", ["polish", "epoxy", "sealer"])
def test_kyles_own_words_after_a_gc_label_always_print(work_type):
    """The GC forms' Texture / Notes lines carry Kyle's words after the label ("Light or Medium or
    Heavy...", "xx"), so they are never candidates, and print as they always did."""
    adj = pw.template_render_adjustments(work_type, "GC")
    assert adj["lines"] == []
    lines = _box_lines(_fill(work_type, "GC"))
    assert sum(1 for t in lines if t.strip().startswith("Notes:") and t.strip() != "Notes:") >= 2
    if work_type == "epoxy":
        assert sum(1 for t in lines if "Light or Medium or Heavy" in t) == 2


def test_the_editor_is_told_which_lines_go_and_on_what():
    got = {(wt, aud): pw.template_render_adjustments(wt, aud)["lines"]
           for wt, aud in (("epoxy", "Direct"), ("polish", "Direct"), ("combo", "Direct"), ("gyp", None))}
    ep = {ln["token"]: ln for ln in got[("epoxy", "Direct")]}
    assert ep["system.texture"]["in_block"] == "system" and ep["system.texture"]["line_key"] == "texture_line"
    assert ep["system.texture"]["label"] == "Texture:"
    assert ep["work_notes"]["in_block"] is None and ep["work_notes"]["label"] == "Notes:"
    assert [ln["token"] for ln in got[("polish", "Direct")]] == ["work_notes"]
    assert sorted(ln["token"] for ln in got[("combo", "Direct")]) == ["texture", "work_notes"]
    gyp = got[("gyp", None)]
    assert len(gyp) == 1 and gyp[0]["sub_item_ids"], gyp
    d = _template("gyp", None)
    texts = {i: t for i, _k, _p, _b, t, _x in pw.iter_editable_blocks(d)}
    assert texts[gyp[0]["sub_item_ids"][0]].startswith("Floor Leveling is NOT included")
    # ...and every id names the "Label: {{token}}" paragraph it describes.
    for (wt, aud), lines in got.items():
        texts = {i: t for i, _k, _p, _b, t, _x in pw.iter_editable_blocks(_template(wt, aud))}
        for ln in lines:
            assert re.match(r"\s*" + re.escape(ln["label"]) + r"\s*\{\{\s*" + re.escape(ln["token"]),
                            texts[ln["id"]]), (wt, ln)
