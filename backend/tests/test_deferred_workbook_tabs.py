"""Estimate Review stops fetching four of the sixteen worksheets before it opens.

THE COST IT REMOVES. `init()` fetched every tab in `estimate_sheet_5.7.xlsx` and waited on all
sixteen before the grid was usable. Measured off the shipped template by serialising exactly what
`/api/sheet` returns and gzipping it the way the middleware does: 6,724,323 bytes of JSON,
493,524 gzipped, over sixteen requests. proposals.wetreadwell.com speaks HTTP/1.1, so six
connections per host means those sixteen requests are THREE waves of round trips; twelve are two.
For an estimator ~250ms from the VPS that wave is the whole point.

WHY FOUR AND NOT ELEVEN. The four are 'Epoxy blank', 'Leveling', 'Specs+Dwgs+Addn' and
'Unit Layouts'. Three separate things have to be true of a tab before the screen can open without
it, and this module checks all three against the shipped artefacts rather than against the
comment above the list:

  1. NO FORMULA ANYWHERE REACHES IT. Checked here by walking every formula in the real workbook
     and collecting the sheet names they cite. This is the assertion that fails the day Kyle
     writes his first `='Epoxy blank'!D85` -- and it fails BEFORE a bid goes out wrong, which is
     the only reason it is derived from the file instead of typed into it.
  2. IT CARRIES NO PRICED ROLE. `roleFor()` answers "other" for all four, so `pricedTabs()`,
     `renderBidOptions()` and `updateTotalBarFromHF()` never read them. Checked by lifting the
     shipped `BASE_ROLE` in the harness, not by retyping it.
  3. IT CANNOT BE THE OPENING TAB. `defaultBaseSheet()` answers Epoxy, Polish or the gyp base.

WHAT DEFERRAL DOES NOT EXCUSE. 'Epoxy blank' and 'Leveling' are written to at load --
`applyJobFlags`, `applyMarkupRates` and `applyRemodelRateOverride` all stamp cells on them. Those
writes land in an engine sheet that is still empty, and the late `setSheetContent` would wipe
them. Every one of those writers records its write in `cellValues` first, so
`loadDeferredIntoEngine` replays this sheet's own keys AFTER the load. The order is the assertion
(see `test_the_late_load_replays_after_it_loads_not_before`); a replay that runs first is
indistinguishable from no replay at all, and the symptom would be a tab quietly charging Kyle's
10% remodel placeholder next to a correct base bid.
"""
import json
import pathlib
import re
import shutil
import subprocess

import pytest

REPO = pathlib.Path(__file__).resolve().parents[2]
FRONTEND = REPO / "frontend"
TEMPLATE = REPO / "backend" / "templates" / "estimate_sheet_5.7.xlsx"
HARNESS = pathlib.Path(__file__).resolve().parent / "js" / "deferred-tabs-harness.js"

needs_node = pytest.mark.skipif(shutil.which("node") is None, reason="node is not installed")


@pytest.fixture(scope="module")
def result():
    proc = subprocess.run(["node", str(HARNESS), str(FRONTEND)],
                          capture_output=True, text=True, encoding="utf-8", timeout=120)
    assert proc.returncode == 0, proc.stderr
    return json.loads(proc.stdout.strip().splitlines()[-1])


@pytest.fixture(scope="module")
def er_src():
    return (FRONTEND / "js" / "estimate-review.js").read_text(encoding="utf-8")


@pytest.fixture(scope="module")
def deferred_names(er_src):
    """The list as the browser has it, read out of the shipped file."""
    m = re.search(r"^const DEFERRABLE_TABS = new Set\((\[[^\]]*\])\);$", er_src, re.M)
    assert m, "DEFERRABLE_TABS is gone from estimate-review.js — rewrite this module"
    return json.loads(m.group(1))


# ── 1. the workbook itself ───────────────────────────────────────────────────

@pytest.fixture(scope="module")
def cross_sheet_citations():
    """{sheet name: the sheets whose formulas cite it}, read out of Kyle's real template.

    Both spellings a formula can use: `'Epoxy blank'!D85` for a name with a space and `Leveling!B77`
    for one without. The bare form is anchored so `Seal!` inside `'Seal (+Jnts)'!` is not counted
    twice and `Gyp (FR)` is not matched by a longer identifier ending in those letters.
    """
    import openpyxl

    wb = openpyxl.load_workbook(TEMPLATE, data_only=False)
    names = wb.sheetnames
    cited = {n: set() for n in names}
    for ws in wb.worksheets:
        for row in ws.iter_rows():
            for cell in row:
                formula = cell.value
                if not isinstance(formula, str) or not formula.startswith("="):
                    continue
                for n in names:
                    if n == ws.title:
                        continue
                    if "'%s'!" % n in formula:
                        cited[n].add(ws.title)
                    elif re.search(r"(?<![A-Za-z0-9_'])%s!" % re.escape(n), formula):
                        cited[n].add(ws.title)
    return cited


def test_the_deferred_tabs_are_real_tabs(deferred_names, cross_sheet_citations):
    """A typo here would defer nothing and look exactly like a working optimisation."""
    assert set(deferred_names) <= set(cross_sheet_citations), (
        "DEFERRABLE_TABS names a sheet the workbook does not have: %s"
        % sorted(set(deferred_names) - set(cross_sheet_citations)))
    assert len(deferred_names) == 4, deferred_names


def test_no_formula_in_the_workbook_reaches_a_deferred_tab(deferred_names, cross_sheet_citations):
    """The load-bearing one. A sheet the engine is never asked about is a sheet it does not need.

    If this goes red, the fix is to take that name OUT of DEFERRABLE_TABS — not to loosen the
    test. A formula pointing at a sheet HyperFormula has not been given resolves to an error or a
    zero, and the estimator sees a price that is simply wrong rather than a page that is broken.
    """
    reached = {n: sorted(cross_sheet_citations[n]) for n in deferred_names
               if cross_sheet_citations[n]}
    assert reached == {}, (
        "a deferred tab is cited by a live formula: %s" % reached)


def test_every_tab_that_is_cited_is_still_loaded_up_front(deferred_names, cross_sheet_citations):
    """The other direction, so the list cannot grow by accident.

    Six sheets in the template are cited by nobody — the four deferred ones plus
    'Gyp (GWorx SC190)' and 'Gyp (FR)'. Those two stay eager because they are PRICED gyp
    variants, which is the second condition, and this test says so out loud so the next person
    to read the citation map does not conclude the list is simply short by two.
    """
    uncited = {n for n, who in cross_sheet_citations.items() if not who}
    assert set(deferred_names) <= uncited
    assert uncited - set(deferred_names) == {"Gyp (GWorx SC190)", "Gyp (FR)"}, sorted(uncited)


# ── 2. the browser's own rules ───────────────────────────────────────────────

@needs_node
def test_no_deferred_tab_carries_a_priced_role(result):
    """Read out of the shipped BASE_ROLE, which is Epoxy, Polish, five gyp variants and two seal
    sheets. A priced tab gets a chip on the bid bar and a total read straight out of the engine."""
    assert result["pricedAmongDeferrable"] == [], result["pricedAmongDeferrable"]


@needs_node
def test_no_deferred_tab_can_be_the_opening_one(result, deferred_names):
    """`defaultBaseSheet()` answers Epoxy, Polish or the gyp base and nothing else, so the tab the
    page paints first is never one that has not been fetched."""
    opening = {result["CANONICAL_SHEET"], "Polish", result["GYP_BASE"]}
    assert opening.isdisjoint(set(deferred_names)), sorted(opening & set(deferred_names))


@needs_node
def test_an_ordinary_draft_defers_all_four(result):
    assert result["plain"] == sorted(["Epoxy blank", "Leveling", "Specs+Dwgs+Addn", "Unit Layouts"])


@needs_node
def test_a_tab_somebody_copied_keeps_its_eager_load(result):
    """init() step 3b rehydrates a copy from `sheetCache[c.source]`. Defer the source and that
    lookup finds nothing, the loop gives up on it as an orphan, and a priced option is left on the
    tab bar with no worksheet behind it — clicking it fetches /api/sheet/Copy1, gets a 404 and
    paints "Failed to load Copy1" where a customer's option should be."""
    assert result["withLevelingCopy"] == ["Epoxy blank", "Specs+Dwgs+Addn", "Unit Layouts"], (
        "a copied source was deferred anyway: %s" % result["withLevelingCopy"])
    assert result["withTwoCopies"] == ["Leveling", "Specs+Dwgs+Addn"], (
        "two copied sources: %s" % result["withTwoCopies"])


# ── 3. the late load ─────────────────────────────────────────────────────────

@needs_node
def test_the_late_load_replays_after_it_loads_not_before(result):
    """`HF.loadSheet` calls setSheetContent, which REPLACES the sheet. Everything
    `applyJobFlags` / `applyMarkupRates` / `applyRemodelRateOverride` stamped in while the tab
    was still empty is inside that replacement, so the replay has to come SECOND. Running it
    first is indistinguishable from not running it at all, and what reaches the estimator is a
    tab quietly holding the template's own 10% remodel placeholder."""
    log = [tuple(row) for row in result["replay"]["log"]]
    assert result["replay"]["returned"] is True
    assert log[0] == ("loadSheet", "Leveling", 2), (
        "the sheet content is no longer loaded first: %r" % (log[:1],))
    assert log[1:] == [
        ("setCellValue", "Leveling!B6", "No"),
        ("setCellValue", "Leveling!D6", "No"),
        ("setCellValue", "Leveling!B77", '=IF(D6="yes",0.07975,0)'),
    ], "the replay is not putting this sheet's saved values back: %r" % (log[1:],)


@needs_node
def test_the_replay_touches_only_the_sheet_that_just_loaded(result):
    """`Epoxy!B6` and `Copy1!B6` are in the same cellValues dict. Replaying either into 'Leveling'
    would write a value onto a cell nobody asked about; replaying them onto their own sheets would
    re-assert edits the estimator may have changed since the page opened."""
    written = [row[1] for row in result["replay"]["log"] if row[0] == "setCellValue"]
    # The EXACT list, not "they all start with Leveling!". The replay passes `name` as the sheet,
    # so a prefix check is true by construction and stays green with the filter deleted -- which
    # is how a mutation run caught this assertion asserting nothing. `Stnd Alts!Z99` in the
    # fixture is the counterexample: unfiltered, it arrives here as `Leveling!Z99`.
    assert written == ["Leveling!B6", "Leveling!D6", "Leveling!B77"], written
    assert result["replay"]["stillDeferred"] == ["Epoxy blank"], (
        "loading one tab dropped the others from the deferred set: %s"
        % result["replay"]["stillDeferred"])


@needs_node
def test_a_tab_that_was_loaded_up_front_is_never_re_loaded(result):
    """setSheetContent replaces a sheet wholesale. Run this on an eager tab — or on a copy, whose
    cells only ever existed client-side — and every edit typed since the page opened is gone."""
    assert result["eagerTabUntouched"] == {"returned": False, "log": []}
    assert result["copyUntouched"] == {"returned": False, "log": []}


@needs_node
def test_the_late_load_happens_once(result):
    """showSheet's fetch branch is skipped once the tab is cached, but a re-render path that
    cleared the cache would call through again — and the second load would discard the first
    load's replayed values."""
    o = result["onceOnly"]
    assert o["first"] is True and o["second"] is False
    assert o["callsAfterFirst"] == o["callsAfterSecond"] == 2, o


@needs_node
def test_an_engine_that_is_not_up_does_not_leave_the_tab_marked_loaded(result):
    """`HF.ready` false means the load cannot happen now, and the honest answer is False. What it
    must not do is claim the work was done while writing nothing."""
    assert result["engineNotReady"] == {"returned": False, "log": []}


# ── 4. the two call sites, out of the shipped source ─────────────────────────
#
# `init()` and `showSheet` both fetch before they get anywhere, so executing them means stubbing
# the network -- which would prove nothing about the ordering that is actually at issue. These
# assert WHERE the calls sit, because position is the whole of what can go wrong with them.

def test_the_deferral_is_actually_applied_to_the_up_front_load(er_src):
    """The Set is inert unless init()'s Promise.all filters on it. A `DEFERRABLE_TABS` that
    nothing consults is a comment, and every test above it would still pass."""
    body = er_src[er_src.index("async function init()"):er_src.index("\nfunction renderTabs()")]
    assign = body.index("deferredTabs = deferrableNow()")
    load = body.index("await Promise.all(sheets")
    assert assign < load, "deferrableNow() is computed after the load it is meant to shorten"
    window = body[load:load + 160]
    assert "deferredTabs.has(name)" in window, (
        "the up-front load no longer filters out the deferred tabs: %r" % window)


def test_the_deferred_tab_reaches_the_engine_before_it_is_rendered(er_src):
    """renderSheet paints formula cells out of HF. Load the tab after the render and the grid
    comes up holding blanks where the template has numbers."""
    body = er_src[er_src.index("async function showSheet("):]
    body = body[:body.index("\n// Map Excel border styles")]
    fetched = body.index("sheetCache[name] = await r.json()")
    engine = body.index("loadDeferredIntoEngine(name)")
    render = body.index("renderSheet(sheetCache[name])")
    assert fetched < engine < render, (
        "the deferred load moved: fetched=%d engine=%d render=%d" % (fetched, engine, render))


@needs_node
def test_a_slow_tab_never_paints_over_the_tab_you_switched_to(result):
    """Deferring four worksheets made showSheet's await reachable by clicking a tab, and with
    it this race. Found by an audit of the merged staging tree on 2026-09-17, not by any PR.

    showSheet sets activeSheet at the top and renders at the bottom. Between them is a fetch,
    measured at 425-925 ms on throttled 4G. Click a deferred tab, switch before it lands, and
    the slow one comes back and paints ITS grid -- while activeSheet, the tab bar and the badge
    all name the tab you switched to.

    THAT IS NOT A COSMETIC MISMATCH. Every structural op reads activeSheet, so right-clicking
    "Delete row 55" on the grid in front of you deletes row 55 from the OTHER sheet -- rekeying
    its cell values and lock overrides and persisting the result. Silent corruption of an
    estimator's priced tab, with the evidence on screen saying it went somewhere else.

    EXECUTED WITH A CONTROLLED FETCH, because the defect lives entirely in the interleaving --
    asserting that the source contains "activeSheet !== name" would pass with the line in the
    wrong place. Mutation-proven: delete the guard and renderedInOrder becomes
    ["Epoxy", "Leveling"] with activeSheet still "Epoxy".
    """
    race = result["sheetSwitchRace"]
    assert race["gridMatchesActiveSheet"], (
        "the grid shows %r while activeSheet is %r -- a structural op would edit the wrong "
        "sheet (render order: %r)"
        % (race["renderedInOrder"][-1:], race["activeSheetAtEnd"], race["renderedInOrder"]))
    assert race["staleNeverPainted"], (
        "the abandoned tab painted anyway: %r" % race["renderedInOrder"])
    assert race["renderedInOrder"] == ["Epoxy"], (
        "the tab actually switched to must be the only thing drawn: %r"
        % race["renderedInOrder"])
