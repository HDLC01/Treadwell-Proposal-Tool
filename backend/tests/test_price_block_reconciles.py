"""The PRICE block follows the estimate sheet's tax answers, adds up, and the screen shows the same one.

THE MODEL (Hanz, 2026-09-25, verbatim): "Remodel Tax should be triggered by remodel tax in the
estimate form. Taxable is where base bid and other options are taxable or not." And: "If remodel and
material sales tax is on then there should be both. If remodel is only the one on then only remodel
tax." And: "if one of the taxes is set to yes then broken out should be the default option."

So the sheet decides WHETHER there is tax, per priced tab (Taxable? → material sales tax, Remodel
Tax? → remodel tax), and the proposal's TAX control decides only the layout:

  * BROKEN OUT: the base line is the pre-tax figure with NO bracket wording; a Material Sales Tax row
    only when taxable; a Remodel Tax row only when remodel; then the Total — and they add up. Hanz's
    own example, Test33: "$6,767 – Epoxy flooring as described above / $72 – Material Sales Tax /
    $6,839 – Total".
  * ONE LINE: the whole bid, and the wording says which taxes are in it — "(Remodel Tax AND material
    sales tax INCLUDED)", "(material sales tax INCLUDED)", "(Remodel Tax INCLUDED)", "(tax exempt)".
    No tax rows and no Total.

ON EVERY TEMPLATE. The GC and Gyp files author their tax rows as plain paragraphs, which used to
print whatever the tax ("$0.00 – Remodel Tax", a $0 Material Sales Tax on an exempt job). The render
now takes out every row the rule does not print, whether the file wraps it in a {{#block}} or not.

THE BID IS TAX-INCLUSIVE — D88 already contains both taxes — so Broken out backs them OUT using the
sheet's own tax cells (never a guessed rate: sales tax compounds through markup). Kyle's 2026-09-08
report is the reason this file adds up printed dollars rather than trusting a helper: "the proposal
says $6,182 and my estimate says $6,307".

WHY IN DOLLARS, ON THE REAL ARTEFACTS. Every document here is filled from a real template through
the real renderer (`_render_documents`, the path Download, Send and the customer's PDF share), and
the printed strings are read back out of the .docx and added up. The screen half runs the page's own
writers in Node (js/price-block-harness.js) over the shapes walked out of the same template files,
and the two are compared figure for figure.
"""
import io
import json
import pathlib
import re
import shutil
import subprocess

import docx
import pytest
from docx import Document
from starlette.requests import Request

import main
import price_rules
import proposal_writer as pw

FRONTEND = pathlib.Path(__file__).resolve().parents[2] / "frontend"
HARNESS = pathlib.Path(__file__).resolve().parent / "js" / "price-block-harness.js"

needs_node = pytest.mark.skipif(shutil.which("node") is None, reason="node is not installed")

_MC = "{http://schemas.openxmlformats.org/markup-compatibility/2006}"

# (total, sales tax cell, remodel tax cell, Taxable?, Remodel Tax?) — the figures always agree with
# the flags, as the sheet's own cells do (sales tax is 0 exactly when Taxable? is No).
FIGURES = {
    "exempt":            (6307.00,   0.00,   0.00, False, False),
    "sales tax only":    (6307.00, 125.00,   0.00, True,  False),    # Kyle's own job
    "remodel tax only":  (6307.00,   0.00, 471.00, False, True),
    "both, to the cent": (6307.50, 125.25, 471.13, True,  True),
    "Test33":            (6839.00,  72.00,   0.00, True,  False),    # Hanz's own example
}
LAYOUTS = ["ONE_LINE", "BROKEN_OUT"]
TEMPLATES = [("epoxy", "Direct"), ("polish", "Direct"), ("combo", "Direct"),
             ("epoxy", "GC"), ("polish", "GC"), ("sealer", "GC"), ("gyp", "Direct")]
PHRASE = {(True, True): "(Remodel Tax AND material sales tax INCLUDED)",
          (True, False): "(material sales tax INCLUDED)",
          (False, True): "(Remodel Tax INCLUDED)",
          (False, False): "(tax exempt)"}


# ── reading the printed price block back off the .docx ───────────────────────
def _lines(docx_bytes):
    """Every non-empty paragraph of the document, once: its OWN text (a paragraph that anchors a
    text box must not report the box's contents) and the mc:Fallback twin dropped."""
    d = Document(io.BytesIO(docx_bytes))
    out = []
    for p in d.element.xpath("//w:p"):
        if any(True for _ in p.iterancestors(f"{_MC}Fallback")):
            continue
        t = pw._own_text(p).strip()
        if t:
            out.append(t)
    return out


def _usd(s):
    """The number in a printed price string. Refuses rather than returning 0.0 — a row read as zero
    is how a broken price block looks balanced."""
    m = re.search(r"\$([\d,]+(?:\.\d+)?)", str(s))
    assert m, f"no dollar figure in {s!r}"
    return float(m.group(1).replace(",", ""))


def _price_rows(docx_bytes):
    """`{base, base_line, material, remodel, total, options, highlights}` — each PRICE row that
    PRINTED, keyed by its own wording. A row the layout takes out is ABSENT, never 0.0.

    `options` is every line under the Options heading, in order, as printed."""
    blob = docx_bytes
    rows = {"options": []}
    lines = _lines(blob)
    in_opts = False
    for line in lines:
        if re.match(r"^Options\b", line):
            in_opts = True
            continue
        if in_opts:
            if re.match(r"^(Pigtail|\*|Excludes|NOTES|Add \$|Treadwell typically)", line):
                in_opts = False
            elif line.startswith("$") or line.startswith("Add ") or line.startswith("Deduct"):
                rows["options"].append(line)
                continue
            else:
                in_opts = False
        if re.search(r"–\s*Material Sales Tax$", line):
            rows.setdefault("material", _usd(line))
        elif re.search(r"–\s*(Kansas\s+)?Remodel Tax$", line):
            rows.setdefault("remodel", _usd(line))
        elif re.search(r"–\s*Total$", line):
            rows.setdefault("total", _usd(line))
        elif "as described above" in line and line.lstrip().startswith("$"):
            rows.setdefault("base", _usd(line))
            rows.setdefault("base_line", line)
    d = Document(io.BytesIO(blob))
    rows["highlights"] = len(list(d.element.body.iter(
        "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}highlight")))
    return rows


def _values(total, sales, remodel, taxable, remodel_on, layout):
    """A payload's `values`, formatted the way the page ships them — through the same `_fmt_usd`
    the document prints with."""
    return {
        "project_name": "Viracor", "job_name": "Viracor", "city_state": "Lees Summit, MO",
        "sqft": "3,000", "epoxy_sf": "3,000", "polish_sf": "3,000", "cove_lf": "0",
        "estimator_name": "Kyle", "bid_date_formatted": "9/8/26", "texture": "OP",
        "system_name": "MACRO", "scope_notes": "scope", "schedule_notes": "~1w",
        "exclusions": "std", "disposal": "d", "state_name": "Kansas",
        "site_visit_phrase": "per site visit on 9/8",
        "total_formatted": main._fmt_usd(total),
        "material_tax_formatted": main._fmt_usd(sales),
        "tax_amount_formatted": main._fmt_usd(remodel),
        "total_label": f"{main._fmt_usd(total)} – Total",
        "area_description": "~3,000 sf of polished concrete flooring",
        # DELIBERATELY THE WRONG BASE: `_generate` must re-derive it by the rule, never trust it.
        "base_bid_formatted": main._fmt_usd(total - sales - remodel),
        "tax_layout": layout, "price_taxable": taxable, "price_remodel_on": remodel_on,
        "gyp_soft_sf": "3,000", "gyp_hard_sf": "0", "gyp_corridor_sf": "0",
        "gyp_soft_thickness": '3/4"', "gyp_hard_thickness": '1"',
        "gyp_corridor_thickness": '3/4"', "mobilizations_line": "1 Mobilization to Site.",
        "work_description": "per plans",
    }


def _render(payload):
    """The real renderer every document goes through (Download, Send, the customer's PDF)."""
    req = Request({"type": "http", "method": "POST", "path": "/t", "headers": [], "query_string": b""})
    return main._render_documents(payload, req, want_estimate=False)["docx"]["content"]


# Two options for the itemised-options check, each off its own tab: one taxable with remodel, one
# remodel only (a seal sheet on an exempt-material job, say). Their figures are their own.
OPTION_ROOMS = [
    {"id": "Epoxy", "name": "Epoxy", "is_base": True,
     "bid": {"total": 6307.5, "sales_tax": 125.25, "remodel": 471.13}},
    {"id": "Copy1", "name": "Epoxy copy", "is_base": False, "show": True, "price_mode": "total",
     "option_desc": "Treadwell 3/16\" Urethane Cement", "system_desc": "Treadwell 3/16\" Urethane Cement",
     "base_total": 6307.5,
     "bid": {"total": 7696, "sales_tax": 96, "remodel": 500, "taxable": True, "remodel_on": True}},
    {"id": "Seal", "name": "Seal", "is_base": False, "show": True, "price_mode": "total",
     "option_desc": "Sealed Concrete", "system_desc": "Sealed Concrete", "base_total": 6307.5,
     "bid": {"total": 1476, "sales_tax": 0, "remodel": 90, "taxable": False, "remodel_on": True}},
]
OPTION_TEMPLATES = [("epoxy", "Direct"), ("polish", "Direct"), ("gyp", "Direct")]

_DOCS: dict = {}


def _need():
    for wt, aud in TEMPLATES:
        for lay in LAYOUTS:
            for fig in FIGURES:
                yield ("block", wt, aud, lay, fig)
    for wt, aud in OPTION_TEMPLATES:
        for lay in LAYOUTS:
            yield ("options", wt, aud, lay, "both, to the cent")


@pytest.fixture
def docs():
    """Every document this module reads, rendered once through the real renderer."""
    if _DOCS:
        return _DOCS
    for kind, wt, aud, lay, fig in _need():
        total, sales, remodel, tx, rm = FIGURES[fig]
        body = {"work_type": wt, "audience": aud,
                "values": _values(total, sales, remodel, tx, rm, lay),
                "remodel": ([{"amount_formatted": main._fmt_usd(remodel)}] if remodel else [])}
        if kind == "options":
            body["rooms"] = OPTION_ROOMS
            body["values"]["base_bid_formatted"] = ""
        _DOCS[(kind, wt, aud, lay, fig)] = _price_rows(_render(body))
    return _DOCS


# ── (1) the document: rows by the sheet, layout by the control, and it adds up ─────────────────
@pytest.mark.parametrize("work_type,audience", TEMPLATES)
@pytest.mark.parametrize("layout", LAYOUTS)
@pytest.mark.parametrize("figure", list(FIGURES))
def test_the_printed_price_block_follows_the_sheet_and_adds_up(docs, work_type, audience, layout, figure):
    total, sales, remodel, tx, rm = FIGURES[figure]
    rows = docs[("block", work_type, audience, layout, figure)]
    where = f"{work_type}/{audience} {layout} {figure}: {rows!r}"
    assert "base" in rows, f"no base line printed — {where}"
    if layout == "BROKEN_OUT":
        # A row prints exactly when its tax applies; the Total always.
        assert ("material" in rows) == tx, f"Material Sales Tax row wrong — {where}"
        assert ("remodel" in rows) == rm, f"Remodel Tax row wrong — {where}"
        assert rows.get("total") == total, f"Total wrong — {where}"
        if tx:
            assert rows["material"] == sales
        if rm:
            assert rows["remodel"] == remodel
        printed = round(rows["base"] + rows.get("material", 0.0) + rows.get("remodel", 0.0), 2)
        assert printed == total, f"the printed rows do not sum to the Total — {where}"
        assert "(" not in rows["base_line"], f"a broken-out base line carries bracket wording — {where}"
    else:
        assert not {"material", "remodel", "total"} & set(rows), f"one line printed a tax row — {where}"
        assert rows["base"] == total, f"one line is not the whole bid — {where}"
        assert rows["base_line"].endswith(PHRASE[(tx, rm)]), f"wrong wording for the sheet — {where}"
    assert rows["highlights"] == 0, f"a review highlight reached the customer — {where}"


def test_hanz_test33_prints_exactly_what_he_sent(docs):
    """His own words for a correct broken-out base: Taxable Yes, Remodel No, $6,839 bid."""
    rows = docs[("block", "epoxy", "Direct", "BROKEN_OUT", "Test33")]
    assert rows["base_line"] == "$6,767 – Epoxy flooring as described above"
    assert rows["material"] == 72.0 and rows["total"] == 6839.0
    assert "remodel" not in rows


def test_kyles_gc_polish_block_reconciles_line_for_line(docs):
    """The exact block Kyle was shown, Broken out: 6,182 + 125 = 6,307, and no "$0.00 – Remodel Tax"
    row between them — a remodel row prints only when Remodel Tax? says Yes."""
    rows = docs[("block", "polish", "GC", "BROKEN_OUT", "sales tax only")]
    assert rows["base"] == 6182.0 and rows["material"] == 125.0
    assert "remodel" not in rows and rows["total"] == 6307.0
    assert "INCLUDED" not in rows["base_line"], rows["base_line"]


def test_an_exempt_gc_job_prints_no_zero_rows(docs):
    """A GC file's tax rows are plain paragraphs, so an exempt job used to print "$0 – Material Sales
    Tax" under "(tax exempt)". One line now: one line."""
    rows = docs[("block", "epoxy", "GC", "ONE_LINE", "exempt")]
    assert rows["base_line"].endswith("(tax exempt)")
    assert not {"material", "remodel", "total"} & set(rows), rows


# ── (2) every option off its own tab ────────────────────────────────────────────────────────────
@pytest.mark.parametrize("work_type,audience", OPTION_TEMPLATES)
def test_each_option_is_itemised_from_its_own_tab_when_broken_out(docs, work_type, audience):
    """Hanz picked this from previews: each option its own pre-tax line, its own tax rows as ITS
    flags say, its own Total — summing. The seal option is remodel-only: no Material Sales Tax row."""
    opts = docs[("options", work_type, audience, "BROKEN_OUT", "both, to the cent")]["options"]
    assert opts == [
        "$7,100 – Treadwell 3/16\" Urethane Cement as described above",
        "$96 – Material Sales Tax",
        "$500 – Remodel Tax",
        "$7,696 – Total",
        "$1,386 – Sealed Concrete as described above",
        "$90 – Remodel Tax",
        "$1,476 – Total",
    ], opts


@pytest.mark.parametrize("work_type,audience", OPTION_TEMPLATES)
def test_each_option_carries_its_own_tax_wording_on_one_line(docs, work_type, audience):
    """One line: the option's whole price and the wording ITS tab's flags call for. These used to
    say "(material sales tax INCLUDED)" whatever the job, so Kyle typed EXCLUDED onto every option
    of every exempt job by hand."""
    opts = docs[("options", work_type, audience, "ONE_LINE", "both, to the cent")]["options"]
    assert opts == [
        "$7,696 – Treadwell 3/16\" Urethane Cement as described above "
        "(Remodel Tax AND material sales tax INCLUDED)",
        "$1,476 – Sealed Concrete as described above (Remodel Tax INCLUDED)",
    ], opts


# ── (3) the discriminator is the wrapper, not the folder ────────────────────
def _synthetic(*texts):
    d = docx.Document()
    for t in texts:
        d.add_paragraph(t)
    return d


def test_free_tax_rows_follows_the_wrapper():
    """THE COUNTEREXAMPLE. `free_tax_rows` claims to answer "does this row print unconditionally",
    and a green test over eight files that happen to split GC-vs-Direct would pass for a function
    that just read the folder name. These documents differ ONLY by the {{#tax_breakout}} /
    {{#remodel}} wrappers, so the answer has to move with them."""
    free = pw.free_tax_rows(_synthetic(
        "Base Bid",
        "{{base_bid_formatted}} – Polished Concrete as described above {{base_tax_phrase}}",
        "{{material_tax_formatted}} – Material Sales Tax",
        "{{tax_amount_formatted}} – Remodel Tax",
        "{{total_formatted}} – Total"))
    assert free == {"material": True, "remodel": True}

    gated = pw.free_tax_rows(_synthetic(
        "Base Bid",
        "{{base_bid_formatted}} – Polished Concrete as described above {{base_tax_phrase}}",
        "{{#tax_breakout}}",
        "{{material_tax_formatted}} – Material Sales Tax",
        "{{/tax_breakout}}",
        "{{#remodel}}",
        "{{remodel.amount_formatted}} – Remodel Tax",
        "{{/remodel}}",
        "{{total_formatted}} – Total"))
    assert gated == {"material": False, "remodel": False}

    assert pw.free_tax_rows(_synthetic(
        "{{#tax_breakout}}", "{{material_tax_formatted}} – Material Sales Tax",
        "{{/tax_breakout}}", "{{tax_amount_formatted}} – Remodel Tax")) == {
            "material": False, "remodel": True}

    assert pw.free_tax_rows(_synthetic("Budget pricing", "{{total_formatted}} – Total")) == {
        "material": False, "remodel": False}


@pytest.mark.parametrize("work_type,audience,expected", [
    ("epoxy", "Direct", False), ("polish", "Direct", False), ("combo", "Direct", False),
    ("budget", "Direct", False),
    ("epoxy", "GC", True), ("polish", "GC", True), ("sealer", "GC", True), ("combo", "GC", True),
    ("gyp", "Direct", True), ("gyp", "GC", True),
])
def test_every_shipped_template_is_classified_from_its_own_file(work_type, audience, expected):
    """What the eight real files say today, cross-checked against an INDEPENDENT walk. The shape
    still matters for one thing: how a draft saved before the TAX control became layout-only is read
    (price_rules.layout_is_broken)."""
    got = pw.template_free_tax_rows(work_type, audience)
    assert got == {"material": expected, "remodel": expected}, pw.pick_template(
        work_type, audience).name
    d = docx.Document(str(pw.pick_template(work_type, audience)))
    walked = [in_block for _i, _k, _p, in_block, text, _x in pw.iter_editable_blocks(d)
              if "material_tax_formatted" in (text or "")]
    if expected:
        assert walked and all(b is None for b in walked), walked
    else:
        assert all(b is not None for b in walked), walked


def test_an_unreadable_template_says_why_instead_of_guessing(monkeypatch, caplog):
    gone = pw.TEMPLATES_ROOT / "Direct" / "not-in-the-image.docx"
    monkeypatch.setattr(pw, "pick_template", lambda work_type, audience: gone)
    with caplog.at_level("WARNING", logger="proposal_tool.proposal_writer"):
        assert pw.template_free_tax_rows("polish", "GC") == {"material": False, "remodel": False}
    msg = caplog.text
    assert "tax-row shape" in msg, msg
    assert "polish" in msg and "GC" in msg, "the warning does not name the template it refused"
    assert "Error" in msg or "Exception" in msg, f"no exception type in the warning: {msg}"


def test_the_shape_read_is_cached_per_file_version():
    pw._free_tax_rows_cached.cache_clear()
    pw.template_free_tax_rows("polish", "GC")
    first = pw._free_tax_rows_cached.cache_info()
    pw.template_free_tax_rows("polish", "GC")
    second = pw._free_tax_rows_cached.cache_info()
    assert second.hits == first.hits + 1 and second.misses == first.misses


# ── (4) the screen and the document print the same price block ──────────────
def _template_blocks(work_type, audience):
    d = docx.Document(str(pw.pick_template(work_type, audience)))
    return [{"id": i, "in_block": in_block, "text": text}
            for i, _k, _p, in_block, text, _x in pw.iter_editable_blocks(d)]


@pytest.fixture(scope="module")
def screen():
    """Run the page's own price-block writers, in Node, over the real template shapes."""
    if shutil.which("node") is None:
        pytest.skip("node is not installed")
    cases = []
    for work_type, audience in TEMPLATES:
        blocks = _template_blocks(work_type, audience)
        for layout in LAYOUTS:
            for figure, (total, sales, remodel, tx, rm) in FIGURES.items():
                cases.append({"name": f"block/{work_type}/{audience}/{layout}/{figure}",
                              "work_type": work_type, "audience": audience, "blocks": blocks,
                              "tax_layout": layout, "taxable": tx, "remodel_on": rm,
                              "total": total, "sales_tax": sales, "remodel_tax": remodel})
    for work_type, audience in OPTION_TEMPLATES:
        blocks = _template_blocks(work_type, audience)
        total, sales, remodel, tx, rm = FIGURES["both, to the cent"]
        for layout in LAYOUTS:
            cases.append({"name": f"options/{work_type}/{audience}/{layout}",
                          "work_type": work_type, "audience": audience, "blocks": blocks,
                          "tax_layout": layout, "taxable": tx, "remodel_on": rm,
                          "total": total, "sales_tax": sales, "remodel_tax": remodel,
                          "rooms": OPTION_ROOMS})
    p = subprocess.run(["node", str(HARNESS), str(FRONTEND)],
                       input=json.dumps(cases), capture_output=True, text=True, encoding="utf-8")
    assert p.returncode == 0, p.stderr
    return {r["name"]: r for r in json.loads(p.stdout)}


def _free_row_kind(text):
    if re.search(r"\{\{\s*material_tax_formatted\s*\}\}", text or ""):
        return "material"
    if re.search(r"\{\{\s*(tax_amount_formatted|remodel\.amount_formatted)\s*\}\}", text or ""):
        return "remodel"
    if re.search(r"\{\{\s*(total_formatted|total_label)\s*\}\}", text or ""):
        return "total"
    return None


@needs_node
@pytest.mark.parametrize("work_type,audience", TEMPLATES)
@pytest.mark.parametrize("layout", LAYOUTS)
@pytest.mark.parametrize("figure", list(FIGURES))
def test_the_screen_and_the_document_print_the_same_price_block(screen, docs, work_type, audience,
                                                                layout, figure):
    """THE PROPERTY BEING FIXED: the figures, the wording and the rows the estimator proofreads are
    the ones the customer reads — every template shape, both layouts, every combination of the two
    taxes. JS off the served blocks, Python off the template file."""
    r = screen[f"block/{work_type}/{audience}/{layout}/{figure}"]
    rows = docs[("block", work_type, audience, layout, figure)]
    where = f"{work_type}/{audience} {layout} {figure}"
    tk = r["tokens"]
    assert tk["tax_layout"] == layout and r["layout"] == layout, where
    # The base bid and its wording, as the document prints them.
    assert rows["base"] == _usd(tk["base_bid_formatted"]), (
        f"{where}: screen {tk['base_bid_formatted']} vs document ${rows['base']:,.2f}")
    assert rows["base_line"].endswith(tk["base_tax_phrase"] or "as described above"), (
        f"{where}: the document's wording {rows['base_line']!r} is not the screen's "
        f"{tk['base_tax_phrase']!r}")
    # Which rows print, three ways: the tokens the payload carries, the rows the page paints (the
    # Direct files mount them), and the free rows the page shows (the GC / Gyp files).
    doc_rows = {k for k in ("material", "remodel", "total") if k in rows}
    screen_rows = {k for k in ("material", "remodel", "total") if tk[f"price_rows_{k}"]}
    assert screen_rows == doc_rows, f"{where}: screen {screen_rows} vs document {doc_rows}"
    free = r["freeRows"]
    blocks = {b["id"]: b for b in _template_blocks(work_type, audience)}
    for bid, shown in free.items():
        kind = _free_row_kind(blocks[int(bid)]["text"])
        assert shown == (kind in doc_rows), f"{where}: free {kind} row shown={shown}, document {doc_rows}"
    p = r["painted"]
    painted_rows = {k for k, on in (("material", p["salesRowShown"]), ("remodel", p["remodelRowShown"]),
                                    ("total", p["totalRowShown"])) if on}
    if not free:
        assert painted_rows == doc_rows, f"{where}: painted {painted_rows} vs document {doc_rows}"
        assert _usd(p["base"]) == rows["base"], f"{where}: painted {p['base']!r}"


@needs_node
@pytest.mark.parametrize("work_type,audience", OPTION_TEMPLATES)
@pytest.mark.parametrize("layout", LAYOUTS)
def test_the_screen_and_the_document_print_the_same_option_lines(screen, docs, work_type, audience, layout):
    """Every option line, as drawn in the editor, is the line the document prints — its own tab's
    figures, its own wording, its own rows."""
    r = screen[f"options/{work_type}/{audience}/{layout}"]
    shown = [ln["text"] for ln in r["options"]]
    printed = docs[("options", work_type, audience, layout, "both, to the cent")]["options"]
    assert shown == printed, f"{work_type}/{audience} {layout}: screen {shown} vs document {printed}"


@needs_node
def test_the_default_layout_follows_the_sheet_until_somebody_picks():
    """"if one of the taxes is set to yes then broken out should be the default option." A draft
    that never chose a layout follows the sheet; one that did keeps it; a draft saved before the
    control became layout-only is read by what it printed (Broken out stays Broken out, Included and
    Exempt were one line — on a GC/Gyp file, whose rows always printed, the default answers)."""
    blocks = _template_blocks("epoxy", "Direct")
    gc = _template_blocks("epoxy", "GC")
    base = {"work_type": "epoxy", "audience": "Direct", "blocks": blocks,
            "total": 6839.0, "sales_tax": 72.0, "remodel_tax": 0.0}
    cases = [
        dict(base, name="new, taxable", taxable=True, remodel_on=False),
        dict(base, name="new, remodel only", taxable=False, remodel_on=True, sales_tax=0.0, remodel_tax=40.0),
        dict(base, name="new, neither", taxable=False, remodel_on=False, sales_tax=0.0),
        dict(base, name="picked one line", taxable=True, remodel_on=False, tax_layout="ONE_LINE"),
        dict(base, name="legacy BROKEN_OUT", taxable=False, remodel_on=False, sales_tax=0.0,
             tax_inclusion="BROKEN_OUT"),
        dict(base, name="legacy INCLUDED", taxable=True, remodel_on=False, tax_inclusion="INCLUDED"),
        dict(base, name="legacy EXEMPT", taxable=True, remodel_on=False, tax_inclusion="EXEMPT"),
        dict(base, name="legacy GC INCLUDED", audience="GC", blocks=gc, taxable=True,
             remodel_on=False, tax_inclusion="INCLUDED"),
        dict(base, name="legacy GC EXEMPT, exempt job", audience="GC", blocks=gc, taxable=False,
             remodel_on=False, sales_tax=0.0, tax_inclusion="EXEMPT"),
    ]
    p = subprocess.run(["node", str(HARNESS), str(FRONTEND)], input=json.dumps(cases),
                       capture_output=True, text=True, encoding="utf-8")
    assert p.returncode == 0, p.stderr
    got = {r["name"]: r["layout"] for r in json.loads(p.stdout)}
    assert got == {
        "new, taxable": "BROKEN_OUT", "new, remodel only": "BROKEN_OUT", "new, neither": "ONE_LINE",
        "picked one line": "ONE_LINE", "legacy BROKEN_OUT": "BROKEN_OUT",
        "legacy INCLUDED": "ONE_LINE", "legacy EXEMPT": "ONE_LINE",
        "legacy GC INCLUDED": "BROKEN_OUT", "legacy GC EXEMPT, exempt job": "ONE_LINE",
    }, got
    # And the document reads a legacy payload the same way (no tax_layout on it).
    for legacy, free, tx, rm, want in [("INCLUDED", False, True, False, False),
                                       ("EXEMPT", False, True, True, False),
                                       ("BROKEN_OUT", False, False, False, True),
                                       ("INCLUDED", True, True, False, True),
                                       ("EXEMPT", True, False, False, False)]:
        assert price_rules.layout_is_broken(None, legacy, free_rows=free, taxable=tx,
                                            remodel_on=rm) is want, (legacy, free, tx, rm)
