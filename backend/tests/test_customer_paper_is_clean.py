"""Two things a customer must never read on a Treadwell proposal: a literal
"&amp;" where an "&" belongs, and a raw "{{token}}" where a number belongs.

Both shipped. Both were invisible to the whole suite for the same reason -- the
existing tests assert on the templates and the payloads the live browser
produces, and the live browser happens to produce clean ones.

(a) THE AMPERSAND. `annotate_templates.py` matches SEARCH strings against raw
    document.xml (so those are XML-escaped, "&amp;") but the engine
    xml_escape()s its REPLACEMENT strings on the way in
    (`inner.replace(search, xml_escape(replacement))`). Three replacements
    still carried "&amp;", which double-escaped to "&amp;amp;" in the file --
    which python-docx faithfully decodes back to the five characters "&amp;"
    and prints. GC Polish and GC Resinous each shipped it, twice per string:
    every floating text box is authored twice, once as modern DrawingML
    (`mc:Choice`) and once as legacy VML (`mc:Fallback`), and the fill pass
    writes both so old Word readers get filled tokens too. This walks EVERY
    part of EVERY template, not the two known files and not just
    document.xml, because a header/footer is the same mistake waiting
    somewhere nobody looked.

(b) THE RAW TOKEN. CLAUDE.md claims "Generation never emits a raw {{token}}".
    `_ensure_value_aliases` is what has to make that true, and for two tokens
    it did not: Direct Polish's Total row is the WHOLE-LINE token
    {{total_label}} and its Area line is {{area_description}}, so the
    amount-token backfills beside them (base_bid_formatted,
    material_tax_formatted) never reached either one. The live browser always
    sends both (`proposal-review.js` computeTokenValues), so a normal Generate
    is clean and the hole only opens on a REPLAY of a frozen payload --
    /api/admin/proposal-pdf, a revision's file links, the To-Dropbox re-file --
    where a payload saved before those tokens existed prints
    "{{total_label}}" where the customer's Total belongs.

    Watch the tax flag in these tests. Direct Polish wraps its Total row in
    {{#tax_breakout}}, so on the DEFAULT "sales tax INCLUDED" setting the row
    is stripped and {{total_label}} is not in the document at all -- a version
    of this test without `tax_inclusion: BROKEN_OUT` passes whether or not the
    bug is fixed, because the line it is about never renders.

Everything here executes the real fill path and reads the produced .docx. A
source assertion could catch neither bug: the "&amp;" lives in a template FILE,
and a missing backfill only exists in output. Note the price and area rows live
inside floating text boxes, so `docx.Document().paragraphs` never reaches them
-- `iter_editable_blocks` / `_own_text` is how you see what the customer sees.
"""
import glob
import html
import io
import os
import re
import zipfile

import docx
import pytest
from fastapi.testclient import TestClient

import main
import proposal_writer as pw

client = TestClient(main.app)

# Every .docx we ship -- the 8 proposal templates, the 7 cover letters and the
# deposit invoice. Globbed, not listed, so a new template is covered the day it
# lands rather than the day someone remembers to add it here.
_ALL_TEMPLATES = sorted(glob.glob(
    os.path.join(os.path.dirname(os.path.abspath(pw.__file__)),
                 "templates", "**", "*.docx"), recursive=True))

assert _ALL_TEMPLATES, "no templates found -- the glob is wrong, not the repo"

# A doubly-escaped ampersand. In raw XML a CORRECT "&" already is "&amp;"; the
# bug is the extra layer, "&amp;amp;", which decodes to the visible text "&amp;".
_DOUBLE_AMP = "&amp;amp;"

_WP = re.compile(r"<w:p[ >].*?</w:p>", re.DOTALL)
_WT = re.compile(r"<w:t\b[^>]*>(.*?)</w:t>", re.DOTALL)
# {{token}}, {{ spaced }}, {{dotted.token}} and the block markers {{#x}}/{{/x}}.
_TOKEN = re.compile(r"\{\{\s*[#/]?\s*([A-Za-z_][A-Za-z0-9_.]*)\s*\}\}")

# Every (work_type, audience) the tool can pick, read straight off the picker so
# a new pair is covered automatically.
_PAIRS = sorted(pw.TEMPLATE_PICKER.keys(), key=lambda k: (k[0], str(k[1])))

# What a frozen payload from the live browser carries: the intake/estimate DATA
# fields, and none of the display tokens derived from them. This is the shape a
# replay actually has, and the shape that used to print "{{total_label}}".
# BROKEN_OUT is required or Direct Polish strips the Total row (see the note up
# top) and the total_label half of this file proves nothing.
_LEGACY_PAYLOAD = {
    "job_name": "Acme Warehouse",
    "city_state": "Olathe, KS",
    "bid_date": "2026-09-08",
    "sqft": "1,600",
    "epoxy_sf": "1,600",
    "polish_sf": "1,600",
    "cove_lf": "500",
    "texture": "Light",
    "system_name": "Standard Sheen",
    "total_formatted": "$6,307",
    "material_tax_formatted": "$125.00",
    "tax_amount_formatted": "$0.00",
    "tax_inclusion": "BROKEN_OUT",
}

# Tokens that survive when the payload carries NO data at all. Every one is a
# value the CALLER supplies -- there is nothing to derive them from, and putting
# a made-up number where a price or a customer name belongs would be worse than
# printing the token. Listed with the reason so a NEW token that nobody wired up
# fails test_empty_payload_leftovers_are_only_caller_data instead of quietly
# joining the crowd.
_CALLER_DATA_TOKENS = {
    "bid_date_formatted": "derived from bid_date; a replay must NOT be re-dated to today",
    "city_state":         "project location, typed at intake",
    "cove_lf":            "cove lineal feet, off the estimate sheet",
    "epoxy_sf":           "epoxy square feet, off the estimate sheet",
    "job_name":           "the customer's name",
    "polish_sf":          "polish square feet, off the estimate sheet",
    "sqft":               "area, off the estimate sheet",
    "system_name":        "the picked system (Epoxy!A22 dropdown)",
    "tax_amount_formatted": "remodel tax -- money, never invented",
    "texture":            "the picked texture",
    "total_formatted":    "the bid -- money, never invented",
    # These two ARE derived (that is what this file fixed), but they are derived
    # FROM the caller data above, so with a payload holding literally nothing
    # they have no input either. test_replay_of_a_legacy_payload_* is what
    # proves the derivation works whenever the inputs exist.
    "area_description":   "derived from sqft + work type; blank when sqft is absent",
    "total_label":        "derived from total_formatted; blank when the total is absent",
}


def _xml_parts(path_or_bytes):
    """(part_name, text) for every XML part of a .docx -- document, headers,
    footers, footnotes and anything else Word tucked in there."""
    src = path_or_bytes if isinstance(path_or_bytes, str) else io.BytesIO(path_or_bytes)
    with zipfile.ZipFile(src) as z:
        for name in z.namelist():
            if name.endswith(".xml"):
                yield name, z.read(name).decode("utf-8", "replace")


def _paragraph_texts(xml):
    """The visible text of each <w:p>, runs joined and entities decoded.

    Joining the runs FIRST is load-bearing: Word splits a token across runs
    freely ("{{tot" + "al_label}}"), so a regex run straight at the raw XML
    would miss exactly the leftovers this file exists to catch.
    """
    for m in _WP.finditer(xml):
        yield html.unescape("".join(_WT.findall(m.group(0))))


def _generate(work_type, audience, values):
    """POST /api/generate and hand back the produced .docx bytes.

    Through the route (via pytest's client, so conftest's prod-write guard and
    auth bypass both apply) because the claim under test is about GENERATION,
    and the route backfills more than `_ensure_value_aliases` does on its own.
    """
    body = {"work_type": work_type, "values": dict(values)}
    if audience:
        body["audience"] = audience
    r = client.post("/api/generate", json=body)
    assert r.status_code == 200, r.text
    return client.get(r.json()["docx_download_url"]).content


def _raw_tokens(data):
    """{token: the paragraph it survived in} for a generated .docx."""
    found = {}
    for name, xml in _xml_parts(data):
        for text in _paragraph_texts(xml):
            for tok in _TOKEN.findall(text):
                found.setdefault(tok, text.strip()[:90])
    return found


def _editable_text(data):
    """Every editable paragraph's own text, text boxes included."""
    d = docx.Document(io.BytesIO(data))
    return [t for _i, _k, _p, _b, t, _tb in pw.iter_editable_blocks(d)]


# ── (a) no literal "&amp;" survives in ANY template part ────────────────
@pytest.mark.parametrize("path", _ALL_TEMPLATES,
                         ids=[os.path.basename(p) for p in _ALL_TEMPLATES])
def test_no_double_escaped_ampersand_in_any_template_part(path):
    offenders = {name: xml.count(_DOUBLE_AMP)
                 for name, xml in _xml_parts(path)
                 if _DOUBLE_AMP in xml}
    assert not offenders, (
        f"{os.path.basename(path)} stores a literal '&amp;' instead of '&' in "
        f"{offenders} -- the customer reads the five characters '&amp;'. Fix the "
        f"run text in the .docx; if it came from annotate_templates.py, that "
        f"rule's REPLACEMENT string must carry a plain '&' (replacements are "
        f"xml_escape()d on the way in, search strings are not).")


def test_gc_price_and_area_lines_render_a_real_ampersand():
    """The three lines that shipped broken, through the real fill path.

    The paired negative/positive matters: deleting the paragraph would satisfy a
    bare "no &amp;" assertion just as well as fixing it, and these lines carry
    the customer's price.
    """
    expected = {
        ("polish", "GC"): [
            "Polished Concrete & Joint Filler as described above",
        ],
        ("epoxy", "GC"): [
            "resinous flooring [RESx] & 500 lf of integral base",
            "Resinous floor & integral cove base as described above",
        ],
    }
    for (work_type, audience), wanted in expected.items():
        out = pw.fill_proposal(work_type=work_type, audience=audience, values={
            "sqft": "1,600", "cove_lf": "500",
            "base_bid_formatted": "$6,307",
            "base_tax_phrase": "(material sales tax INCLUDED)",
        })
        seen = "\n".join(_editable_text(out))
        assert "&amp;" not in seen, (
            f"{work_type}/{audience} rendered a literal '&amp;' to the customer")
        for phrase in wanted:
            assert phrase in seen, (
                f"{work_type}/{audience} lost the line containing {phrase!r}")


# ── (b) a replayed legacy payload prints ZERO raw tokens ────────────────
@pytest.mark.parametrize("work_type,audience", _PAIRS,
                         ids=[f"{w}-{a}" for w, a in _PAIRS])
def test_replay_of_a_legacy_payload_prints_no_raw_token(work_type, audience):
    """Every template, filled from a payload that has the DATA but none of the
    derived display tokens -- i.e. what /api/admin/proposal-pdf, a revision's
    file links and the To-Dropbox re-file actually replay."""
    leftover = _raw_tokens(_generate(work_type, audience, _LEGACY_PAYLOAD))
    assert not leftover, (
        f"{work_type}/{audience} printed raw token(s) {sorted(leftover)} on a "
        f"customer-facing proposal: {leftover}. Derive them in "
        f"main._ensure_value_aliases (see total_label / area_description) or, if "
        f"the value genuinely cannot be derived, the caller has to send it and it "
        f"belongs in _CALLER_DATA_TOKENS with a reason.")


def test_polish_total_and_area_rows_are_really_in_the_document():
    """Guards the guard: assert the two lines are PRESENT and correct, not just
    absent-and-therefore-token-free.

    Direct Polish strips its Total row unless tax is broken out, so without this
    the {{total_label}} assertion above would be vacuous.
    """
    rows = _editable_text(_generate("polish", "Direct", _LEGACY_PAYLOAD))
    stripped = [r.strip() for r in rows]
    assert "$6,307 – Total" in stripped, (
        f"the Total row is missing entirely, so nothing proves total_label is "
        f"derived. PRICE rows seen: {[r for r in stripped if r.startswith('$')]}")
    assert "Area: ~1,600 sf of polished concrete flooring" in stripped, (
        f"the Area line is missing or wrong: "
        f"{[r for r in stripped if r.startswith('Area')]}")


def test_replay_matches_what_the_browser_would_have_sent():
    """A replay must reproduce the ORIGINAL document, not a differently-worded
    one -- so the derived strings have to byte-match the browser's."""
    browser = dict(_LEGACY_PAYLOAD)
    browser["total_label"] = "$6,307 – Total"
    browser["area_description"] = "~1,600 sf of polished concrete flooring"

    replay_rows = _editable_text(_generate("polish", "Direct", _LEGACY_PAYLOAD))
    browser_rows = _editable_text(_generate("polish", "Direct", browser))
    assert replay_rows == browser_rows, (
        "a replayed payload rendered different text than the browser's own")


def test_derived_area_description_follows_the_work_type():
    """The noun comes off the work type, the same three ways
    proposal-review.js picks it."""
    cases = {
        "polish": "~1,600 sf of polished concrete flooring",
        "gyp":    "~1,600 sf of gypsum underlayment",
        "epoxy":  "~1,600 sf of epoxy flooring",
        "combo":  "~1,600 sf of epoxy flooring",
    }
    for work_type, expected in cases.items():
        values = {"work_type": work_type, "sqft": "1,600"}
        main._ensure_value_aliases(values, "Direct")
        assert values["area_description"] == expected, work_type


def test_derivations_never_overwrite_what_the_caller_sent():
    """These are FALLBACKS. The estimator's own wording -- a retitled Total row,
    a hand-edited Area line -- has to win, or the doc editor's edits vanish."""
    values = {
        "work_type": "polish", "sqft": "1,600", "total_formatted": "$6,307",
        "total_label": "$6,307 – Total Base Bid",
        "area_description": "the lobby and corridors only",
    }
    main._ensure_value_aliases(values, "Direct")
    assert values["total_label"] == "$6,307 – Total Base Bid"
    assert values["area_description"] == "the lobby and corridors only"


# ── (c) the residue is caller DATA only, and the list stays honest ──────
@pytest.mark.parametrize("work_type,audience", _PAIRS,
                         ids=[f"{w}-{a}" for w, a in _PAIRS])
def test_empty_payload_leftovers_are_only_caller_data(work_type, audience):
    """With a payload holding nothing but the tax flag, anything still raw must
    be a value the caller was always going to have to supply.

    This is the part that closes the CLASS: a new template token that nobody
    taught `_ensure_value_aliases` to derive lands here as an unknown.
    """
    leftover = _raw_tokens(_generate(work_type, audience,
                                     {"tax_inclusion": "BROKEN_OUT"}))
    unknown = {k: v for k, v in leftover.items() if k not in _CALLER_DATA_TOKENS}
    assert not unknown, (
        f"{work_type}/{audience} printed raw token(s) {sorted(unknown)} that "
        f"nothing derives and nothing documents: {unknown}. Either derive it in "
        f"main._ensure_value_aliases or add it to _CALLER_DATA_TOKENS with the "
        f"reason it has to come from the caller.")


def test_caller_data_list_has_no_dead_entries():
    """A stale allow-list is how this test goes quietly vacuous: once a token IS
    derived, leaving it listed would keep excusing the next regression on it."""
    still_raw = set()
    for work_type, audience in _PAIRS:
        still_raw |= set(_raw_tokens(_generate(work_type, audience,
                                               {"tax_inclusion": "BROKEN_OUT"})))
    dead = sorted(set(_CALLER_DATA_TOKENS) - still_raw)
    assert not dead, (
        f"_CALLER_DATA_TOKENS lists {dead}, which no template leaves raw any "
        f"more. Drop them from the list so it keeps meaning something.")
