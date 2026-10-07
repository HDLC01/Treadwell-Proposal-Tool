"""The Dye and Joint Filler figures the bid prices with reach Kyle's downloaded workbook.

Hanz's rule: ONE PRICE EVERYWHERE -- the tool's bid and the .xlsx must never quote different
figures. Since 2026-09-30 the Polish estimate prices both lines off reserved library_items rows
(`dye`, `joint-filler-kit`; see polish-estimate.js condLine) that an admin edits on the Items tab.
Kyle's Polish tab carried its own C25 0.14, C29 500 and "/3500", so an edited row made the two
disagree. bid-model.js's conditionCellWrites now writes, on every save:

    Polish!C25   one coat of dye's price per SF (the row's unit price over its coverage, plus
    Polish!C26   waste) -- BOTH cells: Kyle's rows 25 and 26 are one coat each (DYE_COATS 2)
    Polish!C29   one kit's price (unit_cost / buy_qty)
    Polish!B29   the kit count, as a FORMULA, only where the row's coverage / waste / roundup /
                 pack differ from the template's =ROUNDUP(IF(E29="yes",(E18/3500),0),0)

Everything here is EXECUTED: the cells come from the real page's own save
(polish-estimate-harness.js, out.reservedWorkbook), the workbook from the real
estimate_writer.fill_estimate, and every dollar is read off that workbook by walking its own
formula text (the same reader test_markup_rate_reaches_the_bid.py proves against Excel's cached
totals). openpyxl writes no cached results, so nothing here trusts a number Excel did not compute
from the formulas in the file.

WHAT IS AND IS NOT CLAIMED. D25 + D26, D29 and B29 are the bid's Dye, Joint Filler line and kit
count to the cent. The tab total (D82) is recomputed by Kyle's own formulas from them. It is NOT
the bid's total, and was not before this change: the beta page does not write the takeoff (E18
included) into the workbook -- see polish-estimate.js's file header -- so E18 is set here the way
the intake -> estimate grid path sets it. Rows 34-63 (labor, tooling, travel) are held at the
template's cached figures; test_the_held_rows_read_none_of_the_written_cells proves that holding
them cannot hide a difference.
"""
import io
import json
import pathlib
import re
import shutil
import subprocess
import sys
import warnings

import openpyxl
import pytest

sys.path.insert(0, str(pathlib.Path(__file__).resolve().parents[1]))
sys.path.insert(0, str(pathlib.Path(__file__).resolve().parent))

import estimate_writer as ew  # noqa: E402
from test_markup_rate_reaches_the_bid import _Eval, _num  # noqa: E402  the proven chain reader

FRONTEND = pathlib.Path(__file__).resolve().parents[2] / "frontend"
HARNESS = pathlib.Path(__file__).resolve().parent / "js" / "polish-estimate-harness.js"
TEMPLATE_B29 = '=ROUNDUP(IF(E29="yes",(E18/3500),0),0)'
LINE_CELLS = ("B25", "C25", "D25", "B26", "C26", "D26", "B29", "C29", "D29")
EVALUATED = ("D25", "D26", "D29", "D31", "D33", "D64", "D82")
needs_node = pytest.mark.skipif(shutil.which("node") is None, reason="node is not installed")


@pytest.fixture(scope="module")
def ran():
    if shutil.which("node") is None:
        pytest.skip("node is not installed")
    proc = subprocess.run(["node", str(HARNESS), str(FRONTEND)],
                          capture_output=True, text=True, encoding="utf-8", timeout=180)
    assert proc.returncode == 0, "the harness itself failed:\n" + proc.stderr
    assert proc.stdout.strip(), "the harness printed nothing:\n" + proc.stderr
    return json.loads(proc.stdout.strip().splitlines()[-1])["reservedWorkbook"]


@pytest.fixture(scope="module")
def cached():
    with warnings.catch_warnings():
        warnings.simplefilter("ignore")
        return openpyxl.load_workbook(ew.TEMPLATE_PATH, data_only=True)


class Book:
    """One generated workbook, read the way Excel will: Polish rows 17-33 (material) and 64-82
    (the markup chain) from their own formula text, rows 34-63 at the template's cached figures,
    literals as written."""

    def __init__(self, xlsx, cached):
        with warnings.catch_warnings():
            warnings.simplefilter("ignore")
            self.f = openpyxl.load_workbook(io.BytesIO(xlsx), data_only=False)
        self.cached = cached
        self._memo = {}

    def value(self, sheet, addr):
        a = addr.replace("$", "")
        key = (sheet, a)
        if key in self._memo:
            return self._memo[key]
        raw = self.f[sheet][a].value
        row = int(re.match(r"[A-Z]{1,3}([0-9]+)$", a).group(1))
        if isinstance(raw, str) and raw.startswith("="):
            if sheet == "Polish" and (17 <= row <= 33 or 64 <= row <= 82):
                self._memo[key] = 0.0                      # cycle guard
                out = _Eval(self, sheet).run(raw[1:])
            else:
                out = self.cached[sheet][a].value
        else:
            out = raw
        self._memo[key] = out
        return out

    def num(self, addr):
        return _num(self.value("Polish", addr))

    def raw(self, addr):
        return self.f["Polish"][addr].value


_BOOKS = {}


def _workbook(scn, cached, cells=None):
    """The real writer, fed the page's own saved cells plus the area the way the intake -> grid
    path writes it. One write per distinct input: each costs a template parse and a save."""
    cv = dict(scn["cells"] if cells is None else cells)
    cv["Polish!E18"] = scn["area"]
    key = json.dumps(cv, sort_keys=True)
    if key not in _BOOKS:
        _BOOKS[key] = Book(ew.fill_estimate({}, cv), cached)
    return _BOOKS[key]


def _material_raw(book):
    return sum(book.num("D%d" % r) for r in range(17, 31))


@needs_node
def test_the_seeded_rows_leave_the_workbook_exactly_as_today(ran, cached):
    """With the seeded rows the save writes C25 = C26 = 0.14 and C29 500 -- the template's own
    figures -- and no quantity formula, so every Dye / Joint Filler cell, the material block
    and the tab total come out identical to the workbook today's save produces. And those
    figures ARE the bid's: D25 + D26 is two coats of 3,501 SF at $0.14, D29 two kits at $500.

    Mutation: write the kit's formula whatever it says (B29 differs), or write dye's per-SF rate
    off the wrong coverage."""
    s = ran["seeded"]
    extra = {k: v for k, v in s["cells"].items() if s["today"].get(k) != v}
    assert extra == {"Polish!C25": 0.14, "Polish!C26": 0.14, "Polish!C29": 500}, extra
    now, today = _workbook(s, cached), _workbook(s, cached, cells=s["today"])
    for addr in LINE_CELLS:
        assert now.raw(addr) == today.raw(addr), (addr, now.raw(addr), today.raw(addr))
    for addr in EVALUATED:
        assert now.num(addr) == pytest.approx(today.num(addr)), addr
    assert now.raw("B29") == TEMPLATE_B29
    assert now.num("D25") + now.num("D26") == pytest.approx(s["dyeCost"]), (
        now.num("D25"), now.num("D26"), s["dyeCost"])
    assert now.num("D29") == pytest.approx(s["jfCost"]), (now.num("D29"), s["jfCost"])


@needs_node
def test_an_edited_library_reaches_the_workbook_line_for_line(ran, cached):
    """$650 a kit, coverage 2,000, dye $0.20 a coat on the rows. The save writes C25 = C26 = 0.2,
    C29 650 and a B29 FORMULA carrying the 2,000; the writer keeps it a formula (_coerce lets a
    safe "=ROUNDUP" through), and the workbook's own arithmetic then gives the bid's Dye lines,
    kit count and Joint Filler line to the cent. The material block moves by exactly what the
    bid's lines moved by, and the tab total is recomputed from it.

    Mutation: leave C29 at the template's 500; hardcode "/3500" in the formula; have the writer
    turn the formula into text."""
    e, s = ran["edited"], ran["seeded"]
    assert e["cells"]["Polish!C25"] == pytest.approx(0.2)
    assert e["cells"]["Polish!C26"] == pytest.approx(0.2)
    assert e["cells"]["Polish!C29"] == 650
    assert e["cells"]["Polish!B29"] == '=ROUNDUP(IF(E29="yes",(E18/2000),0),0)'
    edited, seeded = _workbook(e, cached), _workbook(s, cached)
    cell = edited.f["Polish"]["B29"]
    assert cell.data_type == "f" and cell.value == e["cells"]["Polish!B29"], (
        "B29 did not reach the .xlsx as a formula: %r (%s)" % (cell.value, cell.data_type))
    assert edited.num("B29") == e["jfKits"] == 2
    assert edited.num("D25") + edited.num("D26") == pytest.approx(e["dyeCost"]), (
        edited.num("D25"), edited.num("D26"), e["dyeCost"])
    assert edited.num("D29") == pytest.approx(e["jfCost"]), (edited.num("D29"), e["jfCost"])
    moved = (e["dyeCost"] + e["jfCost"]) - (s["dyeCost"] + s["jfCost"])
    assert _material_raw(edited) - _material_raw(seeded) == pytest.approx(moved), moved
    assert edited.num("D82") > seeded.num("D82"), (edited.num("D82"), seeded.num("D82"))


@needs_node
def test_the_kits_waste_reaches_the_workbooks_kit_count(ran, cached):
    """3,200 SF is one kit at 3,500 SF a kit and two once 10% waste is bought. The template's
    formula would say one; the written one says two, the bid's own count.

    Mutation: leave waste out of the formula."""
    w = ran["wasted"]
    assert w["cells"]["Polish!B29"] == '=ROUNDUP(IF(E29="yes",(E18*(1+10/100)/3500),0),0)'
    book = _workbook(w, cached)
    assert book.num("B29") == w["jfKits"] == 2
    assert book.num("D29") == pytest.approx(w["jfCost"])


@needs_node
def test_a_missing_row_writes_nothing_and_touches_nothing(ran, cached):
    """A database the seed has not reached has no row: the save writes exactly what it wrote
    before, the template's cells stand, and a C25 somebody typed on the estimate grid survives.
    And the template's own two dye lines are what the bid's fallback charges -- both coats -- so
    even with no row the workbook's D25 + D26 is the bid's dye.

    Mutation: write the shipped figures when the row is missing, or delete keys; take the
    fallback back to one coat."""
    m = ran["missing"]
    assert m["cells"] == m["today"], {k: v for k, v in m["cells"].items()
                                      if m["today"].get(k) != v}
    assert m["cells"]["Polish!C25"] == 0.3
    assert "Polish!C29" not in m["cells"] and "Polish!B29" not in m["cells"]
    untyped = {k: v for k, v in m["cells"].items() if k != "Polish!C25"}
    book = _workbook(m, cached, cells=untyped)
    assert book.num("D25") + book.num("D26") == pytest.approx(m["dyeCost"]), (
        book.num("D25"), book.num("D26"), m["dyeCost"])
    assert book.num("D29") == pytest.approx(m["jfCost"])


@needs_node
def test_a_row_that_cannot_price_writes_the_figures_the_bid_fell_back_to(ran, cached):
    """A row that is there but cannot price (coverage or cost blanked) makes the bid fall back to
    the shipped formula; the workbook gets the same shipped figures, so the two still agree."""
    b = ran["blanked"]
    assert b["cells"]["Polish!C25"] == pytest.approx(0.14)
    assert b["cells"]["Polish!C26"] == pytest.approx(0.14)
    assert b["cells"]["Polish!C29"] == 500
    assert "Polish!B29" not in b["cells"]
    book = _workbook(b, cached)
    assert book.num("D25") + book.num("D26") == pytest.approx(b["dyeCost"])
    assert book.num("D29") == pytest.approx(b["jfCost"])


@needs_node
def test_putting_the_coverage_back_replaces_the_formula_an_earlier_save_wrote(ran):
    """An earlier save wrote B29 with a 2,000 coverage. The row is back at 3,500: the next save
    puts the template's own formula back (and the kit's price) instead of leaving the stale one to
    quote a different kit count from the bid.

    Mutation: write B29 only when it differs from the template, never when the key is there."""
    r = ran["reset"]
    assert r["today"]["Polish!B29"] == ran["staleB29"]
    assert r["cells"]["Polish!B29"] == TEMPLATE_B29
    assert r["cells"]["Polish!C29"] == 500


# Rows the seeded pair never exercises: a pack bigger than one, no roundup, and a dye row that
# rounds up or covers more than a square foot a unit. Each is priced by the REAL priceLine and
# written by the REAL conditionCellWrites under node, then read back out of the workbook the real
# writer produced.
BRANCHES = {
    "kit_pack_of_5": {"id": "joint-filler-kit", "unit": "Kit", "buy_qty": 5, "unit_cost": 400,
                      "coverage": 2000, "waste_pct": 0, "roundup": True},
    "kit_no_roundup": {"id": "joint-filler-kit", "unit": "Kit", "buy_qty": 1, "unit_cost": 650,
                       "coverage": 2000, "waste_pct": 0, "roundup": False},
    "dye_rounds_up": {"id": "dye", "unit": "Pail", "buy_qty": 4, "unit_cost": 0.8,
                      "coverage": 2, "waste_pct": 5, "roundup": True},
    "dye_coverage_2": {"id": "dye", "unit": "Gal", "buy_qty": 1, "unit_cost": 0.2,
                       "coverage": 2, "waste_pct": 5, "roundup": False},
}
_NODE = r"""
const B = require(process.argv[1] + '/js/bid-model.js');
const L = require(process.argv[1] + '/js/library-core.js');
const rows = JSON.parse(process.argv[2]), area = Number(process.argv[3]), out = {};
for (const name of Object.keys(rows)) {
  const row = rows[name], key = row.id === 'dye' ? 'dye' : 'joint_filler';
  const p = L.priceLine({ item_id: row.id }, [row], area);
  const lib = {};
  lib[key] = { unit_price: p.unit_price, coverage: p.coverage, waste_pct: p.waste_pct,
               roundup: p.roundup, buy_qty: p.buy_qty };
  // The row is ONE coat; a bid buys B.DYE_COATS of them (Kyle's rows 25 and 26).
  out[name] = { key: key, cost: p.cost * (key === 'dye' ? B.DYE_COATS : 1),
                cells: B.conditionCellWrites({ dye: true, joint_filler: true }, {}, lib) };
}
console.log(JSON.stringify(out));
"""


# Two workbooks carry the four shapes -- a dye row and a kit row each -- because the lines are
# separate cells and every workbook costs a template parse and a save.
PAIRS = (("dye_rounds_up", "kit_pack_of_5"), ("dye_coverage_2", "kit_no_roundup"))


@needs_node
@pytest.mark.parametrize("pair", PAIRS, ids=["+".join(p) for p in PAIRS])
def test_every_row_shape_prices_the_same_in_the_workbook(pair, cached):
    """Whole packs, fractional kits, dye bought by the pail: whatever the row says, the line the
    workbook's own formulas compute is the line priceLine priced.

    Mutation: drop the pack from the kit formula, or the waste from dye's per-SF rate."""
    area = 3501
    proc = subprocess.run(["node", "-e", _NODE, str(FRONTEND), json.dumps(BRANCHES), str(area)],
                          capture_output=True, text=True, encoding="utf-8", timeout=60)
    assert proc.returncode == 0, proc.stderr
    got = json.loads(proc.stdout.strip().splitlines()[-1])
    cells = {}
    for name in pair:
        cells.update(got[name]["cells"])
    book = _workbook({"cells": cells, "area": area}, cached)
    for name in pair:
        lines = ("D25", "D26") if got[name]["key"] == "dye" else ("D29",)
        got_cost = sum(book.num(line) for line in lines)
        assert got_cost == pytest.approx(got[name]["cost"]), (
            name, lines, got_cost, got[name]["cost"],
            {k: v for k, v in cells.items() if k[7:9] in ("B2", "C2")})


def test_a_named_field_on_the_same_cell_loses_to_the_saved_cell(cached):
    """POLISH_CELL_MAP names Polish!C25 "dye_cost_1" and Polish!C29 "joint_filler_qty" -- the
    second is misnamed, C29 is the kit's COST and the quantity is B29. No live caller sends either
    key, and if one did, fill_estimate applies `cell_values` after the named fields, so the figure
    the page saved is the one that lands."""
    assert ew.POLISH_CELL_MAP["dye_cost_1"] == "C25"
    assert ew.POLISH_CELL_MAP["joint_filler_qty"] == "C29"
    with warnings.catch_warnings():
        warnings.simplefilter("ignore")
        wb = openpyxl.load_workbook(io.BytesIO(ew.fill_estimate(
            {"dye_cost_1": 9.99, "joint_filler_qty": 999},
            {"Polish!C25": 0.2, "Polish!C29": 650})))
    assert wb["Polish"]["C25"].value == 0.2
    assert wb["Polish"]["C29"].value == 650


def test_the_held_rows_read_none_of_the_written_cells():
    """Book holds Polish rows 34-63 at the template's cached figures. That is exact only if none
    of them reads a cell this change writes or a line it moves -- checked on the template's own
    formula text, so a future formula that did would fail here rather than hide a difference."""
    with warnings.catch_warnings():
        warnings.simplefilter("ignore")
        ws = openpyxl.load_workbook(ew.TEMPLATE_PATH)["Polish"]
    moved = re.compile(r"(?<![A-Z$])\$?(B|C|D)\$?(25|26|29)(?![0-9])"
                       r"|(?<![A-Z$])\$?D\$?(31|33)(?![0-9])")
    for row in ws.iter_rows(min_row=34, max_row=63):
        for c in row:
            if isinstance(c.value, str) and c.value.startswith("="):
                assert not moved.search(c.value), (c.coordinate, c.value)


@pytest.fixture(scope="module")
def cov():
    """The B2 scenarios -- coverage typed on THIS bid -- from the same real-page harness run."""
    proc = subprocess.run(["node", str(HARNESS), str(FRONTEND)], capture_output=True, text=True,
                          encoding="utf-8", timeout=180)
    assert proc.returncode == 0, proc.stderr
    return json.loads(proc.stdout.strip().splitlines()[-1])["condCoverage"]


@needs_node
def test_a_typed_coverage_reaches_the_workbook_the_same_figures(cov, cached):
    """The download must match the screen. The kit's coverage typed on the bid (3,000 over a
    library 2,000) is the divisor in the B29 formula the save writes, and the workbook's own
    arithmetic gives the bid's kit count and line; the dye coverage typed on the bid lands in both
    coat rates (price over coverage).

    Mutation: conditionLibrary reads the library row instead of the priced (typed) line."""
    jf = cov["libJf"]
    assert jf["cells"]["Polish!B29"] == '=ROUNDUP(IF(E29="yes",(E18/3000),0),0)'
    book = _workbook(jf, cached)
    assert book.num("B29") == jf["jfKits"] == 2
    assert book.num("D29") == pytest.approx(jf["jfCost"])
    dye = cov["libDye"]
    assert dye["cells"]["Polish!C25"] == pytest.approx(0.05)
    assert dye["cells"]["Polish!C26"] == pytest.approx(0.05)
    book = _workbook(dye, cached)
    assert book.num("D25") + book.num("D26") == pytest.approx(dye["dyeCost"])
