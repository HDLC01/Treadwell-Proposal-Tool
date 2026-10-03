"""The NOTES box on page 1: where it prints, where the editor draws it, and what size its lines are.

Hanz, 2026-10-02, on staging, a GC epoxy draft: the first note sat above the box's top rule, every
bullet was drawn left of the box with the text running over the rotated "NOTES" label, and one note
("Treadwell is not responsible for wear & tear...") was drawn much larger than the rest. None of it
came from the release under review; all of it was on production. Three causes, each fixed and each
proven here by running the real code:

  1. THE PDF DRIFTED UP A PAGE. Page 1's boxes hang off empty Cambria paragraphs. Word lays each at
     Cambria's 14.07pt; LibreOffice, with Caladea standing in, at 13.8pt -- so every box anchored n
     paragraphs down printed n x 0.27pt higher than Kyle placed it (the GC/Gyp NOTES box, paragraph
     31, by 8.3pt: its first note above the rule). `_pin_anchor_line_heights` gives those paragraphs
     an at-least line Word already exceeds, and the editor's estimate is that same line.
  2. THE EDITOR IGNORED A WRAP. The GC and Gyp NOTES paragraph runs beside the REGARDS box (square
     wrap), so both renderers start its column after that box: 36.3pt (GC) / 39pt (Gyp) right of
     the margin arithmetic the editor used (`_wrap_column_shift`).
  3. A PLAIN EDIT LOOKED LIKE A CHOSEN SIZE. The editor sends every run's size back as it reads it
     off the page, so editing a NOTES line's words arrived with Kyle's own 7.5pt, and the shrink
     left that line at 7.5pt among 4.5pt neighbours. A size equal to the template's is not a choice.

The positions below were measured on 2026-10-03 from real PDFs: the production LibreOffice image
(backend/ops/terms_render_proof.py's container) rendering these templates filled by this code, each
text box given a visible outline. Word's own PDF export of the same files agrees within 0.55pt, and
Word's output did not change by a point with the line pin in place.
"""
import copy
import io
import re
from unittest import mock

import docx
import pytest
from docx import Document
from docx.oxml.ns import qn

import main
import proposal_writer as pw
from test_editor_fit_parity import (FRONTEND, TEMPLATES, _V, _fit, _docx, _box_runs, _harness, _template,
                                    needs_node)

# {(work_type, audience): {box id: (x_pt, y_pt) LibreOffice prints the box at}}
PRINTED = {
    ("epoxy", "Direct"): {0: (125.2, 36.0), 1: (18.35, 36.0), 2: (162.35, 152.9), 3: (162.25, 495.0),
                          4: (162.25, 321.35), 5: (23.3, 503.45)},
    ("polish", "Direct"): {0: (125.2, 36.0), 1: (18.35, 36.0), 2: (27.35, 504.0), 3: (162.0, 152.9),
                           4: (162.25, 321.35), 5: (162.3, 491.15)},
    ("combo", "Direct"): {0: (125.2, 36.0), 1: (18.35, 36.0), 2: (27.35, 504.0), 3: (162.0, 152.9),
                          4: (162.3, 492.95), 5: (162.25, 321.35)},
    ("polish", "GC"): {0: (125.2, 36.0), 1: (18.35, 36.0), 2: (27.35, 504.0), 3: (160.7, 153.85),
                       4: (161.15, 332.25), 5: (161.3, 509.1)},
    ("epoxy", "GC"): {0: (125.2, 36.0), 1: (18.35, 36.0), 2: (27.35, 504.0), 3: (161.8, 153.4),
                      4: (161.75, 332.8), 5: (161.95, 508.95)},
    ("sealer", "GC"): {0: (125.2, 36.0), 1: (18.35, 36.0), 2: (27.35, 504.0), 3: (160.7, 155.6),
                       4: (160.65, 329.85), 5: (160.65, 508.15)},
    ("gyp", "Direct"): {0: (125.2, 36.0), 1: (18.35, 36.0), 2: (161.35, 153.2), 3: (161.95, 361.1),
                        4: (27.0, 503.45), 5: (161.5, 511.2)},
}
# Where the artwork's NOTES rule ends (its lower edge), and the right edge of the label strip.
NOTES_RULE_BOTTOM = {("epoxy", "Direct"): 488.05, ("polish", "Direct"): 489.45, ("combo", "Direct"): 489.45,
                     ("polish", "GC"): 509.6, ("epoxy", "GC"): 509.6, ("sealer", "GC"): 509.6,
                     ("gyp", "Direct"): 509.6}
LABEL_STRIP_RIGHT = 156.0


def _geometry(wt, aud):
    return pw.template_geometry(docx.Document(str(pw.pick_template(wt, aud))))["boxes"]


def _notes_box_id(wt, aud):
    """The box holding the NOTES text: the one lowest-down 400pt-wide box."""
    return max((b for b in _geometry(wt, aud) if b["w_pt"] > 400), key=lambda b: b["y_pt"])["id"]


@pytest.mark.parametrize("wt,aud", sorted(PRINTED))
def test_the_editor_draws_every_page_one_box_where_the_pdf_prints_it(wt, aud):
    """Every box the editor places (what /api/proposal-template serves it) within 0.2pt of the
    rectangle the PDF prints, both axes. Before: the GC/Gyp NOTES box 36-39pt left, and every box
    hung off a deep paragraph up to 8pt off, in one direction or the other."""
    for b in _geometry(wt, aud):
        x, y = PRINTED[(wt, aud)][b["id"]]
        assert b["x_pt"] == pytest.approx(x, abs=0.2), (wt, aud, b["id"], "x", b["x_pt"], x)
        assert b["y_pt"] == pytest.approx(y, abs=0.2), (wt, aud, b["id"], "y", b["y_pt"], y)


@pytest.mark.parametrize("wt,aud", sorted(PRINTED))
def test_the_notes_box_starts_inside_its_frame(wt, aud):
    """Right of the label strip, and its top not above the artwork's rule by more than the rule is
    thick (the GC/Gyp box top meets the rule; Kyle placed it there in Word)."""
    b = next(g for g in _geometry(wt, aud) if g["id"] == _notes_box_id(wt, aud))
    assert b["x_pt"] > LABEL_STRIP_RIGHT, (wt, aud, b)
    assert b["y_pt"] >= NOTES_RULE_BOTTOM[(wt, aud)] - 1.5, (wt, aud, b)


def _filled(wt, aud):
    return Document(io.BytesIO(pw.fill_proposal(work_type=wt, audience=aud, values=dict(_V),
                                                notes=main._notes_for(wt, []))))


def _deepest_box_paragraph(d):
    body = d.element.body
    tops = [c for c in body if c.tag == qn("w:p")]
    deepest = -1
    for tx in pw._iter_txbx(d):
        a = pw._txbx_anchor(tx)
        if a is None or pw._anchor_offset(a, "positionV")[1] != "paragraph":
            continue
        deepest = max(deepest, tops.index(pw._top_level_of(body, a)))
    return tops, deepest


@pytest.mark.parametrize("wt,aud", sorted(PRINTED))
def test_the_paragraphs_page_one_hangs_off_print_at_the_editors_line(wt, aud):
    """In the generated document, every paragraph above the deepest box carries an AT-LEAST line of
    exactly the editor's estimate (`_ANCHOR_LINE_H_PT`), and nothing from that box down is touched,
    so the Terms pages are not. At-least: Word's taller Cambria line wins there, unchanged."""
    d = _filled(wt, aud)
    tops, deepest = _deepest_box_paragraph(d)
    assert deepest >= 5, "no page-1 box found"
    assert pw._ANCHOR_LINE_TW / 20.0 == pw._ANCHOR_LINE_H_PT
    for i, p in enumerate(tops):
        sp = p.find(qn("w:pPr") + "/" + qn("w:spacing"))
        line = sp.get(qn("w:line")) if sp is not None else None
        if i < deepest:
            assert (line, sp.get(qn("w:lineRule"))) == (str(pw._ANCHOR_LINE_TW), "atLeast"), (wt, aud, i)
    def pinned(p):
        sp = p.find(qn("w:pPr") + "/" + qn("w:spacing"))
        return (sp is not None and sp.get(qn("w:line")) == str(pw._ANCHOR_LINE_TW)
                and sp.get(qn("w:lineRule")) == "atLeast")
    assert not [p for p in tops[deepest:] if pinned(p)], (wt, aud)


# ── 3. a plain edit on a NOTES line shrinks with the lines around it ────────────────────────────
WEAR = "Treadwell is not responsible for wear"


def _gc_body(overrides):
    tpl = _template("epoxy", "GC")
    return {"work_type": "epoxy", "audience": "GC", "template_version": tpl["template_version"],
            "values": dict(_V), "paragraph_overrides": overrides}


def _wear_block():
    return next(b for b in _template("epoxy", "GC")["blocks"] if str(b["text"]).startswith(WEAR))


def _as_the_editor_sends(block, edit=None):
    """The block's runs as the page serializes them after a text edit: every run carries the size
    it is drawn at, which is the template's own (fmtAt reads the span's inline size)."""
    runs = [{"text": r["text"], "size_pt": r["size_pt"]} for r in block["runs"]]
    if edit:
        runs[0]["text"] = runs[0]["text"].replace("Treadwell", edit, 1)
    return {"id": block["id"], "text": "".join(r["text"] for r in runs), "runs": runs}


def test_a_plain_edit_on_a_notes_line_shrinks_with_its_neighbours():
    blk = _wear_block()
    box = blk["txbx"]
    plain = _gc_body([])
    edited = _gc_body([_as_the_editor_sends(blk, edit="TREADWELL")])
    f = next(b for b in _fit(edited) if b["id"] == box)
    assert f["scale"] < 0.999, "the GC NOTES box no longer shrinks, so this proves nothing"
    assert f["exempt"] == [], "a size equal to the template's was taken for the estimator's choice"
    # The document prints the edited line at the same scaled sizes as the unedited one.
    want = [hp for hp, t in _box_runs(_docx(plain))[box] if t and t.startswith(("Treadwell is not", "&", " tear",
                                                                                "ng from", "See Terms"))]
    got = [hp for hp, t in _box_runs(_docx(edited))[box] if t and t.startswith(("TREADWELL is not", "&", " tear",
                                                                                 "ng from", "See Terms"))]
    assert got and got == want and max(got) < 14, (got, want)


def test_a_size_the_estimator_did_choose_on_a_notes_line_is_still_left_alone():
    blk = _wear_block()
    ov = _as_the_editor_sends(blk)
    ov["runs"][0]["size_pt"] = 9                       # not the template's 7.5pt: a real choice
    f = next(b for b in _fit(_gc_body([ov])) if b["id"] == blk["txbx"])
    assert f["exempt"] == [blk["id"]]
    runs = _box_runs(_docx(_gc_body([ov])))[blk["txbx"]]
    assert [hp for hp, t in runs if t.startswith("Treadwell is not")] == [18]


@needs_node
def test_the_editor_draws_the_plainly_edited_notes_line_at_its_neighbours_size():
    """The route's real answer for that payload, handed to the page's real applyBoxFit: the edited
    line's spans get the same printed size as the line above it. Before: none (exempt), so the
    page showed it at 7.5pt among 4.5pt lines."""
    tpl = _template("epoxy", "GC")
    blk = _wear_block()
    body = _gc_body([_as_the_editor_sends(blk, edit="TREADWELL")])
    rep = {b["id"]: b for b in _fit(body)}
    case = _harness([{"name": "wear", "blocks": tpl["blocks"],
                      "fit": {str(k): v for k, v in rep.items()}, "tokens": body["values"]}])[0]
    got = {b["id"]: b for b in case["blocks"]}
    prev = max(i for i in got if i < blk["id"] and got[i]["txbx"] == blk["txbx"] and got[i]["spans"])
    mine = {s["printed"] for s in got[blk["id"]]["spans"] if s["printed"]}
    theirs = {s["printed"] for s in got[prev]["spans"] if s["printed"]}
    assert mine and mine <= theirs | {"4pt", "4.5pt"}, (mine, theirs)


@needs_node
def test_a_shrunk_notes_line_keeps_its_paragraph_mark_on_screen():
    """The writer scales runs, never the paragraph mark, and the mark (with the list bullet drawn at
    its size) sets how tall the line prints: 8.7pt apart for the GC's 4.5pt notes in the PDF. The
    page marks every shrunk paragraph that states its line spacing (`data-tw-mark`), which
    styles.css turns into a mark-sized strut; an unshrunk box and an exempt line get none."""
    tpl = _template("epoxy", "GC")
    body = _gc_body([])
    rep = {b["id"]: b for b in _fit(body)}
    blk = _wear_block()
    case = _harness([{"name": "marks", "blocks": tpl["blocks"],
                      "fit": {str(k): v for k, v in rep.items()}, "tokens": body["values"]}])[0]
    notes = [b for b in case["blocks"] if b["txbx"] == blk["txbx"]]
    stated = {b["id"] for b in tpl["blocks"] if b["txbx"] == blk["txbx"]
              and (b.get("para") or {}).get("spacing", {}).get("line")}
    assert stated, "no NOTES paragraph states its line spacing"
    assert {b["id"] for b in notes if b.get("twMark")} == {b["id"] for b in notes if b["id"] in stated and b["printed"]}
    assert all(not b.get("twMark") for b in case["blocks"] if rep[b["txbx"]]["scale"] >= 0.999)


# ══ FOLLOW-UPS from the review of the fix above (2026-10-03) ═══════════════════════════════════════
# 1. the size rule matched runs by CHARACTER OFFSET, so an edit that changed the line's length put a
#    later run on its neighbour's size and the line printed large again;
# 2. (editor) an untouched paragraph holding a no-break space read as edited -- test_doc_editor_fidelity.py;
# 3. the mark-sized strut and the pin guard had no test that could fail;
# 4. `template_geometry` re-read the whole body's section map for every pair of anchors.
#
# ── the size rule: a size is the estimator's choice only when it is none of the line's own ─────────
_WEAR_EDITS = ["delete a word", "responsible -> liable", "is not -> isn't", "insert text",
               "reword longer", "delete the last word"]


def _wear_edit(kind):
    """The wear & tear note's runs as the editor sends them after ONE edit to its words: every run at
    the size the page draws it at (the template's own: 7.5pt x4, then 7.0pt), the text changed. Under
    the first fix each of these moved a size boundary and the line printed large again."""
    blk = _wear_block()
    runs = [{"text": r["text"], "size_pt": r["size_pt"]} for r in blk["runs"]]

    def sub(old, new):
        hit = [r for r in runs if old in r["text"]]
        assert len(hit) == 1, (old, [r["text"] for r in runs])
        hit[0]["text"] = hit[0]["text"].replace(old, new, 1)

    if kind == "delete a word":
        sub("existing ", "")
    elif kind == "responsible -> liable":
        sub("responsible", "liable")
    elif kind == "is not -> isn't":
        sub("is not", "isn't")
    elif kind == "insert text":
        runs[0]["text"] = "TEST " + runs[0]["text"]
    elif kind == "reword longer":
        sub("wear", "general wear and ordinary")
    elif kind == "delete the last word":
        sub("Conditions.", "")
    else:
        raise KeyError(kind)
    return {"id": blk["id"], "text": "".join(r["text"] for r in runs), "runs": runs}


@pytest.mark.parametrize("edit", _WEAR_EDITS)
def test_an_edit_that_changes_the_length_still_shrinks_with_its_neighbours(edit):
    """Block 161's runs are 7.5pt x4 then 7.0pt, ending at characters 38/39/73/104/127. The first fix
    matched a run to the template run at the SAME OFFSET, so an edit that moved a boundary put the
    7.0pt tail on a 7.5pt offset, took it for the estimator's choice and left the line out of the
    shrink: the PDF and the editor printed it at 7.5/7.0pt beside 4.5pt notes. Now a size is the
    estimator's only when it is none of the line's own, wherever in the line it sits: every run of the
    box prints at the size the unedited document gives it, in the document and (the route is the
    editor's source) in /api/proposal-fit.
    Mutation: the first fix's offset-matching rule back in `_set_paragraph_runs`."""
    blk = _wear_block()
    box = blk["txbx"]
    edited = _gc_body([_wear_edit(edit)])
    f = next(b for b in _fit(edited) if b["id"] == box)
    assert f["scale"] < 0.999 and f["exempt"] == [], f
    assert ([hp for hp, _t in _box_runs(_docx(edited))[box]]
            == [hp for hp, _t in _box_runs(_docx(_gc_body([])))[box]])


def test_a_size_the_estimator_chose_survives_an_edit_that_also_changed_the_length():
    """The other half of the rule: a size the line does not use is a choice, however the words around
    it were edited -- the tail made 9pt (neither the template's 7.5 nor its 7.0) after a word was
    deleted ahead of it. Mutation: `user_sized` never set (every size is taken for the template's)."""
    blk = _wear_block()
    ov = _wear_edit("delete a word")
    ov["runs"][-1]["size_pt"] = 9
    body = _gc_body([ov])
    f = next(b for b in _fit(body) if b["id"] == blk["txbx"])
    assert f["exempt"] == [blk["id"]]
    runs = _box_runs(_docx(body))[blk["txbx"]]
    assert [hp for hp, t in runs if t.startswith("See Terms")] == [18]


@needs_node
def test_the_editor_draws_an_edit_that_changed_the_length_at_its_neighbours_size():
    """The editor takes its exemptions from the route, so fixing the writer fixes the page. The route's
    real answer for the 'delete a word' payload, handed to the page's real applyBoxFit, draws the
    edited line at the size the line above it prints at; exempt, it drew nothing here (7.5pt among
    4.5pt lines on screen)."""
    tpl = _template("epoxy", "GC")
    blk = _wear_block()
    body = _gc_body([_wear_edit("delete a word")])
    rep = {b["id"]: b for b in _fit(body)}
    assert rep[blk["txbx"]]["exempt"] == []
    case = _harness([{"name": "wear", "blocks": tpl["blocks"],
                      "fit": {str(k): v for k, v in rep.items()}, "tokens": body["values"]}])[0]
    got = {b["id"]: b for b in case["blocks"]}
    prev = max(i for i in got if i < blk["id"] and got[i]["txbx"] == blk["txbx"] and got[i]["spans"])
    mine = {s["printed"] for s in got[blk["id"]]["spans"] if s["printed"]}
    theirs = {s["printed"] for s in got[prev]["spans"] if s["printed"]}
    assert mine and mine <= theirs | {"4pt", "4.5pt"}, (mine, theirs)


def _free_box_paragraphs(d):
    for idx, kind, p, in_block, _t, tb in pw.iter_editable_blocks(d):
        if tb is not None and in_block is None and kind == "p":
            yield idx, p


def _chars_of(p):
    """[(character, half-points or None)] of a paragraph's text runs: what the editor reads back."""
    media = (qn("w:drawing"), qn("w:pict"), qn("w:object"))
    out = []
    for r in p.findall(qn("w:r")):
        if any(next(r.iter(t), None) is not None for t in media):
            continue
        sz = r.find(qn("w:rPr") + "/" + qn("w:sz"))
        v = sz.get(qn("w:val")) if sz is not None else None
        hp = int(v) if v and v.isdigit() else None
        for t in r.iter(qn("w:t")):
            out.extend((ch, hp) for ch in (t.text or ""))
    return out


def _runs_of(chars):
    """The editor's runs for these characters: maximal stretches of one size, at the size it draws."""
    runs = []
    for ch, hp in chars:
        if runs and runs[-1][1] == hp:
            runs[-1][0].append(ch)
        else:
            runs.append(([ch], hp))
    return [dict({"text": "".join(cs)}, **({"size_pt": hp / 2.0} if hp is not None else {}))
            for cs, hp in runs]


def _edits_to_the_words(chars):
    """Every kind of edit to a line's words, as characters. A character typed in takes the size of the
    one to its left (the browser extends the span the caret is in); a deleted stretch takes its size
    with it. Placed where the sizes change, which is where an offset-matching rule goes wrong."""
    n = len(chars)
    bounds = [i for i in range(1, n) if chars[i][1] != chars[i - 1][1]]
    for i in sorted({0, n // 2, n - 1} | set(bounds) | {b - 1 for b in bounds}):
        yield "delete one character at %d" % i, chars[:i] + chars[i + 1:]
    text = "".join(c for c, _ in chars)
    for m in re.finditer(r"\S+ *", text):
        a, z = m.start(), m.end()
        if not (a == 0 or z == n or any(a <= b <= z for b in bounds)):
            continue
        left = chars[a][1]
        yield "delete the word at %d" % a, chars[:a] + chars[z:]
        yield "shorten the word at %d" % a, chars[:a] + [("a", left)] + chars[z:]
        yield "lengthen the word at %d" % a, chars[:a] + [(c, left) for c in "approximately "] + chars[z:]
    for i in sorted({0, n // 2, n} | set(bounds)):
        left = chars[i - 1][1] if i else chars[0][1]
        yield "insert at %d" % i, chars[:i] + [(c, left) for c in "TEST "] + chars[i:]
    starts = [0] + bounds + [n]
    for a, z in zip(starts, starts[1:]):
        if z - a < n:
            yield "delete the stretch at %d" % a, chars[:a] + chars[z:]


@pytest.mark.parametrize("wt,aud", TEMPLATES)
def test_no_edit_to_the_words_of_any_text_box_line_makes_it_the_estimators_size(wt, aud):
    """The review's scan, as a test, over EVERY free line of EVERY text box of the template (WORK, PRICE,
    NOTES, the date box): delete a character at each place the size changes, delete / shorten / lengthen
    the words there, insert text, delete a whole stretch -- none of it is a size the estimator chose, so
    `_set_paragraph_runs` reports nothing to exempt. And a size the line does not use, on the whole line,
    on a stretch or on one letter, is a choice and is reported. (Under the first fix this scan exempted
    730 of its 3,140 edits across the eight files -- epoxy GC lines 112/114/115/161, polish GC
    113/115/138/139/163, sealer GC 113/136/162, gyp 132/134/136/142; every one was a plain edit.)
    Mutation: the first fix's offset-matching rule back."""
    d = docx.Document(str(pw.pick_template(wt, aud)))
    scanned = chosen = 0
    for idx, p in _free_box_paragraphs(d):
        chars = _chars_of(p)
        if len(chars) < 3 or all(hp is None for _c, hp in chars):
            continue
        own = {hp for _c, hp in chars if hp is not None}
        for name, new in _edits_to_the_words(chars):
            assert pw._set_paragraph_runs(copy.deepcopy(p), _runs_of(new)) is False, (wt, aud, idx, name)
            scanned += 1
        pick = next(hp for hp in (18, 20, 22, 24, 16, 14) if hp not in own)
        mid = len(chars) // 2
        for name, new in (("the whole line", [(c, pick) for c, _h in chars]),
                          ("a stretch", chars[:mid] + [(c, pick) for c, _h in chars[mid:mid + 3]]
                           + chars[mid + 3:]),
                          ("the first letter", [(chars[0][0], pick)] + chars[1:])):
            assert pw._set_paragraph_runs(copy.deepcopy(p), _runs_of(new)) is True, (wt, aud, idx, name, pick)
            chosen += 1
    assert scanned > 100 and chosen >= 30, (scanned, chosen)      # the scan really ran


def test_a_run_with_no_words_does_not_count_as_a_size_the_line_uses():
    """Word leaves empty runs behind with sizes of their own (a GC PRICE row carries a dozen at 9pt and
    8.5pt after its 10pt text). The editor draws and reads back only runs with words, so those sizes are
    never "the line's own": the estimator picking 9pt on that row has chosen it, and is not overruled.
    Mutation: collect the sizes of every run, empty ones too."""
    d = docx.Document(str(pw.pick_template("epoxy", "GC")))
    row = next(p for _i, p in _free_box_paragraphs(d) if "{{base_bid_formatted}}" in pw._own_text(p))
    empties = {sz.get(qn("w:val")) for r in row.findall(qn("w:r")) if not "".join(t.text or "" for t in r.iter(qn("w:t")))
               for sz in r.findall(qn("w:rPr") + "/" + qn("w:sz"))}
    chars = _chars_of(row)
    own = {hp for _c, hp in chars if hp is not None}
    assert "18" in empties and 18 not in own, (empties, own)
    assert pw._set_paragraph_runs(copy.deepcopy(row), [{"text": "$1 -- a price", "size_pt": 9}]) is True


# The positions `template_geometry` returned at f7b4992, for every box and every picture, as exact reprs.
# (x_pt, y_pt, w_pt, h_pt) per box; (media name, x_pt, y_pt, w_pt, h_pt, para_index) per picture.
GEOMETRY_AT_F7B4992 = {
    ("epoxy", "Direct"): {
        "boxes": [
            (125.2, 36.0, 324.8, 99.0),
            (18.349999999999994, 36.0, 72.0, 18.0),
            (162.35, 152.9, 423.0, 171.0),
            (162.3, 495.0, 422.65, 162.0),
            (162.3, 321.35, 422.65, 164.5),
            (23.31425196850394, 503.45, 90.0, 90.0),
        ],
        "images": [
            ('image1.png', 0.0, -0.6857480314960611, 612.0, 791.0, 0),
            ('image2.png', 0.09999999999999432, 646.8325196850394, 612.0, 791.0, 46),
            ('image2.png', 0.09999999999999432, 885.6825196850394, 612.0, 791.0, 63),
            ('image2.png', 0.09999999999999432, 1110.5325196850395, 612.0, 791.0, 79),
        ],
    },
    ("polish", "Direct"): {
        "boxes": [
            (125.2, 36.0, 324.8, 99.0),
            (18.349999999999994, 36.0, 72.0, 18.0),
            (27.35, 504.0, 90.0, 90.0),
            (162.0, 152.9, 423.0, 171.0),
            (162.3, 321.35, 422.65, 173.5),
            (162.35, 491.21425196850396, 422.65, 172.85),
        ],
        "images": [
            ('image1.png', 0.0, 0.6857480314960611, 612.0, 791.0, 0),
            ('image2.png', 0.09999999999999432, 631.9325196850393, 612.0, 791.0, 45),
            ('image2.png', 0.06748031496063334, 871.65, 612.0, 791.0, 62),
            ('image2.png', 0.06748031496063334, 1096.5, 612.0, 791.0, 78),
        ],
    },
    ("combo", "Direct"): {
        "boxes": [
            (125.2, 36.0, 324.8, 99.0),
            (18.349999999999994, 36.0, 72.0, 18.0),
            (27.35, 504.0, 90.0, 90.0),
            (162.0, 152.9, 423.0, 171.0),
            (162.35, 493.0, 422.65, 172.85),
            (162.3, 321.35, 422.65, 173.5),
        ],
        "images": [
            ('image1.png', 0.0, 0.6857480314960611, 612.0, 791.0, 0),
            ('image2.png', 0.09999999999999432, 631.9325196850393, 612.0, 791.0, 45),
            ('image2.png', 0.06748031496063334, 871.65, 612.0, 791.0, 62),
            ('image2.png', 0.06748031496063334, 1096.5, 612.0, 791.0, 78),
        ],
    },
    ("budget", "Direct"): {
        "boxes": [
            (125.2, 36.0, 324.8, 99.0),
            (18.349999999999994, 36.0, 72.0, 18.0),
            (27.35, 504.0, 90.0, 90.0),
            (162.0, 152.9, 423.0, 171.0),
            (162.35, 493.0, 422.65, 172.85),
            (162.3, 321.35, 422.65, 173.5),
        ],
        "images": [
            ('image1.png', 0.0, 0.6857480314960611, 612.0, 791.0, 0),
        ],
    },
    ("polish", "GC"): {
        "boxes": [
            (125.2, 36.0, 324.8, 99.0),
            (18.349999999999994, 36.0, 72.0, 18.0),
            (27.35, 504.0, 90.0, 90.0),
            (160.7, 153.9, 423.0, 182.15),
            (161.2, 332.25, 422.65, 177.1),
            (161.35, 509.0674803149606, 422.65, 189.0),
        ],
        "images": [
            ('image1.png', 0.0, 0.6857480314960611, 612.0095275590551, 791.0123622047244, 0),
            ('image2.png', 0.09999999999999432, 646.8340157480316, 612.0, 791.0, 46),
            ('image2.png', 0.09999999999999432, 884.9340157480316, 612.0, 791.0, 63),
            ('image2.png', 0.06598425196850144, 1110.5, 612.0, 791.0, 79),
        ],
    },
    ("epoxy", "GC"): {
        "boxes": [
            (125.2, 36.0, 324.8, 99.0),
            (18.349999999999994, 36.0, 72.0, 18.0),
            (27.35, 504.0, 90.0, 90.0),
            (161.8, 153.45, 423.0, 183.75),
            (161.8, 332.84999999999997, 422.65, 176.45),
            (162.0, 508.9628346456693, 422.65, 184.0),
        ],
        "images": [
            ('image1.png', 0.0, 0.6857480314960611, 612.0095275590551, 791.0123622047244, 0),
            ('image2.png', 0.09999999999999432, 646.8659842519686, 612.0, 791.0, 46),
            ('image2.png', 0.09999999999999432, 885.7840157480316, 612.0, 791.0, 63),
            ('image2.png', -0.7000000000000028, 1110.6159842519685, 612.0, 791.0, 79),
        ],
    },
    ("sealer", "GC"): {
        "boxes": [
            (125.2, 36.0, 324.8, 99.0),
            (18.349999999999994, 36.0, 72.0, 18.0),
            (27.35, 504.0, 90.0, 90.0),
            (160.7, 155.6, 423.0, 181.9),
            (160.7, 329.83724409448814, 422.65, 181.9),
            (160.7, 508.1627559055118, 422.65, 180.0),
        ],
        "images": [
            ('image1.png', 0.0, 0.6857480314960611, 612.0095275590551, 791.0123622047244, 0),
            ('image2.png', 0.06598425196850144, 646.0500000000001, 612.0, 791.0, 46),
            ('image2.png', 0.06598425196850144, 885.7, 612.0, 791.0, 63),
            ('image2.png', 0.09999999999999432, 1110.4840157480317, 612.0, 791.0, 79),
        ],
    },
    ("gyp", "Direct"): {
        "boxes": [
            (125.2, 36.0, 324.8, 99.0),
            (18.349999999999994, 36.0, 72.0, 18.0),
            (161.35, 153.2, 423.0, 222.2),
            (162.0, 361.1, 422.65, 154.4),
            (27.0, 503.45, 93.0, 90.0),
            (161.5, 511.25, 423.75, 185.25),
        ],
        "images": [
            ('image1.png', 0.0, 0.6857480314960611, 612.0095275590551, 791.012283464567, 0),
            ('image2.png', 0.0, 645.9825196850394, 612.0, 791.0, 46),
            ('image2.png', 0.0, 941.4749606299213, 612.0, 791.0, 67),
            ('image2.png', 0.0, 1166.535748031496, 611.9, 790.8707086614173, 83),
            ('image2.png', 1.5999999999999943, 1475.335748031496, 609.85, 792.8674803149606, 105),
        ],
    },
}


# ── the geometry the editor is served, pinned exactly ────────────────────────────────────────────────
@pytest.mark.parametrize("wt,aud", sorted(GEOMETRY_AT_F7B4992))
def test_every_templates_geometry_is_exactly_what_it_was_when_the_notes_box_fix_shipped(wt, aud):
    """Not within a tolerance: EXACTLY the numbers `template_geometry` returned at f7b4992 (the NOTES
    box fix), every box and every picture of all eight files, as `==` on the floats. The section-map
    change below touches the arithmetic every one of them goes through, and a position that moved by a
    rounding error would still be inside the 0.2pt of the printed-position test above.
    Mutation: `_section_ordinal` told the wrong map (a section break at every paragraph)."""
    got = pw.template_geometry(docx.Document(str(pw.pick_template(wt, aud))))
    want = GEOMETRY_AT_F7B4992[(wt, aud)]
    assert [(b["x_pt"], b["y_pt"], b["w_pt"], b["h_pt"]) for b in got["boxes"]] == want["boxes"]
    assert [(i["name"], i["x_pt"], i["y_pt"], i["w_pt"], i["h_pt"], i["para_index"])
            for i in got["images"]] == want["images"]


# ── 4. the section map is read once, not once per pair of anchors ────────────────────────────────────
def _section_map_reads(d):
    reads = []
    real = pw._section_break_indices

    def counting(body):
        reads.append(1)
        return real(body)

    with mock.patch.object(pw, "_section_break_indices", counting):
        pw.template_geometry(d)
    return len(reads)


def _add_boxes_the_others_wrap_beside(d, n):
    """n more copies of the box the NOTES paragraph wraps beside (the REGARDS box, square wrap)."""
    src = next(a for a in d.element.body.iter(qn("wp:anchor")) if a.find(qn("wp:wrapSquare")) is not None)
    for _ in range(n):
        src.getparent().append(copy.deepcopy(src))


@pytest.mark.parametrize("wt,aud", sorted(PRINTED))
def test_the_geometry_reads_the_section_map_a_fixed_number_of_times(wt, aud):
    """`_wrap_column_shift` asked `_section_ordinal` about every OTHER box for every box it placed, and
    each answer re-read the whole body's section map: 118 reads per `template_geometry` call on the GC
    resinous file (132 on gyp), 14ms became 91ms, and the route, the shrink and every generate pay it. Now the map is
    read a handful of times (the page metrics, once for the pass), and not a read more for the
    boxes added: ten more square-wrapped boxes, which multiplied the old count, change nothing.
    Mutation: `_section_ordinal` / `_pos_of_anchor` / `_wrap_column_shift` ignoring the map they are
    handed (the per-pair re-read back)."""
    d = docx.Document(str(pw.pick_template(wt, aud)))
    base = _section_map_reads(d)
    assert 1 <= base <= 8, base
    _add_boxes_the_others_wrap_beside(d, 10)
    assert _section_map_reads(d) == base


# ── the line-height pin leaves a line the template states alone ──────────────────────────────────────
def test_a_ruler_paragraph_that_states_its_own_line_spacing_is_left_alone():
    """`_pin_anchor_line_heights` gives every paragraph above the deepest page-1 box an at-least line,
    "unless it already states a line of its own". No shipped template has one that does, so that guard
    could be deleted with every test green. Here one ruler paragraph states a line of its own (360,
    auto) and another states only space before, in the same document: the first keeps its line
    exactly, the second gets the pin and keeps its space before, and the count is the paragraphs
    actually changed. Mutation: drop the guard (the stated line is overwritten with the pin); build the
    spacing fresh (the space before is lost)."""
    d = docx.Document(str(pw.pick_template("epoxy", "GC")))
    tops, deepest = _deepest_box_paragraph(d)
    assert deepest > 6
    own_line, only_before = tops[3], tops[5]
    sp = own_line.get_or_add_pPr().get_or_add_spacing()
    sp.set(qn("w:line"), "360")
    sp.set(qn("w:lineRule"), "auto")
    only_before.get_or_add_pPr().get_or_add_spacing().set(qn("w:before"), "120")

    def spacing(p):
        return p.find(qn("w:pPr") + "/" + qn("w:spacing"))

    assert [p for p in tops[:deepest] if spacing(p) is not None and spacing(p).get(qn("w:line"))] == [own_line]
    n = pw._pin_anchor_line_heights(d)
    assert (spacing(own_line).get(qn("w:line")), spacing(own_line).get(qn("w:lineRule"))) == ("360", "auto")
    s = spacing(only_before)
    assert (s.get(qn("w:before")), s.get(qn("w:line")), s.get(qn("w:lineRule"))) == (
        "120", str(pw._ANCHOR_LINE_TW), "atLeast")
    assert all(spacing(p).get(qn("w:line")) == str(pw._ANCHOR_LINE_TW)
               for p in tops[:deepest] if p is not own_line)
    assert n == deepest - 1, (n, deepest)


# ── the mark-sized strut: the stylesheet rule and the hook the page sets are one thing ───────────────
def _strut_rules():
    """The stylesheet rules (comments stripped) keyed on the `data-tw-mark` hook: [(selector, {property:
    value})]."""
    css = re.sub(r"/\*.*?\*/", "", (FRONTEND / "styles.css").read_text(encoding="utf-8"), flags=re.S)
    out = []
    for m in re.finditer(r"(?P<sel>[^{}]+)\{(?P<body>[^{}]*)\}", css):
        if "[data-tw-mark]" in m.group("sel"):
            decl = dict((k.strip(), v.strip()) for k, v in
                        (d.split(":", 1) for d in m.group("body").split(";") if d.strip()))
            out.append((" ".join(m.group("sel").split()), decl))
    return out


def _selector_matches(selector, el):
    """Does this selector match the paragraph `el` the page built? Exactly the shapes the strut rule is
    written in: a chain of compound selectors made of .class, [attribute] and :not(.class), the last
    one ending in ::after; every ancestor class has to be on the paragraph's box."""
    *ancestors, last = selector.replace("::after", "").split()
    if any(not set(re.findall(r"\.([\w-]+)", a)) <= set(el["boxClasses"]) for a in ancestors):
        return False
    nots = set(re.findall(r":not\(\.([\w-]+)\)", last))
    base = re.sub(r":not\([^)]*\)", "", last)
    return (set(re.findall(r"\.([\w-]+)", base)) <= set(el["classes"])
            and all(a in el["attrs"] for a in re.findall(r"\[([\w-]+)\]", base))
            and not nots & set(el["classes"]))


@needs_node
def test_the_strut_rule_in_the_stylesheet_is_bound_to_the_hook_the_page_sets():
    """The writer scales runs, never a paragraph mark, and the mark sets how tall a line prints: a shrunk
    GC NOTES list prints its 4.5pt notes 8.7pt apart. `applyBoxFit` marks the paragraph (`data-tw-mark`)
    and styles.css draws the mark-sized strut that holds the line open. Nothing tested the rule: delete it
    and every test stayed green while the editor's GC NOTES list collapsed from ~145pt to ~106pt. Here
    the real stylesheet rule is read, and its selector is matched against every paragraph the real page
    code (renderBlock, applyBoxFit, fitTxbx, over the real GC resinous form) produced: it must match the
    shrunk NOTES lines that state a spacing and nothing in a box that was not shrunk, and what it reads
    (`--tw-mark-pt`) must be set on every one of them.
    Mutations: delete the rule; rename `data-tw-mark` in the rule only, or in the page only; drop its
    `font-size: var(--tw-mark-pt)`."""
    rules = _strut_rules()
    assert len(rules) == 1, "styles.css has no single rule keyed on the data-tw-mark hook: %r" % (rules,)
    selector, decl = rules[0]
    assert selector.endswith("::after") and ".tw-txbx" in selector, selector
    assert decl.get("content") == '"\\00a0"' and decl.get("font-size") == "var(--tw-mark-pt)", decl
    assert decl.get("display") == "inline-block" and decl.get("width") == "0", decl

    tpl = _template("epoxy", "GC")
    body = _gc_body([])
    rep = {b["id"]: b for b in _fit(body)}
    blk = _wear_block()
    assert rep[blk["txbx"]]["scale"] < 0.999, "the GC NOTES box no longer shrinks, so this proves nothing"
    case = _harness([{"name": "strut", "blocks": tpl["blocks"],
                      "fit": {str(k): v for k, v in rep.items()}, "tokens": body["values"]}])[0]
    hit = [b for b in case["blocks"] if _selector_matches(selector, b)]
    assert blk["id"] in {b["id"] for b in hit}, "the wear & tear note's mark never gets the strut"
    assert {b["id"] for b in hit if b["txbx"] == blk["txbx"]}, "no line of the NOTES box gets the strut"
    for b in hit:
        assert b["twMark"] and b["printed"] and b["markVar"], b
        assert rep[b["txbx"]]["scale"] < 0.999, ("a box that was not shrunk gets the strut", b["id"], b["txbx"])
