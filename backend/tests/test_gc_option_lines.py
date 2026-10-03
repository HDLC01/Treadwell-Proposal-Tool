"""A pricing option the estimator chose and showed reaches the customer on the GC files too.

Hanz, 2026-10-03, with a screenshot of a GC Epoxy proposal in the editor: the Pricing options
sidebar had "EPOXY 2" ($36,157) ticked "Show as a proposal option" and "Show in proposal", priced
as "total amount". The PRICE box printed the Base Bid line and Kyle's own "Options & Unit Prices"
rows, and NOTHING for EPOXY 2. "The options don't appear in the proposal editor."

THE CAUSE. An option, a manual "Add for" line and a combo's Option 1 / Option 2 breakout all travel
to the document as `{{#price_line}}` rows, and a row can only be expanded where the template has
that region. Kyle's three GC files (Polish, Resinous -- which is also the combo -- and Sealer) have
none: their PRICE box is plain paragraphs, a free heading "Options & Unit Prices" over his static
unit-price rows. So `_expand_named_block` found no region and quietly dropped every row, and the
editor, which mounts its price-line island only through a region, left it in the hidden staging
panel. The Estimate step already knew (it warns "This GC template does not print options") and the
Proposal step, where the option is configured, said nothing. It has been so since the GC forms were
tokenised, on prod as well; #601 re-installed Kyle's files with the same shape.

THE FIX (writer: `_insert_price_lines_under_headings`; editor: `annotatePriceLineAnchor` /
`renderBlockList` / `priceLineRecord`). The rows are inserted as new paragraphs DIRECTLY UNDER the
marked Options heading, before Kyle's static rows, in every copy of the box, each cloned from the
first of Kyle's own money rows under that heading (so the PRICE-list bullet, the indent, the font
and the size are his), and the editor mounts the same island at the same place. The GC FILES ARE
NOT TOUCHED: their bytes and content hash are pinned below, because the real GC drafts carry saved
paragraph edits keyed by editor block id and stamped with that hash, and a paragraph added to the
file would move every id after it.

Every document check here goes through the real `main._generate`, the function Download, Send and
the customer PDF all render through, and reads the .docx itself in BOTH copies of the text box: the
mc:Choice copy Word reads and the mc:Fallback copy LibreOffice renders to the PDF (LibreOffice is
not installed on the dev box). The editor checks run the page's own functions (js/
price-bullets-harness.js: renderBlockList, refreshPriceDisplay, the ribbon's keys) over the real
/api/proposal-template block list, and compare what they draw with the document the same payload
prints.
"""
import copy
import hashlib
import io
import json
import logging
import pathlib
import re
import shutil
import subprocess
import zipfile

import pytest
from docx import Document
from docx.oxml.ns import qn
from fastapi.testclient import TestClient
from starlette.requests import Request

import main
import proposal_writer as pw
import template_versions as tv

client = TestClient(main.app)

BACKEND = pathlib.Path(__file__).resolve().parents[1]
TEMPLATES = BACKEND / "templates"
FRONTEND = BACKEND.parent / "frontend"
_MC = "{http://schemas.openxmlformats.org/markup-compatibility/2006}"

RESINOUS = "GC/xx TREADWELL RESINOUS PROPOSAL - xx.docx"
POLISH = "GC/xx TREADWELL POLISH PROPOSAL - xx.docx"
SEALER = "GC/xx TREADWELL SEALER PROPOSAL - xx.docx"
HEADING = "Options & Unit Prices"

# The three GC files by work type. "combo" has no file of its own: pick_template sends it to the
# Resinous one (TEMPLATE_PICKER), so a GC combo is that file with the breakout lines.
GC_FILES = [
    pytest.param("epoxy", RESINOUS, id="resinous"),
    pytest.param("polish", POLISH, id="polish"),
    pytest.param("sealer", SEALER, id="sealer"),
    pytest.param("combo", RESINOUS, id="combo"),
]

_VALS = {
    "job_name": "GC Options QA", "project_name": "GC Options QA", "city_state": "Olathe, KS",
    "bid_date_formatted": "10/3/26", "total_formatted": "$27,865", "base_bid_formatted": "$27,865",
    "state_name": "Kansas", "system_name": "MACRO", "texture": "OP", "epoxy_sf": "12,000",
    "scope_notes": "s", "schedule_notes": "~5d", "work_description": "w",
    "site_visit_date": "10/3", "disposal": "d", "exclusions": "std",
}

BASE_DESC = "Treadwell micro Flake Double Broadcast"
BASE_TOTAL = 27865


def _room(rid, name, total, *, base=False, desc=None, mode="total", show=True, notes=(),
          base_total=BASE_TOTAL):
    """One entry of `rooms` as the Proposal step's rebuildPricing makes it."""
    desc = desc or name
    return {"id": rid, "name": name, "is_base": base, "show": show, "price_mode": mode,
            "base_total": base_total, "deduct_amount": base_total - total,
            "bid": {"total": total, "sales_tax": 0, "remodel": 0, "taxable": False,
                    "remodel_on": False},
            "system_desc": desc, "option_desc": desc, "base_desc": BASE_DESC,
            "show_system": True, "show_diff": False,
            "notes_auto": [], "notes_manual": list(notes)}


BASE = _room("Epoxy", "Epoxy", BASE_TOTAL, base=True, desc=BASE_DESC)
OPT_TOTAL = _room("Copy1", "EPOXY 2", 36157, desc="Treadwell MACRO Flake")
OPT_DEDUCT = _room("Copy2", "GRIND", 24665, desc="Grind & Seal", mode="deduct")
OPT_ADD = _room("Copy3", "UPGRADE", 36157, desc="Treadwell MACRO Flake", mode="deduct")
OPT_NOTES = _room("Copy1", "EPOXY 2", 36157, desc="Treadwell MACRO Flake",
                  notes=["Includes 2 mobilizations", "Colour: Gray Blend only"])

BULLET = ("3", "0")      # Kyle's PRICE list, level 0: the red square
SUB = ("3", "1")         # level 1: the hollow "o" a line typed under a price line gets

L_TOTAL = ("$36,157 – Treadwell MACRO Flake as described above (tax exempt)", BULLET)
L_DEDUCT = ("Deduct ($3,200) – VE for Grind & Seal, in lieu of " + BASE_DESC + ".", BULLET)
L_ADD = ("Add $8,292 – Treadwell MACRO Flake", BULLET)
L_NOTES = ("$36,157 – Treadwell MACRO Flake as described above (tax exempt) — "
           "Includes 2 mobilizations; Colour: Gray Blend only", BULLET)
L_MANUAL = ("$1,500 – Add dye", BULLET)

COMBO_LINES = [
    {"amount_formatted": "$27,865", "key": "combo:epoxy.flooring",
     "label": "Option 1: Epoxy flooring as described above (tax exempt)"},
    {"amount_formatted": "$13,585", "key": "combo:polish.flooring",
     "label": "Option 2: Polished Concrete flooring as described above (tax exempt)"},
]
L_COMBO1 = ("$27,865 – Option 1: Epoxy flooring as described above (tax exempt)", BULLET)
L_COMBO2 = ("$13,585 – Option 2: Polished Concrete flooring as described above (tax exempt)", BULLET)


def _generate(work_type, audience="GC", **extra):
    """The .docx the real generate path produces, read straight out of its file cache."""
    body = {"work_type": work_type, "audience": audience, "values": dict(_VALS)}
    body.update(extra)
    req = Request({"type": "http", "headers": [], "method": "POST", "path": "/api/generate"})
    out = main._generate(main.GenerateIn(**body), req, persist=False, want_estimate=False)
    return main._FILE_CACHE[out.docx_download_url.rstrip("/").split("/")[-1]]["content"]


def _row(p):
    """(words, (numId, ilvl) or None) of one paragraph: its words and the list it is on."""
    return pw._own_text(p), pw._para_num_ref(p)


def _price_boxes(content, heading=HEADING):
    """{"choice": [row, ...], "fallback": [row, ...]}: every paragraph of the text box that holds
    the Options heading, once per copy of the box. Both copies must exist."""
    d = Document(io.BytesIO(content))
    out = {}
    for tx in d.element.body.iter(qn("w:txbxContent")):
        ps = [p for p in tx if p.tag == qn("w:p")]
        if not any(pw._own_text(p) == heading for p in ps):
            continue
        where = "fallback" if any(True for _ in tx.iterancestors(_MC + "Fallback")) else "choice"
        assert where not in out, f"two {where} boxes hold {heading!r}"
        out[where] = [_row(p) for p in ps]
    assert sorted(out) == ["choice", "fallback"], f"{heading!r} found in {sorted(out)}"
    return out


def _after_heading(rows, heading=HEADING):
    return rows[[r[0] for r in rows].index(heading) + 1:]


def _template_after_heading(rel):
    """Kyle's own paragraphs under the heading, straight out of the pristine file."""
    d = Document(str(TEMPLATES / rel))
    for tx in d.element.body.iter(qn("w:txbxContent")):
        ps = [p for p in tx if p.tag == qn("w:p")]
        if any(pw._own_text(p) == HEADING for p in ps):
            return _after_heading([_row(p) for p in ps])
    raise AssertionError(f"{rel} has no {HEADING!r}")


def _assert_under_heading(content, rel, lines):
    """THE CLAIM, for both copies: the paragraphs directly under the heading are `lines`, in this
    order, and everything after them is Kyle's own rows, as he wrote them."""
    static = _template_after_heading(rel)
    for where, rows in _price_boxes(content).items():
        after = _after_heading(rows)
        assert after[:len(lines)] == lines, (
            f"{where} copy: directly under {HEADING!r}\n  want {lines}\n  got  {after[:len(lines) + 2]}")
        assert after[len(lines):] == static, (
            f"{where} copy: Kyle's rows moved or changed\n  want {static}\n  got  {after[len(lines):]}")


# ── 1. a chosen option prints under the heading, in both copies, on every GC file ────────────────
@pytest.mark.parametrize("work_type,rel", GC_FILES)
def test_a_shown_option_prints_under_the_options_heading(work_type, rel):
    """Hanz's screenshot: Epoxy the base, "EPOXY 2" a shown option priced as a total. The line is
    EXACTLY what the Direct files print for it, as the first row under the heading."""
    _assert_under_heading(_generate(work_type, rooms=[BASE, OPT_TOTAL]), rel, [L_TOTAL])


@pytest.mark.parametrize("work_type,rel", GC_FILES)
def test_an_add_deduct_option_prints_under_the_heading(work_type, rel):
    """Priced as add/deduct: a cheaper option is a Deduct with its VE wording, a costlier one an
    Add. The sign is the line's own, as on every other template."""
    _assert_under_heading(_generate(work_type, rooms=[BASE, OPT_DEDUCT]), rel, [L_DEDUCT])
    _assert_under_heading(_generate(work_type, rooms=[BASE, OPT_ADD]), rel, [L_ADD])


@pytest.mark.parametrize("work_type,rel", GC_FILES)
def test_two_options_print_in_the_order_the_sidebar_lists_them(work_type, rel):
    _assert_under_heading(_generate(work_type, rooms=[BASE, OPT_TOTAL, OPT_DEDUCT]), rel,
                          [L_TOTAL, L_DEDUCT])
    _assert_under_heading(_generate(work_type, rooms=[BASE, OPT_DEDUCT, OPT_TOTAL]), rel,
                          [L_DEDUCT, L_TOTAL])


@pytest.mark.parametrize("work_type,rel", GC_FILES)
def test_an_options_notes_print_on_its_line(work_type, rel):
    _assert_under_heading(_generate(work_type, rooms=[BASE, OPT_NOTES]), rel, [L_NOTES])


@pytest.mark.parametrize("work_type,rel", GC_FILES)
def test_manual_price_lines_print_after_the_options(work_type, rel):
    """The Estimate step's own "Add for" price lines, which were dropped on GC the same way."""
    _assert_under_heading(
        _generate(work_type, rooms=[BASE, OPT_TOTAL],
                  price_lines=[{"label": "Add dye", "amount": 1500}]), rel, [L_TOTAL, L_MANUAL])
    _assert_under_heading(
        _generate(work_type, price_lines=[{"label": "Add dye", "amount": 1500}]), rel, [L_MANUAL])


@pytest.mark.parametrize("work_type,rel", GC_FILES)
def test_a_hidden_or_empty_option_prints_nothing(work_type, rel):
    """The rooms the sidebar did not show still print nothing, and a bid with no option leaves
    Kyle's rows exactly where they were: nothing is added under the heading."""
    hidden = _room("Copy1", "EPOXY 2", 36157, show=False)
    _assert_under_heading(_generate(work_type, rooms=[BASE, hidden]), rel, [])
    _assert_under_heading(_generate(work_type), rel, [])


# ── 2. exactly what the Direct file prints for the same payload ──────────────────────────────────
def _direct_lines(content, trailing_blanks):
    """The lines under the Direct epoxy file's "Options:" heading, without the template's own
    blank paragraphs at the end of the box."""
    return {where: _after_heading(rows, "Options:")[:-trailing_blanks]
            for where, rows in _price_boxes(content, "Options:").items()}


def _template_trailing_blanks(rel):
    d = Document(str(TEMPLATES / rel))
    for tx in d.element.body.iter(qn("w:txbxContent")):
        ps = [p for p in tx if p.tag == qn("w:p")]
        if any("{{#price_line}}" in pw._own_text(p) for p in ps):
            n = 0
            for p in reversed(ps):
                if pw._own_text(p).strip():
                    break
                n += 1
            return n
    raise AssertionError("no price_line region in " + rel)


def test_gc_prints_the_same_lines_direct_prints_for_the_same_payload():
    """The strongest statement of the fix: take the lines the Direct epoxy file prints under its
    Options heading -- text and bullet level of every paragraph, including a re-typed option, the
    lines typed round it and a blank one -- and the GC file prints that same sequence under its
    own heading. The oracle is the shipped Direct path, not a list written for this test."""
    rooms = [BASE, OPT_TOTAL, OPT_DEDUCT, OPT_ADD]
    pov = {"lines2": {"option:Copy1": "⟦amount⟧ – Option two, my words ⟦tax⟧"},
           "after": {"option:Copy1": ["a typed line under the option", ""]},
           "before": {"option:Copy2": ["typed above"]}}
    extra = {"rooms": rooms, "price_overrides": pov,
             "price_lines": [{"label": "Add dye", "amount": 1500}]}
    direct = _generate("epoxy", "Direct", **extra)
    blanks = _template_trailing_blanks("Direct/XX.XX TREADWELL EPOXY PROPOSAL - New Direct.docx")
    want = _direct_lines(direct, blanks)
    # Written out by hand as well, so the comparison is not of two things that could be wrong alike.
    literal = [("$36,157 – Option two, my words (tax exempt)", BULLET),
               ("a typed line under the option", SUB), ("", None),
               ("typed above", SUB), L_DEDUCT, L_ADD, L_MANUAL]
    assert want["choice"] == want["fallback"] == literal
    for work_type, rel in (("epoxy", RESINOUS), ("polish", POLISH), ("sealer", SEALER)):
        _assert_under_heading(_generate(work_type, "GC", **extra), rel, want["choice"])


# ── 3. the lines the estimator retypes and types round an option ─────────────────────────────────
@pytest.mark.parametrize("work_type,rel", GC_FILES)
def test_a_retyped_option_line_keeps_its_words_and_follows_the_estimate(work_type, rel):
    """price_overrides.lines2["option:<tab>"]: his words, with today's amount and tax wording put
    back in at the markers -- the same live shape every other line has."""
    pov = {"lines2": {"option:Copy1": "⟦amount⟧ – Option two, my words ⟦tax⟧"}}
    got = _generate(work_type, rooms=[BASE, OPT_TOTAL], price_overrides=pov)
    _assert_under_heading(got, rel, [("$36,157 – Option two, my words (tax exempt)", BULLET)])
    # The option re-priced: the line follows it.
    dearer = _room("Copy1", "EPOXY 2", 40000, desc="Treadwell MACRO Flake")
    got = _generate(work_type, rooms=[BASE, dearer], price_overrides=pov)
    _assert_under_heading(got, rel, [("$40,000 – Option two, my words (tax exempt)", BULLET)])


@pytest.mark.parametrize("work_type,rel", GC_FILES)
def test_lines_typed_under_and_over_an_option_are_their_own_paragraphs(work_type, rel):
    """Enter at the end of an option line, in the editor, stores price_overrides.after; at its
    start, .before. Each prints as a paragraph of its own: a sub-line on the "o" level, a blank
    one as the blank line it is."""
    pov = {"after": {"option:Copy1": ["a typed line under the option", ""]},
           "before": {"option:Copy1": ["typed above"]}}
    got = _generate(work_type, rooms=[BASE, OPT_TOTAL], price_overrides=pov)
    _assert_under_heading(got, rel, [("typed above", SUB), L_TOTAL,
                                     ("a typed line under the option", SUB), ("", None)])


@pytest.mark.parametrize("work_type,rel", GC_FILES)
def test_the_line_bullet_the_estimator_set_prints(work_type, rel):
    """The ribbon's Bullet / Indent act on a price line and travel as price_overrides.line_props:
    an option line he took the bullet off prints with none, and with his indent, in both copies."""
    pov = {"line_props": {"option:Copy1": {"bullet": False, "indent": 288}}}
    got = _generate(work_type, rooms=[BASE, OPT_TOTAL], price_overrides=pov)
    _assert_under_heading(got, rel, [(L_TOTAL[0], None)])
    d = Document(io.BytesIO(got))
    lefts = []
    for tx in d.element.body.iter(qn("w:txbxContent")):
        for p in tx:
            if p.tag == qn("w:p") and pw._own_text(p) == L_TOTAL[0]:
                ind = p.find(qn("w:pPr")).find(qn("w:ind"))
                lefts.append(ind.get(qn("w:left")) if ind is not None else None)
    assert lefts == ["288", "288"], lefts


# ── 4. a GC combo: the Option 1 / Option 2 breakout ──────────────────────────────────────────────
COMBO_BASE = _room("Epoxy", "Epoxy", 41450, base=True, desc=BASE_DESC, base_total=41450)
COMBO_OPT = _room("Copy1", "EPOXY 2", 36157, desc="Treadwell MACRO Flake", base_total=41450)


def test_a_gc_combo_prints_its_option_1_and_option_2_lines():
    """Direct Combo prints the two option totals as the price. On the GC Resinous file (the combo's
    template) they were dropped, and the page printed one combined base line and nothing else."""
    got = _generate("combo", rooms=[], combo_options=COMBO_LINES)
    _assert_under_heading(got, RESINOUS, [L_COMBO1, L_COMBO2])


def test_a_gc_combo_option_follows_the_breakout_and_prints_no_second_heading():
    """The Direct combo file loses its Options heading with {{#single_bid}} and the writer restores
    one as a row ("Options:"). The GC files keep their own heading, so a second one would print a
    second blank gap and a second title: the lines simply follow the breakout."""
    got = _generate("combo", rooms=[COMBO_BASE, COMBO_OPT], combo_options=COMBO_LINES,
                    price_lines=[{"label": "Add dye", "amount": 1500}])
    _assert_under_heading(got, RESINOUS, [L_COMBO1, L_COMBO2,
                                          ("$36,157 – Treadwell MACRO Flake as described above "
                                           "(tax exempt)", BULLET), L_MANUAL])
    for rows in _price_boxes(got).values():
        assert [r for r in rows if r[0].strip().rstrip(":") == "Options"] == []



# ── 5. a template with nowhere to print them says so instead of dropping them in silence ─────────
BUDGET = "Direct/xx.xx TREADWELL BUDGET PRICING.docx"


def test_a_template_with_neither_a_region_nor_a_heading_warns(caplog):
    """Direct Budget has no {{#price_line}} region and no Options heading. Its lines still cannot
    print, but the writer logs WARNING with the count instead of discarding them without a word."""
    with caplog.at_level(logging.WARNING, logger="proposal_tool.proposal_writer"):
        pw.fill_proposal(work_type="budget", audience="Direct", values=dict(_VALS),
                         price_lines=[{"label": "Add dye", "amount_formatted": "$1,500"},
                                      {"label": "More", "amount_formatted": "$2,000"}])
    hits = [r for r in caplog.records if r.levelno == logging.WARNING and "price line" in r.getMessage()]
    assert hits, [r.getMessage() for r in caplog.records]
    msg = hits[0].getMessage()
    assert "xx.xx TREADWELL BUDGET PRICING.docx" in msg and "2" in msg, msg


def test_nothing_to_print_is_not_a_warning(caplog):
    """No lines is the normal case on every template: no noise."""
    with caplog.at_level(logging.WARNING, logger="proposal_tool.proposal_writer"):
        pw.fill_proposal(work_type="budget", audience="Direct", values=dict(_VALS), price_lines=[])
        pw.fill_proposal(work_type="epoxy", audience="GC", values=dict(_VALS), price_lines=[])
    assert [r.getMessage() for r in caplog.records if "price line" in r.getMessage()] == []


# ── 6. the GC files are untouched, and a saved edit still lands where it was made ────────────────
GC_PINS = {
    # Kyle's 2026-09-11 forms as installed by #601. The sha256 of the bytes, and the content
    # version the editor and the generate guard stamp saved edits with (template_versions).
    POLISH: ("0fae245d7494a255b71107f2b3e219de484cc5ad5b73f9eca4723c340eea1c95",
             "sha256:00df2b19da129697"),
    RESINOUS: ("3337cd2ebc336ec682e921b074c4a1008b03d2835b882a4deb01800a59e24f51",
               "sha256:d6b9ceca38014714"),
    SEALER: ("64579bcb708d4da151a71a61b2ee9ac05c219afc9b24b5f27ffb141a8391c433",
             "sha256:138dfc2c3113df3d"),
}


@pytest.mark.parametrize("rel", [POLISH, RESINOUS, SEALER])
def test_the_gc_template_bytes_and_content_hash_are_unchanged(rel):
    """THE REASON THE FIX LIVES IN THE WRITER. A paragraph added to a GC file moves the id of every
    paragraph after it, and the real GC drafts hold saved edits keyed by id, stamped with this hash
    (PREDECESSOR_VERSIONS exists to keep four of them). Re-saving a file with the lines in it would
    refuse every one of those edits. If this fails, you changed a GC template: prove the id walk
    again and move the pins on purpose."""
    sha, version = GC_PINS[rel]
    path = TEMPLATES / rel
    assert hashlib.sha256(path.read_bytes()).hexdigest() == sha
    assert tv.content_version(path) == version
    assert main._template_proposal_version(path) == version
    assert tv.predecessor_versions(path), "the predecessor entry stopped applying"


def _walk_ids(path):
    d = Document(str(path))
    return [(i, kind, txbx, in_block, text)
            for i, kind, _p, in_block, text, txbx in pw.iter_editable_blocks(d)]


@pytest.mark.parametrize("rel,count", [(POLISH, 171), (RESINOUS, 169), (SEALER, 170)])
def test_the_editor_walk_of_the_gc_files_is_unchanged(rel, count):
    """The ids the editor is served: same number of blocks, and the PRICE box's paragraphs are the
    same ones in the same places (the walk is over the pristine file, and the file is untouched)."""
    walk = _walk_ids(TEMPLATES / rel)
    assert len(walk) == count
    j = client.get("/api/proposal-template?work_type=%s&audience=GC"
                   % {POLISH: "polish", RESINOUS: "epoxy", SEALER: "sealer"}[rel]).json()
    assert [(b["id"], b["kind"], b["txbx"], b["in_block"], b["text"]) for b in j["blocks"]] == walk


def _price_box_ids(rel):
    """The editor ids of the PRICE box: first and last paragraph of the box with the heading."""
    d = Document(str(TEMPLATES / rel))
    walk = list(pw.iter_editable_blocks(d))
    heading = next(w for w in walk if w[4] == HEADING)
    box = [w for w in walk if w[5] == heading[5]]
    return heading[0], box[0][0], box[-1][0]


@pytest.mark.parametrize("work_type,rel", [("epoxy", RESINOUS), ("polish", POLISH),
                                           ("sealer", SEALER), ("combo", RESINOUS)])
def test_saved_paragraph_edits_before_and_after_the_price_box_still_land_where_they_were_made(
        work_type, rel):
    """A real GC draft's saved overrides: edits on the job header, in the WORK box, on Kyle's own
    rows inside the PRICE box (above AND below the new lines), and in the boxes after it, stamped
    with the file's current content hash. With a shown option printing under the heading, every
    one prints on the paragraph it was made on -- the ids are the pristine file's, so nothing
    moved -- and the new lines sit between them."""
    heading_id, first_id, last_id = _price_box_ids(rel)
    walk = _walk_ids(TEMPLATES / rel)
    # Kyle's rows: one above the heading (the Base Bid line), one right under it (a money row
    # that sits BELOW the new lines) and the last one in the box.
    above = next(w[0] for w in walk if first_id <= w[0] < heading_id and w[4].startswith("{{base_bid"))
    below = next(w[0] for w in walk if w[0] > heading_id and w[4].startswith("$x – Add for"))
    edits = {
        99: "GC-EDIT-HEADER",
        above: "$27,865 – GC-EDIT-BASE-ROW",
        below: "$5,000 – GC-EDIT-KYLES-ROW",
    }
    after_box = [w[0] for w in walk if w[0] > last_id and w[1] == "p" and w[2] is not None
                 and w[4].strip() and "{{" not in w[4]]
    assert after_box, "no paragraph in a box after the PRICE box"
    edits[after_box[0]] = "GC-EDIT-AFTER-THE-PRICE-BOX"
    body = dict(paragraph_overrides=[{"id": i, "text": t} for i, t in edits.items()],
                template_version=tv.content_version(TEMPLATES / rel))
    with_option = _generate(work_type, rooms=[BASE, OPT_TOTAL], **body)
    plain = _generate(work_type, rooms=[BASE], **body)

    def words(content):
        return [pw._own_text(p) for p in Document(io.BytesIO(content)).element.body.iter(qn("w:p"))]

    # Every edit printed, on its own paragraph (once in each copy of its box, where it has two).
    for t in edits.values():
        assert words(plain).count(t) >= 1, f"{t!r} did not print at all"
    # THE WHOLE DOCUMENT, paragraph by paragraph: the option's lines are the only difference, so
    # no edit moved, was dropped, or landed on a neighbour.
    assert [w for w in words(with_option) if w != L_TOTAL[0]] == words(plain)
    assert words(with_option).count(L_TOTAL[0]) == 2, "once per copy of the PRICE box"
    # ...and in the PRICE box the new line sits between the heading and Kyle's edited rows. The
    # editor's edits are written to the real (mc:Choice) copy of a box, as they always were; the
    # VML twin keeps the template's words (it is not part of the id walk), and carries the line too.
    boxes = _price_boxes(with_option)
    names = [r[0] for r in boxes["choice"]]
    i = names.index(HEADING)
    assert names[i + 1] == L_TOTAL[0], names[i:i + 4]
    assert names.index("$5,000 – GC-EDIT-KYLES-ROW") > i + 1
    assert "$27,865 – GC-EDIT-BASE-ROW" in names[:i]
    twin = [r[0] for r in boxes["fallback"]]
    assert twin[twin.index(HEADING) + 1] == L_TOTAL[0]


# ── 7. the Direct files and Gyp are not touched ──────────────────────────────────────────────────
DIRECT_AND_GYP = [("epoxy", "Direct"), ("polish", "Direct"), ("combo", "Direct"), ("gyp", "Direct")]


def _parts(content):
    z = zipfile.ZipFile(io.BytesIO(content))
    return {n: z.read(n) for n in z.namelist()}


@pytest.mark.parametrize("work_type,audience", DIRECT_AND_GYP)
def test_the_heading_path_adds_nothing_where_the_file_has_a_price_line_region(
        work_type, audience, monkeypatch):
    """Direct Epoxy / Polish / Combo and Gyp print their lines through {{#price_line}}. The new
    path must be inert there: the same payload generates the same .docx, part for part, with the
    path switched off. (The cross-tree proof against origin/staging is in the commit message.)"""
    extra = {"rooms": [BASE, OPT_TOTAL, OPT_DEDUCT],
             "price_lines": [{"label": "Add dye", "amount": 1500}],
             "price_overrides": {"after": {"option:Copy1": ["typed under"]}}}
    if work_type == "combo":
        extra["combo_options"] = COMBO_LINES
        extra["rooms"] = [COMBO_BASE, COMBO_OPT]
    with_it = _parts(_generate(work_type, audience, **extra))
    monkeypatch.setattr(pw, "_insert_price_lines_under_headings", lambda *a, **k: 0)
    without = _parts(_generate(work_type, audience, **extra))
    assert sorted(with_it) == sorted(without)
    for name in with_it:
        assert with_it[name] == without[name], f"{name} differs"
    # ...and the lines really are in the file, once each, so the comparison was not of two empties.
    text = "\n".join(pw._own_text(p) for p in Document(io.BytesIO(
        _generate(work_type, audience, **extra))).element.body.iter(qn("w:p")))
    assert text.count("$1,500 – Add dye") == 2, "once per copy of the box"


# ── 8. the editor draws the same lines, in the same place ────────────────────────────────────────
# js/price-bullets-harness.js builds the editor's PRICE box through the page's own code over the
# real template's blocks (renderBlockList mounts the islands, refreshPriceDisplay paints them) and
# hands back what it shows and the payload it would send; the payload goes through the real
# _generate and the .docx is read back. Before this fix the editor mounted #price-lines-block only
# through a {{#price_line}} / {{#has_options}} region, so on a GC file the island stayed in the
# hidden staging panel: the editor showed nothing under the heading, and nothing printed either.
HARNESS = BACKEND / "tests" / "js" / "price-bullets-harness.js"
needs_node = pytest.mark.skipif(shutil.which("node") is None, reason="node is not installed")
FILE_OF = {"epoxy": RESINOUS, "combo": RESINOUS, "polish": POLISH, "sealer": SEALER}


def _template_json(work_type, audience="GC"):
    req = Request({"type": "http", "method": "GET", "path": "/api/proposal-template", "headers": [],
                   "query_string": b""})
    return json.loads(main.api_proposal_template(req, work_type=work_type, audience=audience).body)


def _editor_state(work_type, rooms, *, price_lines=(), pov=None, priced_tabs=None):
    """The draft the Proposal step holds: Kyle's figures, the rooms the sidebar lists."""
    return {
        "work_type": work_type, "audience": "GC", "project_name": "GC Options QA",
        "job_name": "GC Options QA", "proposal_lump_sum": BASE_TOTAL, "proposal_sales_tax": 0,
        "proposal_remodel_tax": 0, "proposal_taxable": False, "proposal_remodel_on": False,
        "tax_layout": "ONE_LINE", "priced_tabs": list(priced_tabs or []), "cell_values": {},
        "sheet_area": {"epoxy_sf": 12000, "polish_sf": 12000, "cove_lf": 0},
        "rooms": copy.deepcopy(rooms), "price_lines": copy.deepcopy(list(price_lines)),
        "price_overrides": copy.deepcopy(pov or {}),
    }


def _editor_case(name, work_type, state, actions=()):
    tj = _template_json(work_type)
    box = next(b["txbx"] for b in tj["blocks"] if b.get("txbx") is not None
               and "base_bid_formatted" in (b["text"] or ""))
    return {"name": name, "work_type": work_type, "audience": "GC",
            "blocks": [b for b in tj["blocks"] if b.get("txbx") == box],
            "options_heading_ids": tj.get("options_heading_ids") or [],
            "price_lines_anchor": tj.get("price_lines_anchor"),
            "state": state, "actions": list(actions)}


def _run_editor(cases):
    p = subprocess.run(["node", str(HARNESS), str(FRONTEND)], input=json.dumps(cases),
                       capture_output=True, text=True, encoding="utf-8")
    assert p.returncode == 0, "the harness itself failed:\n" + p.stderr
    return json.loads(p.stdout)


def _lvl(ref):
    return int(ref[1]) if ref else None


def _editor_vs_document(case, run):
    """The editor's lines under the heading and the document's (per copy of the box), each a list
    of (words, bullet level or None), plus the editor's raw lines, where the heading is, how many
    lines are ours, and Kyle's own rows."""
    pl = run["payload"]
    st = case["state"]
    content = _generate(case["work_type"], "GC", values=pl["values"],
                        price_overrides=pl["price_overrides"],
                        paragraph_overrides=pl["paragraph_overrides"],
                        combo_options=pl["combo_options"], rooms=st["rooms"],
                        price_lines=[{"label": l["label"], "amount": l["amount"]}
                                     for l in st["price_lines"]])
    static = _template_after_heading(FILE_OF[case["work_type"]])
    docs = {}
    for where, rows in _price_boxes(content).items():
        after = _after_heading(rows)
        n = len(after) - len(static)
        assert after[n:] == static, f"{where}: Kyle's rows moved"
        docs[where] = [(t, _lvl(ref)) for t, ref in after[:n]]
    ed = run["lines"]
    h = next(i for i, ln in enumerate(ed) if ln["text"] == HEADING)
    n = len(docs["choice"])
    mine = [("" if ln["blank"] else ln["text"], ln["level"] if ln["bullet"] else None)
            for ln in ed[h + 1:h + 1 + n]]
    return mine, docs, ed, h, n, static


def _assert_editor_matches(case, run, want):
    mine, docs, ed, h, n, static = _editor_vs_document(case, run)
    assert mine == docs["choice"] == docs["fallback"], (
        f"{case['name']}: the editor and the document disagree under {HEADING!r}\n  editor {mine}"
        f"\n  choice {docs['choice']}\n  twin   {docs['fallback']}")
    if want is not None:
        assert mine == [(t, _lvl(ref)) for t, ref in want], mine
    # THE ISLAND SITS DIRECTLY UNDER THE HEADING, AHEAD OF KYLE'S OWN ROWS: the line after the last
    # of ours is the first of his.
    assert ed[h + 1 + n]["text"] == static[0][0], (ed[h + n:h + n + 3], static[:1])
    return ed, h, n


def _model_spacing(case, ed):
    """The vertical geometry of Kyle's row the document clones the lines from, as the editor draws it."""
    mid = case["price_lines_anchor"]["model_id"]
    return next(ln["spacing"] for ln in ed if ln.get("id") == mid)


ED_CASES = [
    ("one option (total)", "epoxy", [BASE, OPT_TOTAL], (), None, [L_TOTAL]),
    ("one option (deduct)", "epoxy", [BASE, OPT_DEDUCT], (), None, [L_DEDUCT]),
    ("two options", "epoxy", [BASE, OPT_TOTAL, OPT_DEDUCT], (), None, [L_TOTAL, L_DEDUCT]),
    ("an option with notes", "epoxy", [BASE, OPT_NOTES], (), None, [L_NOTES]),
    ("manual price lines", "epoxy", [BASE, OPT_TOTAL], [{"label": "Add dye", "amount": 1500}], None,
     [L_TOTAL, L_MANUAL]),
    ("polish: one option", "polish", [BASE, OPT_TOTAL], (), None, [L_TOTAL]),
    ("sealer: two options", "sealer", [BASE, OPT_TOTAL, OPT_DEDUCT], (), None, [L_TOTAL, L_DEDUCT]),
    ("a retyped option line", "epoxy", [BASE, OPT_TOTAL], (),
     {"lines2": {"option:Copy1": "⟦amount⟧ – Option two, my words ⟦tax⟧"}},
     [("$36,157 – Option two, my words (tax exempt)", BULLET)]),
    ("a line typed after the option", "epoxy", [BASE, OPT_TOTAL], (),
     {"after": {"option:Copy1": ["a typed line under the option", ""]}},
     [L_TOTAL, ("a typed line under the option", SUB), ("", None)]),
]


@needs_node
@pytest.mark.parametrize("name,work_type,rooms,price_lines,pov,want", ED_CASES,
                         ids=[c[0] for c in ED_CASES])
def test_the_editor_shows_exactly_the_lines_the_document_prints(
        name, work_type, rooms, price_lines, pov, want):
    """THE CLAIM. The Proposal step's PRICE box, drawn by the page's own code, carries directly
    under the Options heading the same lines the .docx prints under it (both copies of its text
    box): the same words in the same order, each with the same bullet level -- and they take the
    spacing of Kyle's row the document clones them from, so the box is as tall on screen as on
    paper. On origin/staging the editor drew nothing there."""
    case = _editor_case(name, work_type, _editor_state(work_type, rooms, price_lines=price_lines, pov=pov))
    run = _run_editor([case])[0]
    ed, h, n = _assert_editor_matches(case, run, want)
    model = _model_spacing(case, ed)
    assert model != [None, None, None], "the model row carries no spacing: nothing to compare"
    for ln in ed[h + 1:h + 1 + n]:
        assert ln["spacing"] == model, (ln["text"], ln["spacing"], model)


@needs_node
def test_a_line_typed_and_a_line_retyped_in_the_editor_reach_the_document():
    """The estimator's own typing, through the page's keys: Enter at the end of an option line and
    a sentence typed (price_overrides.after), and the option line itself re-worded
    (price_overrides.lines2) -- on a GC file as on Direct."""
    st = _editor_state("epoxy", [BASE, OPT_TOTAL])
    case = _editor_case("typed in the editor", "epoxy", st, [
        {"enter_after": {"key": "option:Copy1"}, "type": "typed through the editor"},
    ])
    run = _run_editor([case])[0]
    _assert_editor_matches(case, run, [L_TOTAL, ("typed through the editor", SUB)])
    assert run["payload"]["price_overrides"]["after"] == {"option:Copy1": ["typed through the editor"]}

    st = _editor_state("epoxy", [BASE, OPT_TOTAL])
    case = _editor_case("re-worded in the editor", "epoxy", st, [
        {"settext": {"key": "option:Copy1"}, "text": "$36,157 – My reworded option (tax exempt)"},
    ])
    run = _run_editor([case])[0]
    _assert_editor_matches(case, run, [("$36,157 – My reworded option (tax exempt)", BULLET)])
    assert "option:Copy1" in run["payload"]["price_overrides"]["lines2"]


@needs_node
def test_the_ribbons_bullet_on_an_option_line_prints_on_a_gc_file():
    """Bullet pressed on the option line through the page's own ribbon handler: the editor shows the
    line unbulleted and the document prints it so, in both copies."""
    st = _editor_state("epoxy", [BASE, OPT_TOTAL])
    case = _editor_case("ribbon bullet", "epoxy", st,
                        [{"ribbon": {"key": "option:Copy1"}, "click": "bullet"}])
    run = _run_editor([case])[0]
    mine, docs, *_ = _editor_vs_document(case, run)
    assert mine == docs["choice"] == docs["fallback"]
    assert mine[0][0] == L_TOTAL[0] and mine[0][1] is None, mine


@needs_node
def test_a_gc_combo_breakout_sits_under_the_heading_in_the_editor_too():
    """The Option 1 / Option 2 breakout is drawn into #combo-price-block, which only a
    {{#single_bid}} region used to mount: on the GC Resinous file (the combo's template) it never
    reached the page either. Now it follows the heading, ahead of the options, as printed."""
    tabs = [{"id": "Epoxy", "role": "epoxy", "kind": "base", "total": 27865, "sales_tax": 0,
             "remodel": 0, "taxable": False, "remodel_on": False},
            {"id": "Polish", "role": "polish", "kind": "base", "total": 13585, "sales_tax": 0,
             "remodel": 0, "taxable": False, "remodel_on": False}]
    st = _editor_state("combo", [OPT_TOTAL], priced_tabs=tabs)
    case = _editor_case("combo", "combo", st)
    run = _run_editor([case])[0]
    ed, h, n = _assert_editor_matches(case, run, [L_COMBO1, L_COMBO2, L_TOTAL])
    assert [ln["key"] for ln in ed[h + 1:h + 1 + n]] == [
        "combo:epoxy.flooring", "combo:polish.flooring", "option:Copy1"]
    model = _model_spacing(case, ed)
    for ln in ed[h + 1:h + 1 + n]:
        assert ln["spacing"] == model, (ln["text"], ln["spacing"], model)


def test_the_template_response_names_where_the_editor_mounts_the_lines():
    """/api/proposal-template names the anchor on a template that prints its lines under a free
    heading: the heading block (the first of options_heading_ids) and Kyle's first money row under
    it, the one the writer clones the lines from. None on every template that prints them through a
    region (Direct, Gyp) and on the one with nowhere to print them (Budget)."""
    for wt in ("epoxy", "polish", "sealer", "combo"):
        j = _template_json(wt)
        text = {b["id"]: b["text"] for b in j["blocks"]}
        a = j["price_lines_anchor"]
        assert a and a["heading_id"] == j["options_heading_ids"][0], (wt, a)
        assert text[a["heading_id"]] == HEADING
        assert text[a["model_id"]].startswith("$x – Add for"), (wt, text[a["model_id"]])
        assert a["model_id"] > a["heading_id"]
    for wt, aud in (("epoxy", "Direct"), ("polish", "Direct"), ("combo", "Direct"),
                    ("gyp", "Direct"), ("budget", "Direct")):
        assert _template_json(wt, aud)["price_lines_anchor"] is None, (wt, aud)


@needs_node
def test_a_response_without_the_anchor_still_mounts_the_lines_under_the_heading():
    """A browser holding the response of an older server (no price_lines_anchor, but the heading
    named in options_heading_ids): the lines still go under the heading; only the spacing falls back
    to the page's own."""
    case = _editor_case("older body", "epoxy", _editor_state("epoxy", [BASE, OPT_TOTAL, OPT_DEDUCT]))
    case["price_lines_anchor"] = None
    run = _run_editor([case])[0]
    mine, docs, ed, h, n, static = _editor_vs_document(case, run)
    assert mine == docs["choice"] == [(L_TOTAL[0], 0), (L_DEDUCT[0], 0)]
    assert ed[h + 1 + n]["text"] == static[0][0]


@needs_node
def test_a_template_with_a_price_line_region_ignores_the_heading_list():
    """The mount is for a template with NO {{#price_line}} / {{#has_options}} region. Handed the
    Direct epoxy file's blocks with a heading id (which a Direct response never carries), the page
    still mounts the lines where its region does: the box is drawn exactly as without the id. The id
    is the box's LAST block, a free paragraph drawn after every region: a mount there would be the
    one that wins (a region's own mount comes first and would be moved), so it is the one that shows."""
    tj = _template_json("epoxy", "Direct")
    box = next(b["txbx"] for b in tj["blocks"] if b.get("txbx") is not None
               and "base_bid_formatted" in (b["text"] or ""))
    blocks = [b for b in tj["blocks"] if b.get("txbx") == box]
    heading = blocks[-1]
    assert not (heading["text"] or "").strip() and "{{" not in (heading["text"] or "")
    base = {"name": "d", "work_type": "epoxy", "audience": "Direct", "blocks": blocks,
            "options_heading_ids": [], "price_lines_anchor": None,
            "state": _editor_state("epoxy", [BASE, OPT_TOTAL])}
    base["state"]["audience"] = "Direct"
    forced = dict(base, name="forced", options_heading_ids=[heading["id"]],
                  price_lines_anchor={"heading_id": heading["id"], "model_id": heading["id"]})
    a, b = _run_editor([base, forced])
    strip = lambda r: [{k: v for k, v in ln.items() if k != "spacing"} for ln in r["lines"]]
    assert any(ln["text"].startswith("$36,157") for ln in a["lines"]), "the option is not drawn at all"
    assert strip(a) == strip(b)


def test_a_line_typed_under_the_options_heading_prints_once_on_a_gc_combo():
    """A Direct combo loses its Options heading with {{#single_bid}}, so main.py restores one as a
    row and the lines typed under it ride that row. The GC files keep their own heading and print
    the lines typed under it themselves, FIRST under it: restoring a second row would print the
    typed line twice (and a second heading)."""
    pov = {"after": {"heading_options": ["typed under the heading"]}}
    got = _generate("combo", rooms=[BASE, OPT_TOTAL], combo_options=COMBO_LINES, price_overrides=pov)
    typed = ("typed under the heading", SUB)
    for rows in _price_boxes(got).values():
        after = _after_heading(rows)
        assert [r for r in after if r[0] == "typed under the heading"] == [after[0]], after[:6]
        assert [r[0] for r in after[:4]] == ["typed under the heading", L_COMBO1[0], L_COMBO2[0], L_TOTAL[0]]
    assert typed[0] in [r[0] for r in _price_boxes(got)["choice"]]


def _gc_resinous_after_heading(strip_money):
    """The Resinous file, in memory, with the rows under the Options heading thinned: every money
    row ("$x – Add for ...") out of the first copy of the box when `strip_money`, and then every
    row on the PRICE list as well when it is "all"."""
    d = Document(str(TEMPLATES / RESINOUS))
    head = next(h for h in pw.options_heading_paragraphs(d) if not pw._is_fallback_paragraph(h))
    nxt = head.getnext()
    while nxt is not None:
        cur, nxt = nxt, nxt.getnext()
        if cur.tag != qn("w:p"):
            continue
        listed = pw._para_num_ref(cur) == (pw._PRICE_LIST_NUM_ID, "0")
        if listed and (strip_money == "all" or pw._MONEY_ROW_RE.match(pw._own_text(cur))):
            cur.getparent().remove(cur)
    return head


def test_the_model_row_is_the_first_money_row_else_the_first_listed_row_else_the_heading():
    """The paragraph the lines are cloned from. Kyle's first listed row on the Resinous file is an
    italic sentence (his note on flake sizes), so a money row comes first; with none, the first
    listed row with words; with none of those, the heading itself."""
    d = Document(str(TEMPLATES / RESINOUS))
    head = next(h for h in pw.options_heading_paragraphs(d) if not pw._is_fallback_paragraph(h))
    first = pw._price_line_model(head)
    assert pw._own_text(first).startswith("$x – Add for") and pw._MONEY_ROW_RE.match(pw._own_text(first))
    head = _gc_resinous_after_heading("money")
    sentence = pw._price_line_model(head)
    assert sentence is not head and pw._own_text(sentence).startswith("If a different flake/chip size")
    head = _gc_resinous_after_heading("all")
    assert pw._price_line_model(head) is head


@pytest.mark.parametrize("work_type,audience", [("epoxy", "Direct"), ("polish", "Direct"),
                                                ("combo", "Direct"), ("gyp", "Direct"),
                                                ("epoxy", "GC"), ("polish", "GC"), ("sealer", "GC")])
def test_a_template_that_prints_the_lines_does_not_warn(work_type, audience, caplog):
    """The warning is for a template with nowhere to put the lines. Every other one prints them
    (through its region, or under its heading) and stays quiet."""
    with caplog.at_level(logging.WARNING, logger="proposal_tool.proposal_writer"):
        pw.fill_proposal(work_type=work_type, audience=audience, values=dict(_VALS),
                         price_lines=[{"label": "Add dye", "amount_formatted": "$1,500"}])
    assert [r.getMessage() for r in caplog.records if "price line" in r.getMessage()] == []
