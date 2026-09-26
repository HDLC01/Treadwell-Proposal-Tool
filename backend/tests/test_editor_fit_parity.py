"""The Proposal step shows every text box at the size the PDF prints it.

Hanz, 2026-09-26, clicking through staging: "please follow the text size of what is written in the
proposal PDFs. Its different on the editor and on the output" / "whatever is the font size in the
PDF should also be the same as in the Proposal Editor".

WHAT WAS WRONG. Two shrinks, one in each place, answering to nothing in common. The writer
(`proposal_writer._shrink_overflowing_text_boxes`) estimates a box's content and scales an
overflowing box's runs down to a 0.60 floor; the editor ran its own ladder off a browser
measurement that stopped at 0.75 and then put the box back at FULL size and clipped it behind
"Show all". So a full box showed full-size text on screen and printed smaller. Measured on the
templates as they ship: every GC file's NOTES box prints at the floor with its default text (7.5pt
shown, 4.5pt printed), and so does the GC and Budget date box.

WHAT IT IS NOW. One rule, the writer's. POST /api/proposal-fit runs the real fill for the payload
Continue would send and returns what the shrink decided per box (`fit_report`); the page applies
that scale run by run with `F.fitHp` (proposal-format-core.js), a port of `_scale_txbx_runs`'
per-run arithmetic, and never clips. This file proves each link of that chain by executing it:

  1. the route reports the scale the document is actually printed at -- every run of every text
     box of the real `_generate` output equals `F.fitHp(design size, reported scale)`, run through
     node, over every template with short, full and overflowing content (Python writer vs JS rule);
  2. the editor puts exactly the writer's per-run size on every element it draws (the REAL
     renderBlock / applyBoxFit / fitTxbx, js/editor-fit-harness.js), checked against the REAL
     `_scale_txbx_runs` run on the same sizes (JS page vs Python writer);
  3. with no shrink, every line's base size is the `w:sz` the writer emits: words typed into a
     line, an emptied line's height, the NOTES bullets and the PRICE rows.

LibreOffice is not installed on the dev box, so the document half is checked on the .docx.
"""
import copy
import io
import json
import pathlib
import re
import shutil
import subprocess
import tempfile
from unittest import mock

import pytest
from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from fastapi.testclient import TestClient
from starlette.requests import Request

import main
import proposal_writer as pw

FRONTEND = pathlib.Path(__file__).resolve().parents[2] / "frontend"
CORE = FRONTEND / "js" / "proposal-format-core.js"
HARNESS = pathlib.Path(__file__).resolve().parent / "js" / "editor-fit-harness.js"
NODE = shutil.which("node")
needs_node = pytest.mark.skipif(NODE is None, reason="node is not installed")

# Every template the Proposal step can put on screen: all four Direct files, all three GC files and
# the Gyp file.
TEMPLATES = [("epoxy", "Direct"), ("polish", "Direct"), ("combo", "Direct"), ("budget", "Direct"),
             ("epoxy", "GC"), ("polish", "GC"), ("sealer", "GC"), ("gyp", "Direct")]
KINDS = ["short", "full", "overflow"]

_V = {
    "job_name": "Fit QA", "project_name": "Fit QA", "city_state": "Olathe, KS",
    "bid_date_formatted": "9/26/26", "total_formatted": "$36,763", "base_bid_formatted": "$36,763",
    "state_name": "Kansas", "system_name": "MACRO", "texture": "OP", "epoxy_sf": "12,000",
    "scope_notes": "Grind and coat.", "schedule_notes": "~5 days", "work_description": "Warehouse",
    "site_visit_date": "9/25", "disposal": "d", "exclusions": "Standard exclusions.",
    "estimator_name": "QA Estimator",
}
_LONG = ("Excludes moisture mitigation, joint filling, crack repair beyond the allowance, night or "
         "weekend work, and anything not described above. ") * 8


def _content(kind):
    """short: one note, nothing long. full: the standard notes the backend fills in. overflow: long
    WORK text, eighteen notes and twelve option lines, which puts WORK, PRICE and NOTES over their
    boxes on every Direct template and on Gyp (checked in test_the_fixtures_cover_what_they_claim)."""
    if kind == "short":
        return {"values": dict(_V), "notes": ["One note."]}
    if kind == "full":
        return {"values": dict(_V)}
    v = dict(_V)
    v.update({"exclusions": _LONG, "scope_notes": _LONG[:600], "schedule_notes": _LONG[:300]})
    return {"values": v,
            "notes": ["Note number %d about the job and its conditions." % i for i in range(18)],
            "price_lines": [{"label": "Add option %d for the extra work" % i, "amount": 1000 + i}
                            for i in range(12)]}


def _req(path, method="POST"):
    return Request({"type": "http", "headers": [], "method": method, "path": path, "query_string": b""})


_TPL = {}


def _template(wt, aud):
    if (wt, aud) not in _TPL:
        _TPL[(wt, aud)] = json.loads(
            main.api_proposal_template(_req("/api/proposal-template", "GET"), wt, aud).body)
    return _TPL[(wt, aud)]


def _body(wt, aud, kind, **extra):
    body = {"work_type": wt, "audience": aud, "template_version": _template(wt, aud)["template_version"]}
    body.update(copy.deepcopy(_content(kind)))
    body.update(copy.deepcopy(extra))
    return body


def _fit(body):
    return main.api_proposal_fit(main.GenerateIn(**copy.deepcopy(body)), _req("/api/proposal-fit"))["boxes"]


def _docx(body, unscaled=False):
    """The real generate path's .docx; `unscaled` = the same document with the overflow shrink's
    run scaling switched off, i.e. every run at the size it has before the shrink touches it."""
    def run():
        out = main._generate(main.GenerateIn(**copy.deepcopy(body)), _req("/api/generate"),
                             persist=False, want_estimate=False)
        return main._FILE_CACHE[out.docx_download_url.rstrip("/").split("/")[-1]]["content"]
    if not unscaled:
        return run()
    with mock.patch.object(pw, "_scale_txbx_runs", lambda *a, **k: None):
        return run()


def _sz(el):
    rpr = el.find(qn("w:rPr"))
    sz = rpr.find(qn("w:sz")) if rpr is not None else None
    return int(sz.get(qn("w:val"))) if sz is not None else None


def _text(p):
    return "".join(t.text or "" for t in p.iter(qn("w:t")))


def _box_runs(docx_bytes):
    """Per real text box: [(run half-points or None, run text)] in document order."""
    d = Document(io.BytesIO(docx_bytes))
    return [[(_sz(r), "".join(t.text or "" for t in r.iter(qn("w:t")))) for r in tx.iter(qn("w:r"))]
            for tx in pw._iter_txbx(d)]


def _node_fit_hp(pairs):
    """F.fitHp(hp, scale) for each pair, executed in node from the shipped core."""
    src = ("const F = require(%s); const pairs = JSON.parse(require('fs').readFileSync(0, 'utf8'));"
           "process.stdout.write(JSON.stringify(pairs.map(([h, s]) => F.fitHp(h, s))));"
           % json.dumps(str(CORE)))
    got = subprocess.run([NODE, "-e", src], input=json.dumps(pairs), capture_output=True,
                         encoding="utf-8", errors="replace", timeout=60)
    assert got.returncode == 0, got.stderr
    return json.loads(got.stdout)


def _writer_hp(hps, scale):
    """What the REAL `_scale_txbx_runs` does to runs of these sizes at this scale: a text box of
    one run per size, scaled by the writer's own function, read back."""
    tx = OxmlElement("w:txbxContent")
    p = OxmlElement("w:p")
    tx.append(p)
    for hp in hps:
        r = OxmlElement("w:r")
        rpr = OxmlElement("w:rPr")
        sz = OxmlElement("w:sz")
        sz.set(qn("w:val"), str(hp))
        rpr.append(sz)
        r.append(rpr)
        p.append(r)
    pw._scale_txbx_runs(tx, scale)
    return [_sz(r) for r in p.findall(qn("w:r"))]


def _harness(cases, answers=None):
    return _harness_raw(cases, answers)["cases"]


def _harness_raw(cases, answers=None):
    with tempfile.NamedTemporaryFile("w", suffix=".json", delete=False, encoding="utf-8") as fh:
        json.dump({"cases": cases, "answers": answers}, fh)
        name = fh.name
    try:
        got = subprocess.run([NODE, str(HARNESS), str(FRONTEND), name], capture_output=True,
                             encoding="utf-8", errors="replace", timeout=120)
    finally:
        pathlib.Path(name).unlink(missing_ok=True)
    assert got.returncode == 0, got.stderr
    return json.loads(got.stdout)


def _pt(v):
    """'7.5pt' -> 15 half-points; None stays None."""
    if v is None:
        return None
    assert v.endswith("pt"), v
    return round(float(v[:-2]) * 2)


# ── 1. the route reports the scale the document is printed at ─────────────────
@needs_node
@pytest.mark.parametrize("kind", KINDS)
@pytest.mark.parametrize("wt,aud", TEMPLATES)
def test_every_printed_run_is_the_reported_scale_applied_by_the_editors_rule(wt, aud, kind):
    """For every run of every text box of the real document: a box the route says was left alone
    prints every run at its design size, and a box it says was shrunk prints every run at exactly
    `F.fitHp(design size, reported scale)` -- the shipped JS rule, executed. A run with no size of
    its own is counted at the box's reported `default_hp`, as the writer does."""
    body = _body(wt, aud, kind)
    rep = {b["id"]: b for b in _fit(body)}
    printed = _box_runs(_docx(body))
    design = _box_runs(_docx(body, unscaled=True))
    assert len(printed) == len(design) == len(rep), "the route did not report every text box"
    pairs, want = [], []
    for i, (p_runs, d_runs) in enumerate(zip(printed, design)):
        f = rep[i]
        assert f["exempt"] == [], "no paragraph here has a size the estimator chose"
        assert [t for _, t in p_runs] == [t for _, t in d_runs], "the two renders differ in text"
        if f["scale"] >= 0.999:
            assert f["scale"] == 1.0
            assert [h for h, _ in p_runs] == [h for h, _ in d_runs], (
                "box %d is reported unshrunk and printed shrunk" % i)
            continue
        for (p_hp, _), (d_hp, _) in zip(p_runs, d_runs):
            pairs.append([d_hp if d_hp is not None else f["default_hp"], f["scale"]])
            want.append(p_hp)
    if pairs:
        assert _node_fit_hp(pairs) == want


@needs_node
def test_the_js_rule_is_the_writers_rule_on_every_size():
    """`F.fitHp` against the REAL `_scale_txbx_runs` over every size the templates use and more,
    at scales that include EXACT ties: 20 x 0.625 is 12.5, which the writer's Python `round` sends
    to 12 and a half-up rule to 13. No real document happens to land on a tie, so this is where the
    rounding is pinned; the 4pt floor is exercised here too (6 x 0.6 = 3.6)."""
    scales = [0.6, 0.625, 0.65, 0.7, 0.75, 0.8, 0.816, 0.875, 0.9, 0.95, 0.998]
    sizes = list(range(2, 41))
    pairs = [[hp, sc] for sc in scales for hp in sizes]
    want = [w for sc in scales for w in _writer_hp(sizes, sc)]
    assert any(hp * sc == int(hp * sc) + 0.5 for hp, sc in pairs), "the grid holds no exact tie"
    assert _node_fit_hp(pairs) == want
    # And a box the writer leaves alone keeps every size, even one under the floor.
    assert _node_fit_hp([[6, 1.0], [6, 0.999], [40, 1.0]]) == [6, 6, 40]


def test_the_fixtures_cover_what_they_claim():
    """A parity test over boxes that never shrink proves nothing. The overflowing content shrinks
    WORK, PRICE and NOTES on every Direct template that carries them and on Gyp; at least one box
    reaches the floor; and a short payload shrinks no Direct WORK/PRICE/NOTES box at all."""
    boxes = {  # the WORK / PRICE / NOTES box ids per template, read off where their rows live
        ("epoxy", "Direct"): (2, 4, 3), ("polish", "Direct"): (3, 4, 5),
        ("combo", "Direct"): (3, 5, 4), ("gyp", "Direct"): (2, 3, 5),
    }
    floors = 0
    for (wt, aud), ids in boxes.items():
        over = {b["id"]: b for b in _fit(_body(wt, aud, "overflow"))}
        short = {b["id"]: b for b in _fit(_body(wt, aud, "short"))}
        for i in ids:
            assert over[i]["scale"] < 0.999, (wt, aud, i, over[i])
            if wt != "gyp":
                assert short[i]["scale"] == 1.0, (wt, aud, i, short[i])
        floors += sum(1 for b in over.values() if b["at_floor"])
    assert floors, "no fixture reaches the 0.60 floor, so text past the box is never exercised"


def test_a_size_the_estimator_chose_is_left_alone_and_a_runless_line_takes_the_box_size():
    """The two cases the shrink treats specially, both reported by the route. A paragraph sent with
    run sizes (every formatted edit is) is EXEMPT: its runs keep their size in a shrunk box, and its
    block id is in `exempt`. Words typed into a template line that had no run get a bare run, which
    the shrink gives the box's most common size (`default_hp`) scaled -- not 12pt scaled."""
    tpl = _template("epoxy", "Direct")
    scope = next(b for b in tpl["blocks"] if b["text"].startswith("Scope:"))
    blank = next(b for b in tpl["blocks"] if b["txbx"] == scope["txbx"] and not b["text"]
                 and b["fit"]["typed_sized"] is False and b["fit"]["removable"] and b["id"] > scope["id"])
    body = _body("epoxy", "Direct", "overflow", paragraph_overrides=[
        {"id": scope["id"], "text": "Scope:  EXEMPT-QA",
         "runs": [{"text": "Scope:", "bold": True, "size_pt": 8}, {"text": "  EXEMPT-QA", "size_pt": 8}]},
        {"id": blank["id"], "text": "TYPED-QA"},
    ])
    f = next(b for b in _fit(body) if b["id"] == scope["txbx"])
    assert f["scale"] < 0.999 and f["exempt"] == [scope["id"]]
    runs = _box_runs(_docx(body))[scope["txbx"]]
    assert [h for h, t in runs if t in ("Scope:", "  EXEMPT-QA")] == [16, 16], "the chosen size was shrunk"
    typed = [h for h, t in runs if t == "TYPED-QA"]
    d = Document(io.BytesIO(_docx(body, unscaled=True)))
    design_default = pw._txbx_default_hp(list(pw._iter_txbx(d))[scope["txbx"]])
    assert f["default_hp"] == design_default
    assert typed == _writer_hp([design_default], f["scale"])


# ── 2. the editor draws exactly the writer's per-run size ─────────────────────
@needs_node
@pytest.mark.parametrize("wt,aud", TEMPLATES)
def test_the_editor_draws_every_line_at_the_size_the_writer_prints(wt, aud):
    """The REAL renderBlock over the REAL template blocks, handed the REAL fit report for the
    overflowing payload, then fitTxbx -- the page's own path. Every element that states a size is
    given `--tw-fit-pt` equal to the REAL `_scale_txbx_runs` result for its design size, the box
    itself the page's 9pt scaled, a runless line the box's default size scaled; nothing inside an
    exempt paragraph, and nothing at all in a box the writer left alone. The inline font-size -- what
    fmtAt reads back as the estimator's formatting -- keeps the design size throughout."""
    body = _body(wt, aud, "overflow")
    rep = {b["id"]: b for b in _fit(body)}
    tpl = _template(wt, aud)
    by_id = {b["id"]: b for b in tpl["blocks"]}
    case = _harness([{"name": "%s/%s" % (wt, aud), "blocks": tpl["blocks"],
                      "fit": {str(k): v for k, v in rep.items()}, "tokens": body["values"]}])[0]
    checked = 0
    for box in case["boxes"]:
        f = rep[box["id"]]
        if f["scale"] >= 0.999:
            assert box["printed"] is None and not box["marked"]
        else:
            assert _pt(box["printed"]) == _writer_hp([18], f["scale"])[0]
    for blk in case["blocks"]:
        f = rep[blk["txbx"]]
        rec = by_id[blk["id"]]
        # The design size never moves: the base is what typed words print at, unshrunk.
        assert _pt(blk["base"]) == rec["fit"]["typed_hp"]
        shrunk = f["scale"] < 0.999 and blk["id"] not in f["exempt"]
        if not shrunk:
            assert blk["printed"] is None and all(s["printed"] is None for s in blk["spans"])
            continue
        base_hp = f["default_hp"] if rec["fit"]["typed_sized"] is False else rec["fit"]["typed_hp"]
        assert _pt(blk["printed"]) == _writer_hp([base_hp], f["scale"])[0], blk
        for s in blk["spans"]:
            assert s["design"] is not None, "a span with a printed size and no design size"
            assert _pt(s["printed"]) == _writer_hp([_pt(s["design"])], f["scale"])[0], (blk["id"], s)
            checked += 1
    assert case["marksLeftAfterClear"] == 0, "a size stayed on after the fit was cleared"
    # The overflowing payload shrinks at least one box with sized spans in it on every template
    # (16 to 224 of them), so a pass that compared nothing is a failure, not a vacuous success.
    assert checked, "no span in a shrunk box was compared"


@needs_node
@pytest.mark.parametrize("wt,aud", [("epoxy", "Direct"), ("epoxy", "GC"), ("gyp", "Direct")])
def test_the_page_applies_the_routes_answer_through_its_own_request(wt, aud):
    """The REAL requestFit, executed: it posts the question to /api/proposal-fit, parses the REAL
    report the route returns, and re-fits every box -- ending at exactly the printed sizes the render
    test above proved equal to the writer's. Then the four ways an answer must NOT land: the same
    question twice asks once; an answer computed for another template (a base flip in between) is
    dropped; an older answer that arrives after a newer one is dropped; a failed request leaves the
    last good answer. And an answer that shrinks nothing takes every printed size off again."""
    body = _body(wt, aud, "overflow")
    report = _fit(body)
    tpl = _template(wt, aud)
    direct = _harness([{"name": "direct", "blocks": tpl["blocks"],
                        "fit": {str(b["id"]): b for b in report}, "tokens": body["values"]}])[0]
    want = {str(b["id"]): b["printed"] for b in direct["boxes"]}
    assert any(want.values()), "nothing shrinks, so applying the answer is not being tested"
    got = _harness_raw([], {"blocks": tpl["blocks"], "report": report, "tokens": body["values"]})["answers"]
    assert got["applied"] == want
    assert got["request"]["url"] == "/api/proposal-fit" and got["request"]["method"] == "POST"
    assert got["request"]["contentType"] == "application/json" and got["request"]["auth"]
    assert json.loads(got["request"]["body"]) == {"q": "q1"}
    assert got["calls"] == 1 and got["callsAfterRepeat"] == 1, "an unchanged question was asked again"
    assert all(v is None for v in got["stale"].values()), "an answer for another template was applied"
    assert got["outOfOrder"] == want, "an older answer overwrote a newer one"
    assert got["afterFailure"] == want, "a failed request cleared the last good answer"
    assert all(v is None for v in got["unshrunk"].values())


@needs_node
def test_the_editor_leaves_a_paragraph_the_writer_exempts_at_its_own_size():
    """The real report for a payload with a hand-sized WORK line (exempt) in a shrunk box, handed
    to the real page: that line's spans get no printed size, and the lines around it do."""
    tpl = _template("epoxy", "Direct")
    scope = next(b for b in tpl["blocks"] if b["text"].startswith("Scope:"))
    sched = next(b for b in tpl["blocks"] if b["text"].startswith("Schedule:"))
    body = _body("epoxy", "Direct", "overflow", paragraph_overrides=[
        {"id": scope["id"], "text": "Scope:  x",
         "runs": [{"text": "Scope:", "bold": True, "size_pt": 8}, {"text": "  x", "size_pt": 8}]}])
    rep = {b["id"]: b for b in _fit(body)}
    assert rep[scope["txbx"]]["exempt"] == [scope["id"]] and rep[scope["txbx"]]["scale"] < 0.999
    case = _harness([{"name": "exempt", "blocks": tpl["blocks"],
                      "fit": {str(k): v for k, v in rep.items()}, "tokens": body["values"]}])[0]
    blk = {b["id"]: b for b in case["blocks"]}
    assert blk[scope["id"]]["spans"] and all(s["printed"] is None for s in blk[scope["id"]]["spans"])
    assert blk[scope["id"]]["printed"] is None
    assert blk[sched["id"]]["spans"] and all(s["printed"] for s in blk[sched["id"]]["spans"])


# ── 3. the base sizes, where nothing is shrunk ─────────────────────────────────
@pytest.mark.parametrize("wt,aud", TEMPLATES)
def test_words_typed_into_a_line_print_at_the_size_the_editor_shows(wt, aud):
    """Every free text-box line of the template rewritten as plain text, the way a restored draft
    replays it and the way words land after a line's own spans are gone: the writer prints each at
    `fit.typed_hp`, which is the size renderBlock gives the line itself (asserted in the harness
    test above). On a line that had no run, that is 12pt, the document default -- not the 9pt the
    editor used to show."""
    tpl = _template(wt, aud)
    # The free Remodel Tax row (GC, Gyp) is taken out by the writer on a job with no remodel tax,
    # whatever it says, so there is no printed line to measure.
    lines = [b for b in tpl["blocks"] if b["txbx"] is not None and not b["in_block"]
             and not (b.get("para") or {}).get("marker")
             and "tax_amount_formatted" not in b["text"] and "remodel.amount" not in b["text"]]
    body = _body(wt, aud, "short",
                 paragraph_overrides=[{"id": b["id"], "text": "QA-%d" % b["id"]} for b in lines])
    d = Document(io.BytesIO(_docx(body, unscaled=True)))
    dd = d.styles.element.find(qn("w:docDefaults"))
    doc_default = int(dd.find(".//" + qn("w:rPrDefault") + "/" + qn("w:rPr") + "/" + qn("w:sz"))
                      .get(qn("w:val")))
    found = {}
    for tx in pw._iter_txbx(d):
        for p in tx.iter(qn("w:p")):
            t = _text(p)
            if t.startswith("QA-") and t[3:].isdigit():
                found[int(t[3:])] = [_sz(r) for r in p.findall(qn("w:r")) if "".join(
                    x.text or "" for x in r.iter(qn("w:t")))]
    assert set(found) == {b["id"] for b in lines}
    for b in lines:
        (hp,) = found[b["id"]]
        if b["fit"]["typed_sized"]:
            assert hp == b["fit"]["typed_hp"], (b["id"], hp, b["fit"])
        else:
            # A bare run: no size of its own, so it prints at what it inherits, which on every
            # template is the document default -- and that is the size the editor was told.
            assert hp is None, (b["id"], hp, b["fit"])
            assert b["fit"]["typed_hp"] == doc_default, (b["id"], b["fit"], doc_default)


@pytest.mark.parametrize("wt,aud", TEMPLATES)
def test_the_shrink_never_scales_a_paragraph_mark(wt, aud):
    """An EMPTY line prints as tall as its paragraph mark, and `_scale_txbx_runs` scales runs, never
    marks -- so in a shrunk box an empty line keeps its full height. That is why the editor draws a
    `.tw-empty` line at `fit.hp` (styles.css, --tw-mark-pt) and leaves it out of the shrink. Executed
    on the overflowing payload: every paragraph mark of every box is the same in the printed
    document and in the same document with run scaling switched off."""
    body = _body(wt, aud, "overflow")
    marks = []
    for unscaled in (False, True):
        d = Document(io.BytesIO(_docx(body, unscaled=unscaled)))
        marks.append([[_mark_hp(p) for p in tx.iter(qn("w:p"))] for tx in pw._iter_txbx(d)])
    assert marks[0] == marks[1]
    assert any(f["scale"] < 0.999 for f in _fit(body)) or wt == "budget", "nothing was shrunk"


def test_an_emptied_line_prints_at_its_mark_size():
    """The Direct epoxy Exclusions line is the case that shows the difference: its words are 8pt
    (`typed_hp` 16) and its mark is 9pt (`hp` 18). Emptied, it prints at the mark's 9pt, which is
    what the editor now draws for it -- not at the 8pt its words were."""
    tpl = _template("epoxy", "Direct")
    excl = next(b for b in tpl["blocks"] if b["text"].startswith("Exclusions:"))
    assert (excl["fit"]["hp"], excl["fit"]["typed_hp"]) == (18, 16)
    base = Document(io.BytesIO(_docx(_body("epoxy", "Direct", "short"))))
    work = [p for p in list(pw._iter_txbx(base))[excl["txbx"]].iter(qn("w:p"))]
    i = next(k for k, p in enumerate(work) if _text(p).startswith("Exclusions:"))
    d = Document(io.BytesIO(_docx(_body("epoxy", "Direct", "short",
                                        paragraph_overrides=[{"id": excl["id"], "text": ""}]))))
    got = list(list(pw._iter_txbx(d))[excl["txbx"]].iter(qn("w:p")))[i]
    assert _text(got) == "" and _mark_hp(got) == excl["fit"]["hp"]


@needs_node
def test_the_editor_draws_an_empty_line_at_its_mark_and_outside_the_shrink():
    """renderBlock puts every text-box line's mark size on it (`data-mark-pt`, --tw-mark-pt), and the
    stylesheet rule that applies it to `.tw-empty` is more specific than the shrink's rule, so an
    empty line in a shrunk box keeps its full height, as it prints."""
    tpl = _template("epoxy", "Direct")
    case = _harness([{"name": "marks", "blocks": tpl["blocks"], "fit": {}, "tokens": {}}])[0]
    by_id = {b["id"]: b for b in tpl["blocks"]}
    assert case["blocks"]
    for blk in case["blocks"]:
        hp = by_id[blk["id"]]["fit"]["hp"]
        assert _pt(blk["markVar"]) == hp, blk
        assert float(blk["markPt"]) == hp / 2, blk
    css = (FRONTEND / "styles.css").read_text(encoding="utf-8")
    mark = re.search(r"(?m)^\.tw-txbx \.tw-block\.tw-empty\[data-mark-pt\]\s*\{([^}]*)\}", css)
    fit = re.search(r"(?m)^\.tw-txbx \[data-tw-fit\]\s*\{([^}]*)\}", css)
    assert mark and "var(--tw-mark-pt) !important" in mark.group(1)
    assert fit and "var(--tw-fit-pt) !important" in fit.group(1)
    # (0,4,0) against (0,2,0): the mark wins in a shrunk box whichever rule comes first.
    assert css.index(".tw-txbx .tw-block.tw-empty[data-mark-pt]") > 0


def _mark_hp(p):
    ppr = p.find(qn("w:pPr"))
    rpr = ppr.find(qn("w:rPr")) if ppr is not None else None
    sz = rpr.find(qn("w:sz")) if rpr is not None else None
    return int(sz.get(qn("w:val"))) if sz is not None else None


@needs_node
@pytest.mark.parametrize("wt", ["epoxy", "polish", "combo", "gyp"])
def test_the_notes_bullets_are_drawn_at_the_notes_row_size(wt):
    """The {{#notes}} bullets are the editor's own lines, and they stated no size, so they inherited
    the page's 9pt while every note printed at the row's 7.5pt (Epoxy) or 8pt (Polish, Combo). Now
    each bullet, blank ones included, carries the size of the `{{notes.text}}` run, and the
    document's printed notes are exactly that size when the box is not shrunk."""
    aud = "Direct"
    tpl = _template(wt, aud)
    body = _body(wt, aud, "short", notes=["First note.", "", "Third note."])
    case = _harness([{"name": wt, "blocks": tpl["blocks"], "fit": {}, "tokens": body["values"]}])[0]
    row_hp = _pt(str(case["notesRowPt"]) + "pt")
    assert case["notes"] and all(_pt(n["fontSize"]) == row_hp for n in case["notes"]), case["notes"]
    d = Document(io.BytesIO(_docx(body, unscaled=True)))
    sizes = {_sz(r) for tx in pw._iter_txbx(d) for p in tx.iter(qn("w:p"))
             if _text(p) in ("First note.", "Third note.") for r in p.findall(qn("w:r")) if _sz(r)}
    assert sizes == {row_hp}, (sizes, row_hp)


@pytest.mark.parametrize("wt", ["epoxy", "polish", "combo"])
def test_the_price_rows_the_editor_draws_itself_print_at_the_pages_size(wt):
    """The PRICE rows on the Direct files are the editor's own lines (#base-bid-row, the tax rows,
    the option lines), drawn without a size and inheriting `.tw-page`'s 9pt (PAGE_HP = 18). That is
    right only while the writer prints them at 9pt too; this says so, run by run."""
    tpl = _template(wt, "Direct")
    price_box = next(b["txbx"] for b in tpl["blocks"] if b["in_block"] in ("price_line", "tax_breakout"))
    body = _body(wt, "Direct", "short", price_lines=[{"label": "Add dye", "amount": 1500}])
    runs = _box_runs(_docx(body, unscaled=True))[price_box]
    assert runs and {h for h, t in runs if t.strip()} == {18}, runs


# ── the route itself ───────────────────────────────────────────────────────────
def test_the_fit_route_needs_a_signed_in_caller(monkeypatch, real_verify_token):
    """Behind the same sign-in gate as /api/generate: it runs the document fill."""
    monkeypatch.setattr(main.supabase_client, "verify_token", real_verify_token)
    client = TestClient(main.app)
    assert client.post("/api/proposal-fit", json={"work_type": "epoxy"}).status_code == 401
    assert client.post("/api/generate", json={"work_type": "epoxy"}).status_code == 401


def test_the_fit_route_writes_nothing(monkeypatch):
    """It runs the real fill and stops: no cached file, no audit event, no draft write."""
    for name in ("_cache_file",):
        monkeypatch.setattr(main, name, mock.Mock(side_effect=AssertionError(name)))
    monkeypatch.setattr(main.drafts, "log_event", mock.Mock(side_effect=AssertionError("log_event")))
    monkeypatch.setattr(main.drafts, "save_draft", mock.Mock(side_effect=AssertionError("save_draft")))
    rep = _fit(_body("epoxy", "Direct", "overflow"))
    assert [b["id"] for b in rep] == [0, 1, 2, 3, 4, 5]


def test_asking_changes_no_document():
    """`fit_report` only records. A document built with it is byte-for-byte the document built
    without it, so no saved payload prints differently because the editor asked a question."""
    kw = dict(work_type="epoxy", audience="Direct", values=dict(_V, exclusions=_LONG))
    plain = pw.fill_proposal(**kw)
    report = []
    asked = pw.fill_proposal(fit_report=report, **kw)
    assert report and _xml(plain) == _xml(asked)


def _xml(docx_bytes):
    import zipfile
    return zipfile.ZipFile(io.BytesIO(docx_bytes)).read("word/document.xml")


def test_the_route_says_which_template_it_answered_for():
    """The page throws away an answer computed for a template no longer on screen (a base flip in
    between; executed in test_the_page_applies_the_routes_answer_through_its_own_request), which it
    can only do because the route names the template it filled -- the same version string
    /api/proposal-template served, for each template."""
    for wt in ("polish", "epoxy"):
        got = main.api_proposal_fit(main.GenerateIn(**_body(wt, "Direct", "short")),
                                    _req("/api/proposal-fit"))
        assert got["template_version"] == _template(wt, "Direct")["template_version"]
    assert _template("polish", "Direct")["template_version"] != _template("epoxy", "Direct")["template_version"]


def test_notes_that_arrive_after_the_first_fit_ask_for_the_size_again():
    """A brand-new project's boilerplate notes, and the notes a base flip re-seeds, come back from
    /api/default-notes after the editor's first fit question was answered: the NOTES box kept the
    size of notes it no longer held. The page's own prefill and reseedNotesForWorkType now ask again
    (js/editor-fit-harness.js notesFit)."""
    got = _harness_raw([])["notesFit"]
    assert got["seeded"] == "Scope.\nSchedule."
    assert got["afterPrefill"] == 1
    assert got["afterReseed"] == 2


# ── the shrink counts indents: the reviewer's probe (2026-09-26 editor release) ──────────────────
_LONG_O = ("Includes one mobilization, protection of adjacent finishes, all grinding dust control and "
           "the moisture test described above; pricing holds for thirty days from the date shown.")


def _probe(n):
    """Epoxy Direct, Broken out, and `n` long lines typed under the Total: each prints Kyle's
    hollow "o", its text one inch in (level 1, 1440 twips)."""
    b = _body("epoxy", "Direct", "full")
    b["values"].update({"tax_layout": "BROKEN_OUT", "price_taxable": True, "price_remodel_on": False,
                        "material_tax_formatted": "$1,200", "total_formatted": "$36,763",
                        "base_bid_formatted": "$35,563"})
    b["price_overrides"] = {"after": {"total": [_LONG_O] * n}}
    return b


def _price_box(docx_bytes):
    """(index, Document, txbxContent) of the PRICE box: the one holding "Base Bid"."""
    d = Document(io.BytesIO(docx_bytes))
    for i, tx in enumerate(pw._iter_txbx(d)):
        if any(_text(p).strip() == "Base Bid" for p in tx.iter(qn("w:p"))):
            return i, d, tx
    raise AssertionError("no PRICE box")


def _usable(d, tx, i):
    geo = pw.template_geometry(d)["boxes"][i]
    l, r, t, b = pw._txbx_insets(tx)
    return geo["w_pt"] - (l + r) / pw._EMU_PER_PT, geo["h_pt"] - (t + b) / pw._EMU_PER_PT


def test_the_shrink_counts_the_price_boxs_indents():
    """Review of the bullets branch, finding 3: the writer measured every paragraph at the box's
    full width, so the REBID indents (288 twips on a money line, 1440 on an "o" line) were free.
    Three long "o" lines under the Total wrapped into more lines than the estimate counted: the box
    was shrunk to 0.93 and still printed ~17pt past its bottom edge (laid out with Zetta Serif's
    own glyphs, the test below). The estimate now counts each paragraph at its own width, so the
    box it reports is the box that prints, and the editor, which applies that report, draws it the
    same size.

    Pinned: the PRICE box's reported content is the estimate's glyph model applied to each paragraph
    AT ITS OWN WIDTH, read off the unscaled .docx (a model that ignores the indents cannot match),
    the "o" lines are really one inch in, and more of them never shrink the box less. Mutation:
    _fit_line_widths_pt handing back the box's width."""
    scales = {}
    for n in (0, 1, 2, 3):
        body = _probe(n)
        i, _, _ = _price_box(_docx(body))
        box = next(b for b in _fit(body) if b["id"] == i)
        _, du, txu = _price_box(_docx(body, unscaled=True))
        uw, _ = _usable(du, txu, i)
        want, inch_in = 0.0, 0
        for p in txu.iter(qn("w:p")):
            text = _text(p)
            hp = pw._fit_hp(p)
            fpt = hp / 2.0 if hp else 9.0
            left = pw._effective_left_tw(du, p)
            body_w = uw - left / 20.0
            first_w = body_w
            if pw._para_num_ref(p) is None:
                first_w += ((pw._effective_ind_tw(du, p, "hanging") or 0)
                            - (pw._effective_ind_tw(du, p, "firstLine") or 0)) / 20.0
            if text.strip() and left >= 1440:
                inch_in += 1
            fc = first_w / (pw._TXBX_GLYPH_W * fpt)
            bc = body_w / (pw._TXBX_GLYPH_W * fpt)
            lines = 1 + (-(-(len(text) - fc) // bc) if len(text) > fc else 0)
            want += lines * pw._TXBX_LINE_H * fpt
        assert inch_in == n, (n, inch_in)
        assert abs(box["content_pt"] - round(want, 3)) < 0.01, (n, box["content_pt"], want)
        scales[n] = box["scale"]
    assert scales[0] >= scales[1] >= scales[2] >= scales[3] and scales[3] < 0.8, scales


_ZETTA_BOOK = pathlib.Path(pw.__file__).resolve().parent / "fonts" / "Zetta Serif-Book.otf"


def _real_font_height(body):
    """The PRICE box laid out with Zetta Serif Book's own advances (FreeType, through PIL) at each
    paragraph's printed size and its own width, greedy word wrap, its w:spacing over the font's
    ascent + descent. Independent of the estimate's glyph model. Returns (height, usable height)."""
    from PIL import ImageFont
    font = ImageFont.truetype(str(_ZETTA_BOOK), size=1000)
    asc, desc = font.getmetrics()
    i, d, tx = _price_box(_docx(body))
    uw, uh = _usable(d, tx, i)
    total = 0.0
    for p in tx.iter(qn("w:p")):
        runs = [(_text(rr), _sz(rr)) for rr in p.iter(qn("w:r"))]
        text = "".join(s for s, _ in runs)
        hp = next((h for s, h in runs if s.strip() and h), None) or pw._fit_hp(p) or 18
        pt = hp / 2.0
        ppr = p.find(qn("w:pPr"))
        sp = ppr.find(qn("w:spacing")) if ppr is not None else None
        before = int(sp.get(qn("w:before")) or 0) / 20.0 if sp is not None else 0.0
        after = int(sp.get(qn("w:after")) or 0) / 20.0 if sp is not None else 0.0
        line = sp.get(qn("w:line")) if sp is not None else None
        rule = sp.get(qn("w:lineRule")) if sp is not None else None
        one = (asc + desc) * pt / 1000.0
        lh = one * int(line) / 240.0 if line and rule in (None, "auto") else (int(line) / 20.0 if line else one)
        width = uw - pw._effective_left_tw(d, p) / 20.0
        lines, cur = 1, ""
        for w in text.split(" "):
            cand = (cur + " " + w) if cur else w
            if font.getlength(cand) * pt / 1000.0 <= width or not cur:
                cur = cand
            else:
                lines, cur = lines + 1, w
        total += before + lines * lh + after
    return total, uh


@pytest.mark.skipif(not _ZETTA_BOOK.exists(), reason="the licensed Zetta Serif files are not on this machine")
def test_three_long_o_lines_under_the_total_no_longer_print_past_the_box():
    """The reviewer's probe, measured with the real font rather than the estimate's model: long "o"
    lines under the Total in epoxy Direct, Broken out. Before the shrink counted indents, three of
    them printed ~17.5pt past the PRICE box and two ~1pt past; now the box holds them. Skipped where
    the licensed font is absent (CI): the test above pins the estimate itself."""
    for n in (2, 3):
        height, usable = _real_font_height(_probe(n))
        assert height <= usable, (n, height, usable)
