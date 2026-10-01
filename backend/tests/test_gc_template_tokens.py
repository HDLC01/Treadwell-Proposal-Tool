"""Kyle's re-saved GC proposal forms, and the intake fields that fill what they leave as "xx".

Hanz, 2026-10-02: Kyle made new GC proposal templates; they are the formats the app uses for a
General Contractor proposal, and the intake form collects more of what they leave as "xx" --
intake -> estimate sheet -> proposal editor / document.

WHAT THIS FILE HOLDS TO:

  * THE INSTALLED FILES ARE THE ANNOTATOR'S OUTPUT ON KYLE'S RAW FILES. His three forms live,
    untokenized, in docs/GC Templates/; `annotate_templates.annotate_from` reproduces every part
    of every installed GC file from them. A hand edit to a GC template, or a rule that stopped
    matching, fails here -- not on a customer's proposal.
  * A BLANK INTAKE FIELD PRINTS KYLE'S OWN WORDS. Each new token stands where his placeholder was
    ("033543", "xx Architects", "8/1/26", "[PC]", "A900", "0"), and with the field blank the
    document prints exactly the line his raw file prints. Rendered through the real fill path and
    compared with the raw file, line for line -- that is what pins every literal in
    `proposal_writer.TEMPLATE_TOKEN_DEFAULTS` to Kyle's file.
  * NO SAVED EDIT IS LOST. Four real GC drafts carried editor edits stamped with the old files'
    content hashes. The paragraph walk of the new files is identical to the old one -- proven here
    against the old bytes, kept in tests/fixtures/gc_predecessors/ -- so
    `template_versions.PREDECESSOR_VERSIONS` lets those stamps replay, at generate and in the
    editor, and nothing else.
  * THE CUSTOMER SEES ONLY WHAT WAS MEANT. With the new fields blank, a generated GC document
    differs from one built on the old file in the job header (job name and city, as the Gyp form
    has always printed them), in Kyle's own spec-line date, and in his Resinous flake wording.
  * INTAKE -> ESTIMATE SHEET: Drawings Dated, and the Specs+Dwgs+Addn tab.
"""
import hashlib
import io
import json
import pathlib
import re
import shutil
import subprocess
import zipfile

import docx
import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook

import annotate_templates as at
import estimate_writer as ew
import main
import proposal_writer as pw
import template_versions as tv

client = TestClient(main.app)

BACKEND = pathlib.Path(__file__).resolve().parents[1]
REPO = BACKEND.parent
RAW_DIR = REPO / "docs" / "GC Templates"
PREDECESSORS = pathlib.Path(__file__).resolve().parent / "fixtures" / "gc_predecessors"
FRONTEND_JS = REPO / "frontend" / "js"

POLISH = "GC/xx TREADWELL POLISH PROPOSAL - xx.docx"
RESINOUS = "GC/xx TREADWELL RESINOUS PROPOSAL - xx.docx"
SEALER = "GC/xx TREADWELL SEALER PROPOSAL - xx.docx"
GC_FILES = (POLISH, RESINOUS, SEALER)
# Which work type picks each file. Sealer is option-only in the app (no base work type reaches
# it), but fill_proposal still renders it by name, and the file is replaced all the same.
WORK_TYPE = {POLISH: "polish", RESINOUS: "epoxy", SEALER: "sealer"}

# The tokens the intake fields fill, plus the job header's two (the Gyp form's treatment).
NEW_TOKENS = {"job_name", "city_state", "spec_section", "architect", "drawings_dated_formatted",
              "finish_tag", "plan_sheet", "addenda_count"}
_TOKEN = re.compile(r"\{\{\s*([#/]?[A-Za-z_][A-Za-z0-9_.]*)\s*\}\}")
_WP = re.compile(r"<w:p[ >].*?</w:p>", re.DOTALL)
_WT = re.compile(r"<w:t\b[^>]*>([^<]*)</w:t>")
_MC_FALLBACK = "{http://schemas.openxmlformats.org/markup-compatibility/2006}Fallback"

# The editor ids whose WORDS changed from the old file to the new one (the walk is otherwise the
# same paragraph by paragraph): the three job-header lines, the spec line, the area line, the
# addenda line -- and on Resinous, Kyle's System line ("Decorative Flake (macro sizes: 1/4”+)").
CHANGED_IDS = {POLISH: {99, 100, 101, 113, 116, 159},
               RESINOUS: {99, 100, 101, 112, 114, 117, 157},
               SEALER: {99, 100, 101, 113, 116, 158}}


def _parts(path):
    """Every part of a .docx, decompressed, in package order. Compared instead of the file's
    bytes: the deflate stream depends on the zlib that wrote it, the parts do not."""
    with zipfile.ZipFile(path) as z:
        return [(n, z.read(n)) for n in z.namelist()]


def _walk(path_or_bytes):
    src = path_or_bytes if isinstance(path_or_bytes, bytes) else str(path_or_bytes)
    d = docx.Document(io.BytesIO(src) if isinstance(src, bytes) else src)
    return [{"id": i, "kind": k, "in_block": b, "txbx": x, "text": t}
            for i, k, _p, b, t, x in pw.iter_editable_blocks(d)]


def _tokens(path):
    """{token: number of paragraphs it appears in} over document.xml, runs joined per paragraph
    (Word splits a token across runs freely), both the mc:Choice and mc:Fallback copies."""
    with zipfile.ZipFile(path) as z:
        xml = z.read("word/document.xml").decode("utf-8")
    out = {}
    for m in _WP.finditer(xml):
        for tok in _TOKEN.findall("".join(_WT.findall(m.group(0)))):
            out[tok] = out.get(tok, 0) + 1
    return out


def _rendered_lines(docx_bytes):
    """The text of each paragraph Word renders (mc:Choice), skipping the VML fallback copy."""
    d = docx.Document(io.BytesIO(docx_bytes))
    out = []
    for p in d.element.xpath("//w:p"):
        if any(True for _ in p.iterancestors(_MC_FALLBACK)):
            continue
        out.append("".join(t.text or "" for t in p.xpath("./w:r/w:t")))
    return out


# ── 1. the installed files are the annotator's output on Kyle's raw files ─────────────────────
def test_kyles_raw_files_are_in_the_repo_and_untokenized():
    for rel in GC_FILES:
        raw = RAW_DIR / pathlib.Path(rel).name
        assert raw.exists(), f"Kyle's raw file for {rel} is missing from {RAW_DIR}"
        assert not _tokens(raw), f"{raw.name} already carries tokens -- it is not Kyle's raw file"


def test_annotating_kyles_raw_files_reproduces_the_installed_ones(tmp_path):
    """Re-run the annotator on copies of Kyle's files and compare every part with what ships."""
    before = {rel: hashlib.sha256((RAW_DIR / pathlib.Path(rel).name).read_bytes()).hexdigest()
              for rel in GC_FILES}
    done = at.annotate_from(RAW_DIR, list(GC_FILES), out_root=tmp_path)
    assert set(done) == set(GC_FILES)
    for rel in GC_FILES:
        assert _parts(tmp_path / rel) == _parts(BACKEND / "templates" / rel), (
            f"{rel} is not what annotate_templates.py makes of Kyle's raw file. Re-run "
            f"`python annotate_templates.py --from \"../docs/GC Templates\" \"{rel}\"` or fix the "
            f"rule that drifted; never hand-edit the installed file.")
        after = hashlib.sha256((RAW_DIR / pathlib.Path(rel).name).read_bytes()).hexdigest()
        assert after == before[rel], f"annotating changed Kyle's raw file {rel}"


def test_a_rule_that_no_longer_matches_refuses_instead_of_shipping_a_placeholder(tmp_path):
    """The GC header/span rules are strict: a re-saved form whose wording moved raises, and the
    file being annotated is left exactly as it was (no half-written template, no .tmp)."""
    rel = POLISH
    dst = tmp_path / rel
    dst.parent.mkdir(parents=True)
    raw = (RAW_DIR / pathlib.Path(rel).name).read_bytes()
    # Kyle "re-saves" the form with the spec number edited.
    with zipfile.ZipFile(io.BytesIO(raw)) as zin, zipfile.ZipFile(dst, "w", zipfile.ZIP_DEFLATED) as zout:
        for item in zin.infolist():
            data = zin.read(item.filename)
            if item.filename == "word/document.xml":
                data = data.replace(b"3543</w:t>", b"3544</w:t>")
            zout.writestr(item, data)
    edited = dst.read_bytes()
    with pytest.raises(ValueError):
        at.annotate_one(dst, at.TEMPLATE_RULES[rel], rel)
    assert dst.read_bytes() == edited
    assert not dst.with_suffix(dst.suffix + ".tmp").exists()


def test_the_header_rule_is_anchored_on_the_header_box():
    """Only the box whose first three lines ARE "xx" / "xx, KS " / "xx, MO " -- never "Spec 03xx",
    "Notes: xx" or the "xx gallons/kits" scope step, which keep Kyle's text."""
    for rel in GC_FILES:
        walk = _walk(BACKEND / "templates" / rel)
        assert [b["text"] for b in walk[99:102]] == ["{{job_name}}", "{{city_state}}", ""], rel
        assert walk[99]["txbx"] == walk[100]["txbx"] == walk[101]["txbx"] == 0, rel
        texts = [b["text"] for b in walk]
        assert sum(t.count("{{job_name}}") for t in texts) == 1, rel
    resinous = [b["text"] for b in _walk(BACKEND / "templates" / RESINOUS)]
    sealer = [b["text"] for b in _walk(BACKEND / "templates" / SEALER)]
    assert "Notes: xx" in resinous and "Notes: xx" in sealer


# ── 2. the token set: everything the old files carried, plus the intake's ─────────────────────
@pytest.mark.parametrize("rel", GC_FILES)
def test_the_new_file_carries_every_old_token_plus_the_intake_ones(rel):
    new = _tokens(BACKEND / "templates" / rel)
    old = _tokens(PREDECESSORS / pathlib.Path(rel).name)
    assert set(new) == set(old) | NEW_TOKENS, (set(new) ^ (set(old) | NEW_TOKENS))
    for tok in old:
        assert new[tok] == old[tok], (rel, tok, old[tok], new[tok])
    for tok in NEW_TOKENS:
        assert new[tok] == 2, f"{rel}: {{{{{tok}}}}} in {new[tok]} paragraph(s); expected the " \
                              f"mc:Choice and mc:Fallback copy of one line"


def test_every_token_the_gc_rules_add_has_a_default_and_nothing_else_does():
    """A span token with no default would print raw on a blank field; a default for a token the
    file does not print is dead weight that hides a renamed token."""
    assert set(pw.TEMPLATE_TOKEN_DEFAULTS) == set(at.GC_SPAN_RULES) == set(GC_FILES)
    for rel, rules in at.GC_SPAN_RULES.items():
        added = {t for _anchor, spans in rules for _s, r in spans for t in _TOKEN.findall(r)}
        assert added == set(pw.TEMPLATE_TOKEN_DEFAULTS[rel]), rel
        assert added <= set(_tokens(BACKEND / "templates" / rel)), rel


@pytest.mark.parametrize("rel", GC_FILES)
def test_each_default_is_the_text_its_rule_replaced(rel):
    """The literal in TEMPLATE_TOKEN_DEFAULTS, put back into the rule's replacement, gives back the
    exact text the rule took out of Kyle's file ("[PC]" for "[{{finish_tag}}]")."""
    d = pw.TEMPLATE_TOKEN_DEFAULTS[rel]
    for _anchor, spans in at.GC_SPAN_RULES[rel]:
        for search, repl in spans:
            back = _TOKEN.sub(lambda m: d[m.group(1)], repl)
            assert back == search, (rel, search, repl, back)


# ── 3. a blank field prints the line Kyle's raw file prints ───────────────────────────────────
@pytest.mark.parametrize("rel", GC_FILES)
def test_a_blank_field_prints_exactly_what_kyles_file_prints(rel):
    """THE PIN: render the installed file with every new field blank (and the two old tokens on
    those lines set to the numbers Kyle typed), and the spec, area and addenda lines are his raw
    file's own lines, character for character."""
    out = pw.fill_proposal(work_type=WORK_TYPE[rel], audience="GC",
                           values={"sqft": "1,600", "cove_lf": "500", "architect": "",
                                   "spec_section": None, "addenda_count": None})
    got = [b["text"] for b in _walk(out)]
    raw = [b["text"] for b in _walk(RAW_DIR / pathlib.Path(rel).name)]
    for marker in ("Drawings by", "Finish Schedule", "Addenda Acknowledged"):
        want = [t for t in raw if marker in t]
        assert len(want) == 1, (rel, marker, want)
        assert [t for t in got if marker in t] == want, (rel, marker)


@pytest.mark.parametrize("rel", GC_FILES)
def test_intake_values_land_on_the_spec_finish_and_addenda_lines(rel):
    out = pw.fill_proposal(work_type=WORK_TYPE[rel], audience="GC", values={
        "sqft": "1,600", "cove_lf": "500", "job_name": "Acme Warehouse", "city_state": "Olathe, KS",
        "spec_section": "09 67 23", "architect": "Gould Evans", "drawings_dated_formatted": "8/15/26",
        "finish_tag": "RES-1", "plan_sheet": "A601", "addenda_count": 3})
    got = [b["text"] for b in _walk(out)]
    spec = [t for t in got if "Drawings by" in t]
    assert len(spec) == 1 and "per Spec 09 67 23 & Drawings by Gould Evans dated 8/15/26 (NO spec)" in spec[0]
    area = [t for t in got if "Finish Schedule" in t]
    assert len(area) == 1 and "[RES-1]" in area[0] and area[0].endswith("per the Finish Schedule on A601")
    assert "Addenda Acknowledged: 3" in got
    assert got[99:102] == ["Acme Warehouse", "Olathe, KS", ""]
    assert not any("{{" in t for t in spec + area), "a raw token on the spec / area line"


def test_the_defaults_never_reach_the_callers_values():
    """fill_proposal fills a COPY: /api/generate writes `values` back onto the draft and fills the
    estimate sheet from them, so "xx Architects" must never become the job's architect."""
    values = {"sqft": "1,600", "architect": "", "spec_section": None}
    snapshot = dict(values)
    pw.fill_proposal(work_type="polish", audience="GC", values=values)
    assert values == snapshot


def test_a_zero_addenda_count_is_a_value_not_a_blank():
    out = pw.fill_proposal(work_type="polish", audience="GC", values={"addenda_count": 0})
    assert "Addenda Acknowledged: 0" in [b["text"] for b in _walk(out)]
    assert pw._with_token_defaults({"addenda_count": 0}, {"addenda_count": "x"}) == {"addenda_count": 0}
    assert pw._with_token_defaults({"addenda_count": " "}, {"addenda_count": "x"}) == {"addenda_count": "x"}


def test_only_the_gc_files_have_defaults():
    for (wt, aud), rel in pw.TEMPLATE_PICKER.items():
        want = pw.TEMPLATE_TOKEN_DEFAULTS.get(rel, {})
        assert pw.template_token_defaults(wt, aud) == want, (wt, aud)
    assert pw.template_token_defaults("gyp", "GC") == {}, "the Gyp form is unchanged and has none"


def test_the_drawings_date_is_formatted_like_the_header_date():
    for raw, want in (("2026-08-15", "8/15/26"), ("2026-12-01", "12/1/26"), ("8/1/26", "8/1/26"),
                      ("", None), (None, None)):
        values = {"drawings_dated": raw}
        main._ensure_value_aliases(values, "GC")
        assert values.get("drawings_dated_formatted") == want, (raw, values)
    # What the caller already formatted wins, as bid_date_formatted's backfill does.
    values = {"drawings_dated": "2026-08-15", "drawings_dated_formatted": "Aug 15"}
    main._ensure_value_aliases(values, "GC")
    assert values["drawings_dated_formatted"] == "Aug 15"


# ── 4. through the route: no raw token, and the customer sees only what was meant ─────────────
def _generate(body):
    r = client.post("/api/generate", json=body)
    assert r.status_code == 200, r.text
    return client.get(r.json()["docx_download_url"]).content


@pytest.mark.parametrize("work_type", ["polish", "epoxy", "combo"])
def test_a_generated_gc_proposal_with_the_new_fields_blank_prints_no_raw_token(work_type):
    # `texture` is sent blank, as the browser sends it (computeTokenValues): it is caller data
    # (test_customer_paper_is_clean._CALLER_DATA_TOKENS), not one of these fields.
    out = _generate({"work_type": work_type, "audience": "GC", "values": {
        "project_name": "Blank Drawings QA", "bid_date": "2026-10-02", "city_state": "Olathe, KS",
        "sqft": "1,600", "cove_lf": "500", "texture": "", "total_formatted": "$6,307"}})
    lines = _rendered_lines(out)
    assert not [t for t in lines if "{{" in t], [t for t in lines if "{{" in t]
    assert "Blank Drawings QA" in lines and "Olathe, KS" in lines
    assert "Addenda Acknowledged: 0" in lines


@pytest.mark.parametrize("rel", GC_FILES)
def test_what_changed_for_the_customer(rel, monkeypatch):
    """The same values rendered on the OLD file and on the new one. With the new intake fields
    blank, every printed line is the same except: the job header (the job name and city, which the
    old file printed as "xx" / "xx, KS " / "xx, MO "), the spec line's date (Kyle's 8/1/26 for his
    4/1/26) and, on Resinous, Kyle's flake wording. Nothing else a customer reads moved."""
    values = {"job_name": "Acme Warehouse", "city_state": "Olathe, KS", "bid_date_formatted": "10/2/26",
              "estimator_name": "Kyle Loseke", "sqft": "1,600", "cove_lf": "500", "texture": "Light",
              "scope_notes": "s", "schedule_notes": "t", "exclusions": "e",
              "base_bid_formatted": "$6,307", "base_tax_phrase": "(material sales tax INCLUDED)",
              "material_tax_formatted": "$0.00", "tax_amount_formatted": "$0.00",
              "total_formatted": "$6,307"}
    new = _rendered_lines(pw.fill_proposal(work_type=WORK_TYPE[rel], audience="GC", values=dict(values)))
    old_path = PREDECESSORS / pathlib.Path(rel).name
    monkeypatch.setattr(pw, "pick_template", lambda work_type, audience: old_path)
    old = _rendered_lines(pw.fill_proposal(work_type=WORK_TYPE[rel], audience="GC", values=dict(values)))
    assert len(old) == len(new)
    diff = [(a, b) for a, b in zip(old, new) if a != b]
    got_header = [d for d in diff if d[0] in ("xx", "xx, KS ", "xx, MO ")]
    assert got_header == [("xx", "Acme Warehouse"), ("xx, KS ", "Olathe, KS"), ("xx, MO ", "")], diff
    rest = [d for d in diff if d not in got_header]
    spec = [d for d in rest if "Drawings by" in d[0]]
    assert len(spec) == 1 and spec[0][0].replace("4/1/26", "8/1/26") == spec[0][1], spec
    rest = [d for d in rest if d not in spec]
    if rel == RESINOUS:
        assert len(rest) == 1, rest
        a, b = rest[0]
        assert "(sizes 1/4” -or- 1/2”)" in a and b == a.replace("(sizes 1/4” -or- 1/2”)", "(macro sizes: 1/4”+)")
    else:
        assert rest == [], rest


# ── 5. no saved edit is lost: the predecessor versions ────────────────────────────────────────
@pytest.mark.parametrize("rel", GC_FILES)
def test_the_predecessor_entry_is_pinned_to_the_walk(rel):
    """The claim PREDECESSOR_VERSIONS makes -- an edit saved against the old file lands on the same
    paragraph of the new one -- re-proven against the old bytes: same number of blocks, and block
    by block the same kind, text box and region; the words differ only where the swap meant them
    to. The entry names the file's CURRENT content, so it stops applying if the file changes."""
    current, preds = tv.PREDECESSOR_VERSIONS[rel]
    assert tv.content_version(BACKEND / "templates" / rel) == current, (
        f"{rel} changed since its PREDECESSOR_VERSIONS entry was written: re-prove the walk and "
        f"write a new entry, or delete it")
    fixture = PREDECESSORS / pathlib.Path(rel).name
    assert preds == (tv.content_version(fixture),), "the fixture is not the predecessor's bytes"
    old, new = _walk(fixture), _walk(BACKEND / "templates" / rel)
    assert len(old) == len(new)
    shape = lambda w: [(b["id"], b["kind"], b["txbx"], b["in_block"]) for b in w]
    assert shape(old) == shape(new)
    assert {a["id"] for a, b in zip(old, new) if a["text"] != b["text"]} == CHANGED_IDS[rel]
    geo = lambda p: pw.template_geometry(docx.Document(str(p)))
    assert geo(fixture) == geo(BACKEND / "templates" / rel), "the boxes moved"


def test_no_other_template_has_a_predecessor():
    assert set(tv.PREDECESSOR_VERSIONS) == set(GC_FILES)
    for rel in pw.TEMPLATE_PICKER.values():
        if rel not in GC_FILES:
            assert tv.predecessor_versions(BACKEND / "templates" / rel) == ()


@pytest.mark.parametrize("work_type,rel", [("polish", POLISH), ("epoxy", RESINOUS),
                                           ("combo", RESINOUS), ("sealer", SEALER)])
def test_an_edit_saved_against_the_old_file_still_prints(work_type, rel):
    """polish:GC, epoxy:GC and combo:GC are the keys the real drafts' edits are saved under."""
    old_hash = tv.PREDECESSOR_VERSIONS[rel][1][0]
    spec_id = sorted(CHANGED_IDS[rel] - {99, 100, 101})[0]          # the spec line
    body = {"work_type": work_type, "audience": "GC",
            "values": {"project_name": "Predecessor QA", "sqft": "1,600", "cove_lf": "500"},
            "paragraph_overrides": [{"id": 99, "text": "OLD-HEADER-EDIT"},
                                    {"id": spec_id, "text": "Polished: per Spec 1 OLD-SPEC-EDIT"}],
            "template_version": old_hash}
    lines = _rendered_lines(_generate(body))
    assert "OLD-HEADER-EDIT" in lines and any("OLD-SPEC-EDIT" in t for t in lines)
    body["template_version"] = "sha256:0000000000000000"
    lines = _rendered_lines(_generate(body))
    assert "OLD-HEADER-EDIT" not in lines, "an unrelated stamp was accepted"
    assert not tv.accepts("sha256:0000000000000000", BACKEND / "templates" / rel)
    assert tv.accepts(old_hash, BACKEND / "templates" / rel)
    assert tv.accepts(tv.content_version(BACKEND / "templates" / rel), BACKEND / "templates" / rel)


def test_a_predecessor_stamp_is_refused_once_the_file_changes_again(tmp_path, monkeypatch):
    """The entry is tied to the content it was proven against: the same file with one more byte is
    a different file, and the old stamps are refused there, as a legacy entry's are."""
    rel = POLISH
    copy = tmp_path / "templates" / rel
    copy.parent.mkdir(parents=True)
    copy.write_bytes((BACKEND / "templates" / rel).read_bytes() + b"\0")
    monkeypatch.setattr(tv, "TEMPLATES_ROOT", tmp_path / "templates")
    assert tv.predecessor_versions(copy) == ()
    assert not tv.accepts(tv.PREDECESSOR_VERSIONS[rel][1][0], copy)


@pytest.mark.parametrize("work_type,rel", [("polish", POLISH), ("epoxy", RESINOUS),
                                           ("combo", RESINOUS), ("sealer", SEALER)])
def test_the_editor_is_told_the_predecessors_and_the_defaults(work_type, rel):
    j = client.get(f"/api/proposal-template?work_type={work_type}&audience=GC").json()
    assert j["template_version"] == tv.content_version(BACKEND / "templates" / rel)
    assert j["template_version_predecessors"] == list(tv.PREDECESSOR_VERSIONS[rel][1])
    assert j["token_defaults"] == pw.TEMPLATE_TOKEN_DEFAULTS[rel]


def test_a_direct_template_has_neither():
    j = client.get("/api/proposal-template?work_type=polish&audience=Direct").json()
    assert j["template_version_predecessors"] == [] and j["token_defaults"] == {}


def _lift(path, name):
    src = path.read_text(encoding="utf-8")
    m = re.search(r"^  function " + re.escape(name) + r"\b", src, re.M)
    assert m, f"{name} not found in {path.name}"
    depth, j = 0, src.index("{", m.start())
    for k in range(j, len(src)):
        depth += {"{": 1, "}": -1}.get(src[k], 0)
        if depth == 0:
            return src[m.start():k + 1]
    raise AssertionError("unterminated " + name)


def _node(script):
    proc = subprocess.run(["node", "-e", script], capture_output=True, text=True,
                          encoding="utf-8", timeout=60)
    assert proc.returncode == 0, proc.stderr
    return json.loads(proc.stdout.strip().splitlines()[-1])


@pytest.mark.skipif(shutil.which("node") is None, reason="node is not installed")
@pytest.mark.parametrize("rel", GC_FILES)
def test_the_editor_restores_what_the_backend_would_apply(rel):
    """The SHIPPED savedVersionMatches, given what /api/proposal-template sends for the file,
    answers every stamp the way template_versions.accepts does."""
    path = BACKEND / "templates" / rel
    cur, old = tv.content_version(path), tv.PREDECESSOR_VERSIONS[rel][1][0]
    stamps = [cur, old, "sha256:0000000000000000", "", "1788900501000000000", "STALE"]
    js = (FRONTEND_JS / "proposal-review.js")
    got = _node(_lift(js, "savedVersionMatches") + f"""
let templateVersion = {json.dumps(cur)}; let templateLegacyFloorS = {tv.legacy_floor_s(path)};
let templatePredecessors = {json.dumps(list(tv.predecessor_versions(path)))};
console.log(JSON.stringify({json.dumps(stamps)}.map(savedVersionMatches)));""")
    assert got == [tv.accepts(s, path) if s else False for s in stamps]
    assert got == [True, True, False, False, False, False]


@pytest.mark.skipif(shutil.which("node") is None, reason="node is not installed")
def test_the_editor_picks_both_up_from_the_template_response():
    """The load-time lines that read the two new fields, RUN against the endpoint's real answer: a
    renamed key on either side leaves the editor refusing the drafts' edits and drawing blanks."""
    src = (FRONTEND_JS / "proposal-review.js").read_text(encoding="utf-8").replace("\r\n", "\n")
    m = re.search(r"^ *templatePredecessors = [^;]+;\n *templateTokenDefaults = [^;]+;", src, re.M)
    assert m, "initDocumentEditor no longer reads template_version_predecessors / token_defaults"
    body = client.get("/api/proposal-template?work_type=polish&audience=GC").json()
    j = {"template_version_predecessors": body["template_version_predecessors"],
         "token_defaults": body["token_defaults"]}
    got = _node("let templatePredecessors = []; let templateTokenDefaults = {}; const j = %s;\n%s\n"
                "console.log(JSON.stringify([templatePredecessors, templateTokenDefaults]));"
                % (json.dumps(j), m.group(0)))
    assert got == [body["template_version_predecessors"], body["token_defaults"]]


@pytest.mark.skipif(shutil.which("node") is None, reason="node is not installed")
@pytest.mark.parametrize("work_type,rel", [("polish", POLISH), ("epoxy", RESINOUS), ("sealer", SEALER)])
def test_the_editor_shows_what_the_document_prints(work_type, rel):
    """PARITY, executed on both sides. The editor's real fill (withTokenDefaults + fillPlain)
    over the blocks /api/proposal-template serves, against the paragraphs fill_proposal prints for
    the same values -- blank new fields, and filled ones."""
    tpl = client.get(f"/api/proposal-template?work_type={work_type}&audience=GC").json()
    ids = sorted(CHANGED_IDS[rel] - ({114} if rel == RESINOUS else set()))
    blocks = {b["id"]: b["text"] for b in tpl["blocks"] if b["id"] in ids}
    js = FRONTEND_JS / "proposal-review.js"
    for values in ({"job_name": "Acme", "city_state": "Olathe, KS", "sqft": "1,600", "cove_lf": "500",
                    "spec_section": "", "architect": None},
                   {"job_name": "Acme", "city_state": "Olathe, KS", "sqft": "1,600", "cove_lf": "500",
                    "spec_section": "033543-A", "architect": "Gould Evans",
                    "drawings_dated_formatted": "8/15/26", "finish_tag": "PC-2", "plan_sheet": "A1.01",
                    "addenda_count": 4}):
        screen = _node(
            "const DOC_TOKEN_RE = /\\{\\{\\s*([a-zA-Z_][a-zA-Z0-9_]*)\\s*\\}\\}/g;\n"
            + _lift(js, "fillPlain") + "\n" + _lift(js, "withTokenDefaults") + "\n"
            + "let templateTokenDefaults = %s;\n" % json.dumps(tpl["token_defaults"])
            + "const tokens = withTokenDefaults(%s);\n" % json.dumps(values)
            + "const blocks = %s;\n" % json.dumps(blocks)
            + "const out = {}; for (const id in blocks) out[id] = fillPlain(blocks[id], tokens);\n"
            + "console.log(JSON.stringify(out));")
        doc = [b["text"] for b in _walk(pw.fill_proposal(work_type=work_type, audience="GC",
                                                         values=dict(values)))]
        # The job header and the WORK box come before anything the fill adds or removes, so their
        # ids are the template's; the addenda line sits below the PRICE box, whose Options gap the
        # fill redraws, so it is found by its words.
        for bid in ids:
            if bid < 128:
                assert screen[str(bid)] == doc[bid], (rel, bid, screen[str(bid)], doc[bid])
            else:
                assert screen[str(bid)] in doc, (rel, bid, screen[str(bid)])


@pytest.mark.skipif(shutil.which("node") is None, reason="node is not installed")
def test_the_editor_and_the_backend_format_the_drawings_date_the_same():
    """computeTokenValues' drawings_dated_formatted against main._ensure_value_aliases."""
    js = (FRONTEND_JS / "proposal-review.js").read_text(encoding="utf-8").replace("\r\n", "\n")
    m = re.search(r"drawings_dated_formatted: \(\(\) => \{\n(.*?)\n      \}\)\(\),", js, re.S)
    assert m, "computeTokenValues no longer derives drawings_dated_formatted inline"
    cases = ["2026-08-15", "2026-12-01", "2026-01-09T00:00", "8/1/26", " 2026-03-04 ", ""]
    got = _node("const out = %s.map((x) => { const mergedValues = { drawings_dated: x };\n%s\n });\n"
                "console.log(JSON.stringify(out));" % (json.dumps(cases), m.group(1)))
    for raw, js_value in zip(cases, got):
        values = {"drawings_dated": raw}
        main._ensure_value_aliases(values, "GC")
        assert (values.get("drawings_dated_formatted") or "") == js_value, (raw, values, js_value)


# ── 6. intake -> the estimate sheet ───────────────────────────────────────────────────────────
INTAKE = {"project_name": "Specs QA", "architect": "Gould Evans", "drawings_dated": "2026-08-15",
          "spec_section": "033543", "finish_tag": "PC", "plan_sheet": "A900", "addenda_count": 3}


def _wb(values, **kw):
    return load_workbook(io.BytesIO(ew.fill_estimate(values, **kw)))


def test_the_specs_tab_gets_the_drawing_spec_and_addenda_rows():
    ws = _wb(INTAKE)[ew.SPECS_SHEET]
    assert (ws["A3"].value, ws["B3"].value, ws["C3"].value, ws["D3"].value) == (
        "A900", "Finish Schedule", None, "PC")
    assert ws["A14"].value == "033543" and ws["A14"].data_type == "s", "the spec section lost its zero"
    assert ws["B14"].value is None, "a description was invented"
    assert [ws[f"A{r}"].value for r in range(25, 29)] == [1, 2, 3, None]
    assert all(ws[f"B{r}"].value is None for r in range(25, 29)), "an addendum's impact was invented"
    # The tables' own headers are untouched.
    assert (ws["A2"].value, ws["A13"].value, ws["A24"].value, ws["B24"].value) == (
        "Drawing", "Section", "#", "Impact to Scope")


def test_a_blank_intake_writes_nothing_on_the_specs_tab():
    blank = _wb({"project_name": "Specs QA"})[ew.SPECS_SHEET]
    template = load_workbook(str(ew.TEMPLATE_PATH))[ew.SPECS_SHEET]
    for row in template.iter_rows():
        for cell in row:
            assert blank[cell.coordinate].value == cell.value, cell.coordinate
    assert ew.specs_tab_cells({"addenda_count": None, "spec_section": "  "}) == {}


def test_the_addenda_rows_stop_at_the_table():
    cells = ew.specs_tab_cells({"addenda_count": 40})
    assert sorted(int(a[1:]) for a in cells) == list(range(25, 54))
    assert cells["A53"] == 29
    for bad in ("2.5", -1, "two", True):
        assert ew.specs_tab_cells({"addenda_count": bad}) == {}, bad


def test_a_typed_formula_is_text_on_the_specs_tab():
    cells = ew.specs_tab_cells({"plan_sheet": "=HYPERLINK(\"x\")", "spec_section": "+1"})
    assert cells["A3"].startswith("'=") and cells["A14"] == "'+1"


def test_the_specs_tab_only_writes_cells_the_template_leaves_empty():
    template = load_workbook(str(ew.TEMPLATE_PATH))[ew.SPECS_SHEET]
    for addr in ew.specs_tab_cells(dict(INTAKE, addenda_count=29)):
        assert template[addr].value is None, addr


def test_a_cell_the_template_already_fills_is_never_overwritten():
    """The emptiness guard, made to fire: today every target cell is empty in Kyle's workbook, so
    the guard is only proven on a sheet where it is not -- his own words and a formula stay."""
    wb = load_workbook(str(ew.TEMPLATE_PATH))
    ws = wb[ew.SPECS_SHEET]
    ws["A3"] = "Kyle's drawing"
    ws["A25"] = "=1+1"
    n = ew._write_specs_tab(ws, INTAKE)
    assert ws["A3"].value == "Kyle's drawing" and ws["A25"].value == "=1+1"
    assert ws["A14"].value == "033543" and ws["A26"].value == 2
    assert n == len(ew.specs_tab_cells(INTAKE)) - 2


def test_the_estimators_grid_edit_still_wins():
    ws = _wb(INTAKE, cell_values={"Specs+Dwgs+Addn!A3": "A901"})[ew.SPECS_SHEET]
    assert ws["A3"].value == "A901"


def test_drawings_dated_reaches_the_base_tabs_and_never_a_formula():
    wb = _wb(INTAKE)
    assert wb["Epoxy"]["B9"].value == "2026-08-15"
    assert wb["Polish"]["B9"].value == "=Epoxy!B9", "Polish!B9 is Kyle's mirror formula"
    assert wb['Gyp (USG 1-8")']["B11"].value == "2026-08-15"
    assert wb['Gyp (USG N12ULTRA)']["B11"].value == "='Gyp (USG 1-8\")'!B11"
    assert wb["Epoxy"]["B8"].value == "Gould Evans"


def test_the_estimate_grid_seeds_drawings_dated_where_it_seeds_architect():
    src = (FRONTEND_JS / "estimate-review.js").read_text(encoding="utf-8")
    assert re.search(r'^\s*drawings_dated:\s*"Epoxy!B9",', src, re.M)
    assert re.search(r'^\s*drawings_dated:\s*`\$\{GYP_BASE\}!B11`,', src, re.M)
    assert ew.EPOXY_CELL_MAP["drawings_dated"] == "B9" and ew.GYP_CELL_MAP["drawings_dated"] == "B11"


def test_generate_fills_the_sheet_from_the_intake_values():
    r = client.post("/api/generate", json={"work_type": "polish", "audience": "GC",
                                           "values": dict(INTAKE, sqft="1,600")})
    assert r.status_code == 200, r.text
    wb = load_workbook(io.BytesIO(client.get(r.json()["xlsx_download_url"]).content))
    ws = wb[ew.SPECS_SHEET]
    assert (ws["A3"].value, ws["A14"].value, ws["A27"].value) == ("A900", "033543", 3)
    assert wb["Epoxy"]["B8"].value == "Gould Evans" and wb["Epoxy"]["B9"].value == "2026-08-15"
    blank = client.post("/api/generate", json={"work_type": "polish", "audience": "GC",
                                               "values": {"project_name": "Specs QA"}})
    wb = load_workbook(io.BytesIO(client.get(blank.json()["xlsx_download_url"]).content))
    assert wb["Epoxy"]["B8"].value is None, "the proposal's 'xx Architects' leaked onto the sheet"
    assert wb[ew.SPECS_SHEET]["A14"].value is None


# ── 7. the intake form ────────────────────────────────────────────────────────────────────────
def test_the_intake_form_asks_for_them_all_optionally():
    html = (REPO / "frontend" / "index.html").read_text(encoding="utf-8")
    box = html[html.index('<fieldset id="drawings-specs-box">'):]
    box = box[:box.index("</fieldset>")]
    want = {"drawings_dated": "date", "spec_section": "text", "finish_tag": "text",
            "plan_sheet": "text", "addenda_count": "number"}
    for name, typ in want.items():
        m = re.search(r'<input\b[^>]*\bname="%s"[^>]*>' % name, box)
        assert m, f"{name} is not in the Drawings & specs box"
        assert f'type="{typ}"' in m.group(0), (name, m.group(0))
        assert "required" not in m.group(0), f"{name} must stay optional"
    addenda = re.search(r'<input\b[^>]*\bname="addenda_count"[^>]*>', box).group(0)
    assert 'min="0"' in addenda and 'step="1"' in addenda
    # Right after the Project Info box, which ends with Architect / GC: nothing but a comment
    # between the two.
    before = html[:html.index('<fieldset id="drawings-specs-box">')]
    before = re.sub(r"<!--.*?-->", "", before, flags=re.S).rstrip()
    assert before.endswith("</fieldset>") and 'name="architect"' in before[-400:]
